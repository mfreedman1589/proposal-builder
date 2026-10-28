"""The dev-only QA page and the QA loop's helpers (qa_ledger.py, deck_text.py,
db's qa_findings functions). The table itself is stubbed with an in-memory
fake -- a test never writes real rows, the same rule test_feedback.py follows.

    python tests/test_qa_page.py
"""
import os
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
os.chdir(REPO)
os.environ["PROPOSAL_BUILDER_TEST_MODE"] = "1"
os.environ.pop("PROPOSAL_BUILDER_DEV_MODE", None)

import db  # noqa: E402
db.fetch_audiences = lambda: (None, "stubbed for test isolation")

import deck_text  # noqa: E402
import qa_ledger  # noqa: E402
from streamlit.testing.v1 import AppTest  # noqa: E402

failures = []


def check(label, condition, detail=""):
    if condition:
        print(f"  PASS  {label}")
    else:
        print(f"  FAIL  {label}{'  -- ' + str(detail) if detail else ''}")
        failures.append(label)


class FakeQA:
    def __init__(self, rows):
        self.rows = [dict(r) for r in rows]
        self.updates = []
        self.submitted = []

    def fetch(self, limit=1000):
        return sorted(self.rows, key=lambda r: db._qa_number(r["qa_id"])), None

    def submit(self, fields, state=None, build_stamp=None, created_by=None):
        row = dict(fields, qa_id=db.next_qa_id(self.rows), status="Open",
                   build_stamp=build_stamp, created_by=created_by, state_json=state)
        self.rows.append(row)
        self.submitted.append(row)
        return row, None

    def update(self, qa_id, **changes):
        self.updates.append((qa_id, changes))
        for row in self.rows:
            if row["qa_id"] == qa_id:
                row.update(changes)
        return True, None


SEED = [{"qa_id": "QA-008", "title": "Regency flight end", "severity": "Low", "status": "By design",
         "triage_verdict": "answer-key error"},
        {"qa_id": "QA-002", "title": "Earlier fix", "severity": "Medium", "status": "Fixed",
         "fix_commit": "abc1234"}]


def check_helpers():
    print("\nHelpers -- IDs, the merge gate, the ledger export, deck text")
    check("first new finding is QA-009 when only the first session's IDs exist",
          db.next_qa_id(SEED) == "QA-009", db.next_qa_id(SEED))
    check("IDs continue past the highest, never reuse", db.next_qa_id(SEED + [{"qa_id": "QA-014"}])
          == "QA-015")
    rows = SEED + [{"qa_id": "QA-009", "severity": "High", "status": "Open"},
                   {"qa_id": "QA-010", "severity": "Critical", "status": "Triaged"},
                   {"qa_id": "QA-011", "severity": "Critical", "status": "Fixed"},
                   {"qa_id": "QA-012", "severity": "Medium", "status": "Open"}]
    check("the merge gate names only Critical/High that are Open or Triaged",
          db.merge_blocking_findings(rows) == ["QA-009", "QA-010"], db.merge_blocking_findings(rows))
    on_main = rows + [{"qa_id": "QA-013", "severity": "High", "status": "Open", "also_on_main": True}]
    check("...and never one marked 'also on main' (the public app already has it)",
          db.merge_blocking_findings(on_main) == ["QA-009", "QA-010"],
          db.merge_blocking_findings(on_main))
    check("the ledger shows the 'also on main' column",
          "Also on main" in qa_ledger.ledger_markdown(on_main))
    md = qa_ledger.ledger_markdown(SEED)
    check("the ledger export lists every finding and says it's generated",
          "QA-008" in md and "QA-002" in md and "do not edit by hand" in md)

    original = db.fetch_qa_findings
    try:
        db.fetch_qa_findings = lambda limit=1000: (rows, None)
        check("`qa_ledger.py --gate` exits 1 while something blocks", qa_ledger.main(["--gate"]) == 1)
        db.fetch_qa_findings = lambda limit=1000: (SEED, None)
        check("...and 0 when nothing does", qa_ledger.main(["--gate"]) == 0)
        db.fetch_qa_findings = lambda limit=1000: (None, "offline")
        check("...and 2 (never 0) when the table can't be read", qa_ledger.main(["--gate"]) == 2)
    finally:
        db.fetch_qa_findings = original

    template = REPO / "REPORT_MASTER_v0_15.pptx"
    if template.exists():
        text = deck_text.deck_text(str(template))
        check("deck text names each slide with its key and renders tables",
              "`report:analyst_inventory`" in text and "| Model | Viewed | Sold | Look-to-Book |"
              in text, text[:300])
    else:
        print("  SKIP  REPORT_MASTER_v0_15.pptx not present")


class _FakeTable:
    """Just enough of the supabase client for db.update_qa_finding."""
    def __init__(self, rows):
        self.rows, self._op, self._row = rows, None, None

    def select(self, *_a):
        self._op = "select"
        return self

    def update(self, row):
        self._op, self._row = "update", row
        return self

    def eq(self, column, value):
        self._match = (column, value)
        return self

    def limit(self, _n):
        return self

    def execute(self):
        column, value = self._match
        hits = [r for r in self.rows if r.get(column) == value]
        if self._op == "update":
            for r in hits:
                r.update(self._row)
        return type("Result", (), {"data": [dict(r) for r in hits]})()


class _FakeClient:
    def __init__(self, rows):
        self.rows = rows

    def table(self, _name):
        return _FakeTable(self.rows)


def check_history():
    print("\nA finding's history -- changes and rulings are appended, never removed")
    entries = db.qa_history_entries({"finding_type": "Answer-key question", "status": "Open",
                                     "severity": "Low"},
                                    {"finding_type": "Bug", "status": "Open"}, by="Matt",
                                    note="Even split must always be offered.", at="2026-10-23T14:05")
    check("one entry per field that really changed, plus the ruling",
          [(e["kind"], e.get("field")) for e in entries] == [("change", "finding_type"),
                                                            ("note", None)], entries)
    rows = [{"qa_id": "QA-025", "status": "Open", "finding_type": "Answer-key question",
             "history": [{"at": "2026-10-22T10:00", "by": "Muse", "kind": "note", "note": "first"}]},
            {"qa_id": "QA-030", "status": "Open"}]
    original = db.get_client
    try:
        db.get_client = lambda: _FakeClient(rows)
        ok, error = db.update_qa_finding("QA-025", by="Matt", note="Ruling.", finding_type="Bug")
        history = rows[0]["history"]
        check("the update appends to the existing history, keeping the old entries",
              ok and not error and history[0]["note"] == "first"
              and [h["kind"] for h in history] == ["note", "change", "note"]
              and rows[0]["finding_type"] == "Bug", (ok, error, history))
        ok, error = db.update_qa_finding("QA-030", by="Matt", status="Fixed", fix_commit="abc")
        check("before the Stage 23 column exists the change still saves, and says the history "
              "didn't", ok and "Stage 23" in (error or "") and rows[1]["status"] == "Fixed",
              (ok, error))
    finally:
        db.get_client = original
    md = qa_ledger.ledger_markdown(rows)
    check("the ledger shows each finding's history, dated and signed",
          "2026-10-22 10:00 · Muse · Ruling / note: first" in md
          and "· Matt · Type: Answer-key question -> Bug" in md, md[-600:])


def _new_app():
    at = AppTest.from_file(str(REPO / "app.py"), default_timeout=300)
    at.session_state["authed"] = True
    at.session_state["current_user"] = "Muse"
    return at


def check_page():
    print("\nQA page -- dev only")
    fake = FakeQA(SEED)
    originals = (db.fetch_qa_findings, db.submit_qa_finding, db.update_qa_finding)
    db.fetch_qa_findings, db.submit_qa_finding, db.update_qa_finding = (
        fake.fetch, fake.submit, fake.update)
    try:
        at = _new_app()
        at.run()
        nav = [r for r in at.radio if r.key == "nav_section"]
        check("public app (no DEV_MODE): no QA page in the nav",
              nav and "QA" not in nav[0].options, nav and nav[0].options)

        os.environ["PROPOSAL_BUILDER_DEV_MODE"] = "1"
        at = _new_app()
        at.session_state["page_choice"] = "QA"
        at.run()
        nav = [r for r in at.radio if r.key == "nav_section"]
        check("dev app: QA page in the nav", nav and "QA" in nav[0].options)
        check("page renders without an exception", not at.exception, at.exception)
        texts = " ".join(str(getattr(m, "value", "")) for m in at.markdown)
        captions = " ".join(str(getattr(c, "value", "")) for c in at.caption)
        check("shows the running build stamp", "Running build" in captions, captions[:200])
        check("renders QA_RULES.md and the answer key",
              "QA rules" in texts and "answer key" in texts.lower(), texts[:200])
        check("ledger lists the seeded findings",
              any("QA-008" in str(getattr(e, "label", "")) for e in at.expander))

        verify = [b for b in at.button if b.key == "qapage_verify_QA-002"]
        check("a Fixed finding has a 'Mark verified' control", bool(verify))
        if verify:
            verify[0].click().run()
            check("...which marks it Verified, as the person who did it",
                  ("QA-002", {"status": "Verified", "by": "Muse"}) in fake.updates, fake.updates)
        reopen = [b for b in at.button if b.key == "qapage_reopen_QA-008"]
        check("a By-design finding can be reopened", bool(reopen))

        # Updating a filed finding: Type / Severity / Status and a ruling.
        at.selectbox(key="qapage_upd_type_QA-008").set_value("Bug")
        at.selectbox(key="qapage_upd_status_QA-008").set_value("Open")
        at.text_area(key="qapage_upd_note_QA-008").set_value(
            "Matt's ruling 2026-10-23: the end date is right.").run()
        at.button(key="qapage_upd_save_QA-008").click().run()
        update = next((c for q, c in fake.updates if q == "QA-008"), {})
        check("a filed finding's Type and Status change, with a dated ruling and who made it",
              update.get("finding_type") == "Bug" and update.get("status") == "Open"
              and update.get("note", "").startswith("Matt's ruling") and update.get("by") == "Muse"
              and "severity" not in update, update)

        title = [t for t in at.text_input if (t.key or "").startswith("qapage_title_")]
        actual = [t for t in at.text_area if (t.key or "").startswith("qapage_actual_")]
        title[0].set_value("Plan title doubles the client name")
        actual[0].set_value("WAEPA WAEPA CTV Strategy").run()
        submit = [b for b in at.button if (b.key or "").startswith("qapage_submit_")]
        submit[0].click().run()
        filed = fake.submitted[-1] if fake.submitted else {}
        check("filing assigns QA-009 and stores the build stamp and session state",
              filed.get("qa_id") == "QA-009" and filed.get("build_stamp")
              and (filed.get("state_json") or {}).get("page") == "QA", filed)
        check("...and records who filed it", filed.get("created_by") == "Muse", filed.get("created_by"))
    finally:
        os.environ.pop("PROPOSAL_BUILDER_DEV_MODE", None)
        db.fetch_qa_findings, db.submit_qa_finding, db.update_qa_finding = originals


MONEY = "$20,000 over $6,667/mo"


def _unescaped_dollars(texts):
    import re
    return [t for t in texts if re.search(r"(?<!\\)\$", t)]


def check_dollars():
    print("\nDollar amounts survive st.markdown (Streamlit reads $...$ as LaTeX)")
    import app_shell
    check("md_dollars escapes both signs", app_shell.md_dollars(MONEY)
          == r"\$20,000 over \$6,667/mo", app_shell.md_dollars(MONEY))
    check("... and wrapping twice changes nothing",
          app_shell.md_dollars(app_shell.md_dollars(MONEY)) == app_shell.md_dollars(MONEY))
    row = {"qa_id": "QA-031", "title": "Budget", "severity": "High", "status": "Open",
           "actual": f"The plan says {MONEY}.", "triage_note": f"Kept {MONEY}.",
           "history": [{"at": "2026-10-23 14:22", "by": "Muse", "kind": "note",
                        "note": f"Re-verify: '{MONEY}' rendered as '20,000'."}]}
    fake = FakeQA([row])
    original = db.fetch_qa_findings
    db.fetch_qa_findings = fake.fetch
    os.environ["PROPOSAL_BUILDER_DEV_MODE"] = "1"
    try:
        at = _new_app()
        at.session_state["page_choice"] = "QA"
        at.run()
        check("renders", not at.exception, at.exception)
        shown = [str(getattr(e, "value", "")) for e in
                 list(at.markdown) + list(at.caption) + list(at.info) + list(at.warning)]
        body = " ".join(shown)
        check("the finding's Actual, Triage and history notes keep both $",
              body.count(r"\$20,000 over \$6,667/mo") >= 3, body[:400])
        check("no markdown on the page carries an unescaped $",
              not _unescaped_dollars(shown), [t[:120] for t in _unescaped_dollars(shown)])
    finally:
        os.environ.pop("PROPOSAL_BUILDER_DEV_MODE", None)
        db.fetch_qa_findings = original


LONG_NOTE = ("2026-10-28, Muse -- re-verify. " + " ".join(
    f"Step {i}: checked line {i} of the plan, $10,000/mo held." for i in range(1, 80))
    + " END-OF-NOTE")


def _long_deck(path):
    """A slide built to be hard to show in full: a multi-line first box,
    tiles (a value box and a label box each), a long paragraph, a table."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    def box(text, top):
        frame = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(9), Inches(0.4)).text_frame
        lines = text.split("\n")
        frame.text = lines[0]
        for line in lines[1:]:
            frame.add_paragraph().text = line
    box("Executive Summary\nOf the 1,696 vehicles our audience viewed, 939 have since sold.", 0.2)
    box("$39,168,431", 1.0)
    box("Est. value sold", 1.3)
    box("73.1%", 1.6)
    box("Look-to-Book, used", 1.9)
    box("Long footnote: " + " ".join(f"clause {i} of the source line" for i in range(200))
        + " FOOTNOTE-END", 2.2)
    table = slide.shapes.add_table(2, 2, Inches(0.5), Inches(3), Inches(4), Inches(1)).table
    for (r, c), value in {(0, 0): "Model", (0, 1): "Sold", (1, 0): "Ford F-150",
                          (1, 1): "121"}.items():
        table.cell(r, c).text = value
    prs.save(path)


def _deck_text_page(path):
    import os
    os.environ["PROPOSAL_BUILDER_DEV_MODE"] = "1"
    import app
    app.render_dev_deck_text(path)


def check_full_text():
    print("\nLong text renders in full: deck text and the QA ledger (Muse, 2026-10-28)")
    import tempfile
    path = str(Path(tempfile.mkdtemp()) / "long.pptx")
    _long_deck(path)
    text = deck_text.deck_text(path)
    check("a multi-line first box keeps every line (the second isn't dropped)",
          "### Slide 1: Executive Summary" in text
          and "Of the 1,696 vehicles our audience viewed, 939 have since sold." in text, text[:300])
    check("each tile value is its own paragraph, beside its own label",
          "\n\n$39,168,431\n\nEst. value sold\n\n" in text
          and "\n\n73.1%\n\nLook-to-Book, used\n\n" in text, text)
    check("a long paragraph and the table come through whole",
          "FOOTNOTE-END" in text and "| Ford F-150 | 121 |" in text)
    os.environ["PROPOSAL_BUILDER_DEV_MODE"] = "1"
    try:
        at = AppTest.from_function(_deck_text_page, args=(path,), default_timeout=300)
        at.run()
        shown = " ".join(str(getattr(m, "value", "")) for m in at.markdown)
        check("on the page: the expander shows the whole slide, nothing cut",
              not at.exception and "FOOTNOTE-END" in shown and "939 have since sold" in shown
              and "Look-to-Book, used" in shown, at.exception or shown[-300:])
    finally:
        os.environ.pop("PROPOSAL_BUILDER_DEV_MODE", None)

    row = {"qa_id": "QA-031", "title": "Per-month budget priced as campaign total",
           "severity": "High", "status": "Open", "triage_note": "Triage: " + LONG_NOTE,
           "history": [{"at": "2026-10-28T17:14:13", "by": "Matt", "kind": "note",
                        "note": LONG_NOTE + "\nSecond paragraph of the note. SECOND-END"}]}
    fake = FakeQA([row])
    original = db.fetch_qa_findings
    db.fetch_qa_findings = fake.fetch
    os.environ["PROPOSAL_BUILDER_DEV_MODE"] = "1"
    try:
        at = _new_app()
        at.session_state["page_choice"] = "QA"
        at.run()
        shown = " ".join(str(getattr(m, "value", "")) for m in at.markdown)
        check("the ledger is a wrapping table, not a clipping grid", not at.exception
              and not list(at.dataframe) and "| QA-031 |" in shown, at.exception)
        check("... its triage note in full", shown.count("END-OF-NOTE") >= 2, shown[:300])
        check("the history note in full, both paragraphs",
              "Step 79: checked line 79" in shown and "SECOND-END" in shown, shown[-400:])
    finally:
        os.environ.pop("PROPOSAL_BUILDER_DEV_MODE", None)
        db.fetch_qa_findings = original


def main():
    check_helpers()
    check_history()
    check_page()
    check_dollars()
    check_full_text()
    print()
    print(f"{len(failures)} failure(s)" if failures else "All checks passed")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
