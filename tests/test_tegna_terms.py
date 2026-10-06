"""The media plan's terms block is TEGNA's (Standard Advertising Terms
3.2, 3.3, 9.1), and build_master_tegna_terms.py is what puts it there.

Offline: a synthetic plan slide carrying the master's own terms-box shape
(bold label + body runs, the URL as its own hyperlinked run) goes through
the real build script. The expected wording below is copied from the
request that set it, not read from the script, so a drift in the script's
constant is a failure here rather than something this agrees with.

    python tests/test_tegna_terms.py
"""
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

import assembly  # noqa: E402
import build_master_tegna_terms as build  # noqa: E402

URL = "https://bit.ly/TegnaStandardAdTerms"
EXPECTED = [
    "TEGNA Terms & Conditions: Your purchase of advertising is subject to the TEGNA "
    "Standard Advertising Terms and Conditions (“Standard Terms”), available at: " + URL,
    "Payment: Payment in full is due no later than 5 business days before the campaign "
    "start date, unless TEGNA has granted you credit terms, in which case payment is due "
    "within 30 days of invoice.",
    "Cancellation: You may cancel an Insertion Order at any time with 30 days’ prior "
    "written notice. Your campaign will continue to run during the notice period, and you "
    "are responsible for all fees for that period.",
]

failures = []


def check(label, condition, detail=""):
    print(f"  {'PASS' if condition else 'FAIL'}  {label}"
          f"{'' if condition or not detail else '  -- ' + str(detail)}")
    if not condition:
        failures.append(label)


def old_master_slide(path):
    """One plan slide shaped like the v10 master's standard variant."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.6), Inches(0.6), Inches(11), Inches(0.7)) \
        .text_frame.text = "{{CLIENT_NAME}} {{PLAN_TITLE}}"
    slide.shapes.add_table(2, 2, Inches(0.4), Inches(1.9), Inches(12.5), Inches(2))
    box = slide.shapes.add_textbox(Inches(0.39), Inches(6.70), Inches(10.17), Inches(0.72))
    frame = box.text_frame
    frame.word_wrap = True
    old = [("Premion Terms & Conditions: ", "Your purchase of advertising from Premion is "
            "subject to the Premion Standard Advertising Terms and Conditions, available at: "),
           ("Payment: ", "Payment for all Campaigns is due no later than 5 business days in "
            "advance unless you have established credit with us. "),
           ("Cancellation: ", "You may terminate an Insertion Order at any time with 10 days’ "
            "prior written notice.")]
    for i, (label, body) in enumerate(old):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        for text, bold in ((label, True), (body, False)):
            run = para.add_run()
            run.text, run.font.bold, run.font.size = text, bold, Pt(9)
            run.font.name = "Proxima Nova Light"
        if i == 0:
            link = para.add_run()
            link.text = "http://premionmedia.com/advertising-terms-and-conditions/"
            link.font.size, link.font.name = Pt(9), "Proxima Nova Light"
            link.hyperlink.address = link.text
    prs.save(path)


def terms_box(path):
    slide = Presentation(path).slides[0]
    return slide, assembly._text_shape_containing(slide, assembly._TERMS_ANCHOR)


with tempfile.TemporaryDirectory() as tmp:
    src, out, again = (str(Path(tmp) / n) for n in ("src.pptx", "out.pptx", "again.pptx"))
    old_master_slide(src)
    build.main(src, out)

    slide, box = terms_box(out)
    paragraphs = box.text_frame.paragraphs
    check("the box is still found by the co-viewing anchor", box is not None)
    check("three paragraphs, worded as TEGNA's terms",
          [p.text for p in paragraphs] == EXPECTED, [p.text for p in paragraphs])
    check("no 10-day notice anywhere", "10 days" not in box.text_frame.text)

    links = [r for p in paragraphs for r in p.runs if r.hyperlink.address]
    check("exactly one link, on the URL text, pointing at it",
          [(r.text, r.hyperlink.address) for r in links] == [(URL, URL)],
          [(r.text, r.hyperlink.address) for r in links])
    external = [r.target_ref for r in slide.part.rels.values() if r.is_external]
    check("the old link's relationship is gone", external == [URL], external)

    check("labels stay bold, bodies stay regular",
          all(p.runs[0].font.bold and not p.runs[1].font.bold for p in paragraphs))
    check("every run is set to the fitted size",
          {r.font.size.pt for p in paragraphs for r in p.runs} == {build.FONT_SIZE.pt})
    check("the box clears the Total TV station logos (ends before 10.87in)",
          (box.left + box.width) / 914400 < 10.87, (box.left + box.width) / 914400)

    build.main(out, again)
    check("re-running the build changes nothing",
          [p.text for p in terms_box(again)[1].text_frame.paragraphs] == EXPECTED)

    warning = assembly.add_coviewing_footnote(slide, "1.4x factor applied (TVision, State of "
                                              "Streaming 2025; Nielsen, 2025) -- effective CPM "
                                              "at estimated exposure $20.25.")
    check("with the co-viewing line added, the block still fits the slide", warning is None,
          warning)

print(f"\n{'FAILED: ' + ', '.join(failures) if failures else 'All checks passed.'}")
sys.exit(1 if failures else 0)
