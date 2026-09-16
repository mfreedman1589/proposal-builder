"""report:summary / report:case_study -- the one-slide summary and the
vault-bound case study (ATTRIBUTION_REPORT_PLAN.md Phase 6, REPORT_MASTER_
v0_9's two new standalone slides).

Most checks here run against a small, fully SYNTHETIC AttributionExport
(no gitignored real client file needed, never SKIPs) -- these two slides'
own fill logic doesn't depend on any real-fixture quirk, unlike the rest of
report_assembly.py's checks, which do need the real MW/Cardinal pairs. A
few checks that specifically verify the real-fixture / storage-round-trip
path use the real files and SKIP when they're absent, same convention as
test_report_assembly.py.

    python tests/test_report_summary_case_study.py
"""
import dataclasses
import sys
from datetime import date
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

import market_lookup  # noqa: E402
import package_check  # noqa: E402
from pptx import Presentation  # noqa: E402

import assembly  # noqa: E402
import attribution_import as ai  # noqa: E402
import db  # noqa: E402
import report_assembly as ra  # noqa: E402
from attribution_import import AttributionRow  # noqa: E402

TEMPLATE = REPO / "REPORT_MASTER_v0_9.pptx"
ATTRIBUTION_MW = REPO / "MW attribution excel.xlsx"
DELIVERY_MW = REPO / "MW delivery.xlsx"
PROPOSAL_MASTER = REPO / "TEGNA_MASTER_DECK_v1_1.pptx"

market_lookup.install()


class Report:
    def __init__(self):
        self.passed, self.failed, self.skipped = 0, [], []

    def check(self, label, ok, detail=""):
        if ok:
            print(f"  PASS  {label}")
            self.passed += 1
        else:
            print(f"  FAIL  {label}{'  -- ' + str(detail) if detail else ''}")
            self.failed.append(label)

    def skip(self, reason):
        print(f"  SKIP  {reason}")
        self.skipped.append(reason)


def _all_text(slide):
    """Every string a viewer could actually read on this slide -- every
    text-frame run plus every table cell -- concatenated, lowercased. The
    one shared surface every "does X ever leak/appear" check below scans."""
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text_frame.text)
    return " ".join(parts)


def make_synthetic_attribution(client_name="Acme Test Co", by_device=None, few_intent=False,
                               no_outperforming_zip=False):
    by_url = ({"https://example.com/": 100} if few_intent else {
        "https://example.com/": 100,
        "https://example.com/collections/mattresses": 300,
        "https://example.com/store-locator": 50,
        "https://example.com/about": 20,
    })
    by_zip = [
        AttributionRow(label="20002", delivered_impressions=40_000,
                      attributed_impressions=600, attributed_rate=0.015),
    ] if no_outperforming_zip else [
        AttributionRow(label="20001", delivered_impressions=50_000,
                      attributed_impressions=2000, attributed_rate=0.04),
        AttributionRow(label="20002", delivered_impressions=40_000,
                      attributed_impressions=600, attributed_rate=0.015),
    ]
    return ai.AttributionExport(
        client_name=client_name,
        flight_start=date(2026, 6, 1), flight_end=date(2026, 6, 30),
        delivered_impressions=1_000_000,
        attributed_rate=0.02,
        attributed_unique_visitors=500,
        attributed_unique_visitor_rate=0.0005,
        attributed_conversions=25, has_conversions=True,
        by_url=by_url,
        by_zip=by_zip,
        by_recency={"0 - 3 Days": 300, "4 - 7 Days": 150, "8+ Days": 50},
        by_referral_domain={"Direct": 250, "Google": 150, "Facebook": 100},
        by_device={"smartphone": 400, "desktop": 100} if by_device is None else by_device,
    )


ONE_THREAD = [{
    "head": "Fast responders", "anchor": "goal", "goal_ref": "in-store visits",
    "finding": "61% of attributed visitors responded within 3 days.",
    "meaning": "The campaign drove quick, measurable in-store intent.",
    "action": "Maintain current flight pacing next month.",
}]
FOUR_THREADS = ONE_THREAD + [
    {"head": "Strong ZIP concentration", "anchor": "signal",
     "finding": "20001 led at 2x the campaign average.",
     "meaning": "Delivery concentrated in top ZIPs paid off.", "action": ""},
    {"head": "Direct traffic dominant", "anchor": "signal",
     "finding": "50% of attributed visits arrived direct.",
     "meaning": "Brand awareness is already strong in this market.", "action": ""},
    {"head": "Weekend softness", "anchor": "signal",
     "finding": "Saturday lagged the rest of the week.",
     "meaning": "Weekend dayparting may be worth revisiting.", "action": "Review weekend pacing."},
]
ACCEPTED_OPT = [{"dimension": "zip", "value": "20099", "decision": "accepted",
                "final_text": "Reduced delivery in ZIP 20099."}]
DECLINED_OPT = [{"dimension": "creative", "value": "Banner B", "decision": "declined",
                "reason": "Still testing", "final_text": "Reduce delivery in the Banner B creative."}]


def check_standalone_slides_drop_from_normal_build(rep):
    print("\nbuild_report_deck drops both v0_9 standalone slides unconditionally")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present")
        return
    attribution = make_synthetic_attribution()
    out, warnings = ra.build_report_deck(
        str(TEMPLATE), attribution, None, str(REPO / "tests" / "_scratch_normal.pptx"),
        client_name="Acme Test Co", goals_bullets=["Grow visits"],
        whats_next_bullets=["Keep monitoring"], audience_bullets=["Adults 25-54"])
    prs = Presentation(out)
    keys = {slide_key for slide in prs.slides
           for slide_key in ([ra.slide_map.notes_key(slide)] if slide.has_notes_slide else [])}
    rep.check("report:summary is not in the built deck", "report:summary" not in keys, keys)
    rep.check("report:case_study is not in the built deck", "report:case_study" not in keys, keys)
    rep.check("built deck passes package_check", package_check.check_package(out) == [],
             package_check.check_package(out))
    Path(out).unlink(missing_ok=True)


def check_build_single_slide_isolates_each(rep):
    print("\nbuild_single_slide produces exactly one slide, correctly keyed")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present")
        return
    for key in ("report:summary", "report:case_study"):
        out = REPO / "tests" / f"_scratch_single_{key.split(':')[1]}.pptx"
        ra.build_single_slide(str(TEMPLATE), key, str(out), lambda slide: None)
        prs = Presentation(out)
        rep.check(f"{key}: exactly one slide", len(prs.slides) == 1, len(prs.slides))
        rep.check(f"{key}: it's the right slide", ra.slide_map.notes_key(prs.slides[0]) == key,
                 ra.slide_map.notes_key(prs.slides[0]))
        rep.check(f"{key}: passes package_check", package_check.check_package(str(out)) == [],
                 package_check.check_package(str(out)))
        out.unlink(missing_ok=True)


def check_summary_all_tokens_resolved(rep):
    print("\nfill_summary_slide leaves no literal {{TOKEN}} anywhere on the slide")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present")
        return
    attribution = make_synthetic_attribution()
    out = REPO / "tests" / "_scratch_summary_tokens.pptx"
    ra.build_summary_slide(str(TEMPLATE), str(out), attribution=attribution, delivery=None,
                          client_name="Acme Test Co", threads=FOUR_THREADS,
                          accepted_optimizations=ACCEPTED_OPT, include_conversions=True)
    prs = Presentation(out)
    text = _all_text(prs.slides[0])
    rep.check("no literal {{ left on the slide", "{{" not in text, text)
    rep.check("client name filled", "Acme Test Co" in text, text)
    rep.check("passes package_check", package_check.check_package(str(out)) == [])
    out.unlink(missing_ok=True)


def check_summary_takeaway_reflow_fewer_than_three(rep):
    print("\nfewer than 3 surviving threads deletes the trailing row(s)/divider(s), no reflow gap")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present")
        return
    attribution = make_synthetic_attribution()
    out = REPO / "tests" / "_scratch_summary_one_thread.pptx"
    ra.build_summary_slide(str(TEMPLATE), str(out), attribution=attribution, delivery=None,
                          client_name="Acme Test Co", threads=ONE_THREAD, accepted_optimizations=[])
    prs = Presentation(out)
    names = {sh.name for sh in prs.slides[0].shapes}
    rep.check("row 1 kept", "SummaryTakeaway1Head" in names)
    rep.check("row 2 deleted (only one thread)", "SummaryTakeaway2Head" not in names, names)
    rep.check("row 3 deleted", "SummaryTakeaway3Head" not in names, names)
    rep.check("divider 1 deleted", "SummaryTakeawayDivider1" not in names, names)
    rep.check("divider 2 deleted", "SummaryTakeawayDivider2" not in names, names)
    out.unlink(missing_ok=True)


def check_summary_tile4_conversions_vs_rate(rep):
    print("\nSUMMARY_TILE_4 shows conversions only when BOTH include_conversions and "
         "has_conversions are true -- no half-states")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present")
        return
    attribution = make_synthetic_attribution()
    out_conv = REPO / "tests" / "_scratch_tile4_conv.pptx"
    ra.build_summary_slide(str(TEMPLATE), str(out_conv), attribution=attribution, delivery=None,
                          client_name="Acme", threads=ONE_THREAD, accepted_optimizations=[],
                          include_conversions=True)
    prs = Presentation(out_conv)
    label = next(sh.text_frame.text for sh in prs.slides[0].shapes if sh.name == "SummaryTile4Label")
    rep.check("include_conversions=True + has_conversions=True -> Conversions tile", label == "Conversions", label)
    out_conv.unlink(missing_ok=True)

    out_norate = REPO / "tests" / "_scratch_tile4_rate.pptx"
    ra.build_summary_slide(str(TEMPLATE), str(out_norate), attribution=attribution, delivery=None,
                          client_name="Acme", threads=ONE_THREAD, accepted_optimizations=[],
                          include_conversions=False)
    prs2 = Presentation(out_norate)
    label2 = next(sh.text_frame.text for sh in prs2.slides[0].shapes if sh.name == "SummaryTile4Label")
    rep.check("include_conversions=False (rep's toggle off) -> rate tile, not conversions",
             label2 == "Unique Visitor Rate", label2)
    out_norate.unlink(missing_ok=True)


def check_sidebar_reflow_with_one_intent_class(rep):
    print("\nfewer than 4 intent classes deletes the unused SidebarSub rows")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present")
        return
    attribution = make_synthetic_attribution(few_intent=True)
    out = REPO / "tests" / "_scratch_sidebar_subs.pptx"
    ra.build_summary_slide(str(TEMPLATE), str(out), attribution=attribution, delivery=None,
                          client_name="Acme", threads=ONE_THREAD, accepted_optimizations=[])
    prs = Presentation(out)
    names = {sh.name for sh in prs.slides[0].shapes}
    rep.check("stat row kept (there is one intent class)", "SidebarStatValue" in names, names)
    rep.check("sub rows deleted (only one class total)", "SidebarSub1Value" not in names, names)
    out.unlink(missing_ok=True)


def check_device_split_present_and_absent(rep):
    print("\nsidebar second block (device split) fills when present, deletes cleanly when absent")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present")
        return
    with_device = make_synthetic_attribution()
    out1 = REPO / "tests" / "_scratch_device_present.pptx"
    ra.build_summary_slide(str(TEMPLATE), str(out1), attribution=with_device, delivery=None,
                          client_name="Acme", threads=ONE_THREAD, accepted_optimizations=[])
    prs1 = Presentation(out1)
    text1 = _all_text(prs1.slides[0])
    rep.check("device split text present", "smartphone" in text1, text1)
    out1.unlink(missing_ok=True)

    no_device = make_synthetic_attribution(by_device={})
    out2 = REPO / "tests" / "_scratch_device_absent.pptx"
    ra.build_summary_slide(str(TEMPLATE), str(out2), attribution=no_device, delivery=None,
                          client_name="Acme", threads=ONE_THREAD, accepted_optimizations=[])
    prs2 = Presentation(out2)
    names2 = {sh.name for sh in prs2.slides[0].shapes}
    rep.check("second-block shapes deleted, not left blank", "SidebarSecondHeader" not in names2, names2)
    rep.check("passes package_check even with the block gone",
             package_check.check_package(str(out2)) == [])
    out2.unlink(missing_ok=True)


def check_case_study_white_label_redacts_everywhere(rep):
    print("\nwhite_label=True: the client name never survives ANYWHERE on the slide, "
         "even when a thread's own head/finding/meaning names it")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present")
        return
    attribution = make_synthetic_attribution(client_name="Mattress Warehouse")
    leaky_threads = [{
        "head": "Mattress Warehouse saw strong response", "anchor": "goal",
        "finding": "Mattress Warehouse attributed visits skewed toward product pages.",
        "meaning": "Mattress Warehouse benefited from concentrated ZIP-level delivery.",
        "action": "",
    }]
    out = REPO / "tests" / "_scratch_cs_white_label.pptx"
    ra.build_case_study_slide(str(TEMPLATE), str(out), attribution=attribution, delivery=None,
                             threads=leaky_threads, accepted_optimizations=[],
                             client_name="Mattress Warehouse", vertical="retail",
                             white_label=True, budget=125_000)
    prs = Presentation(out)
    text = _all_text(prs.slides[0]).lower()
    rep.check("client name is gone from the white-labeled slide", "mattress warehouse" not in text, text)
    rep.check("passes package_check", package_check.check_package(str(out)) == [])
    out.unlink(missing_ok=True)

    out2 = REPO / "tests" / "_scratch_cs_named.pptx"
    ra.build_case_study_slide(str(TEMPLATE), str(out2), attribution=attribution, delivery=None,
                             threads=leaky_threads, accepted_optimizations=[],
                             client_name="Mattress Warehouse", vertical="retail",
                             white_label=False, budget=None)
    prs2 = Presentation(out2)
    text2 = _all_text(prs2.slides[0]).lower()
    rep.check("named mode DOES show the real client name (proves the redaction check isn't vacuous)",
             "mattress warehouse" in text2, text2)
    out2.unlink(missing_ok=True)


def check_cost_per_visitor_tile_conditional(rep):
    print("\nCsTile3 (cost per visitor) appears only with a real budget, and reflows away without one")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present")
        return
    attribution = make_synthetic_attribution()
    out_budget = REPO / "tests" / "_scratch_cs_budget.pptx"
    ra.build_case_study_slide(str(TEMPLATE), str(out_budget), attribution=attribution, delivery=None,
                             threads=ONE_THREAD, accepted_optimizations=[], client_name="Acme",
                             vertical="retail", white_label=True, budget=25_000)
    prs = Presentation(out_budget)
    names = {sh.name for sh in prs.slides[0].shapes}
    rep.check("CsTile3 present with a budget", "CsTile3Value" in names, names)
    out_budget.unlink(missing_ok=True)

    out_nobudget = REPO / "tests" / "_scratch_cs_nobudget.pptx"
    ra.build_case_study_slide(str(TEMPLATE), str(out_nobudget), attribution=attribution, delivery=None,
                             threads=ONE_THREAD, accepted_optimizations=[], client_name="Acme",
                             vertical="retail", white_label=True, budget=None)
    prs2 = Presentation(out_nobudget)
    names2 = {sh.name for sh in prs2.slides[0].shapes}
    rep.check("CsTile3 gone with no linked proposal budget", "CsTile3Value" not in names2, names2)
    rep.check("survivors still pass package_check after reflow",
             package_check.check_package(str(out_nobudget)) == [])
    out_nobudget.unlink(missing_ok=True)


def check_declined_optimization_never_surfaces(rep):
    print("\na DECLINED optimization appears in neither artifact; an ACCEPTED one appears in both")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present")
        return
    attribution = make_synthetic_attribution()
    log = ACCEPTED_OPT + DECLINED_OPT
    accepted_only = ra.accepted_optimizations_from_report_json({"optimizations": log})
    rep.check("accepted_optimizations_from_report_json excludes the declined entry",
             len(accepted_only) == 1 and accepted_only[0]["value"] == "20099", accepted_only)

    out_sum = REPO / "tests" / "_scratch_decline_summary.pptx"
    ra.build_summary_slide(str(TEMPLATE), str(out_sum), attribution=attribution, delivery=None,
                          client_name="Acme", threads=ONE_THREAD, accepted_optimizations=accepted_only)
    text_sum = _all_text(Presentation(out_sum).slides[0])
    rep.check("summary: accepted optimization text present", "20099" in text_sum, text_sum)
    rep.check("summary: declined optimization text absent", "Banner B" not in text_sum, text_sum)
    out_sum.unlink(missing_ok=True)

    out_cs = REPO / "tests" / "_scratch_decline_cs.pptx"
    ra.build_case_study_slide(str(TEMPLATE), str(out_cs), attribution=attribution, delivery=None,
                             threads=ONE_THREAD, accepted_optimizations=accepted_only,
                             client_name="Acme", vertical="retail", white_label=True)
    text_cs = _all_text(Presentation(out_cs).slides[0])
    rep.check("case study: accepted optimization text present", "20099" in text_cs, text_cs)
    rep.check("case study: declined optimization text absent", "Banner B" not in text_cs, text_cs)
    out_cs.unlink(missing_ok=True)


def check_regenerate_from_rehydrated_matches_fresh(rep):
    print("\nregenerate-later path: a report_json-round-tripped attribution/delivery produces "
         "BYTE-IDENTICAL slide text to the freshly-parsed export -- no second draft call needed")
    if not (TEMPLATE.exists() and ATTRIBUTION_MW.exists() and DELIVERY_MW.exists()):
        rep.skip("REPORT_MASTER_v0_9.pptx or the real MW fixture pair not present")
        return
    attribution = ai.parse_attribution_export(str(ATTRIBUTION_MW))
    delivery = ai.parse_delivery_export(str(DELIVERY_MW))
    attr_dict = db._json_safe(dataclasses.asdict(attribution))
    deliv_dict = db._json_safe(dataclasses.asdict(delivery))
    rep.check("flight_start round-trips through storage as a plain string (the thing "
             "rehydrate_attribution has to undo)", isinstance(attr_dict["flight_start"], str),
             attr_dict["flight_start"])

    rehydrated_attr = ra.rehydrate_attribution(attr_dict)
    rehydrated_deliv = ra.rehydrate_delivery(deliv_dict)
    rep.check("rehydrated flight_start is a real date again",
             hasattr(rehydrated_attr.flight_start, "year"), rehydrated_attr.flight_start)

    fresh_out = REPO / "tests" / "_scratch_regen_fresh.pptx"
    stored_out = REPO / "tests" / "_scratch_regen_stored.pptx"
    ra.build_summary_slide(str(TEMPLATE), str(fresh_out), attribution=attribution, delivery=delivery,
                          client_name="Mattress Warehouse", threads=ONE_THREAD,
                          accepted_optimizations=[])
    ra.build_summary_slide(str(TEMPLATE), str(stored_out), attribution=rehydrated_attr,
                          delivery=rehydrated_deliv, client_name="Mattress Warehouse",
                          threads=ONE_THREAD, accepted_optimizations=[])
    text_fresh = _all_text(Presentation(fresh_out).slides[0])
    text_stored = _all_text(Presentation(stored_out).slides[0])
    rep.check("fresh-export and rehydrated-from-storage renders are text-identical",
             text_fresh == text_stored,
             (text_fresh, text_stored) if text_fresh != text_stored else "")
    fresh_out.unlink(missing_ok=True)
    stored_out.unlink(missing_ok=True)


def check_real_cardinal_fixture_end_to_end(rep):
    print("\nreal Cardinal fixture -- both slides build clean, both pass package_check")
    attribution_path = REPO / "Premion Website Attribution Cardinal.xlsx"
    delivery_path = REPO / "Premion OTT.xlsx"
    if not (TEMPLATE.exists() and attribution_path.exists() and delivery_path.exists()):
        rep.skip("REPORT_MASTER_v0_9.pptx or the real Cardinal fixture pair not present")
        return
    attribution = ai.parse_attribution_export(str(attribution_path))
    delivery = ai.parse_delivery_export(str(delivery_path))
    out_sum = REPO / "tests" / "_scratch_cardinal_summary.pptx"
    out_cs = REPO / "tests" / "_scratch_cardinal_cs.pptx"
    ra.build_summary_slide(str(TEMPLATE), str(out_sum), attribution=attribution, delivery=delivery,
                          client_name="Cardinal Plumbing", threads=FOUR_THREADS,
                          accepted_optimizations=[])
    ra.build_case_study_slide(str(TEMPLATE), str(out_cs), attribution=attribution, delivery=delivery,
                             threads=FOUR_THREADS, accepted_optimizations=[],
                             client_name="Cardinal Plumbing", vertical="home_improvement",
                             white_label=True, budget=40_000)
    text_sum = _all_text(Presentation(out_sum).slides[0])
    text_cs = _all_text(Presentation(out_cs).slides[0]).lower()
    rep.check("summary: no literal token left", "{{" not in text_sum, text_sum)
    rep.check("summary: passes package_check", package_check.check_package(str(out_sum)) == [])
    rep.check("case study: no literal token left", "{{" not in text_cs, text_cs)
    rep.check("case study: white-labeled, no client name leak", "cardinal plumbing" not in text_cs, text_cs)
    rep.check("case study: passes package_check", package_check.check_package(str(out_cs)) == [])
    out_sum.unlink(missing_ok=True)
    out_cs.unlink(missing_ok=True)


def check_white_label_redacts_informal_name_reference(rep):
    print("\nwhite_label_text catches an informal/shortened reference to the client, "
         "not just the full canonical phrase (2026-09-15 real find)")
    rep.check("'Cardinal' alone, from 'Cardinal Plumbing', is redacted -- a real "
             "vault-insertion check first caught this surviving a full-phrase-only pass",
             ra.white_label_text("Cardinal drove strong direct response.", "Cardinal Plumbing")
             == "the campaign drove strong direct response.",
             ra.white_label_text("Cardinal drove strong direct response.", "Cardinal Plumbing"))
    rep.check("a single-word client name (already fully covered by the first pass) is unaffected",
             ra.white_label_text("WAEPA benefited from concentrated delivery.", "WAEPA")
             == "the campaign benefited from concentrated delivery.")
    rep.check("short generic connector words in the name are never redacted on their own "
             "(no 4+ length false trigger from 'of'/'the'/'co' etc.)",
             ra.white_label_text("The Company saw growth.", "The Company")
             == "the campaign saw growth.",
             ra.white_label_text("The Company saw growth.", "The Company"))
    rep.check("blank client_name passes through unchanged",
             ra.white_label_text("Cardinal saw growth.", "") == "Cardinal saw growth.")
    rep.check("a doubled 'the the' left by redacting a name right after an article is "
             "collapsed -- caught live walking a real Cardinal case study through the app "
             "(2026-09-15): '...into the Cardinal services funnel' became '...into the the "
             "campaign services funnel' before this cleanup",
             ra.white_label_text(
                 "...moving qualified prospects into the Cardinal services funnel.",
                 "Cardinal Plumbing")
             == "...moving qualified prospects into the campaign services funnel.",
             ra.white_label_text(
                 "...moving qualified prospects into the Cardinal services funnel.",
                 "Cardinal Plumbing"))


def check_case_study_inserts_into_proposal_deck(rep):
    print("\na report-generated, white-labeled case study slide inserts into a REAL proposal "
         "master deck via the existing cross-deck copy mechanism (append_case_studies) -- "
         "the exact path the vault picker on Build a proposal already uses for every other "
         "case study, and the scenario that first caught the informal-name leak above")
    if not (TEMPLATE.exists() and PROPOSAL_MASTER.exists()):
        rep.skip(f"{TEMPLATE.name} or {PROPOSAL_MASTER.name} not present")
        return
    attribution = make_synthetic_attribution(client_name="Cardinal Plumbing")
    leaky_threads = [{
        "head": "Strong local response", "anchor": "goal", "finding": "62% direct.",
        "meaning": "Cardinal drove strong direct response.", "action": "",
    }]
    cs_path = REPO / "tests" / "_scratch_insert_case_study.pptx"
    ra.build_case_study_slide(str(TEMPLATE), str(cs_path), attribution=attribution, delivery=None,
                             threads=leaky_threads, accepted_optimizations=[],
                             client_name="Cardinal Plumbing", vertical="home_improvement",
                             white_label=True, budget=None)

    prs = Presentation(str(PROPOSAL_MASTER))
    before = len(prs.slides)
    n = assembly.append_case_studies(prs, [{"path": str(cs_path), "slides": [0]}])
    rep.check("exactly one slide was appended", n == 1, n)
    rep.check("the deck grew by exactly one slide", len(prs.slides) == before + 1,
             (before, len(prs.slides)))

    out = REPO / "tests" / "_scratch_master_with_case_study.pptx"
    prs.save(out)
    rep.check("the combined deck passes package_check", package_check.check_package(str(out)) == [],
             package_check.check_package(str(out)))

    prs2 = Presentation(str(out))
    leak_sites = []
    for index, slide in enumerate(prs2.slides):
        for shape in slide.shapes:
            if shape.has_text_frame and "cardinal" in shape.text_frame.text.lower():
                leak_sites.append((index, shape.name, shape.text_frame.text))
    rep.check("the client's name appears NOWHERE in the assembled deck, including the "
             "informal thread reference that used to leak", not leak_sites, leak_sites)
    cs_path.unlink(missing_ok=True)
    out.unlink(missing_ok=True)


if __name__ == "__main__":
    rep = Report()
    check_standalone_slides_drop_from_normal_build(rep)
    check_build_single_slide_isolates_each(rep)
    check_summary_all_tokens_resolved(rep)
    check_summary_takeaway_reflow_fewer_than_three(rep)
    check_summary_tile4_conversions_vs_rate(rep)
    check_sidebar_reflow_with_one_intent_class(rep)
    check_device_split_present_and_absent(rep)
    check_case_study_white_label_redacts_everywhere(rep)
    check_cost_per_visitor_tile_conditional(rep)
    check_declined_optimization_never_surfaces(rep)
    check_regenerate_from_rehydrated_matches_fresh(rep)
    check_real_cardinal_fixture_end_to_end(rep)
    check_white_label_redacts_informal_name_reference(rep)
    check_case_study_inserts_into_proposal_deck(rep)
    total = rep.passed + len(rep.failed)
    print(f"\n{rep.passed} passed, {len(rep.failed)} failed, {len(rep.skipped)} skipped "
         f"out of {total}")
    sys.exit(1 if rep.failed else 0)
