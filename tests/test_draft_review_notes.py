"""The draft's "Before sending" checklist never contradicts the table.

QA session, Visit Hershey & Harrisburg, build a7e4dad (2026-10-02): the plan
was right, but the review notes said a segment was both on the avails table
and "dropped", that TRAVEL Family was "not on the avails table" when it had
6 plan rows, and that Wilkes-Barre had no rows when it did. One wrong note
makes a rep re-check every other one.

Two fixes, both checked here:
- the app's own presence notes compare a drafted name against every TERM
  of every group (normalized), so a stacked group's audience is found;
- every note claiming something is missing is checked against the table
  and plan the draft leaves behind, and dropped (and logged) when the state
  contradicts it -- `app.drop_contradicted_draft_notes`.

Expected values come from the real Hershey avails document's own groups
(tests/group_scenario_fixtures.py) and Muse's own filed notes, never from
the helper under test.

    python tests/test_draft_review_notes.py
"""
import copy
import os
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(Path(__file__).resolve().parent))
os.chdir(REPO)
os.environ.setdefault("PROPOSAL_BUILDER_TEST_MODE", "1")

import db                                      # noqa: E402
db.fetch_audiences = lambda: (None, "stubbed for test isolation")

import app                                     # noqa: E402
import targeting_groups as tg                  # noqa: E402
import group_scenario_fixtures as gsf          # noqa: E402

failures = []


def check(label, condition, detail=""):
    if condition:
        print(f"  PASS  {label}")
    else:
        print(f"  FAIL  {label}{'  -- ' + str(detail) if detail else ''}")
        failures.append(label)


class _StubSt:
    def __init__(self):
        self.session_state = {}
        self.secrets = {}


def apply_draft(draft, preset_state):
    real = app.st
    stub = _StubSt()
    stub.session_state.update(preset_state)
    app.st = stub
    try:
        app.apply_draft_to_form(copy.deepcopy(draft))
    finally:
        app.st = real
    return stub.session_state


def hershey_groups():
    if not gsf.HERSHEY_PDF.exists():
        return None
    groups = [dict(g, include_in_plan=True, include_locked=False)
              for g in gsf.build_hershey()["groups"]]
    for g in groups:
        g.pop("_flight_label", None)
    # Muse's session also carried a Wilkes-Barre market (Harrisburg's own
    # DMA neighbour); the real document doesn't, so one group says so by name.
    wb = dict(groups[3], id="wilkesbarre", name="Wilkes Barre-Scranton-Hazleton")
    return groups + [wb]


def plan_with(groups, audience_terms):
    rows = [{"Tactic": "Premion Streaming TV", "Targeting": tg.audience_label(g),
             "Geo": tg.geo_label(g, label_for=app._market_display_name)}
            for g in groups if set(audience_terms) & set(g["terms"])]
    return [{"name": "Option A", "rows": rows}]


def check_state_check(groups):
    print("\nNotes checked against the table and plan")
    plan = plan_with(groups, ["TRAVEL Family", "AFIRST Travel Buffs and Sightseers"])
    muse = [
        "Unrecognized audience segment(s) dropped: AFIRST Travel Buffs and Sightseers, DEMO Age A21-44",
        "The notes mention TRAVEL Family, which is not on the avails table -- add it in Section D2 "
        "(Audiences & avails) and tick Plan if it should be sold.",
        "Wilkes-Barre has no rows on the avails table.",
    ]
    true_notes = [
        "The AFIRST Travel Buffs and Sightseers segment appears in the avails table as a labeled "
        "avails row.",
        "The notes mention TRAVEL Golf, which is not on the avails table -- add it in Section D2.",
        "TRAVEL Golf in Philadelphia isn't on the avails table yet.",
        "The $0 line for TRAVEL Family was dropped -- no-charge items belong in Included with Campaign.",
        "No split was stated between the two audiences, so the budget was divided evenly. Confirm "
        "the intended split.",
        "Pittsburgh has no rows -- pull avails for it if it's in scope.",
    ]
    kept, dropped = app.drop_contradicted_draft_notes(muse + true_notes, groups, plan)
    check("all three of Muse's contradicted notes are dropped",
          [d["note"] for d in dropped] == muse, dropped)
    check("every true note is kept, including a true 'not on the avails table' about an absent "
          "audience, a market named beside it, and a dropped $0 LINE",
          kept == true_notes, kept)
    check("each drop says what contradicted it (logged, never silent)",
          all(d["contradicted_by"] for d in dropped)
          and "travel family" in dropped[1]["contradicted_by"]
          and "wilkes" in dropped[2]["contradicted_by"], dropped)
    off_plan = app.drop_contradicted_draft_notes(
        ['"Option A" -- available but not on the plan: DEMO Age A25 Plus.'], groups,
        plan_with(groups, ["AFIRST Travel Buffs and Sightseers"]))
    check("'available but not on the plan' stays when the audience really isn't on the plan "
          "(on the table is a different claim)", off_plan[0] and not off_plan[1], off_plan)
    on_plan = app.drop_contradicted_draft_notes(
        ["TRAVEL Family isn't on the plan."], groups, plan)
    check("...and goes when the plan carries it", not on_plan[0] and on_plan[1], on_plan)


def check_apply_draft(groups):
    print("\nThe app's own presence notes, through apply_draft_to_form")
    scn = gsf.build_hershey()
    draft = {
        "client_name": "QA-TEST-Visit Hershey & Harrisburg", "vertical": "travel",
        "total_budget": 20000, "breakout": "full_flight", "media_plan_lines": [],
        "group_selection": {"mode": "all"}, "group_allocation": None,
        "group_selection_reason": "",
        "audiences": [{"segment": "AFIRST Travel Buffs and Sightseers"},
                      {"segment": "TRAVEL Family"},
                      {"segment": "DEMO Age A21-44"},
                      {"segment": "AFIRST Travel Buffs and Sightseers, DEMO Age A21-44"},
                      {"segment": "TRAVEL Golf"}],
        "attribution": [], "sports": [],
        "unresolved": [],
        "unresolved_internal": ["The notes mention TRAVEL Family, which is not on the avails table."],
    }
    state = apply_draft(draft, {"targeting_groups": groups, "flight_start": scn["flight_start"],
                                "flight_end": scn["flight_end"], "avails_mode": True})
    notes = (state.get("draft_unresolved") or []) + (state.get("draft_unresolved_internal") or [])
    joined = " ".join(notes)
    check("no audience on the table is reported 'dropped' (AFIRST, DEMO Age A21-44, their stack)",
          not any("dropped" in n and ("AFIRST" in n or "A21-44" in n) for n in notes), notes)
    check("TRAVEL Family (one term of a stacked group) is never 'not on the avails table'",
          not any("TRAVEL Family" in n and "not on the avails table" in n for n in notes), notes)
    check("an audience genuinely absent (TRAVEL Golf) is still reported",
          "TRAVEL Golf" in joined, notes)
    check("the model's own contradicted note was dropped and recorded for the feedback capture",
          any("TRAVEL Family" in d["note"] for d in state.get("draft_notes_dropped") or []),
          state.get("draft_notes_dropped"))
    check("the groups themselves are untouched by the draft (same ids, same avails)",
          [(g["id"], g["avails_monthly"]) for g in state["targeting_groups"]]
          == [(g["id"], g["avails_monthly"]) for g in groups])


def check_strategy_summary(groups):
    print("\nStrategy summary line: written only when it matches the plan")
    plan = plan_with(groups, ["TRAVEL Family"])     # only TRAVEL Family is sold
    good = "Targeted CTV strategy to reach TRAVEL Family households in Philadelphia and New York."
    check("a line naming only what the plan carries passes",
          app.strategy_summary_conflict(good, groups, plan) is None,
          app.strategy_summary_conflict(good, groups, plan))
    off = "Targeted CTV strategy to reach AFIRST Travel Buffs and Sightseers in Philadelphia."
    check("a line naming an audience on the table but not the plan is caught",
          "afirst travel buffs and sightseers" in (app.strategy_summary_conflict(off, groups, plan) or ""),
          app.strategy_summary_conflict(off, groups, plan))
    wb_plan = [{"name": "A", "rows": [r for r in plan[0]["rows"] if "Wilkes" not in r["Geo"]]}]
    market = "Targeted CTV strategy to reach TRAVEL Family households in Wilkes-Barre."
    check("a line naming a market no plan line runs in is caught",
          "wilkes" in (app.strategy_summary_conflict(market, groups, wb_plan) or ""),
          app.strategy_summary_conflict(market, groups, wb_plan))

    scn = gsf.build_hershey()
    base = {"client_name": "QA-TEST-Visit Hershey", "vertical": "travel", "total_budget": 20000,
            "breakout": "full_flight", "media_plan_lines": [],
            "group_selection": {"mode": "all"}, "group_allocation": None,
            "group_selection_reason": "", "audiences": [], "attribution": [], "sports": [],
            "unresolved": [], "unresolved_internal": []}
    preset = {"targeting_groups": groups, "flight_start": scn["flight_start"],
              "flight_end": scn["flight_end"], "avails_mode": True}
    ok_state = apply_draft(dict(base, strategy_summary="Targeted CTV strategy to reach travel "
                                                        "families across the Mid-Atlantic."), preset)
    check("a matching drafted line lands in the strategy box",
          ok_state.get("strategy_summary_text") == "Targeted CTV strategy to reach travel families "
                                                   "across the Mid-Atlantic.",
          ok_state.get("strategy_summary_text"))
    unselected = [dict(g, include_in_plan=False) for g in groups]
    bad_state = apply_draft(dict(base, group_selection={"mode": "ids", "ids": [groups[3]["id"]]},
                                 strategy_summary="CTV to reach AFIRST Travel Buffs and Sightseers."),
                            dict(preset, targeting_groups=unselected))
    check("a line the plan doesn't back is not written, and the seller is told why",
          "strategy_summary_text" not in bad_state
          and any("strategy line" in n for n in bad_state.get("draft_unresolved_internal") or []),
          (bad_state.get("strategy_summary_text"), bad_state.get("draft_unresolved_internal")))

    print("\nStrategy summary on the slide (template shape StrategySummary)")
    from pptx import Presentation
    from pptx.util import Inches
    import assembly
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(9), Inches(0.35))
    box.name = "StrategySummary"
    box.text_frame.text = "{{STRATEGY_SUMMARY}}"
    table = slide.shapes.add_table(2, 2, Inches(0.5), Inches(1.5), Inches(9), Inches(1)).table
    assembly.apply_strategy_summary(slide, "Targeted CTV strategy to reach travel families.")
    check("with text: the token is filled",
          box.text_frame.text == "Targeted CTV strategy to reach travel families.", box.text_frame.text)
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    box2 = slide2.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(9), Inches(0.35))
    box2.name = "StrategySummary"
    box2.text_frame.text = "{{STRATEGY_SUMMARY}}"
    frame = slide2.shapes.add_table(2, 2, Inches(0.5), Inches(1.5), Inches(9), Inches(1))
    assembly.apply_strategy_summary(slide2, None)
    check("toggle off: the line is removed and the table moves up by its height",
          not any(s.name == "StrategySummary" for s in slide2.shapes)
          and frame.top == Inches(1.15), frame.top)
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    check("a template without the shape is untouched",
          assembly.apply_strategy_summary(slide3, "anything") is False)


def main():
    groups = hershey_groups()
    if groups is None:
        print(f"  SKIP  {gsf.HERSHEY_PDF.name} not present")
        return 0
    check_state_check(groups)
    check_apply_draft(groups)
    check_strategy_summary(groups)
    print()
    print(f"{len(failures)} failure(s)" if failures else "All checks passed")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
