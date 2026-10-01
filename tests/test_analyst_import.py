"""Auto-Sales Analyst facts JSON -> the report builder.

Covers analyst_import.py (parse, period gate, visitor cross-check, the slide
and payload helpers), report:url_report filled from it, and the facts
payload + facts-only checker accepting the Analyst's own numbers.

Real fixtures (gitignored client data; the real-file checks SKIP without
them): Ted Britt's August 2026 website attribution export and the Analyst's
facts export for the same group and month. Expected values are taken from
the Analyst file's own totals block and the attribution export's own
advertiser tab -- never from the helpers under test.

    python tests/test_analyst_import.py
"""
import json
import re
import sys
from datetime import date
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

import market_lookup  # noqa: E402
import package_check  # noqa: E402
from pptx import Presentation  # noqa: E402

import analyst_import as an  # noqa: E402
import app  # noqa: E402
import attribution_import as ai  # noqa: E402
import report_assembly as ra  # noqa: E402
import slide_map  # noqa: E402

TEMPLATE = REPO / "REPORT_MASTER_v0_14.pptx"
ATTRIBUTION_TB930 = REPO / "Premion Website Attribution and Reach Extension TB930.xlsx"
# Schema v2 (the Analyst's 2026-10-01 export); the older v1 file still parses
# and is used where the checks are about v1 compatibility.
ANALYST_TB_AUG = REPO / "ted_britt_aug2026_facts.json"
ANALYST_TB_AUG_V1 = REPO / "auto-group-5-sites-12-07-pm-et_2026-08-01_2026-08-31_facts.json"
OUT_DIR = REPO / "tests" / "_manual_output"

market_lookup.install()

failures = []
skipped = []


def check(label, condition, detail=""):
    if condition:
        print(f"  PASS  {label}")
    else:
        print(f"  FAIL  {label}{'  -- ' + str(detail) if detail else ''}")
        failures.append(label)


def skip(label):
    print(f"  SKIP  {label}")
    skipped.append(label)


def _synthetic(**overrides):
    data = {
        "schema_version": 1,
        "meta": {"group_name": "Test Group", "is_group": True, "site_count": 2,
                 "analysis_period": {"start": "2026-08-01", "end": "2026-08-31",
                                     "source": "rep_entered_report_month"},
                 "inventory_scanned_at": "2026-09-30", "lookback_days": 30},
        "sites": [{"site_id": "a", "dealer_name": "Alpha Ford", "domain": "alpha.com"},
                  {"site_id": "b", "dealer_name": "Beta Chevy", "domain": "beta.com"}],
        "totals": {
            "vehicles_shopped": 200, "vehicles_sold": 80, "look_to_book_pct": 40.0,
            "est_revenue_sold": 3_200_000, "est_pipeline_value": 8_000_000,
            "visits_total": 1000, "unique_visitors": 400,
            "traffic_mix": [
                {"category": "NEW_VDP", "label": "New VDP", "visits": 300},
                {"category": "OTHER", "label": "Other", "visits": 250},
                {"category": "USED_VDP", "label": "Used VDP", "visits": 200},
                {"category": "HOMEPAGE", "label": "Homepage", "visits": 100},
                {"category": "SERVICE", "label": "Service", "visits": 80},
                {"category": "NEW_CAR_SEARCH", "label": "New Car Search", "visits": 70},
            ],
            "sold_by_make_model": [
                {"make": "FORD", "model": "MUSTANG MACH-E", "label": "FORD MUSTANG MACH-E", "count": 7},
                {"make": "GMC", "model": "SIERRA 1500", "label": "GMC SIERRA 1500", "count": 12},
                {"make": "FORD", "model": "F-150", "label": "FORD F-150", "count": 20},
            ],
            "missed_opportunities": {"count": 1, "vehicles": [
                {"make": "FORD", "model": "BRONCO", "label": "FORD BRONCO", "visits": 44,
                 "vin": "1FMEE9BP3SLB32625"}]},
        },
        "by_site": {"b": {"visits_total": 5, "traffic_mix": []}},
    }
    data.update(overrides)
    return json.dumps(data).encode()


def _raises(source):
    try:
        an.parse_analyst_facts(source)
    except an.AnalystParseError as exc:
        return str(exc)
    return None


def check_parser():
    print("\nParser -- shape, errors a rep can act on")
    parsed = an.parse_analyst_facts(_synthetic())
    check("a valid export parses, dates kept as ISO strings",
          parsed["period_start"] == "2026-08-01" and parsed["period_end"] == "2026-08-31",
          (parsed["period_start"], parsed["period_end"]))
    check("dealer names carried from the sites list",
          [s["dealer_name"] for s in parsed["sites"]] == ["Alpha Ford", "Beta Chevy"])
    check("the parsed dict is JSON-safe (it's logged with the report)",
          json.loads(json.dumps(parsed)) == parsed)

    msg = _raises(b"{not json")
    check("unreadable JSON -> plain message pointing at the Analyst's facts download",
          msg is not None and "facts" in msg and "Traceback" not in msg, msg)
    msg = _raises(json.dumps({"hello": 1}).encode())
    check("some other JSON file -> named as not an Analyst facts export",
          msg is not None and "isn't an Auto-Sales Analyst" in msg, msg)
    msg = _raises(_synthetic(schema_version=3))
    check("a schema version newer than v2 is refused, never half-read",
          msg is not None and "3" in msg, msg)
    bad_period = json.loads(_synthetic())
    bad_period["meta"]["analysis_period"]["start"] = "August"
    msg = _raises(json.dumps(bad_period).encode())
    check("an unreadable analysis period is refused (the period gate depends on it)",
          msg is not None and "August" in msg, msg)
    empty = json.loads(_synthetic())
    empty["totals"] = {"visits_total": 0}
    msg = _raises(json.dumps(empty).encode())
    check("a file with no traffic or vehicle data is refused", msg is not None, msg)
    bom = b"\xef\xbb\xbf" + _synthetic()
    check("a UTF-8 BOM (Windows save) still parses", _raises(bom) is None, _raises(bom))


def check_gates():
    print("\nPeriod overlap and visitor cross-check")
    parsed = an.parse_analyst_facts(_synthetic())
    check("overlapping window (Jul 27 - Aug 31) -> used",
          an.periods_overlap(parsed, date(2026, 7, 27), date(2026, 8, 31)))
    check("one shared day at the edge counts as overlap",
          an.periods_overlap(parsed, date(2026, 8, 31), date(2026, 9, 30)))
    check("June report vs August Analyst -> not used",
          not an.periods_overlap(parsed, date(2026, 6, 1), date(2026, 6, 30)))
    check("ISO-string bounds work the same as dates (manual period entry stores strings)",
          not an.periods_overlap(parsed, "2026-06-01", "2026-06-30")
          and an.periods_overlap(parsed, "2026-08-10", "2026-08-12"))
    check("a missing report bound never blocks the file",
          an.periods_overlap(parsed, None, date(2026, 6, 30)))
    check("matching visitor counts -> no mismatch", an.visitor_count_mismatch(parsed, 400) is None)
    check("differing counts -> both named, Analyst first",
          an.visitor_count_mismatch(parsed, 390) == (400, 390), an.visitor_count_mismatch(parsed, 390))
    check("no export count -> nothing to compare", an.visitor_count_mismatch(parsed, 0) is None)

    check("a normal file (80 sold) is usable", an.unusable_reason(parsed) is None)
    zero = json.loads(_synthetic())
    zero["totals"]["vehicles_sold"] = 0
    reason = an.unusable_reason(an.parse_analyst_facts(json.dumps(zero).encode()))
    check("0 sold with vehicles shopped -> set aside, per the Sales Assist's own rule",
          reason is not None and "0 vehicles sold" in reason, reason)
    check("scan 30 days after period end (the recommended window) -> no advisory",
          an.scan_delay_days(parsed) == 30 and an.lookback_advisory(parsed) is None,
          an.scan_delay_days(parsed))
    late = json.loads(_synthetic())
    late["meta"]["inventory_scanned_at"] = "2026-11-24"
    late_parsed = an.parse_analyst_facts(json.dumps(late).encode())
    advisory = an.lookback_advisory(late_parsed)
    check("scanned 85 days after period end -> advisory naming the days (scan date, not "
          "lookback_days, decides)", advisory is not None and "85 days" in advisory, advisory)
    facts = {"analyst": an.payload_facts(late_parsed)}
    items = app.attr_actionable_review_items(facts)
    check("...and the same fact reaches Review before sending",
          any("85 days" in i for i in items), items)
    no_scan = json.loads(_synthetic())
    no_scan["meta"].pop("inventory_scanned_at")
    no_scan["meta"]["lookback_days"] = 90
    check("no scan date -> falls back to lookback_days",
          an.scan_delay_days(an.parse_analyst_facts(json.dumps(no_scan).encode())) == 90)

    check("franchise makes come from the client and dealer names",
          an.franchise_makes(parsed, "Alpha Ford Lincoln") == ["CHEVROLET", "FORD", "LINCOLN"],
          an.franchise_makes(parsed, "Alpha Ford Lincoln"))
    stafford = {"sites": [{"dealer_name": "Stafford Honda"}, {"dealer_name": "Dominion Motors"}]}
    check("short makes need a whole word (Stafford isn't Ford, Dominion isn't Mini)",
          an.franchise_makes(stafford) == ["HONDA"], an.franchise_makes(stafford))
    check("run-together domain names still match a long make (Tedbrittchevrolet)",
          an.franchise_makes({"sites": [{"dealer_name": "Tedbrittchevrolet"}]}) == ["CHEVROLET"])
    unverified = json.loads(_synthetic())
    unverified["totals"]["shopped_vehicle_status"] = {"available": 100, "sold": 80, "unverified": 20}
    payload = an.payload_facts(an.parse_analyst_facts(json.dumps(unverified).encode()))
    check("'Inventory Unavailable' vehicles ride separately, never inside the sold count",
          payload["vehicles_status_unconfirmed"] == 20 and payload["vehicles_sold_since"] == 80,
          payload)


def check_slide_helpers():
    print("\nSlide helpers -- traffic mix and top models")
    parsed = an.parse_analyst_facts(_synthetic())
    rows = an.traffic_mix_rows(parsed)
    labels = [r["label"] for r in rows]
    check("named categories largest-first, capped, remainder row last",
          labels == ["New VDP", "Used VDP", "Homepage", "Service", "All other pages"], labels)
    check("the Analyst's own 'Other' never appears beside 'All other pages'",
          "Other" not in labels, labels)
    check("the table accounts for every visit (sums to visits_total = 1000)",
          sum(r["visits"] for r in rows) == 1000, sum(r["visits"] for r in rows))
    check("visit shares sum to 100%", abs(sum(r["share"] for r in rows) - 1) < 1e-9)
    models = an.top_model_rows(parsed)
    check("models sorted by units, not file order",
          [m["count"] for m in models] == [20, 12, 7], models)
    check("labels title-cased, model codes and acronyms kept",
          [m["label"] for m in models] == ["Ford F-150", "GMC Sierra 1500", "Ford Mustang Mach-E"],
          [m["label"] for m in models])
    check("model share is of vehicles sold (20 of 80 = 25%)", abs(models[0]["share"] - 0.25) < 1e-9)
    check("VDP share counts new + used VDP visits only (500 of 1000)",
          abs(an.vdp_visit_share(parsed) - 0.5) < 1e-9, an.vdp_visit_share(parsed))

    payload = an.payload_facts(parsed)
    check("payload carries no VINs", "1FMEE9BP3SLB32625" not in json.dumps(payload))
    check("payload keeps the full, unfolded mix (Service and Other both present)",
          {m["label"] for m in payload["traffic_mix"]} >= {"Service", "Other"})
    check("payload names what sold as 'vehicles_sold_since' (the wording the prompt teaches)",
          payload["vehicles_sold_since"] == 80)


def check_real_ted_britt():
    print("\nReal Ted Britt, August 2026 -- export + Analyst for the same group and month")
    for path in (ATTRIBUTION_TB930, ANALYST_TB_AUG, TEMPLATE):
        if not path.exists():
            skip(f"{path.name} not present")
            return
    attribution = ai.parse_attribution_export(str(ATTRIBUTION_TB930))
    analyst = an.parse_analyst_facts(str(ANALYST_TB_AUG))
    raw = json.loads(ANALYST_TB_AUG.read_text(encoding="utf-8"))

    check("export's own unique-visitor count is 2,671 (advertiser tab)",
          attribution.attributed_unique_visitors == 2671, attribution.attributed_unique_visitors)
    check("...and the Analyst's matches it, so no mismatch warning fires",
          an.visitor_count_mismatch(analyst, attribution.attributed_unique_visitors) is None)
    check("August Analyst overlaps the Jul 27 - Aug 31 report period",
          an.periods_overlap(analyst, attribution.flight_start, attribution.flight_end),
          (attribution.flight_start, attribution.flight_end))
    mix = an.traffic_mix_rows(analyst)
    check("slide's traffic table sums to the file's visits_total (8,694)",
          sum(r["visits"] for r in mix) == raw["totals"]["visits_total"] == 8694,
          sum(r["visits"] for r in mix))
    check("a truck-shop site with only visits in by_site doesn't break parsing",
          "tedbritttruckshop" in analyst["by_site"])

    OUT_DIR.mkdir(exist_ok=True)
    with_path = OUT_DIR / "TB930_with_analyst.pptx"
    _, warnings = ra.build_report_deck(
        str(TEMPLATE), attribution, None, str(with_path),
        client_name="Ted Britt Ford & Ted Britt Chantilly", goals_bullets=[],
        whats_next_bullets=["Analyst check: what's-next placeholder"], analyst=analyst)
    check("deck builds with no fit/column warnings", not warnings, warnings)
    prs = Presentation(str(with_path))
    url = next(s for s in prs.slides if s.has_notes_slide
               and slide_map.notes_key(s) == "report:url_report")
    texts = [sh.text_frame.text for sh in slide_map.iter_all_shapes(url.shapes) if sh.has_text_frame]
    tables = {sh.name: sh.table for sh in slide_map.iter_all_shapes(url.shapes)
              if getattr(sh, "has_table", False) and sh.has_table}
    all_text = " ".join(texts) + " ".join(
        c.text for t in tables.values() for r in t.rows for c in r.cells)
    check("no token left unfilled on the slide", "{{" not in all_text)
    check("section headers rewritten for the Analyst data",
          ra._ANALYST_MIX_HEADER in texts and ra._ANALYST_MODELS_HEADER in texts, texts)
    check("old URL-report headers gone", "VISITORS BY INTENT" not in texts, texts)
    mix_table = tables["IntentSummaryTable"]
    check("traffic table header is Page type | Visits | % of visits (no Converted column)",
          [c.text for c in mix_table.rows[0].cells] == list(ra._ANALYST_MIX_COLUMNS),
          [c.text for c in mix_table.rows[0].cells])
    check("first traffic row is New VDP, 2,818 visits (the file's largest category)",
          [c.text for c in mix_table.rows[1].cells][:2] == ["New VDP", "2,818"],
          [c.text for c in mix_table.rows[1].cells])
    model_table = tables["TopUrlTable"]
    check("models table's top row is the F-150 at 100 units (file's top_sellers[0])",
          [c.text for c in model_table.rows[1].cells][:2] == ["Ford F-150", "100"],
          [c.text for c in model_table.rows[1].cells])
    check("subtitle is the approved inventory-movement shape, counts from the file (1,696 / 939)",
          "Of the 1,696 vehicles our audience viewed, 939 have since sold." in texts,
          texts)
    check("section headers use the Analyst's own names (Traffic Mix, Top Sold Models)",
          "TRAFFIC MIX" in texts and "TOP SOLD MODELS" in texts, texts)
    check("footnote defines 'sold' as inventory movement and names the source and month",
          any("Auto-Sales Analyst, August 2026" in t
              and "not purchases by our visitors" in t
              for t in texts), texts)
    check("package is structurally clean", not package_check.check_package(str(with_path)),
          package_check.check_package(str(with_path)))

    without_path = OUT_DIR / "TB930_without_analyst.pptx"
    ra.build_report_deck(
        str(TEMPLATE), attribution, None, str(without_path),
        client_name="Ted Britt Ford & Ted Britt Chantilly", goals_bullets=[],
        whats_next_bullets=["Analyst check: what's-next placeholder"])
    prs2 = Presentation(str(without_path))
    url2 = next(s for s in prs2.slides if s.has_notes_slide
                and slide_map.notes_key(s) == "report:url_report")
    texts2 = [sh.text_frame.text for sh in slide_map.iter_all_shapes(url2.shapes) if sh.has_text_frame]
    check("without the Analyst (and no vertical) the fallback URL report still builds, "
          "relabelled for visits", "TOP PAGES BY ATTRIBUTED VISITS" in texts2, texts2)

    facts = ra.build_facts_payload(attribution, None, analyst=analyst,
                                   client_name="Ted Britt Ford & Ted Britt Chantilly")
    check("facts payload carries the Analyst slice", facts["analyst"] is not None)
    check("Ted Britt's franchise makes are Ford, Chevrolet and Lincoln -- so Ford-heavy "
          "movement there is inventory mix, not audience proof",
          facts["analyst"]["franchise_makes"] == ["CHEVROLET", "FORD", "LINCOLN"],
          facts["analyst"]["franchise_makes"])
    check("scan was 31 days after period end (Oct 1 vs Aug 31) -- no timing item",
          facts["analyst"]["days_scanned_after_period_end"] == 31
          and not any("Auto-Sales Analyst" in i for i in app.attr_actionable_review_items(facts)))
    check("...and None without it", ra.build_facts_payload(attribution, None)["analyst"] is None)
    sentence = ("939 of the 1,696 vehicles our audience viewed have since sold, a 55.4% "
                "look-to-book rate worth an estimated $39.2M; 57% of visits hit a vehicle "
                "detail page and the F-150 led with 100 units, with Ford at 537.")
    violations = app._attr_draft_number_violations([("thread", sentence)], facts)
    check("the facts-only checker traces every Analyst number a narrative would cite",
          not violations, violations)
    invented = app._attr_draft_number_violations(
        [("thread", "an estimated $41.2M in inventory sold")], facts)
    check("...and still flags one the file doesn't contain", bool(invented), invented)
    prompt = app.build_attr_draft_prompt(facts)
    check("the drafting prompt carries the Analyst section and its data",
          "Auto-Sales Analyst" in prompt and '"vehicles_sold_since": 939' in prompt)


REACH_FIXTURES = [
    "MW attribution excel.xlsx", "Premion Website Attribution Cardinal.xlsx",
    "Premion Website Attribution and Reach Extension (13).xlsx",
    "Premion Website Attribution and Reach Extension (15).xlsx",
    "Premion Website Attribution and Reach Extension TB.xlsx",
    "Premion Website Attribution and Reach Extension TB930.xlsx",
]


def check_fallback_url_report():
    """The fallback "Where Visitors Went" -- what every report without the
    Analyst file gets. Two fixes guarded here: (1) no share on it exceeds
    100% (the export's per-page unique-visitor counts were being summed
    into class "reach", 285% for Ted Britt's "Other pages"), and (2) an auto
    site's pages are named the Analyst's way, never by raw path."""
    print("\nFallback URL slide -- shares bounded, automotive names")
    present = [REPO / f for f in REACH_FIXTURES if (REPO / f).exists()]
    if not present:
        skip("no real attribution exports present")
    for path in present:
        attribution = ai.parse_attribution_export(str(path))
        for vertical in (None, "auto"):
            rows = (ra.intent_summary_rows(attribution, vertical=vertical)
                    + ra.top_url_rows(attribution, vertical=vertical))
            classes = ra.intent_facts(attribution, vertical=vertical)["classes"]
            worst = max([r["_visit_share_raw"] for r in rows]
                        + [r["reach"] for r in rows + classes if r.get("reach") is not None]
                        + [c["visit_share"] for c in classes])
            check(f"{path.name} ({vertical or 'no vertical'}): no share or reach over 100% "
                  f"(max {worst:.0%})", worst <= 1)
            share_sum = sum(r["_share_raw"] for r in ra.intent_summary_rows(attribution, vertical=vertical))
            check(f"{path.name} ({vertical or 'no vertical'}): Traffic Mix shares sum to 100%",
                  abs(share_sum - 1) < 1e-6, share_sum)
    try:
        ra.assert_shares_within_bounds([{"label": "Other pages", "_share_raw": 2.85}])
        raised = False
    except ValueError:
        raised = True
    check("the bounds assertion itself fires on the old 285% figure", raised)

    if not (ATTRIBUTION_TB930.exists() and ANALYST_TB_AUG.exists() and TEMPLATE.exists()):
        skip("Ted Britt August fixtures not present")
        return
    attribution = ai.parse_attribution_export(str(ATTRIBUTION_TB930))
    old = {r["intent"]: r for r in ra.intent_summary_rows(attribution)}
    check("Ted Britt with no vertical: 'Other pages' is now a share of visits, well under 100% "
          "(was 285% of visitors)", old["other"]["_share_raw"] < 1, old["other"]["share"])
    raw = json.loads(ANALYST_TB_AUG.read_text(encoding="utf-8"))
    analyst_mix = {m["label"]: m["visits"] for m in raw["totals"]["traffic_mix"]}
    ours = {c["label"]: c["visits"] for c in ra.intent_facts(attribution, vertical="auto")["classes"]}
    shared = [label for label in analyst_mix if label != "Other"]
    mismatched = {label: (ours.get(label), analyst_mix[label]) for label in shared
                  if ours.get(label) != analyst_mix[label]}
    check("auto categories reproduce the Analyst file's own Traffic Mix, same names, same "
          "visit counts (New VDP 2,818, Used VDP 2,113, ...)", not mismatched, mismatched)
    check("the Analyst's 'Other' equals our Other + Finance + Trade-In (the two classes we "
          "split out of it)",
          ours.get("Other pages", 0) + ours.get("Finance / Credit App", 0)
          + ours.get("Trade-In", 0) == analyst_mix["Other"],
          (ours.get("Other pages"), ours.get("Finance / Credit App"), ours.get("Trade-In"),
           analyst_mix["Other"]))

    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / "TB930_fallback_auto.pptx"
    _, warnings = ra.build_report_deck(
        str(TEMPLATE), attribution, None, str(out), client_name="Ted Britt",
        goals_bullets=[], whats_next_bullets=["fallback check"], vertical="auto")
    prs = Presentation(str(out))
    url = next(s for s in prs.slides if s.has_notes_slide
               and slide_map.notes_key(s) == "report:url_report")
    text = " ".join(
        [sh.text_frame.text for sh in slide_map.iter_all_shapes(url.shapes) if sh.has_text_frame]
        + [c.text for sh in slide_map.iter_all_shapes(url.shapes)
           if getattr(sh, "has_table", False) and sh.has_table
           for r in sh.table.rows for c in r.cells])
    check("no raw path ('Searchnew.Aspx', '.aspx') anywhere on the fallback slide",
          "searchnew" not in text.lower() and ".aspx" not in text.lower(), text[:300])
    check("vehicle rows are named as vehicles (2026 Ford F-150 Lariat)",
          "2026 Ford F-150 Lariat" in text, text[:400])
    check("no percentage over 100 on the slide",
          all(float(p) <= 100 for p in re.findall(r"(\d+(?:\.\d+)?)%", text)),
          re.findall(r"(\d+(?:\.\d+)?)%", text))
    check("fallback deck builds without fit warnings", not warnings, warnings)

    # One basis per column: every "% of visits" cell is that row's visits over
    # ALL attributed page visits -- recomputed here from the export's own URL
    # tab, not from the helper that filled the cell.
    total_visits = sum(int(v or 0) for u, v in attribution.by_url.items()
                       if not ra._is_noise_url(u))
    bad = []
    for sh in slide_map.iter_all_shapes(url.shapes):
        if not (getattr(sh, "has_table", False) and sh.has_table):
            continue
        for row in list(sh.table.rows)[1:]:
            cells = [c.text for c in row.cells]
            visits, pct = int(cells[1].replace(",", "")), float(cells[2].rstrip("%"))
            if abs(visits / total_visits * 100 - pct) > 0.51:
                bad.append(cells)
    check("every '% of visits' cell on both tables is visits / all page visits (one basis, "
          "additive)", not bad, bad)
    check("footnote says what the column is and that rows add up",
          "% of visits = each row's share of all attributed page visits, so rows add up" in text)

    nwfcu = REPO / "Premion Website Attribution and Reach Extension (15).xlsx"
    if nwfcu.exists():
        n_att = ai.parse_attribution_export(str(nwfcu))
        home = next(r for r in ra.intent_summary_rows(n_att) if r["intent"] == "homepage")
        check("NWFCU homepage: the table shows its share of visits (36%); its genuine 72% "
              "reach stays out of the column, available to the narrative as 'reach'",
              home["share"] == "36%" and round(home["reach"] * 100) == 72,
              (home["share"], home["reach"]))

    stacked = {"threads": [{"head": "Inventory", "meaning":
               "An estimated $38,063,252 in revenue sold and an estimated $75,200,770 "
               "Pipeline Value across all shopped inventory."}]}
    alone = {"threads": [{"head": "Inventory", "meaning":
             "The campaign drove traffic to an estimated $75,200,770 Pipeline Value."}]}
    check("the pipeline check flags Revenue Sold and Pipeline Value side by side (the real "
          "Ted Britt draft's phrasing)", bool(app._pipeline_stack_violations(stacked)))
    check("...and passes Pipeline Value cited alone",
          not app._pipeline_stack_violations(alone), app._pipeline_stack_violations(alone))
    plus = {"url_intent_narrative": "$38M sold plus $37M still in the pipeline."}
    check("...and flags a '$X sold plus $Y pipeline' sentence with no 'revenue' word",
          bool(app._pipeline_stack_violations(plus)))
    representing = {"url_intent_narrative": (
        "Of the 1,696 vehicles those visitors shopped, 918 have since sold, representing an "
        "estimated Pipeline Value of $75,200,770 across all vehicles attributed visitors viewed.")}
    check("...and flags Pipeline tied to the sold count with one dollar figure (a real draft: "
          "'918 have since sold, representing an estimated Pipeline Value')",
          bool(app._pipeline_stack_violations(representing)))
    # v2 vocabulary (2026-10-01): the same rule under the new name. "Est." is
    # an abbreviation -- the splitter used to cut there, so the sold count and
    # the total landed in different "sentences" and both of these passed.
    v2_stacked = {"url_intent_narrative": (
        "An estimated $39,168,431 in Est. value sold and an Est. total value viewed of "
        "$75,200,770.")}
    v2_tied = {"url_intent_narrative": (
        "Of the 1,696 vehicles our audience viewed, 939 have since sold, with an Est. total "
        "value viewed of $75,200,770.")}
    v2_alone = {"url_intent_narrative": (
        "Of the 1,696 vehicles our audience viewed, 939 have since sold. The Est. total value "
        "viewed reached an estimated $75,200,770.")}
    check("v2: Est. value sold beside Est. total value viewed is flagged",
          bool(app._pipeline_stack_violations(v2_stacked)))
    check("v2: Est. total value viewed tied to the sold count is flagged ('Est.' doesn't end "
          "the sentence)", bool(app._pipeline_stack_violations(v2_tied)))
    check("v2: Est. total value viewed in a sentence of its own passes",
          not app._pipeline_stack_violations(v2_alone), app._pipeline_stack_violations(v2_alone))
    shopper = app._enforce_draft_rules(
        {"analyst_watchlist_narrative": "Keeps high-intent shoppers moving; vehicles they shopped.",
         "url_intent_narrative": ("The campaign drove 8,694 visits from in-market shoppers. "
                                  "Shoppers viewed 1,696 vehicles and 939 have since sold."),
         "threads": [{"head": "Website", "meaning": "Streaming TV drove shoppers to the site."}]},
        {"analyst": {"period_end": "2026-08-31", "vehicles_viewed": 1696,
                     "vehicles_sold_since": 939, "visits_total": 8694}})
    check("'shopped'/'shoppers' become 'viewed'/'visitors' in Analyst text (a real v2 draft "
          "wrote 'high-intent shoppers')",
          shopper["analyst_watchlist_narrative"] == "Keeps high-intent visitors moving; vehicles they viewed.",
          shopper["analyst_watchlist_narrative"])
    check("...but website-attribution language is left alone: 'drove ... visits from in-market "
          "shoppers' stays, only the sentence citing Analyst sales changes (scope correction, "
          "2026-10-01)",
          shopper["url_intent_narrative"] == ("The campaign drove 8,694 visits from in-market "
                                              "shoppers. Visitors viewed 1,696 vehicles and 939 "
                                              "have since sold.")
          and shopper["threads"][0]["meaning"] == "Streaming TV drove shoppers to the site.",
          (shopper["url_intent_narrative"], shopper["threads"][0]["meaning"]))
    check("...and nothing is swapped without the Analyst",
          app._swap_analyst_vocabulary({"url_intent_narrative": "shoppers"}, {"analyst": None})
          ["url_intent_narrative"] == "shoppers")

    # The deterministic backstops after the retry -- real sentences from the
    # Ted Britt drafts that survived a corrective retry.
    # vdp_visit_share carries the real 56.7% the kept sentence quotes -- the
    # number-trace rule (QA-007) would otherwise remove it too, correctly.
    facts_ott = {"analyst": {"period_end": "2026-08-31", "vdp_visit_share": 0.567},
                 "ott_retargeting": {"impressions": 1}}
    survived = {
        "url_intent_narrative": (
            "56.7% of all attributed page visits landed on New or Used VDP pages. An estimated "
            "$38,063,252 in vehicles that attributed visitors shopped have since sold, with an "
            "additional estimated $75,200,770 in Pipeline Value across all shopped inventory."),
        "threads": [
            {"head": "Idea", "action_tier": 2, "meaning": "m",
             "action": "Add OTT Retargeting to build multi-screen frequency."},
            {"head": "Site", "action_tier": 2, "meaning": "m",
             "action": "Consider Site Retargeting for lead-page visitors."}],
        "goal_alignment_notes": []}
    fixed = app._enforce_draft_rules(survived, facts_ott)
    check("a stacked Pipeline sentence that survived the retry is removed; the rest stays",
          fixed["url_intent_narrative"] == "56.7% of all attributed page visits landed on New "
                                           "or Used VDP pages."
          and not app._pipeline_stack_violations(fixed), fixed["url_intent_narrative"])
    check("an 'add OTT Retargeting' idea is dropped when retargeting is already running; "
          "Site Retargeting (a different product) is kept",
          fixed["threads"][0]["action"] is None
          and fixed["threads"][1]["action"] == "Consider Site Retargeting for lead-page visitors.",
          fixed["threads"])
    check("both removals are recorded in the draft notes, never silent",
          len(fixed["goal_alignment_notes"]) == 2, fixed["goal_alignment_notes"])
    only_stacked = {"url_intent_narrative": "An estimated $38M sold and a $75M Pipeline Value."}
    check("a field left empty becomes None, so the slide uses its computed sentence",
          app._strip_pipeline_stacking(only_stacked, facts_ott)["url_intent_narrative"] is None)


ANALYST_DECK = REPO / "Auto Group (5 Sites) - 10_51 AM ET_Summary (1).pptx"


def check_cross_facts_and_order():
    """Ted Britt review, 2026-10-01: richer facts joined from the export's
    own vehicle pages, and Takeaways closing the deck."""
    print("\nCross-source Analyst facts, and Takeaways last")
    check("model families match across the two sources' spellings",
          [an.model_family_key(m) for m in ("F-150", "F 150 Lariat", "GX-460", "Gx 460 Premium",
                                            "SILVERADO 2500 HD", "Q7 55 Prestige")]
          == ["F150", "F150", "GX460", "GX460", "SILVERADO", "Q7"])
    check("a short letter code before a model number stays a code (GX 460, not Gx 460)",
          an._title_label("2020 LEXUS GX 460 PREMIUM PKG") == "2020 Lexus GX 460 Premium Pkg",
          an._title_label("2020 LEXUS GX 460 PREMIUM PKG"))
    if not (ATTRIBUTION_TB930.exists() and ANALYST_TB_AUG.exists() and TEMPLATE.exists()):
        skip("Ted Britt August fixtures not present")
        return
    attribution = ai.parse_attribution_export(str(ATTRIBUTION_TB930))
    analyst = an.parse_analyst_facts(str(ANALYST_TB_AUG))
    raw = json.loads(ANALYST_TB_AUG.read_text(encoding="utf-8"))["totals"]
    x = ra.analyst_cross_facts(analyst, attribution)
    nu = x["new_vs_used"]
    check("the export's vehicle pages total the Analyst's own vehicles viewed (1,696)",
          nu["new_vehicles_viewed"] + nu["used_vehicles_viewed"] == raw["vehicles_shopped"] == 1696)
    check("new/used viewed counts reproduce the Analyst's OWN Look-to-Book from its sold counts "
          "(430/1,000 = 43.0%, 509/696 = 73.1%)",
          round(raw["vehicles_sold_new"] / nu["new_vehicles_viewed"] * 100, 1) == raw["look_to_book_pct_new"]
          and round(raw["vehicles_sold_used"] / nu["used_vehicles_viewed"] * 100, 1)
          == raw["look_to_book_pct_used"], nu)
    check("used is the faster side and the gap is material", nu["faster_side"] == "used"
          and nu["look_to_book_gap_material"], nu)
    check("no model is emitted with more sold than viewed",
          all(m["sold_since"] <= m["viewed"] for m in x["models_viewed_vs_sold"]))
    f150_viewed = sum(m["count"] for m in raw["shopped_by_make_model"]
                      if an.model_family_key(m["model"]) == "F150")
    f150_sold = sum(m["count"] for m in raw["sold_by_make_model"]
                    if an.model_family_key(m["model"]) == "F150")
    top = x["models_viewed_vs_sold"][0]
    check(f"F-150 is the most-viewed family, {f150_viewed} viewed / {f150_sold} sold (the file's "
          "own by-model rows, every F-150 trim folded together)",
          (top["model"], top["viewed"], top["sold_since"]) == ("Ford F-150", f150_viewed, f150_sold), top)
    gap = x["store_gap"]
    check("store gap: Chantilly 59.5% vs Chevrolet 43.0%, material, by the Analyst's own store names",
          (gap["strongest"], gap["strongest_pct"], gap["weakest"], gap["weakest_pct"], gap["material"])
          == ("Ted Britt Ford of Chantilly", 59.5, "Ted Britt Chevrolet", 43.0, True), gap)
    check("stores are keyed on their domain (a store's identity; site ids are name slugs)",
          x["stores"][0]["domain"] == "tedbrittchantilly.com", x["stores"][:1])
    check("a store with traffic but no vehicles (the truck shop) isn't in the scoreboard",
          not any("truckshop" in s["store"].lower() for s in x["stores"]), x["stores"])
    check("every Missed Opportunity is joined to its page: year/trim, new/used and Est. price",
          all(m["new_or_used"] for m in x["missed_opportunities"])
          and x["missed_opportunities"][0] == {"vehicle": "2026 Ford Mustang Dark Horse SC",
                                               "new_or_used": "new", "visits": 58, "est_value": 45000},
          x["missed_opportunities"][:2])
    lexus = [m["vehicle"] for m in x["missed_opportunities"] if "Lexus" in m["vehicle"]]
    check("a trim word the model name already carries isn't repeated (GX-460, not 'Gx-460 460')",
          lexus == ["2020 Lexus GX-460 Premium Pkg"], lexus)
    tiers = {t["tier"]: t for t in x["price_tiers"]}
    viewed_total = sum(t["count"] for t in raw["shopped_by_price_tier"])
    check("price tiers carry % of viewed from the file's own shopped_by_price_tier",
          abs(tiers["Budget (<$30k)"]["share_of_viewed"] - 305 / viewed_total) < 1e-9
          and abs(sum(t["share_of_sold"] for t in x["price_tiers"]) - 1) < 1e-9, tiers)
    check("Budget over-indexes (24% of sold vs 18% of viewed); Core and Premium don't",
          [t["tier"] for t in x["price_tiers"] if t["over_indexes"]] == ["Budget (<$30k)"],
          x["price_tiers"])
    facts = ra.build_facts_payload(attribution, None, analyst=analyst, client_name="Ted Britt Ford")
    blob = json.dumps(facts["analyst"])
    vins = [v["vin"] for v in raw["missed_opportunities"]["vehicles"] if v.get("vin")]
    check(f"none of the file's {len(vins)} VINs reach the model-facing facts",
          vins and not any(v in blob for v in vins))
    check("the model never sees 'shopped', 'pipeline', a benchmark, an influence count or the "
          "visit bands (v2 fields the report deliberately leaves out)",
          not re.findall(r"shopped|pipeline|benchmark|influence|visit_band", blob.lower()),
          re.findall(r"\w*(?:shopped|pipeline|benchmark|influence|visit_band)\w*", blob.lower()))
    period = ra.period_facts_for_report(attribution, analyst=analyst)["analyst"]
    check("visit bands ride in period_facts only (for later months), never the draft payload",
          period["visit_bands"] and "look_to_book_by_visit_band" not in blob, period)

    for label, deck in (("without", None), ("with", ANALYST_DECK)):
        if deck is not None and not deck.exists():
            skip(f"{deck.name} not present")
            continue
        out = OUT_DIR / f"TB930_order_{label}_append.pptx"
        ra.build_report_deck(str(TEMPLATE), attribution, None, str(out),
                             client_name="Ted Britt", goals_bullets=[],
                             whats_next_bullets=["order check"], analyst=analyst,
                             extra_deck_path=str(deck) if deck else None)
        prs = Presentation(str(out))
        last = prs.slides[len(prs.slides._sldIdLst) - 1]
        check(f"Takeaways is the last slide ({label} the Analyst deck appended)",
              last.has_notes_slide and slide_map.notes_key(last) == "report:takeaways")
        check(f"...and the package is structurally clean ({label} append)",
              not package_check.check_package(str(out)))


TEMPLATE_V015 = REPO / "REPORT_MASTER_v0_15.pptx"


def _slide_texts(prs, key):
    slide = next((s for s in prs.slides if s.has_notes_slide and slide_map.notes_key(s) == key), None)
    if slide is None:
        return None, None
    texts, tables = [], {}
    for sh in slide_map.iter_all_shapes(slide.shapes):
        if sh.has_text_frame:
            texts.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            tables[sh.name] = [[c.text for c in r.cells] for r in sh.table.rows]
    return texts, tables


def check_native_slides():
    """REPORT_MASTER_v0_15's native Analyst slides (Matt's build_v0_15.py).
    A v2 file fills every column; a v1 file hides the v2-only columns."""
    print("\nNative Analyst slides (v0_15)")
    v2 = an.parse_analyst_facts(json.dumps(dict(json.loads(_synthetic()), schema_version=2)).encode())
    check("a schema_version 2 file is accepted", v2["schema_version"] == 2)
    draft = {"threads": [{"head": "W", "meaning": "m",
                          "action": "Prioritize the Missed Opportunities watch list below -- check photos."}],
             "url_intent_narrative": "The Store Scoreboard above shows it."}
    cleaned = app._strip_list_position_words(draft)
    check("'watch list below' / 'Scoreboard above' lose the position word, keep the name",
          cleaned["threads"][0]["action"] == "Prioritize the Missed Opportunities watch list -- check photos."
          and cleaned["url_intent_narrative"] == "The Store Scoreboard shows it.", cleaned)
    for path in (ATTRIBUTION_TB930, ANALYST_TB_AUG, TEMPLATE_V015):
        if not path.exists():
            skip(f"{path.name} not present")
            return
    attribution = ai.parse_attribution_export(str(ATTRIBUTION_TB930))
    analyst = an.parse_analyst_facts(str(ANALYST_TB_AUG))
    client = "Ted Britt Ford & Ted Britt Chantilly"
    check("client-store pre-guess finds the store named in the client's own name, by domain",
          ra.guess_client_stores(analyst, client) == ["tedbrittchantilly.com"],
          ra.guess_client_stores(analyst, client))

    def build(name, data, stores=None):
        out = OUT_DIR / name
        _, w = ra.build_report_deck(str(TEMPLATE_V015), attribution, None, str(out),
                                    client_name=client, goals_bullets=[],
                                    whats_next_bullets=["native check"], vertical="auto",
                                    analyst=data, analyst_client_stores=stores)
        return Presentation(str(out)), out, w

    prs, out, w = build("TB930_v015_v2.pptx", analyst, ["tedbrittchantilly.com"])
    keys = [slide_map.notes_key(s) for s in prs.slides]
    check("all three native slides present, after zip/ott and before Takeaways (last)",
          keys[-4:] == ["report:analyst_inventory", "report:analyst_watchlist",
                        "report:analyst_group", "report:takeaways"], keys)
    check("no fit warnings, package clean", not w and not package_check.check_package(str(out)),
          w)
    all_text = " ".join(sh.text_frame.text for s in prs.slides for sh in slide_map.iter_all_shapes(s.shapes)
                        if sh.has_text_frame)
    all_cells = " ".join(c.text for s in prs.slides for sh in slide_map.iter_all_shapes(s.shapes)
                         if getattr(sh, "has_table", False) and sh.has_table
                         for r in sh.table.rows for c in r.cells)
    check("no token left anywhere", "{{" not in all_text + all_cells)
    check("no 'shopped', 'pipeline' or influence wording anywhere in the deck",
          not re.findall(r"shopp|pipeline|influenc", (all_text + all_cells).lower()),
          re.findall(r"\w*(?:shopp|pipeline|influenc)\w*", (all_text + all_cells).lower()))
    texts, tables = _slide_texts(prs, "report:analyst_inventory")
    check("inventory tiles carry the file's own figures (939 / $39,168,431 / $75,200,770 / "
          "55.4% / 43.0% / 73.1%) under the v2 labels",
          all(v in texts for v in ("939", "$39,168,431", "$75,200,770", "55.4%", "43.0%", "73.1%",
                                   "Est. value sold", "Viewed vehicles sold", "Est. total value viewed")),
          texts)
    check("the two dollar tiles are never side by side (the units tile sits between them)",
          ra._ANALYST_ROW1_TILES == ("AnalystRevenueTile", "AnalystUnitsTile", "AnalystPipelineTile"),
          ra._ANALYST_ROW1_TILES)
    check("no influence tile, whatever the file carries (sold_above_benchmark is ignored)",
          not any("campaign visits" in t for t in texts), texts)
    check("tier table is Tier | Units sold | % of sold | % of viewed",
          tables["AnalystTierTable"][0] == ["Tier", "Units sold", "% of sold", "% of viewed"],
          tables["AnalystTierTable"][0])
    check("models table header and first row: Ford F-150, 213 viewed / 121 sold",
          tables["AnalystModelsTable"][0][:3] == ["Model", "Viewed", "Sold"]
          and tables["AnalystModelsTable"][1][:3] == ["Ford F-150", "213", "121"],
          tables["AnalystModelsTable"][:2])
    texts, tables = _slide_texts(prs, "report:analyst_watchlist")
    check("watch list: Est. price column, 10 rows, no VIN",
          tables["AnalystWatchlistTable"][0] == ["Vehicle", "New/Used", "Visits", "Est. price"]
          and tables["AnalystWatchlistTable"][1][-1] == "$45,000"
          and len(tables["AnalystWatchlistTable"]) == 11
          and "1FA6P8GJ3T5551409" not in str(tables), tables["AnalystWatchlistTable"][:2])
    texts, tables = _slide_texts(prs, "report:analyst_group")
    rows = tables["AnalystGroupTable"][1:]
    check("scoreboard: the Analyst's own dealer names, sorted by Look-to-Book, the client "
          "marked by domain",
          [r[0] for r in rows] == ["Ted Britt Ford of Chantilly", "Ted Britt Chantilly Lincoln",
                                   "Ted Britt Ford of Fairfax", "Ted Britt Chevrolet"]
          and rows[0][-1] == "Client" and all(r[-1] == "" for r in rows[1:]), rows)
    check("the truck shop (traffic, no vehicles) is named in the footnote by its domain -- its "
          "dealer_name is only the slug run together",
          any("tedbritttruckshop.com" in t for t in texts)
          and not any("Tedbritttruckshop" in t for t in texts), texts)
    check("store_label keeps a real dealer name and falls back to the domain only for a slug",
          an.store_label({"dealer_name": "Ted Britt Chevrolet", "domain": "tedbrittchevrolet.com"})
          == "Ted Britt Chevrolet"
          and an.store_label({"dealer_name": "Tedbritttruckshop", "domain": "tedbritttruckshop.com"})
          == "tedbritttruckshop.com")

    if ANALYST_TB_AUG_V1.exists():
        prs1, out1, w1 = build("TB930_v015_v1.pptx", an.parse_analyst_facts(str(ANALYST_TB_AUG_V1)))
        _t1, t1 = _slide_texts(prs1, "report:analyst_inventory")
        _t2, t2 = _slide_texts(prs1, "report:analyst_watchlist")
        check("a v1 file still builds: no '% of viewed' tier column, no 'Est. price' (v2-only)",
              t1["AnalystTierTable"][0] == ["Tier", "Units sold", "% of sold"]
              and t2["AnalystWatchlistTable"][0] == ["Vehicle", "New/Used", "Visits"]
              and not w1 and not package_check.check_package(str(out1)),
              (t1["AnalystTierTable"][0], t2["AnalystWatchlistTable"][0], w1))
    else:
        skip(f"{ANALYST_TB_AUG_V1.name} not present")

    prs3, _o3, _w3 = build("TB930_v015_none.pptx", None)
    keys3 = [slide_map.notes_key(s) for s in prs3.slides]
    check("no Analyst file: every analyst_set slide dropped, Takeaways still last",
          not any(k.startswith("report:analyst_") for k in keys3 if k) and keys3[-1] == "report:takeaways",
          keys3)


def main():
    check_native_slides()
    check_parser()
    check_gates()
    check_slide_helpers()
    check_real_ted_britt()
    check_fallback_url_report()
    check_cross_facts_and_order()
    print()
    print(f"{len(failures)} failure(s), {len(skipped)} skipped" if failures
          else f"All checks passed ({len(skipped)} skipped)")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
