"""notes_file_import.py -- pure, offline, no gitignored fixtures required for
.txt/.docx (both generated on the fly); .pdf reuses whichever real avails PDF
already sits at the repo root for other tests and SKIPs cleanly without one.

    python tests/test_notes_file_import.py
"""
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

import notes_file_import as nfi   # noqa: E402
from docx import Document          # noqa: E402

failures = []


def check(label, condition, detail=""):
    if condition:
        print(f"  PASS  {label}")
    else:
        print(f"  FAIL  {label}{'  -- ' + str(detail) if detail else ''}")
        failures.append(label)


def main():
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)

        print("=" * 78)
        print("SCENARIO  .txt notes")
        print("=" * 78)
        txt_path = tmp / "notes.txt"
        txt_path.write_text("Client wants a CTV campaign.\nBudget is $50,000.", encoding="utf-8")
        text, error = nfi.extract_notes_text(str(txt_path), "notes.txt")
        check("no error", error is None, error)
        check("real text extracted", text == "Client wants a CTV campaign.\nBudget is $50,000.", text)

        print("\nSCENARIO  .txt with bytes that aren't valid UTF-8 -- replaced, not raised")
        print("=" * 78)
        bad_path = tmp / "bad.txt"
        bad_path.write_bytes(b"Budget is $50,000 \xff\xfe more notes")
        text, error = nfi.extract_notes_text(str(bad_path), "bad.txt")
        check("no error -- bad bytes are replaced, not fatal", error is None, error)
        check("the readable parts survive", "Budget is $50,000" in text and "more notes" in text, text)

        print("\nSCENARIO  .docx notes")
        print("=" * 78)
        docx_path = tmp / "notes.docx"
        doc = Document()
        doc.add_paragraph("Client wants a CTV campaign.")
        doc.add_paragraph("Budget is $50,000.")
        doc.save(str(docx_path))
        text, error = nfi.extract_notes_text(str(docx_path), "notes.docx")
        check("no error", error is None, error)
        check("both paragraphs extracted, in order",
              text == "Client wants a CTV campaign.\nBudget is $50,000.", text)

        print("\nSCENARIO  an unsupported extension is reported, not raised")
        print("=" * 78)
        bad_ext = tmp / "notes.xlsx"
        bad_ext.write_bytes(b"not really an xlsx")
        text, error = nfi.extract_notes_text(str(bad_ext), "notes.xlsx")
        check("text is None", text is None, text)
        check("a plain-language error names the file", error and "notes.xlsx" in error, error)

        print("\nSCENARIO  a file that fails to open as the format its extension claims")
        print("=" * 78)
        corrupt_pdf = tmp / "corrupt.pdf"
        corrupt_pdf.write_bytes(b"this is not a real pdf")
        text, error = nfi.extract_notes_text(str(corrupt_pdf), "corrupt.pdf")
        check("text is None", text is None, text)
        check("a plain-language error names the file", error and "corrupt.pdf" in error, error)

        corrupt_docx = tmp / "corrupt.docx"
        corrupt_docx.write_bytes(b"this is not a real docx")
        text, error = nfi.extract_notes_text(str(corrupt_docx), "corrupt.docx")
        check("text is None", text is None, text)
        check("a plain-language error names the file", error and "corrupt.docx" in error, error)

    print("\nSCENARIO  a real avails PDF, read as notes text (not parsed for structure -- "
          "just proving real-world PDF text extraction works)")
    print("=" * 78)
    real_pdfs = list(REPO.glob("*.pdf"))
    if not real_pdfs:
        print("  SKIP -- no real PDF at the repo root")
    else:
        pdf_path = real_pdfs[0]
        text, error = nfi.extract_notes_text(str(pdf_path), pdf_path.name)
        check("no error", error is None, error)
        check("real text extracted, non-trivial length", bool(text) and len(text) > 200,
              len(text) if text else 0)

    print()
    if failures:
        print(f"{len(failures)} FAILED: {failures}")
        return 1
    print("Meeting notes uploaded as .txt, .pdf or .docx all extract to plain text; an "
          "unsupported extension or a corrupt file is reported in plain language, never raised.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
