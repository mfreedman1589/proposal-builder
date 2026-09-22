"""The Attribution Reports page (ATTRIBUTION_REPORT_PLAN.md Phase 2) --
through the real form via AppTest + test_mode_upload injection, the same
technique test_avails_pdf_wiring.py/test_wideorbit.py use for a
file_uploader AppTest can't drive directly. db.py's advertiser/proposal/
report functions are stubbed with a small in-memory fake store -- the real
`advertisers`/`attribution_reports` tables don't exist in Supabase until the
migration is pasted in by hand, and a test must never write real rows
either way (same reason test_draft_regression.py stubs db.log_proposal).

    python tests/test_attribution_reports_page.py
"""
import os
import sys
from datetime import date
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
os.chdir(REPO)
os.environ["PROPOSAL_BUILDER_TEST_MODE"] = "1"

import db  # noqa: E402
db.fetch_audiences = lambda: (None, "stubbed for test isolation")

import app  # noqa: E402
from streamlit.testing.v1 import AppTest  # noqa: E402

MW_FIXTURE = REPO / "MW attribution excel.xlsx"
DELIVERY_MW = REPO / "MW delivery.xlsx"
WAEPA_FIXTURE = REPO / "Premion Website Attribution and Reach Extension (13).xlsx"
POLK_FIXTURE = REPO / "Polk Dashboard.xlsx"

failures = []


def check(label, condition, detail=""):
    if condition:
        print(f"  PASS  {label}")
    else:
        print(f"  FAIL  {label}{'  -- ' + str(detail) if detail else ''}")
        failures.append(label)


# A pre-seeded `attr_draft` value for Generate-click checks that don't
# care about drafted CONTENT, only that Generate builds a deck -- setting
# this in session_state BEFORE clicking Generate makes `draft_to_use =
# attr_draft` non-None, so the self-sufficient "draft it myself" branch
# (2026-09-10 UX fix) never fires and no live Claude call happens.
# `apply_attr_draft` reads every field via `.get(...)` with an `or []`/
# `or None` fallback, so `{}` is a completely safe, empty draft.
#
# Real gotcha, found building this: `app.call_claude_attr_draft = stub`
# does NOT intercept the call `render_attribution_reports_page` makes when
# driven through `AppTest.from_file(...)` -- confirmed directly (a stub
# instance sitting in `app.__dict__["call_claude_attr_draft"]`, verified
# still there after the run, was never invoked; a real, billed API call
# happened instead). AppTest re-executes app.py's source into its own
# script-run context, and a bare-name call to a function DEFINED directly
# in app.py resolves against THAT execution's own globals, not the
# separately-`import app`-ed module object a test file holds a reference
# to -- a sibling of DECISIONS.md's "stubbing a function that is someone
# else's default argument is a silent no-op" trap, just triggered by
# AppTest's own re-exec model instead of a captured default parameter.
# `db.*` stubbing above works fine BECAUSE `db` is cached in `sys.modules`
# and shared regardless of how many times app.py's script body re-runs;
# a function defined directly in app.py has no such shared home to patch.
# This is exactly why test_draft_regression.py (Tier 1) never clicks a
# live-call button through AppTest either -- it calls `apply_draft_to_form`
# directly and injects the result into session_state, the same shape used
# here.
_ATTR_STUB_DRAFT = {}


class FakeStore:
    """Records every call and hands back plausible rows -- enough of
    db.py's real contract for app.py's logic to exercise, without touching
    Supabase.
    """
    def __init__(self):
        self.advertisers = []
        self.proposals = []
        self.reports = []
        self.case_studies = []
        self.create_advertiser_calls = []
        self.log_report_calls = []
        self.link_calls = []
        self.upload_report_file_calls = []
        self.delete_report_calls = []
        self.set_vertical_calls = []
        self.upload_case_study_calls = []
        self._next = 1

    def _id(self):
        self._next += 1
        return f"fake-{self._next}"

    def fetch_advertisers(self, limit=1000, active_only=False):
        rows = [r for r in self.advertisers if (not active_only or r.get("active", True))]
        return list(rows), None

    def create_advertiser(self, name):
        self.create_advertiser_calls.append(name)
        display = " ".join((name or "").split())
        key = display.lower()
        for row in self.advertisers:
            if row["canonical_name"].lower() == key:
                return row, None
        row = {"id": self._id(), "canonical_name": display}
        self.advertisers.append(row)
        return row, None

    def fetch_proposals(self, limit=500):
        return list(self.proposals), None

    def fetch_proposal(self, proposal_id):
        for row in self.proposals:
            if row["id"] == proposal_id:
                return row, None
        return None, "That proposal no longer exists."

    def link_proposal_advertiser(self, proposal_id, advertiser_id):
        self.link_calls.append((proposal_id, advertiser_id))
        return True, None

    def log_attribution_report(self, advertiser_id, proposal_id, report_json,
                               status="parsed", created_by=None, series_id=None):
        call = dict(advertiser_id=advertiser_id, proposal_id=proposal_id,
                   report_json=report_json, status=status, created_by=created_by,
                   series_id=series_id)
        self.log_report_calls.append(call)
        # Also kept as a real "logged report" row, so a LATER page render's
        # own prior_periods fetch (Highlights/Takeaways rework) can read it
        # back the same way `db.fetch_attribution_reports` would. ONE id
        # computed here and reused for both the row and the return value --
        # a real, pre-existing fixture bug (two separate `self._id()` calls)
        # meant the id handed back to the caller never matched any row in
        # `self.reports`, which `upload_report_file` below now depends on.
        report_id = self._id()
        self.reports.append({"id": report_id, "advertiser_id": advertiser_id,
                             "proposal_id": proposal_id, "report_json": report_json,
                             "created_by": created_by, "series_id": series_id,
                             "created_at": "2026-01-01T00:00:00"})
        return report_id, None

    def fetch_attribution_reports(self, advertiser_id=None, limit=500):
        rows = [r for r in self.reports if not advertiser_id or r.get("advertiser_id") == advertiser_id]
        return list(rows), None

    def set_report_series(self, report_id, series_id):
        for row in self.reports:
            if row["id"] == report_id:
                row["series_id"] = series_id
                return True, None
        return False, "report not found"

    def fetch_series_reports(self, series_id, advertiser_id=None):
        if not series_id:
            return [], None
        rows = [r for r in self.reports if r.get("series_id") == series_id
               and (not advertiser_id or r.get("advertiser_id") == advertiser_id)]
        return list(rows), None

    def update_attribution_report_json(self, report_id, report_json):
        for row in self.reports:
            if row["id"] == report_id:
                row["report_json"] = report_json
                return True, None
        return False, "report not found"

    def set_advertiser_vertical(self, advertiser_id, vertical):
        self.set_vertical_calls.append((advertiser_id, vertical))
        for row in self.advertisers:
            if row["id"] == advertiser_id:
                row["vertical"] = vertical
        return True, None

    def upload_report_file(self, report_id, local_path, filename):
        self.upload_report_file_calls.append((report_id, local_path, filename))
        storage_path = f"{report_id}/{filename}"
        for row in self.reports:
            if row["id"] == report_id:
                row["storage_path"] = storage_path
                return row, {}, None
        return None, None, "report not found"

    def report_file(self, report_id, storage_path):
        # Never actually read in these tests (nothing here clicks
        # "Download"); a real, existing local path is enough to satisfy
        # the type the caller expects if it ever is.
        return str(REPO / "app.py")

    def delete_attribution_report(self, report_id):
        self.delete_report_calls.append(report_id)
        using = [cs for cs in self.case_studies if cs.get("source_report_id") == report_id]
        if using:
            names = ", ".join(cs.get("title") or "(untitled)" for cs in using)
            return False, f"Can't delete -- used by {len(using)} case study/studies ({names})."
        before = len(self.reports)
        self.reports = [r for r in self.reports if r["id"] != report_id]
        return (len(self.reports) < before), None

    def upload_case_study(self, local_path, filename, title, verticals, products, summary,
                          added_by, optimize=True, advertiser_id=None, source_report_id=None):
        call = dict(local_path=local_path, filename=filename, title=title,
                   verticals=list(verticals), products=list(products), summary=summary,
                   added_by=added_by, advertiser_id=advertiser_id, source_report_id=source_report_id)
        self.upload_case_study_calls.append(call)
        row = {"id": self._id(), "title": title, "advertiser_id": advertiser_id,
              "source_report_id": source_report_id}
        self.case_studies.append(row)
        return row, {"source_size": 1000, "dst_size": 900}, None


def _fake_proposal_row(rid="prop-1", client_name="Mattress Warehouse", target_dmas=None,
                       goals="Drive qualified leads for the fall promotion",
                       audience="Adults 25-54 in-market for home improvement",
                       flight_label="Jun-Jul 2026", budget_cost="$50,000"):
    """Phase 5's field-map shape (ATTRIBUTION_REPORT_PLAN.md) -- real market
    KEYS (`washington_hagerstown`/`baltimore`, the same two the real WAEPA
    proposal uses) so `target_market_labels` resolves them against the real
    market-profile catalog, `campaign_specs.goals`/`.audience` (the goals-
    prefill bug's own real field), and a `deck_payload.media_plan_options`
    shape matching what a real proposal's Generate actually writes."""
    return {
        "id": rid, "client_name": client_name, "vertical": "retail", "market": "DC",
        "generated_at": "2026-08-01T12:00:00", "deck_version_id": 5,
        "parent_proposal_id": None, "revision_label": None,
        "file_storage_path": None, "logo_storage_path": None, "created_by": "Matt",
        "target_dmas": target_dmas if target_dmas is not None else ["washington_hagerstown", "baltimore"],
        "form_json": {
            "proposal_title": "CTV Strategy", "plan_options": [],
            # "start"/"end" match "setup" below -- a real proposal's Generate
            # writes both sections from the same source dates (app.py's own
            # form_json["flight"] construction), and `linked_proposal_report_
            # fields` reads THIS "flight" dict for its own flight_start/
            # flight_end, not "setup" -- Part 2 rework's report-type
            # inference is the first thing in this test file to depend on
            # that field actually being populated.
            "flight": {"label": flight_label, "start": "2026-06-01", "end": "2026-07-17"},
            "setup": {"originating_market": "DC", "flight_start": "2026-06-01",
                     "flight_end": "2026-07-17", "plan_basis": "monthly",
                     "avails_mode": True, "total_tv": False},
            "campaign_specs": {"goals": goals, "audience": audience},
            "deck_payload": {"media_plan_options": [
                {"full_flight_total": {"cost": budget_cost, "impressions": "1,000,000"}}]},
            "targeting_groups": [],
        },
    }


def _ss(at, key, default=None):
    """at.session_state has no .get() -- AppTest's SafeSessionState only
    proxies __getitem__/__contains__, and .get resolves as an attribute
    lookup for a literal key named "get" instead. Every other test file
    here works around this with an explicit `in` check; this is that
    check, factored out since this file needs it repeatedly.
    """
    return at.session_state[key] if key in at.session_state else default


def new_app():
    at = AppTest.from_file(str(REPO / "app.py"), default_timeout=300)
    at.session_state["authed"] = True
    at.session_state["current_user"] = "Test"
    return at


def _confirm_advertiser(at):
    """Clicks through section 2's advertiser confirm, whichever of its two
    shapes actually rendered. `store` (FakeStore) is ONE instance shared
    across every check in `main()`'s sequence, so `create_advertiser` calls
    accumulate real entries across tests -- a later test whose export's
    client name exactly matches an advertiser an EARLIER test already
    created (a real, common case: several tests share MW/WAEPA) hits
    section 2's own exact-match shortcut (a lone `st.success` + a
    `key="attr_confirm_exact_advertiser"` button), not the radio +
    "Confirm client" button a brand-new name gets. Found the hard way:
    a test assuming only the radio shape failed silently, since neither
    branch raises when the WRONG button is clicked (or none is) -- it just
    never sets `attr_advertiser_id`."""
    exact_buttons = [b for b in at.button if b.key == "attr_confirm_exact_advertiser"]
    if exact_buttons:
        exact_buttons[0].click().run()
        return
    radios = [r for r in at.radio if r.key == "attr_advertiser_choice"]
    if radios:
        radios[0].set_value(radios[0].options[-1]).run()
        confirm_buttons = [b for b in at.button if b.label == "Confirm client"]
        if confirm_buttons:
            confirm_buttons[0].click().run()


def check_upload_first_new_advertiser_no_proposal(store):
    print("\nUpload-first door: brand-new advertiser, no matching proposal")
    if not MW_FIXTURE.exists():
        print("  SKIP  MW attribution excel.xlsx not present")
        return

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_attribution_upload_path"] = str(MW_FIXTURE)
    at.run()
    check("no exception after upload", not at.exception, at.exception)

    parsed = _ss(at, "attr_parsed_attribution")
    check("the export parsed", bool(parsed), parsed)
    if parsed:
        check("client name resolved from the export", parsed.get("client_name"), "Mattress Warehouse")
        check("client name is exactly 'Mattress Warehouse'",
             parsed.get("client_name") == "Mattress Warehouse", parsed.get("client_name"))

    radios = at.radio
    check("an advertiser radio is shown", len(radios) >= 1, len(radios))
    if not radios:
        return
    advertiser_radio = radios[0]
    check("its only option is 'create new'", "Create new client" in advertiser_radio.options[-1],
         advertiser_radio.options)
    advertiser_radio.set_value(advertiser_radio.options[-1]).run()

    confirm_buttons = [b for b in at.button if b.label == "Confirm client"]
    check("a Confirm advertiser button is present", bool(confirm_buttons), [b.label for b in at.button])
    if confirm_buttons:
        confirm_buttons[0].click().run()
    check("no exception after confirming the advertiser", not at.exception, at.exception)
    check("attr_advertiser_id got set", _ss(at, "attr_advertiser_id"),
         _ss(at, "attr_advertiser_id"))
    check("create_advertiser was called with the parsed client name",
         store.create_advertiser_calls == ["Mattress Warehouse"], store.create_advertiser_calls)

    no_proposal_buttons = [b for b in at.button
                           if b.label == "No proposal -- build this report standalone"]
    check("the no-matching-proposal path offers a standalone button",
         bool(no_proposal_buttons), [b.label for b in at.button])
    if no_proposal_buttons:
        no_proposal_buttons[0].click().run()
    check("attr_no_proposal is set", _ss(at, "attr_no_proposal") is True,
         _ss(at, "attr_no_proposal"))

    # "Log this report" is gone as its own step (2026-09-10 UX fix) --
    # generating a report logs it, with no separate click.
    check("no standalone 'Log this report' button exists any more",
         not any(b.label == "Log this report" for b in at.button), [b.label for b in at.button])

    print("\n  Generate the report deck (Phase 3 walking skeleton)")
    template = REPO / "REPORT_MASTER_v0_6.pptx"
    if not template.exists():
        print("  SKIP  REPORT_MASTER_v0_6.pptx not present")
        return
    goals_areas = [t for t in at.text_area if t.key == "attr_goals_input"]
    whats_next_areas = [t for t in at.text_area if t.key == "attr_whats_next_input"]
    check("goals text area is present", bool(goals_areas), [t.key for t in at.text_area])
    check("what's-next text area is present", bool(whats_next_areas), [t.key for t in at.text_area])
    if goals_areas and whats_next_areas:
        goals_areas[0].set_value("Drive online engagement for the sale").run()
        whats_next_areas[0].set_value("Expand into the top-performing markets").run()
    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    check("a Generate report deck button is present", bool(generate_buttons),
         [b.label for b in at.button])
    if generate_buttons:
        # A pre-seeded (empty) draft means "already previewed" -- Generate
        # uses it as-is rather than drafting one itself, so this click
        # never makes a live Claude call. See _ATTR_STUB_DRAFT's own
        # comment for why stubbing the drafting function itself doesn't
        # work through AppTest.
        at.session_state["attr_draft"] = _ATTR_STUB_DRAFT
        generate_buttons[0].click().run()
    check("no exception after generating", not at.exception, at.exception)
    download_buttons = [b for b in at.download_button if b.label == "⬇ Download report .pptx"]
    check("a download button appears after a successful generate",
         bool(download_buttons), [b.label for b in at.download_button])

    check("Generate logged the report itself, with no separate click",
         len(store.log_report_calls) == 1, store.log_report_calls)
    if store.log_report_calls:
        call = store.log_report_calls[0]
        check("logged with the created advertiser id",
             call["advertiser_id"] == _ss(at, "attr_advertiser_id"), call)
        check("logged with no proposal", call["proposal_id"] is None, call)
        check("report_json carries the parsed attribution facts",
             call["report_json"].get("attribution", {}).get("client_name") == "Mattress Warehouse",
             call["report_json"])
        check("report_json's delivery half is None (no delivery file uploaded)",
             call["report_json"].get("delivery") is None, call["report_json"])


def check_rfpid_confirm_gate(store):
    print("\nMulti-RFPID confirm gate (WAEPA: 2 RFPIDs, a real split IO)")
    if not WAEPA_FIXTURE.exists():
        print("  SKIP  WAEPA fixture not present")
        return

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_attribution_upload_path"] = str(WAEPA_FIXTURE)
    at.run()
    check("no exception after upload", not at.exception, at.exception)

    rfpid_checkboxes = [c for c in at.checkbox if c.key == "attr_rfpid_confirmed"]
    check("the RFPID confirm checkbox is shown", bool(rfpid_checkboxes),
         [c.key for c in at.checkbox])
    if not rfpid_checkboxes:
        return
    check("2 RFPIDs defaults to confirmed (the split-IO shape, not the rollup one)",
         rfpid_checkboxes[0].value is True, rfpid_checkboxes[0].value)

    # Default confirmed -- the page proceeds past the gate into advertiser
    # matching (a radio, since this fixture's advertiser doesn't yet exist
    # in the fake store).
    check("advertiser matching renders past the gate (default confirmed)",
         len(at.radio) >= 1, len(at.radio))

    # Uncheck it -- the gate blocks, advertiser matching disappears. (Some
    # radio may still exist elsewhere on the page -- e.g. sidebar nav --
    # so this checks for the SPECIFIC advertiser-matching radio's own
    # option text going away, not a bare app-wide radio count.)
    rfpid_checkboxes[0].set_value(False).run()
    check("no exception after unchecking", not at.exception, at.exception)
    check("unchecking the confirm blocks advertiser matching from rendering",
         not any("Create new client" in opt for r in at.radio for opt in r.options),
         [r.options for r in at.radio])
    check("an info message explains what to do",
         any("Confirm above" in str(m) for m in _all_markdown_text(at)), _all_markdown_text(at))


def check_conversions_toggle(store):
    print("\nConversions toggle (WAEPA has them, MW doesn't)")
    if not WAEPA_FIXTURE.exists():
        print("  SKIP  WAEPA fixture not present")
        return

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_attribution_upload_path"] = str(WAEPA_FIXTURE)
    at.run()
    # Default-confirmed RFPID gate lets the rest of the page through with
    # no extra interaction (see check_rfpid_confirm_gate).
    radios = at.radio
    check("advertiser radio present", bool(radios), [r.label for r in radios])
    if not radios:
        return
    radios[0].set_value(radios[0].options[-1]).run()
    confirm_buttons = [b for b in at.button if b.label == "Confirm client"]
    if confirm_buttons:
        confirm_buttons[0].click().run()
    no_proposal_buttons = [b for b in at.button
                           if b.label == "No proposal -- build this report standalone"]
    if no_proposal_buttons:
        no_proposal_buttons[0].click().run()

    conv_checkboxes = [c for c in at.checkbox if c.key == "attr_include_conversions"]
    check("the Include conversions checkbox is shown for a conversions export",
         bool(conv_checkboxes), [c.key for c in at.checkbox])
    if conv_checkboxes:
        check("it defaults ON", conv_checkboxes[0].value is True, conv_checkboxes[0].value)

    template = REPO / "REPORT_MASTER_v0_6.pptx"
    if not template.exists():
        print("  SKIP  REPORT_MASTER_v0_6.pptx not present -- can't test Generate")
        return
    goals_areas = [t for t in at.text_area if t.key == "attr_goals_input"]
    whats_next_areas = [t for t in at.text_area if t.key == "attr_whats_next_input"]
    if goals_areas and whats_next_areas:
        goals_areas[0].set_value("Drive qualified insurance leads").run()
        whats_next_areas[0].set_value("Expand DC/Baltimore targeting").run()
    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    check("a Generate report deck button is present", bool(generate_buttons),
         [b.label for b in at.button])
    if generate_buttons:
        at.session_state["attr_draft"] = _ATTR_STUB_DRAFT  # see its own comment -- no live call
        generate_buttons[0].click().run()
    check("no exception generating WITH conversions on", not at.exception, at.exception)

    # Turn it off and generate again -- must still build cleanly (the
    # graceful-degrade path when the template hasn't been widened yet is
    # exercised by report_assembly's own tests; this just proves the app
    # wiring survives the toggle both ways).
    conv_checkboxes = [c for c in at.checkbox if c.key == "attr_include_conversions"]
    if conv_checkboxes:
        conv_checkboxes[0].set_value(False).run()
    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    if generate_buttons:
        at.session_state["attr_draft"] = _ATTR_STUB_DRAFT
        generate_buttons[0].click().run()
    check("no exception generating WITH conversions off", not at.exception, at.exception)


def check_mw_has_no_conversions_toggle(store):
    print("\nMW (no conversions) shows no toggle at all")
    if not MW_FIXTURE.exists():
        print("  SKIP  MW attribution excel.xlsx not present")
        return
    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_attribution_upload_path"] = str(MW_FIXTURE)
    at.run()
    conv_checkboxes = [c for c in at.checkbox if c.key == "attr_include_conversions"]
    check("no Include conversions checkbox for a non-conversions export",
         not conv_checkboxes, [c.key for c in at.checkbox])


def check_polk_upload_and_projection_toggle(store):
    """Phase 7 -- the Polk upload slot and its "Project for match rate"
    toggle. The toggle only appears once a Polk file is actually uploaded
    (a toggle with nothing to project against is clutter), defaults off,
    and Generate builds cleanly whether or not the current template has
    caught up to report:automotive_registrations yet (report_assembly's
    own tests cover the drop-if-absent branch in detail; this just proves
    the app-level wiring survives end to end through a real Generate
    click)."""
    print("\nPolk automotive match-back upload + 'Project for match rate' toggle")
    if not MW_FIXTURE.exists() or not POLK_FIXTURE.exists():
        print("  SKIP  MW attribution excel.xlsx or Polk Dashboard.xlsx not present")
        return

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_attribution_upload_path"] = str(MW_FIXTURE)
    at.run()
    _confirm_advertiser(at)
    no_proposal_buttons = [b for b in at.button
                           if b.label == "No proposal -- build this report standalone"]
    if no_proposal_buttons:
        no_proposal_buttons[0].click().run()

    check("no Polk file yet -> no 'Project for match rate' checkbox",
         not [c for c in at.checkbox if c.key == "attr_polk_projected"],
         [c.key for c in at.checkbox])

    at.session_state["attr_polk_upload_path"] = str(POLK_FIXTURE)
    at.run()
    polk_checkboxes = [c for c in at.checkbox if c.key == "attr_polk_projected"]
    check("a Polk file uploaded -> the toggle appears", bool(polk_checkboxes),
         [c.key for c in at.checkbox])
    if not polk_checkboxes:
        return
    check("it defaults OFF", polk_checkboxes[0].value is False, polk_checkboxes[0].value)

    # Phase 8: Generate now blocks on a Polk file with no "months covered"
    # entered (see check_polk_months_gate below for the block itself) --
    # every Generate click in THIS check is testing something else, so it
    # sets months=1 (a no-op under normalize_for_months) to get past the
    # gate rather than re-proving the gate here too.
    months_inputs = [n for n in at.number_input if n.key == "attr_polk_months"]
    if months_inputs:
        months_inputs[0].set_value(1).run()

    template = REPO / "REPORT_MASTER_v0_6.pptx"
    if not template.exists():
        print("  SKIP  REPORT_MASTER_v0_6.pptx not present -- can't test Generate")
        return
    goals_areas = [t for t in at.text_area if t.key == "attr_goals_input"]
    whats_next_areas = [t for t in at.text_area if t.key == "attr_whats_next_input"]
    if goals_areas and whats_next_areas:
        goals_areas[0].set_value("Drive incremental vehicle sales").run()
        whats_next_areas[0].set_value("Expand the winning audience segment").run()

    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    check("a Generate report deck button is present", bool(generate_buttons),
         [b.label for b in at.button])
    if generate_buttons:
        # No live Claude call -- see _ATTR_STUB_DRAFT's own comment above.
        at.session_state["attr_draft"] = _ATTR_STUB_DRAFT
        generate_buttons[0].click().run()
    check("no exception generating with a Polk file present (whether or not the "
         "current template has report:automotive_registrations yet)",
         not at.exception, at.exception)

    # This drives the real app end to end against whatever template is
    # LIVE in Supabase right now (db.report_master_deck isn't stubbed here
    # -- only the advertisers/attribution_reports tables are, per this
    # file's own header) -- so once v0_12 is uploaded and activated, the
    # freshest .pptx in the scratch dir this Generate click just wrote
    # should genuinely carry the filled slide, not just build without
    # raising. Scans by mtime rather than reconstructing app.py's own
    # out_path naming (client_name-derived), so this doesn't drift if that
    # naming ever changes.
    from pptx import Presentation as _Presentation
    scratch = db.scratch_dir("attribution_reports")
    pptx_files = sorted(scratch.glob("*.pptx"), key=lambda p: p.stat().st_mtime)
    if pptx_files:
        built = _Presentation(str(pptx_files[-1]))
        keys = [s.notes_slide.notes_text_frame.text.splitlines()[0].replace("key: ", "")
               for s in built.slides if s.has_notes_slide]
        if "report:automotive_registrations" in keys:
            auto_slide = built.slides[keys.index("report:automotive_registrations")]
            text = "\n".join(sh.text_frame.text for sh in auto_slide.shapes if sh.has_text_frame)
            check("the live template's automotive_registrations slide is actually filled "
                 "with this fixture's real matched-households figure, through a real "
                 "Generate click -- not just 'no exception'",
                 "44,756" in text, text[:1500])
            check("no unfilled {{TOKEN}} survived on the live-built slide",
                 "{{" not in text, text[:300])

    # Toggle it on and generate again -- must still build cleanly either way.
    polk_checkboxes = [c for c in at.checkbox if c.key == "attr_polk_projected"]
    if polk_checkboxes:
        polk_checkboxes[0].set_value(True).run()
    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    if generate_buttons:
        at.session_state["attr_draft"] = _ATTR_STUB_DRAFT
        generate_buttons[0].click().run()
    check("no exception generating with the projection toggle ON",
         not at.exception, at.exception)

    # "Clear / new report" sweeps every attr_-prefixed key by construction
    # (clear_attribution_report_state's own prefix-sweep docstring) -- the
    # new Polk keys need no special-casing, but this proves it rather than
    # assuming it.
    clear_buttons = [b for b in at.button if "Clear" in b.label]
    if clear_buttons:
        clear_buttons[0].click().run()
        confirm_clear = [b for b in at.button if b.key == "attr_confirm_clear_yes"]
        if confirm_clear:
            confirm_clear[0].click().run()
        check("Clear / new report also sweeps the Polk keys",
             "attr_polk_path" not in at.session_state
             and "attr_parsed_polk" not in at.session_state)


def check_polk_phase8_wiring(store):
    """Phase 8 -- the months-covered gate (blocks Generate with no value),
    the ROI/Cost Per Visit toggles surfacing their own fields, and that a
    real Generate click with the gate satisfied still builds cleanly.
    report_assembly's own test suite covers the pure math (normalize_
    for_months, compute_roi, compute_cost_per_visit, polk_dealer_table_
    rows) in detail; this proves the app-level wiring, same division of
    labor as check_polk_upload_and_projection_toggle above."""
    print("\nPolk Phase 8 -- months gate, ROI/Cost Per Visit toggles")
    if not MW_FIXTURE.exists() or not POLK_FIXTURE.exists():
        print("  SKIP  MW attribution excel.xlsx or Polk Dashboard.xlsx not present")
        return

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_attribution_upload_path"] = str(MW_FIXTURE)
    at.run()
    _confirm_advertiser(at)
    no_proposal_buttons = [b for b in at.button
                           if b.label == "No proposal -- build this report standalone"]
    if no_proposal_buttons:
        no_proposal_buttons[0].click().run()
    at.session_state["attr_polk_upload_path"] = str(POLK_FIXTURE)
    at.run()

    months_inputs = [n for n in at.number_input if n.key == "attr_polk_months"]
    check("the 'Months covered' input appears once a Polk file is uploaded",
         bool(months_inputs), [n.key for n in at.number_input])
    if months_inputs:
        check("it starts blank (None), not defaulted to 1",
             months_inputs[0].value is None, months_inputs[0].value)

    goals_areas = [t for t in at.text_area if t.key == "attr_goals_input"]
    whats_next_areas = [t for t in at.text_area if t.key == "attr_whats_next_input"]
    if goals_areas and whats_next_areas:
        goals_areas[0].set_value("Drive incremental vehicle sales").run()
        whats_next_areas[0].set_value("Expand the winning audience segment").run()

    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    if generate_buttons:
        at.session_state["attr_draft"] = _ATTR_STUB_DRAFT
        generate_buttons[0].click().run()
    error_texts = [str(getattr(e, "value", "")) for e in at.error]
    check("Generate is BLOCKED with a Polk file present and no months entered",
         any("Months covered by this Polk report" in t for t in error_texts), error_texts)

    months_inputs = [n for n in at.number_input if n.key == "attr_polk_months"]
    if months_inputs:
        months_inputs[0].set_value(2).run()

    roi_checkboxes = [c for c in at.checkbox if c.key == "attr_polk_roi_on"]
    check("the ROI toggle is present once a Polk file is uploaded", bool(roi_checkboxes),
         [c.key for c in at.checkbox])
    if roi_checkboxes:
        check("ROI defaults off", roi_checkboxes[0].value is False, roi_checkboxes[0].value)
        roi_checkboxes[0].set_value(True).run()
        check("ticking it reveals profit-per-vehicle and cost inputs",
             any(n.key == "attr_polk_profit_per_vehicle" for n in at.number_input)
             and any(n.key == "attr_polk_roi_cost" for n in at.number_input),
             [n.key for n in at.number_input])

    cpv_checkboxes = [c for c in at.checkbox if c.key == "attr_cpv_on"]
    check("the Cost Per Visit toggle is present regardless of Polk (a separate feature)",
         bool(cpv_checkboxes), [c.key for c in at.checkbox])
    if cpv_checkboxes:
        check("Cost Per Visit defaults off", cpv_checkboxes[0].value is False)
        cpv_checkboxes[0].set_value(True).run()
        check("ticking it reveals separate CTV and retargeting cost inputs -- never one "
             "blended field",
             any(n.key == "attr_cpv_ctv_cost" for n in at.number_input)
             and any(n.key == "attr_cpv_retargeting_cost" for n in at.number_input),
             [n.key for n in at.number_input])

    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    if generate_buttons:
        at.session_state["attr_draft"] = _ATTR_STUB_DRAFT
        generate_buttons[0].click().run()
    check("with months entered and ROI/CPV toggles on, Generate builds cleanly",
         not at.exception, at.exception)


def check_polk_phase8_persistence(store):
    """Phase 8 follow-up (2026-09-21): every confirmed rep input -- months
    covered, sales-through date, ROI toggle+profit/cost, CPV toggle+costs
    -- is persisted onto the logged report's own `report_json["phase8_
    inputs"]`, and a case study built LATER from that report reuses the
    confirmed CTV-only cost automatically, never the proposal's own
    blended budget and never re-asking. Uses a linked proposal with a
    REAL, different blended budget ($50,000) specifically so the
    precedence is actually being tested -- a case study with no proposal
    at all would pass even with the precedence rule backwards, since
    there'd be nothing to prefer the confirmed figure over."""
    print("\nPolk Phase 8 -- persisted inputs, reused without asking")
    if not MW_FIXTURE.exists() or not POLK_FIXTURE.exists():
        print("  SKIP  MW attribution excel.xlsx or Polk Dashboard.xlsx not present")
        return

    fake_row = _fake_proposal_row(rid="prop-phase8-1", budget_cost="$50,000")
    store.proposals = [fake_row]
    before_report_count = len(store.reports)

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_prelinked_proposal_id"] = "prop-phase8-1"
    at.session_state["attr_attribution_upload_path"] = str(MW_FIXTURE)
    at.session_state["attr_polk_upload_path"] = str(POLK_FIXTURE)
    at.run()
    check("no exception with a linked proposal + attribution + Polk upload",
         not at.exception, at.exception)

    months_inputs = [n for n in at.number_input if n.key == "attr_polk_months"]
    if months_inputs:
        months_inputs[0].set_value(3).run()
    sales_through_inputs = [d for d in at.date_input if d.key == "attr_polk_sales_through"]
    if sales_through_inputs:
        sales_through_inputs[0].set_value(date(2026, 7, 31)).run()

    roi_checkboxes = [c for c in at.checkbox if c.key == "attr_polk_roi_on"]
    if roi_checkboxes:
        roi_checkboxes[0].set_value(True).run()
    profit_inputs = [n for n in at.number_input if n.key == "attr_polk_profit_per_vehicle"]
    if profit_inputs:
        profit_inputs[0].set_value(2500.0).run()
    roi_cost_inputs = [n for n in at.number_input if n.key == "attr_polk_roi_cost"]
    if roi_cost_inputs:
        roi_cost_inputs[0].set_value(12000.0).run()

    cpv_checkboxes = [c for c in at.checkbox if c.key == "attr_cpv_on"]
    if cpv_checkboxes:
        cpv_checkboxes[0].set_value(True).run()
    # Deliberately NOT $50,000 (the proposal's own blended budget) -- the
    # whole point of this check is proving the confirmed figure wins, not
    # just that some number shows up.
    ctv_cost_inputs = [n for n in at.number_input if n.key == "attr_cpv_ctv_cost"]
    if ctv_cost_inputs:
        ctv_cost_inputs[0].set_value(9000.0).run()
    retargeting_cost_inputs = [n for n in at.number_input if n.key == "attr_cpv_retargeting_cost"]
    if retargeting_cost_inputs:
        retargeting_cost_inputs[0].set_value(0.0).run()

    goals_areas = [t for t in at.text_area if t.key == "attr_goals_input"]
    whats_next_areas = [t for t in at.text_area if t.key == "attr_whats_next_input"]
    if goals_areas and whats_next_areas:
        goals_areas[0].set_value("Drive incremental vehicle sales").run()
        whats_next_areas[0].set_value("Expand the winning audience segment").run()

    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    check("a Generate report deck button is present", bool(generate_buttons),
         [b.label for b in at.button])
    if not generate_buttons:
        return
    at.session_state["attr_draft"] = _ATTR_STUB_DRAFT
    generate_buttons[0].click().run()
    check("no exception generating with months/ROI/CPV all confirmed", not at.exception, at.exception)
    check("exactly one new report was logged", len(store.reports) == before_report_count + 1,
         len(store.reports))
    if len(store.reports) != before_report_count + 1:
        return

    logged = store.log_report_calls[-1]["report_json"].get("phase8_inputs") or {}
    check("months covered persisted", logged.get("polk_months") == 3, logged)
    check("sales-through date persisted as an ISO string",
         logged.get("polk_sales_through") == "2026-07-31", logged)
    check("ROI toggle + profit per vehicle + cost persisted",
         logged.get("polk_roi_on") is True and logged.get("polk_profit_per_vehicle") == 2500.0
         and logged.get("polk_roi_cost") == 12000.0, logged)
    check("CPV toggle + both confirmed cost figures persisted",
         logged.get("cpv_on") is True and logged.get("cpv_ctv_cost") == 9000.0
         and logged.get("cpv_retargeting_cost") == 0.0, logged)

    report_row = store.reports[-1]
    rid = report_row["id"]
    print("  Report history row: create a case study -- must reuse the confirmed CTV cost, "
         "never the proposal's blended budget, with no re-asking")
    cs_buttons = [b for b in at.button if b.key == f"rpt_cs_button_{rid}"]
    check("the row's own 'Create case study' button is present", bool(cs_buttons),
         [b.key for b in at.button if b.key and b.key.startswith("rpt_cs_")])
    if not cs_buttons:
        return
    at.session_state[f"rpt_cs_tags_{rid}"] = {"verticals": ["auto"]}
    cs_buttons[0].click().run()
    check("no exception opening the case study review", not at.exception, at.exception)

    from pptx import Presentation as _Presentation
    cs_files = sorted(db.scratch_dir("attribution_reports").glob(f"case_study_{rid}.pptx"),
                      key=lambda p: p.stat().st_mtime)
    check("a case study .pptx was actually built", bool(cs_files))
    if not cs_files:
        return
    built = _Presentation(str(cs_files[-1]))
    text = "\n".join(sh.text_frame.text for sl in built.slides for sh in sl.shapes
                     if sh.has_text_frame)
    mw_attribution = app.attribution_import.parse_attribution_export(str(MW_FIXTURE))
    expected_from_ctv_cost = f"${9000.0 / mw_attribution.attributed_unique_visitors:,.0f}"
    expected_from_blended_budget = f"${50000.0 / mw_attribution.attributed_unique_visitors:,.0f}"
    check("the case study's Cost Per Visitor tile uses the PERSISTED, confirmed CTV cost "
         "($9,000), not the proposal's own blended budget ($50,000)",
         expected_from_ctv_cost in text and (expected_from_ctv_cost == expected_from_blended_budget
                                             or expected_from_blended_budget not in text),
         (expected_from_ctv_cost, expected_from_blended_budget, text[:1200]))


def check_clear_button(store):
    print("\nClear / new report")
    if not MW_FIXTURE.exists():
        print("  SKIP  MW attribution excel.xlsx not present")
        return
    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_attribution_upload_path"] = str(MW_FIXTURE)
    at.run()
    check("something got parsed before clearing", bool(_ss(at, "attr_parsed_attribution")),
         _ss(at, "attr_parsed_attribution"))

    clear_buttons = [b for b in at.button if b.key == "attr_clear_button"]
    check("the Clear / new report button is present", bool(clear_buttons),
         [b.key for b in at.button])
    if not clear_buttons:
        return
    clear_buttons[0].click().run()
    confirm_buttons = [b for b in at.button if b.key == "attr_confirm_clear_yes"]
    check("a confirm step appears rather than clearing immediately",
         bool(confirm_buttons), [b.key for b in at.button])
    if not confirm_buttons:
        return
    confirm_buttons[0].click().run()
    check("no exception after confirming", not at.exception, at.exception)
    check("the parsed attribution export is gone",
         _ss(at, "attr_parsed_attribution") is None, _ss(at, "attr_parsed_attribution"))
    check("the attribution upload path is gone",
         _ss(at, "attr_attribution_path") is None, _ss(at, "attr_attribution_path"))


def check_prelinked_door(store):
    print("\nPre-linked door: 'Build report' from Proposal History")
    fake_row = _fake_proposal_row(rid="prop-history-1", client_name="Cardinal Plumbing")
    store.proposals = [fake_row]

    print("  History row exposes a 'Build report' button that sets the right flags")
    at = new_app()
    at.session_state["page_choice"] = "Proposal history"
    at.run()
    check("no exception rendering History", not at.exception, at.exception)
    build_report_buttons = [b for b in at.button if b.label == "Build report"]
    check("a 'Build report' button is present on the row", bool(build_report_buttons),
         [b.label for b in at.button])
    if build_report_buttons:
        build_report_buttons[0].click().run()
    check("no exception after clicking it", not at.exception, at.exception)
    check("it navigated to Attribution reports",
         _ss(at, "page_choice") == "Attribution reports",
         _ss(at, "page_choice"))
    check("the proposal id was recorded",
         _ss(at, "attr_prelinked_proposal_id") == "prop-history-1",
         _ss(at, "attr_prelinked_proposal_id"))

    print("  The page shows the linked proposal and does not re-search")
    check("the linked proposal's client name is shown",
         any("Cardinal Plumbing" in str(m) for m in _all_markdown_text(at)), None)


def check_phase5_proposal_link(store):
    """ATTRIBUTION_REPORT_PLAN.md Phase 5, through the real page --
    complements test_report_assembly.py's own Phase 5 check (which drives
    report_assembly.py's functions directly, offline) by proving the app
    WIRING actually reaches them: the goals-prefill fix, and that
    Generate's real deck carries the linked proposal's flight/geography/
    budget-derived facts. The advertiser name is deliberately made to
    DISAGREE with the export's own ("WAEPA Insurance" vs. the export's
    "WAEPA") -- proving the linked proposal's name wins throughout,
    including the client name in the finished deck's own CLIENT_NAME/
    recap tokens and the download filename.
    """
    print("\nPhase 5: proposal-linked fields reach the real page and the generated deck")
    if not WAEPA_FIXTURE.exists():
        print("  SKIP  WAEPA fixture not present")
        return
    template = REPO / "REPORT_MASTER_v0_6.pptx"
    if not template.exists():
        print("  SKIP  REPORT_MASTER_v0_6.pptx not present -- can't test Generate")
        return

    fake_row = _fake_proposal_row(
        rid="prop-phase5-1", client_name="WAEPA Insurance",
        goals="Evaluate OTT performance beyond awareness\nTrack engaged visits and cost per engaged visit",
        audience="Federal government employees researching benefits",
        flight_label="Oct 2026 - Dec 2026", budget_cost="$74,970")
    store.proposals = [fake_row]

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_prelinked_proposal_id"] = "prop-phase5-1"
    at.session_state["attr_attribution_upload_path"] = str(WAEPA_FIXTURE)
    at.run()
    check("no exception with a linked proposal + upload", not at.exception, at.exception)

    goals_areas = [t for t in at.text_area if t.key == "attr_goals_input"]
    check("goals text area is present", bool(goals_areas), [t.key for t in at.text_area])
    if goals_areas:
        check("goals prefilled from campaign_specs.goals -- the real prefill bug's own fix "
             "(it used to read a top-level 'goals' key that doesn't exist on any real "
             "proposal, silently prefilling '')",
             "Evaluate OTT performance" in goals_areas[0].value, goals_areas[0].value)

    whats_next_areas = [t for t in at.text_area if t.key == "attr_whats_next_input"]
    check("what's-next text area is present", bool(whats_next_areas),
         [t.key for t in at.text_area])
    if whats_next_areas:
        whats_next_areas[0].set_value("Expand DC/Baltimore targeting").run()

    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    check("a Generate report deck button is present", bool(generate_buttons),
         [b.label for b in at.button])
    if not generate_buttons:
        return
    at.session_state["attr_draft"] = _ATTR_STUB_DRAFT  # see its own comment -- no live call
    generate_buttons[0].click().run()
    check("no exception after generating", not at.exception, at.exception)
    download_buttons = [b for b in at.download_button if b.label == "⬇ Download report .pptx"]
    check("a download button appears after a successful generate",
         bool(download_buttons), [b.label for b in at.download_button])
    if not download_buttons:
        return

    import db as _db
    from pptx import Presentation
    out_path = _db.scratch_dir("attribution_reports") / "WAEPA Insurance.pptx"
    check(f"the generated deck exists on disk ({out_path.name})", out_path.exists(), out_path)
    if not out_path.exists():
        return
    prs = Presentation(str(out_path))
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text.append(cell.text_frame.text)
    alltext = "\n".join(text)
    check("the linked proposal's own client name wins over the export's own "
         "(CLIENT_NAME reads 'WAEPA Insurance', not bare 'WAEPA')",
         "WAEPA Insurance" in alltext, alltext[:300])
    check("FLIGHT_LABEL reads the linked proposal's own flight",
         "Oct 2026 - Dec 2026" in alltext, None)
    check("GEOGRAPHY_LABEL reads the target-market names ('Washington, DC' and 'Baltimore')",
         "Washington, DC" in alltext and "Baltimore" in alltext, None)
    check("AUDIENCE_BULLETS reads the linked proposal's own Campaign Specs audience",
         "Federal government employees researching benefits" in alltext, None)


def check_report_history_tab_and_followup(store):
    """ATTRIBUTION_REPORT_PLAN.md Phase 6, through the real page: the tab
    split leaves the wizard's own behavior unchanged (already proven by
    every other check_* function above still passing), the Report history
    tab shows a report just generated (grouped under its client, with the
    report_type/whats_next this run wrote), "Build follow-up proposal"
    lands the report's what's-next text in the notes box, and Delete
    refuses while a case study references the report.
    """
    print("\nPhase 6: Report history tab, follow-up proposal, delete")
    if not WAEPA_FIXTURE.exists():
        print("  SKIP  WAEPA fixture not present")
        return
    template = REPO / "REPORT_MASTER_v0_6.pptx"
    if not template.exists():
        print("  SKIP  REPORT_MASTER_v0_6.pptx not present -- can't test Generate")
        return

    fake_row = _fake_proposal_row(
        rid="prop-phase6-1", client_name="WAEPA Phase 6",
        flight_label="Oct 2026 - Dec 2026", budget_cost="$74,970")
    store.proposals = [fake_row]
    before_report_count = len(store.reports)

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_prelinked_proposal_id"] = "prop-phase6-1"
    at.session_state["attr_attribution_upload_path"] = str(WAEPA_FIXTURE)
    at.run()
    check("no exception with a linked proposal + upload", not at.exception, at.exception)

    whats_next_areas = [t for t in at.text_area if t.key == "attr_whats_next_input"]
    if whats_next_areas:
        whats_next_areas[0].set_value("Expand DC/Baltimore targeting next quarter").run()

    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    if not generate_buttons:
        check("a Generate report deck button is present", False, [b.label for b in at.button])
        return
    at.session_state["attr_draft"] = _ATTR_STUB_DRAFT
    generate_buttons[0].click().run()
    check("no exception after generating", not at.exception, at.exception)

    check("exactly one new report was logged", len(store.reports) == before_report_count + 1,
         len(store.reports))
    if len(store.reports) != before_report_count + 1:
        return
    report_row = store.reports[-1]
    rid = report_row["id"]
    # Part 2 rework (2026-09-21): report_type is now INFERRED, not a bare
    # selectbox default -- this WAEPA export spans 2 calendar months (Jul,
    # Aug) and ends 2026-08-31, after the linked proposal's own flight_end
    # (2026-07-17), so the correct inferred call is Wrap-up, not "monthly."
    check("report_type is inferred as 'wrap' (export ends after the "
         "proposal's own flight_end)",
         report_row.get("report_json", {}).get("report_type") == "wrap", report_row)
    check("whats_next was stored as the FINAL list Generate used",
         report_row.get("report_json", {}).get("whats_next") == ["Expand DC/Baltimore targeting next quarter"],
         report_row)
    check("upload_report_file was called with the id log_attribution_report actually returned",
         store.upload_report_file_calls and store.upload_report_file_calls[-1][0] == rid,
         store.upload_report_file_calls)

    print("  The Report history tab shows the report, grouped under its client")
    check("the client name appears somewhere on the page (the group header)",
         any("WAEPA Phase 6" in str(m) for m in _all_markdown_text(at)), None)
    check("the report's own what's-next text appears in its row",
         any("Expand DC/Baltimore targeting" in str(m) for m in _all_markdown_text(at)), None)

    print("  Delete refuses while a case study references the report")
    store.case_studies = [{"id": "cs-blocks-delete", "title": "A case study built from this report",
                           "source_report_id": rid}]
    typed_inputs = [t for t in at.text_input if t.key == f"rpt_del_confirm_{rid}"]
    check("the typed-DELETE confirm field is present", bool(typed_inputs),
         [t.key for t in at.text_input if t.key and t.key.startswith("rpt_del")])
    if typed_inputs:
        typed_inputs[0].set_value("DELETE").run()
    delete_buttons = [b for b in at.button if b.key == f"rpt_del_{rid}"]
    check("the delete-permanently button is present and enabled once DELETE is typed",
         bool(delete_buttons) and not delete_buttons[0].disabled,
         [(b.key, b.disabled) for b in at.button if b.key and "del" in b.key])
    if delete_buttons:
        delete_buttons[0].click().run()
    check("no exception after the blocked delete attempt", not at.exception, at.exception)
    check("the report row still exists -- refused, not silently dropped",
         any(r["id"] == rid for r in store.reports), store.reports)
    check("the refusal names the referencing case study",
         any("A case study built from this report" in str(e.value) for e in at.error), None)

    print("  Delete succeeds once nothing references the report")
    store.case_studies = []
    typed_inputs = [t for t in at.text_input if t.key == f"rpt_del_confirm_{rid}"]
    if typed_inputs:
        typed_inputs[0].set_value("DELETE").run()
    delete_buttons = [b for b in at.button if b.key == f"rpt_del_{rid}"]
    if delete_buttons:
        delete_buttons[0].click().run()
    check("no exception after the successful delete", not at.exception, at.exception)
    check("the report row is actually gone", all(r["id"] != rid for r in store.reports), store.reports)


def check_build_follow_up_proposal(store):
    """The "Build follow-up proposal" button (item 6): with a linked
    proposal, it restores the proposal into the Build form (rehydrate_
    proposal_into_form's own job, already tested elsewhere) AND appends the
    report's own what's-next text to the notes box -- the one piece
    rehydrate alone doesn't supply.
    """
    print("\nBuild follow-up proposal: linked-proposal path")
    fake_row = _fake_proposal_row(rid="prop-followup-1", client_name="Followup Co")
    store.proposals = [fake_row]
    report_row = {"id": "report-followup-1", "advertiser_id": None, "proposal_id": "prop-followup-1",
                 "report_json": {"headline_facts": {"period_start": "2026-06-01", "period_end": "2026-06-30"},
                                 "whats_next": ["Increase geofencing budget", "Add a retargeting line"]},
                 "created_by": "Matt", "created_at": "2026-07-01T00:00:00"}
    store.reports = [report_row]

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.run()
    check("no exception rendering the page", not at.exception, at.exception)

    followup_buttons = [b for b in at.button if b.key == "rpt_followup_report-followup-1"]
    check("the 'Build follow-up proposal' button is present", bool(followup_buttons),
         [b.key for b in at.button if b.key and "followup" in b.key])
    if not followup_buttons:
        return
    followup_buttons[0].click().run()
    check("no exception after clicking it", not at.exception, at.exception)
    check("it navigated to Build a proposal",
         _ss(at, "page_choice") == "Build a proposal", _ss(at, "page_choice"))
    check("the client name carried over from the linked proposal",
         _ss(at, "client_name") == "Followup Co", _ss(at, "client_name"))
    notes = _ss(at, "draft_notes_input") or ""
    check("the report's own what's-next items landed in the notes box",
         "Increase geofencing budget" in notes and "Add a retargeting line" in notes, notes)
    check("the notes name which report period they came from",
         "2026-06-01" in notes, notes)


def check_goals_prefill_via_search_and_pick(store):
    """Bug 1 (Matt's WAEPA walkthrough, 2026-09-10): the in-page search-
    and-pick door (this page's own section 3) never prefilled anything --
    Phase 5's field map was only ever wired off `prelinked_row` (Proposal
    History's "Build report" door), which is the ONLY door check_phase5_
    proposal_link above ever drove. This is the Phase 5 acceptance test
    that should have been run through the actual UI the first time: link
    a real-shaped WAEPA proposal via the SEARCH flow (never `attr_
    prelinked_proposal_id`) and assert the goals box holds Michele's real
    goals before the rep types anything.
    """
    print("\nGoals/audience prefill via the in-page search-and-pick door (not pre-linked)")
    if not WAEPA_FIXTURE.exists():
        print("  SKIP  WAEPA fixture not present")
        return

    michele_goals = ("Evaluate OTT performance beyond awareness — measure whether "
                     "concentrating spend into fewer core federal markets and increasing "
                     "share of voice drives more meaningful engagement signals\n"
                     "Track engaged visits to site and application entry point, cost per "
                     "engaged visit, application activity within test markets, and overall "
                     "lift versus 2025 historical performance")
    fake_row = _fake_proposal_row(
        rid="09e61e0b-a945-4022-851d-d3c31b2acbd0", client_name="WAEPA",
        goals=michele_goals,
        audience="Civilian federal government employees, excluding current WAEPA members, "
                 "plus federal job seekers researching benefits",
        flight_label="Oct 2026 - Dec 2026", budget_cost="$74,970")
    store.proposals = [fake_row]

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_attribution_upload_path"] = str(WAEPA_FIXTURE)
    at.run()
    check("no exception after upload", not at.exception, at.exception)

    # Advertiser confirm first (section 2) -- shape depends on whether an
    # earlier check in this same run already created "WAEPA" as an
    # advertiser (the shared FakeStore accumulates across checks), so this
    # goes through `_confirm_advertiser` rather than assuming the radio
    # shape specifically.
    check("an advertiser radio or the exact-match shortcut is shown",
         bool(at.radio) or bool(at.button), [w.key for w in list(at.radio) + list(at.button)])
    _confirm_advertiser(at)
    check("advertiser confirmed", _ss(at, "attr_advertiser_id") is not None,
         _ss(at, "attr_advertiser_id"))

    # Section 3's own search-and-pick radio -- NOT the pre-linked door.
    proposal_radios = [r for r in at.radio if r.key == "attr_proposal_choice"]
    check("the proposal picker radio is shown", bool(proposal_radios),
         [r.key for r in at.radio])
    if not proposal_radios:
        return
    real_option = next((o for o in proposal_radios[0].options
                        if o != "No proposal -- build this report standalone"), None)
    check("a real proposal candidate is offered", real_option is not None,
         proposal_radios[0].options)
    if real_option is None:
        return
    proposal_radios[0].set_value(real_option).run()
    proposal_confirm = [b for b in at.button if b.label == "Confirm"]
    check("a Confirm button is present for the proposal pick", bool(proposal_confirm),
         [b.label for b in at.button])
    if proposal_confirm:
        proposal_confirm[0].click().run()
    check("no exception after picking the proposal", not at.exception, at.exception)
    check("attr_proposal_id was set by the picker",
         _ss(at, "attr_proposal_id") == "09e61e0b-a945-4022-851d-d3c31b2acbd0",
         _ss(at, "attr_proposal_id"))
    check("this really is the search-and-pick door, not the pre-linked one "
         "(otherwise this test would prove nothing new)",
         _ss(at, "attr_prelinked_proposal_id") is None, _ss(at, "attr_prelinked_proposal_id"))

    goals_areas = [t for t in at.text_area if t.key == "attr_goals_input"]
    check("goals text area is present", bool(goals_areas), [t.key for t in at.text_area])
    if goals_areas:
        check("the goals box holds Michele's real goals BEFORE the rep types anything -- "
             "the actual bug: this stayed empty when linked via search-and-pick",
             "concentrating spend into fewer core federal markets" in goals_areas[0].value,
             goals_areas[0].value)
    audience_areas = [t for t in at.text_area if t.key == "attr_audience_input"]
    check("audience text area is present", bool(audience_areas), [t.key for t in at.text_area])
    if audience_areas:
        check("the audience box holds the proposal's own Campaign Specs audience",
             "federal government employees" in audience_areas[0].value.lower(),
             audience_areas[0].value)

    # The validation half of bug 1: goals are optional to TYPE when a
    # proposal supplies them -- Generate must not block on "enter a goal"
    # just because the rep never touched the (already-prefilled) box.
    # What's-next is still untyped and the pre-seeded (empty) draft has no
    # whats_next_bullets of its own to fall back to, so Generate DOES block
    # here -- on the what's-next message specifically, never on the goals
    # one. Pre-seeding also means this click never makes a live Claude
    # call -- see _ATTR_STUB_DRAFT's own comment.
    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    if generate_buttons:
        at.session_state["attr_draft"] = _ATTR_STUB_DRAFT
        generate_buttons[0].click().run()
    error_texts = [str(getattr(e, "value", "")) for e in at.error]
    check("Generate does NOT block on 'enter a campaign goal' -- the linked proposal's "
         "own goals already satisfy it",
         not any("campaign goal" in t for t in error_texts), error_texts)


def check_proposal_picker_shows_identifying_details(store):
    """Bug 2 (Matt's WAEPA walkthrough, 2026-09-10): the picker used to
    show only the client name and a bare match score ("WAEPA (score
    1.00)") -- a rep with several same-client proposals can't tell an
    April test from a July extension from that alone. Every candidate must
    show enough to confirm it's the right one: title, flight dates, date
    logged, who logged it.
    """
    print("\nProposal picker shows identifying details, not just a name and a score")
    if not MW_FIXTURE.exists():
        print("  SKIP  MW attribution excel.xlsx not present")
        return

    fake_row = _fake_proposal_row(
        rid="prop-picker-1", client_name="Mattress Warehouse",
        flight_label="Apr 2026 - Jun 2026", budget_cost="$30,000")
    fake_row["form_json"]["proposal_title"] = "Spring Sale CTV Test"
    # The picker reads the STRUCTURED setup dates (resolve_setup), never
    # the free-text flight.label -- match them so this fixture is internally
    # consistent, the same "structured data over narrative text" rule
    # geo_column_default's own 2026-09-08 fix already established elsewhere.
    fake_row["form_json"]["setup"]["flight_start"] = "2026-04-01"
    fake_row["form_json"]["setup"]["flight_end"] = "2026-06-30"
    fake_row["generated_at"] = "2026-04-02T09:00:00"
    fake_row["created_by"] = "Pat"
    store.proposals = [fake_row]

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_attribution_upload_path"] = str(MW_FIXTURE)
    at.run()
    # See _confirm_advertiser's own comment -- the shared FakeStore may
    # already hold "Mattress Warehouse" from an earlier check this run.
    _confirm_advertiser(at)

    proposal_radios = [r for r in at.radio if r.key == "attr_proposal_choice"]
    check("the proposal picker radio is shown", bool(proposal_radios),
         [r.key for r in at.radio])
    if not proposal_radios:
        return
    options = proposal_radios[0].options
    real_option = next((o for o in options
                        if o != "No proposal -- build this report standalone"), None)
    check("a real proposal candidate is offered", real_option is not None, options)
    if real_option is None:
        return
    check("the candidate names the proposal's own title",
         "Spring Sale CTV Test" in real_option, real_option)
    check("the candidate names the flight dates",
         "2026-04-01" in real_option and "2026-06-30" in real_option, real_option)
    check("the candidate names when it was logged (2026-04-02)",
         "2026-04-02" in real_option, real_option)
    check("the candidate names who logged it (Pat)",
         "Pat" in real_option, real_option)


def _all_markdown_text(at):
    texts = []
    for kind in ("markdown", "caption", "info", "success", "warning", "error", "title", "header", "subheader"):
        try:
            texts.extend(str(getattr(el, "value", "")) for el in getattr(at, kind))
        except Exception:                                                 # noqa: BLE001
            pass
    return texts


def check_pixel_issue_window(store):
    """The known tracking-pixel under-count window (June 9 - July 21, 2026,
    Pansophic memo, 2026-09-08 correction) -- a non-blocking, rep-facing
    warning when an uploaded export's flight overlaps it. Real fixtures on
    both sides: MW's real flight (Jun 1 - Jul 13) genuinely overlaps (this
    is the SAME defect MW's own 2.15x projection factor turned out to be,
    per DECISIONS.md); WAEPA's real flight (Jul 27 - Aug 31) genuinely
    starts after the fix. Pure edge cases (boundary days, missing dates)
    are covered directly against `app._overlaps_pixel_issue_window` --
    this is the "does it actually reach the page" half.
    """
    print("\nPixel-issue-window warning (2026-09-08 correction)")
    cases = [
        (date(2026, 6, 1), date(2026, 6, 15), True, "starts before, ends inside"),
        (date(2026, 7, 15), date(2026, 8, 1), True, "starts inside, ends after"),
        (date(2026, 5, 1), date(2026, 8, 1), True, "fully contains the window"),
        (date(2026, 1, 1), date(2026, 3, 1), False, "well before"),
        (date(2026, 8, 1), date(2026, 9, 1), False, "well after"),
        (date(2026, 6, 9), date(2026, 6, 9), True, "exact start boundary"),
        (date(2026, 7, 21), date(2026, 7, 21), True, "exact end boundary"),
        (date(2026, 7, 22), date(2026, 8, 1), False, "starts the day after the fix"),
        (None, date(2026, 6, 15), False, "missing start"),
    ]
    for start, end, expected, label in cases:
        got = app._overlaps_pixel_issue_window(start, end)
        check(f"overlap logic: {label}", got == expected, (got, expected))

    if not MW_FIXTURE.exists():
        print("  SKIP  MW attribution excel.xlsx not present")
    else:
        at = new_app()
        at.session_state["page_choice"] = "Attribution reports"
        at.session_state["attr_attribution_upload_path"] = str(MW_FIXTURE)
        at.run()
        check("no exception after upload", not at.exception, at.exception)
        texts = _all_markdown_text(at)
        check("MW's real flight overlaps the window -- the warning shows",
             any("known tracking issue" in t for t in texts), texts)

    if not WAEPA_FIXTURE.exists():
        print("  SKIP  WAEPA fixture not present")
    else:
        at = new_app()
        at.session_state["page_choice"] = "Attribution reports"
        at.session_state["attr_attribution_upload_path"] = str(WAEPA_FIXTURE)
        at.run()
        check("no exception after upload", not at.exception, at.exception)
        texts = _all_markdown_text(at)
        check("WAEPA's real flight starts after the fix -- no warning",
             not any("known tracking issue" in t for t in texts), texts)


def check_vertical_conversion_definition_and_plan_toggle(store):
    """Highlights/Takeaways rework's own page additions: the vertical
    selectbox prefills from a linked proposal's own row-level `vertical`
    (never guessed); the conversion-definition text input is present; and
    the plan-vs-actual toggle stays ABSENT until the template's own
    DeliveryByGeoTable is actually widened to 5 columns -- today's
    REPORT_MASTER_v0_6.pptx has 3, confirmed directly in
    test_report_assembly.py's own `named_table_column_count` check, so a
    rep must never see a toggle that would change nothing on the slide.
    """
    print("\nVertical/conversion-definition fields, and the plan-vs-actual toggle's own gate")
    if not MW_FIXTURE.exists():
        print("  SKIP  MW attribution excel.xlsx not present")
        return
    fake_row = _fake_proposal_row(rid="prop-vertical-1", client_name="Mattress Warehouse")
    store.proposals = [fake_row]

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_prelinked_proposal_id"] = "prop-vertical-1"
    at.session_state["attr_attribution_upload_path"] = str(MW_FIXTURE)
    at.session_state["attr_delivery_upload_path"] = str(DELIVERY_MW) if DELIVERY_MW.exists() else None
    at.run()
    check("no exception with a linked proposal + attribution/delivery upload",
         not at.exception, at.exception)

    vertical_boxes = [s for s in at.selectbox if s.key == "attr_vertical_input"]
    check("the vertical selectbox is present", bool(vertical_boxes),
         [s.key for s in at.selectbox])
    if vertical_boxes:
        check("prefilled from the linked proposal's own row-level vertical "
             "('retail' -> 'Retail'), never guessed",
             vertical_boxes[0].value == "Retail", vertical_boxes[0].value)

    conversion_inputs = [t for t in at.text_input if t.key == "attr_conversion_definition_input"]
    check("the conversion-definition text input is present", bool(conversion_inputs),
         [t.key for t in at.text_input])

    toggle_keys = [t.key for t in at.toggle]
    check("the plan-vs-actual toggle is ABSENT -- today's template has only 3 columns "
         "in DeliveryByGeoTable, not the 5 the toggle needs to change anything",
         "attr_show_plan_vs_actual" not in toggle_keys, toggle_keys)


def check_optimization_controls(store):
    """The optimization engine's own controls (ATTRIBUTION_REPORT_PLAN.md
    Phase 6, built after the advertiser spine): a Level selectbox and a
    dimension multiselect, publisher defaulting OFF (Matt's ruling,
    2026-09-14) while every other dimension defaults on, and Generate
    still completes cleanly with them present -- the real threshold/cap
    logic and the model's own behavior are covered elsewhere (tests/
    test_optimization_engine.py offline, tests/test_attribution_draft_
    live.py's "mw_optimization" scenario live); this is just the widgets
    existing and not crashing the page.
    """
    print("\nOptimization recommendation controls: defaults and no-crash")
    if not MW_FIXTURE.exists():
        print("  SKIP  MW attribution excel.xlsx not present")
        return
    template = REPO / "REPORT_MASTER_v0_6.pptx"
    if not template.exists():
        print("  SKIP  REPORT_MASTER_v0_6.pptx not present -- can't test Generate")
        return

    # Self-contained, not inherited from whatever an earlier check_* left in
    # `store` -- an earlier test's own "Mattress Warehouse" proposal row
    # left `store.proposals` non-empty, which turned this run's "no matching
    # proposal" step into a candidate-picker (a radio + Confirm) instead of
    # the bare standalone button this check expected, and the real failure
    # then got masked by an unrelated console-encoding crash trying to print
    # the mismatch. Every other check_* in this file resets what it needs;
    # this one hadn't.
    store.proposals = []

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_attribution_upload_path"] = str(MW_FIXTURE)
    at.run()
    check("no exception after upload", not at.exception, at.exception)

    # The optimization controls render in step 4 (goals/notes/etc.), past
    # the advertiser-confirm and proposal-link steps -- have to get through
    # both before the widgets exist to check.
    _confirm_advertiser(at)
    no_proposal_buttons = [b for b in at.button
                           if b.label == "No proposal -- build this report standalone"]
    if no_proposal_buttons:
        no_proposal_buttons[0].click().run()

    level_boxes = [s for s in at.selectbox if s.key == "attr_optimization_level"]
    check("the Level selectbox is present", bool(level_boxes), [s.key for s in at.selectbox])
    if level_boxes:
        check("it defaults to None (a client with no stored level, per the optimization-"
             "sequence work) -- reporting only unless a rep or the Clients page says otherwise",
             level_boxes[0].value == "None", level_boxes[0].value)

    dim_multiselects = [m for m in at.multiselect if m.key == "attr_optimization_dimensions"]
    check("the dimensions multiselect is present", bool(dim_multiselects),
         [m.key for m in at.multiselect])
    if dim_multiselects:
        defaults = set(dim_multiselects[0].value)
        check("publisher defaults OFF", "publisher" not in defaults, defaults)
        check("zip/market/creative/day_of_week all default ON",
             {"zip", "market", "creative", "day_of_week"} <= defaults, defaults)

    goals_areas = [t for t in at.text_area if t.key == "attr_goals_input"]
    if goals_areas:
        goals_areas[0].set_value("Drive in-store visits").run()

    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    check("a Generate button is present", bool(generate_buttons), [b.label for b in at.button])
    if not generate_buttons:
        return
    at.session_state["attr_draft"] = _ATTR_STUB_DRAFT
    generate_buttons[0].click().run()
    check("no exception generating with optimization dimensions enabled",
         not at.exception, at.exception)


def check_optimization_checklist_accept_edit_decline(store):
    """VERIFY (ATTRIBUTION_REPORT_PLAN.md Phase 6, "the memory between
    monthly reports"; ZIP handling rewritten for round 3, 2026-09-22
    review, item 3): MW at Moderate renders real ZIP GROUPS as a
    checklist (round 3 moved ZIP-level optimization off the flat
    candidate list entirely -- `optimization_candidates` no longer
    produces "zip" dimension candidates at all, see its own docstring --
    so this scenario, which real MW data only ever populated via ZIP
    candidates at Moderate, now exercises `zip_optimization_groups`'s own
    checklist instead of the flat one). Decline one group (with a
    reason), edit another (dropping one ZIP out of it via the "editable
    lists" multiselect), leave any remaining group at the default Accept;
    Generate completes; `report_json["optimizations"]` holds one entry
    PER ZIP across every group, each with its own decision -- "the
    declines are the interesting half for learning later," same
    contract the flat engine's own log already had. The declined group's
    zips carry `final_text=None` (never reach the model); the edited
    group's KEPT zips carry the rep's own recomputed wording, while the
    ONE DROPPED zip is logged "declined" individually (reviewed and
    excluded, not acted on) even though the rest of its group was
    edited/accepted.
    """
    print("\nOptimization checklist: accept/edit/decline, and what gets logged")
    if not MW_FIXTURE.exists():
        print("  SKIP  MW attribution excel.xlsx not present")
        return
    template = REPO / "REPORT_MASTER_v0_6.pptx"
    if not template.exists():
        print("  SKIP  REPORT_MASTER_v0_6.pptx not present -- can't test Generate")
        return

    # Deterministic -- seed a prior MW report explicitly (opens the timing
    # gate) rather than relying on suite ordering to have logged one.
    mw_advertiser = next((a for a in store.advertisers
                         if a["canonical_name"] == "Mattress Warehouse"), None)
    if mw_advertiser is None:
        mw_advertiser, _ = store.create_advertiser("Mattress Warehouse")
    # This prior report's own "accepted" entry for ZIP 27525 (a real,
    # confirmed MW outlier at Moderate) also lets this test prove the
    # in-effect EXCLUSION end to end through the real app wiring --
    # test_optimization_engine.py already proves the underlying function
    # excludes it given the right inputs; this proves app.py actually
    # threads `_prior_reports`/`attribution_obj` into it correctly.
    store.reports.append({
        "id": store._id(), "advertiser_id": mw_advertiser["id"], "proposal_id": None,
        "report_json": {"headline_facts": {"period_start": "2026-01-01", "period_end": "2026-01-31"},
                        "attribution": {"attributed_rate": 0.01, "by_zip": [
                            {"label": "27525", "delivered_impressions": 5000,
                             "attributed_impressions": 5, "attributed_rate": 0.001}]},
                        "optimizations": [{"dimension": "zip", "value": "27525",
                                          "metric": "attributed_rate", "decision": "accepted",
                                          "final_text": "Cut 27525 (from a seeded prior report)"}]}})
    store.proposals = []

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_attribution_upload_path"] = str(MW_FIXTURE)
    at.run()
    check("no exception after upload", not at.exception, at.exception)

    _confirm_advertiser(at)
    no_proposal_buttons = [b for b in at.button
                           if b.label == "No proposal -- build this report standalone"]
    if no_proposal_buttons:
        no_proposal_buttons[0].click().run()

    level_boxes = [s for s in at.selectbox if s.key == "attr_optimization_level"]
    check("level defaults to None even though evidence_periods >= 2 (a seeded prior report "
         "exists) -- level and timing are independent gates",
         bool(level_boxes) and level_boxes[0].value == "None", level_boxes[0].value if level_boxes else None)
    check("at the default None level, nothing renders on the checklist",
         not [r for r in at.radio if r.key and (r.key.startswith("attr_opt_decision_")
                                                or r.key.startswith("attr_zipgrp_decision_"))], None)

    if level_boxes:
        level_boxes[0].set_value("Moderate").run()

    decision_radios = [r for r in at.radio if r.key and r.key.startswith("attr_zipgrp_decision_")]
    check("real zip groups rendered as a checklist", bool(decision_radios),
         [r.key for r in at.radio])
    if len(decision_radios) < 2:
        check("more than one zip group rendered (otherwise decline+edit can't both be exercised)",
             False, len(decision_radios))
        return

    # Expose every group's own editable ZIP list (via Edit) before committing
    # to a scenario, both to size each group for the log-count assertions
    # below and to prove ZIP 27525 -- accepted in the seeded PRIOR report --
    # never appears in any group this month (in-effect exclusion, wired end
    # to end through the real app, same proof the flat-candidate version of
    # this test made before round 3 moved zips off that engine).
    for r in decision_radios:
        r.set_value("Edit").run()
    edit_multiselects = [m for m in at.multiselect if m.key and m.key.startswith("attr_zipgrp_zips_")]
    check("every rendered group exposes its own editable ZIP list",
         len(edit_multiselects) == len(decision_radios),
         (len(edit_multiselects), len(decision_radios)))
    group_zip_lists = {m.key[len("attr_zipgrp_zips_"):]: list(m.options) for m in edit_multiselects}
    all_group_zips = {z for zips in group_zip_lists.values() for z in zips}
    check("ZIP 27525 does NOT reappear in any group this month (in-effect exclusion)",
         "27525" not in all_group_zips, sorted(all_group_zips))

    decline_suffix = decision_radios[0].key[len("attr_zipgrp_decision_"):]
    edit_suffix = decision_radios[1].key[len("attr_zipgrp_decision_"):]

    [r for r in at.radio if r.key == f"attr_zipgrp_decision_{decline_suffix}"][0].set_value("Decline").run()
    [r for r in at.radio if r.key == f"attr_zipgrp_decision_{edit_suffix}"][0].set_value("Edit").run()
    # Every OTHER group goes back to the default Accept -- the probing loop
    # above left them all sitting on Edit.
    for r in [r for r in at.radio if r.key and r.key.startswith("attr_zipgrp_decision_")
             and r.key not in (f"attr_zipgrp_decision_{decline_suffix}",
                              f"attr_zipgrp_decision_{edit_suffix}")]:
        r.set_value("Accept").run()

    reason_inputs = [t for t in at.text_input if t.key == f"attr_zipgrp_reason_{decline_suffix}"]
    if reason_inputs:
        reason_inputs[0].set_value("Client asked to keep these ZIPs live").run()

    edit_group_multiselects = [m for m in at.multiselect if m.key == f"attr_zipgrp_zips_{edit_suffix}"]
    check("the edit group's ZIP multiselect is present once Edit is picked", bool(edit_group_multiselects),
         [m.key for m in at.multiselect])
    dropped_zip, kept_zips = None, []
    if edit_group_multiselects:
        edit_group_zips = group_zip_lists.get(edit_suffix, list(edit_group_multiselects[0].options))
        check("the edit group has at least 2 ZIPs (so one can be dropped and others kept)",
             len(edit_group_zips) >= 2, edit_group_zips)
        dropped_zip, kept_zips = edit_group_zips[0], edit_group_zips[1:]
        edit_group_multiselects[0].set_value(kept_zips).run()

    goals_areas = [t for t in at.text_area if t.key == "attr_goals_input"]
    if goals_areas:
        goals_areas[0].set_value("Improve efficiency of underperforming zips").run()
    # The stub draft (see its own module-level comment) carries no threads,
    # so drafted_whats_next is empty -- Generate blocks on "enter at least
    # one what's-next item" unless one is typed, same as every other
    # Generate-clicking check in this file that uses the stub.
    whats_next_areas = [t for t in at.text_area if t.key == "attr_whats_next_input"]
    if whats_next_areas:
        whats_next_areas[0].set_value("Keep monitoring performance").run()

    reports_before = len(store.reports)
    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    check("a Generate button is present", bool(generate_buttons), [b.label for b in at.button])
    if not generate_buttons:
        return
    at.session_state["attr_draft"] = _ATTR_STUB_DRAFT
    generate_buttons[0].click().run()
    check("no exception generating with a mixed accept/edit/decline zip-group checklist",
         not at.exception, at.exception)
    check("exactly one new report was logged", len(store.reports) == reports_before + 1,
         len(store.reports))
    if len(store.reports) != reports_before + 1:
        return

    logged = store.reports[-1]["report_json"].get("optimizations") or []
    expected_total_zips = sum(len(zips) for zips in group_zip_lists.values())
    check("every zip across every group was logged, not just the touched groups' -- one entry "
         "per zip, matching the flat engine's own log shape",
         len(logged) == expected_total_zips, (len(logged), expected_total_zips))
    by_zip = {e["value"]: e for e in logged}

    declined_zips = group_zip_lists.get(decline_suffix, [])
    declined_entries = [by_zip.get(z) for z in declined_zips]
    check("every zip in the declined group is logged with decision='declined'",
         all(e and e["decision"] == "declined" for e in declined_entries), declined_entries)
    check("none of the declined group's final_text reaches the model",
         all(e and e.get("final_text") is None for e in declined_entries), declined_entries)
    check("the decline reason was captured on every zip in that group",
         all(e and e.get("reason") == "Client asked to keep these ZIPs live" for e in declined_entries),
         declined_entries)

    dropped_entry = by_zip.get(dropped_zip) if dropped_zip else None
    check("the one ZIP dropped out of the edited group via the multiselect is logged "
         "'declined' individually -- reviewed and excluded, not acted on -- even though the "
         "rest of its group was edited",
         bool(dropped_entry) and dropped_entry["decision"] == "declined", dropped_entry)
    check("the dropped zip's final_text is None too",
         bool(dropped_entry) and dropped_entry.get("final_text") is None, dropped_entry)

    kept_entries = [by_zip.get(z) for z in kept_zips]
    check("every ZIP kept in the edited group is logged with decision='edited'",
         all(e and e["decision"] == "edited" for e in kept_entries), kept_entries)
    kept_final_texts = {e.get("final_text") for e in kept_entries if e}
    check("the kept zips all carry the SAME recomputed group wording, and it's not None",
         len(kept_final_texts) == 1 and None not in kept_final_texts, kept_final_texts)

    touched_zips = set(declined_zips) | set(kept_zips) | {dropped_zip}
    accepted_count = sum(1 for e in logged if e["decision"] == "accepted")
    untouched_zip_count = len(all_group_zips - touched_zips)
    check("every zip in an untouched group defaulted to accepted",
         accepted_count == untouched_zip_count, (accepted_count, untouched_zip_count))


def check_dev_warnings_reach_feedback_export():
    """2026-09-12 follow-up to the warnings-triage rework: confirm
    `capture_feedback_state` actually carries BOTH dev-warning keys
    through to a filed bug report -- that's the whole reason they're
    routed to session_state instead of just printed and dropped. Two
    separate keys (not one merged list) because they have different
    lifetimes: `attr_dev_warnings` describes the Generate click that just
    ran; `attr_upload_dev_warnings` describes the currently uploaded
    export and must survive a later, clean Generate."""
    print("\nAttribution Report Builder's dev warnings reach the feedback export")

    class _StStub:
        def __init__(self, session_state):
            self.session_state = session_state

    real_st, real_active_deck_version = app.st, db.active_deck_version
    db.active_deck_version = lambda: (None, None)
    app.st = _StStub({
        "attr_dev_warnings": ["BreakdownTable has 4 columns but 5 were supplied -- "
                             "not showing: conv_rate."],
        "attr_upload_dev_warnings": ["Atlanta reaches AL, GA, which have no ZCTA "
                                    "boundary file built."],
    })
    try:
        state = app.capture_feedback_state("Attribution reports")
    finally:
        app.st = real_st
        db.active_deck_version = real_active_deck_version

    check("attr_dev_warnings (the per-Generate-click bucket) rides along in the captured state",
         state.get("attr_dev_warnings") == [
             "BreakdownTable has 4 columns but 5 were supplied -- not showing: conv_rate."],
         state)
    check("attr_upload_dev_warnings (the per-export bucket) rides along too, as its own key",
         state.get("attr_upload_dev_warnings") == [
             "Atlanta reaches AL, GA, which have no ZCTA boundary file built."],
         state)

    exported = app._feedback_export_markdown([
        {"category": "Bug", "created_at": "2026-09-12T10:00:00", "created_by": "Andrea",
         "page": "Attribution reports", "notes": "Zip map looked wrong.", "state_json": state},
    ])
    check("both dev warnings actually appear in the exported markdown report",
         "conv_rate" in exported and "no ZCTA boundary file built" in exported,
         exported[:2000])


def check_summary_and_case_study_buttons(store):
    """ATTRIBUTION_REPORT_PLAN.md Phase 6's "one-slide summary and case
    study" -- through the real page against the real, live-Supabase active
    report template (REPORT_MASTER_v0_9, same as every other check_* in
    this file that clicks Generate). The Generate-time toggle produces a
    second download in the SAME run; the Report history row's own "One-
    slide summary" button regenerates one LATER from the logged report
    with no second draft call (report_json["draft"] is reused verbatim --
    `_ATTR_STUB_DRAFT` here is `{}`, so `threads` is empty either way, but
    the point of this check is that no exception/API call happens, not
    what the empty draft renders); "Create case study" opens a review that
    defaults white-labeled ("Name the client" OFF) and writes advertiser_
    id/source_report_id onto the saved vault row.
    """
    print("\nOne-slide summary + case study, through the real page")
    if not WAEPA_FIXTURE.exists():
        print("  SKIP  WAEPA fixture not present")
        return

    fake_row = _fake_proposal_row(
        rid="prop-summary-cs-1", client_name="WAEPA Summary Test",
        flight_label="Oct 2026 - Dec 2026", budget_cost="$74,970")
    store.proposals = [fake_row]
    before_report_count = len(store.reports)

    at = new_app()
    at.session_state["page_choice"] = "Attribution reports"
    at.session_state["attr_prelinked_proposal_id"] = "prop-summary-cs-1"
    at.session_state["attr_attribution_upload_path"] = str(WAEPA_FIXTURE)
    at.run()
    check("no exception with a linked proposal + upload", not at.exception, at.exception)

    summary_toggles = [c for c in at.checkbox if c.key == "attr_build_summary_toggle"]
    check("the 'Also build a one-slide summary' toggle is present", bool(summary_toggles),
         [c.key for c in at.checkbox])
    if summary_toggles:
        summary_toggles[0].set_value(True).run()

    whats_next_areas = [t for t in at.text_area if t.key == "attr_whats_next_input"]
    if whats_next_areas:
        whats_next_areas[0].set_value("Keep monitoring DC/Baltimore performance").run()

    generate_buttons = [b for b in at.button if b.label == "✨ Generate report deck"]
    check("a Generate report deck button is present", bool(generate_buttons),
         [b.label for b in at.button])
    if not generate_buttons:
        return
    at.session_state["attr_draft"] = _ATTR_STUB_DRAFT
    generate_buttons[0].click().run()
    check("no exception after generating (with the summary toggle on)", not at.exception, at.exception)
    check("exactly one new report was logged", len(store.reports) == before_report_count + 1,
         len(store.reports))
    if len(store.reports) != before_report_count + 1:
        return
    report_row = store.reports[-1]
    rid = report_row["id"]

    # st.download_button is its own AppTest element type, never at.button
    # (found live writing this check -- download_button elements simply
    # don't appear in at.button at all).
    generate_summary_dl = [b for b in at.download_button if b.key == "attr_summary_download"]
    check("Generate-time toggle produced a one-slide-summary download button in the same run",
         bool(generate_summary_dl),
         [b.key for b in at.download_button if b.key and "summary" in b.key])

    # Real bug, found live 2026-09-22: both download buttons used to be
    # rendered INLINE inside `if st.button("Generate report deck")`, which
    # is only True on the exact run it was clicked. `st.download_button`
    # always forces a rerun on click -- so clicking "Download report
    # .pptx" reran the script, the Generate button read False again, and
    # the one-slide summary's OWN download button (and the review/notes
    # below it) vanished from under the rep mid-download. Reproduced here
    # by clicking the report download button and re-running, exactly what
    # a rep's own click does -- both buttons, and the report count, must
    # survive that rerun untouched (no duplicate log/upload either).
    report_dl = [b for b in at.download_button if b.key == "attr_download"]
    check("the report .pptx download button is present in the same run", bool(report_dl),
         [b.key for b in at.download_button])
    if report_dl:
        report_dl[0].click().run()
        check("no exception after clicking Download report .pptx", not at.exception, at.exception)
        check("the report download button survives that rerun",
             any(b.key == "attr_download" for b in at.download_button),
             [b.key for b in at.download_button])
        check("the one-slide summary download button survives that rerun too",
             any(b.key == "attr_summary_download" for b in at.download_button),
             [b.key for b in at.download_button])
        check("clicking Download did not re-log a duplicate report",
             len(store.reports) == before_report_count + 1, len(store.reports))
        # Real bug, found live 2026-09-22 (Matt): once the vanishing-
        # download-button bug above was fixed, its OWN fix exposed a
        # second one -- the Preview section (above "7. Generate") and the
        # persistent post-Generate block (below the downloads) BOTH render
        # "Review before sending"/"Draft notes" for the same draft, so
        # after a real Generate + download click a rep saw the whole panel
        # twice, once above the Generate button and once below it.
        review_markdowns = [m for m in at.markdown if "Review before sending" in (m.value or "")]
        check("'Review before sending' renders at most once after Generate + a download click",
             len(review_markdowns) <= 1, len(review_markdowns))
        draft_notes_expanders = [e for e in at.expander if e.label == "Draft notes"]
        check("'Draft notes' expander renders at most once too",
             len(draft_notes_expanders) <= 1, len(draft_notes_expanders))

    print("  Report history row: regenerate the summary later (no second draft call)")
    summary_buttons = [b for b in at.button if b.key == f"rpt_summary_{rid}"]
    check("the row's own 'One-slide summary' button is present", bool(summary_buttons),
         [b.key for b in at.button if b.key and b.key.startswith("rpt_summary")])
    if summary_buttons:
        summary_buttons[0].click().run()
        check("no exception building the summary from the logged report (no draft call)",
             not at.exception, at.exception)
        row_summary_dl = [b for b in at.download_button if b.key == f"rpt_summary_dl_{rid}"]
        check("the regenerated summary offers its own download button", bool(row_summary_dl),
             [b.key for b in at.download_button if b.key and "summary_dl" in b.key])

    print("  Report history row: create a case study, white-labeled by default, then save")
    cs_buttons = [b for b in at.button if b.key == f"rpt_cs_button_{rid}"]
    check("the row's own 'Create case study' button is present", bool(cs_buttons),
         [b.key for b in at.button if b.key and b.key.startswith("rpt_cs_")])
    if not cs_buttons:
        return
    # Pre-seed the tag-suggestion cache (2026-09-15 rework: tags are now
    # auto-suggested via a live Claude call the moment this review opens)
    # so opening it here never fires a real, billed API call -- the same
    # "avoid the branch that calls out" shape `attr_draft` pre-seeding
    # already uses elsewhere in this file, and the ONLY shape that works:
    # this file's own top comment documents that stubbing a bare-name
    # app.py function doesn't intercept AppTest's re-executed script.
    #
    # Deliberately a DISAGREEING suggestion (this proposal's own vertical
    # is "retail," per _fake_proposal_row) -- a real live walkthrough found
    # that a stale/wrong Claude suggestion could otherwise silently push
    # the rep's own eyebrow-vertical choice out of the vault tags entirely.
    at.session_state[f"rpt_cs_tags_{rid}"] = {"verticals": ["education"]}
    cs_buttons[0].click().run()
    check("no exception opening the case study review", not at.exception, at.exception)

    # Keyed as f"rpt_cs_verticals_{rid}_{eyebrow_vertical}" (a real 2026-09-
    # 15 fix -- the key must change when the vertical does, or `default=`
    # is ignored on the rerun that follows changing it), so a prefix match
    # rather than an exact one.
    verticals_multiselects_pre = [m for m in at.multiselect
                                  if m.key and m.key.startswith(f"rpt_cs_verticals_{rid}_")]
    check("the eyebrow's own vertical (retail) is in the vault tags even though Claude's "
         "own suggestion disagreed (education) -- the rep's explicit choice is never silently "
         "dropped, only added to",
         bool(verticals_multiselects_pre) and "retail" in verticals_multiselects_pre[0].value
         and "education" in verticals_multiselects_pre[0].value,
         [(m.key, m.value) for m in at.multiselect if m.key and "cs_verticals" in m.key])
    check("the eyebrow's own vertical is listed FIRST, ahead of Claude's own suggestion",
         bool(verticals_multiselects_pre) and verticals_multiselects_pre[0].value[0] == "retail",
         verticals_multiselects_pre[0].value if verticals_multiselects_pre else None)

    named_checkboxes = [c for c in at.checkbox if c.key == f"rpt_cs_named_{rid}"]
    check("the 'Name the client' checkbox is present and OFF by default (white-labeled)",
         bool(named_checkboxes) and named_checkboxes[0].value is False,
         [(c.key, c.value) for c in at.checkbox if c.key and "cs_named" in c.key])

    added_by_inputs = [t for t in at.text_input if t.key == f"rpt_cs_addedby_{rid}"]
    if added_by_inputs:
        added_by_inputs[0].set_value("Test Rep").run()
    save_buttons = [b for b in at.button if b.key == f"rpt_cs_save_{rid}"]
    check("the primary 'Save to case study vault' button is present", bool(save_buttons),
         [b.key for b in at.button if b.key and "cs_save" in b.key])
    download_only_buttons = [b for b in at.download_button if b.key == f"rpt_cs_dl_{rid}"]
    check("a secondary 'Download only' action is also present", bool(download_only_buttons),
         [b.key for b in at.download_button if b.key and "cs_dl" in b.key])
    if save_buttons:
        save_buttons[0].click().run()
    check("no exception saving the case study", not at.exception, at.exception)
    check("upload_case_study was called exactly once", len(store.upload_case_study_calls) == 1,
         store.upload_case_study_calls)
    if store.upload_case_study_calls:
        call = store.upload_case_study_calls[-1]
        check("source_report_id was set to this report's own id",
             call["source_report_id"] == rid, call)
        check("advertiser_id was carried through from the report row",
             call["advertiser_id"] == report_row.get("advertiser_id"), call)


def check_vault_browser_report_source_label(store):
    """The vault browser's own "generated from a report" label (2026-09-15:
    "a rep picking case studies for a pitch should be able to see which
    ones come from real attribution data") -- `app._report_source_labels`
    directly, against the same FakeStore already patched onto `db.
    fetch_attribution_reports`/`db.fetch_advertisers`. A case study whose
    `source_report_id` doesn't resolve to any report (orphaned) degrades to
    a bare label rather than a crash or a blank line.
    """
    print("\nVault browser: 'generated from a report' label")
    store.advertisers = [{"id": "adv-label-1", "canonical_name": "Label Co", "active": True}]
    store.reports = [
        {"id": "report-label-1", "advertiser_id": "adv-label-1",
         "report_json": {"headline_facts": {"period_start": "2026-06-01",
                                            "period_end": "2026-06-30"}}},
    ]
    rows = [
        {"id": "cs-1", "title": "From a real report", "source_report_id": "report-label-1"},
        {"id": "cs-2", "title": "Hand uploaded", "source_report_id": None},
        {"id": "cs-3", "title": "Orphaned reference", "source_report_id": "report-does-not-exist"},
    ]
    labels = app._report_source_labels(rows)
    check("the linked report gets a label naming its period and client",
         labels.get("report-label-1") == "📊 generated from a report — 2026-06-01 to 2026-06-30, "
                                        "Label Co",
         labels)
    check("a case study with no source_report_id gets no entry at all",
         "None" not in labels and None not in labels, labels)
    check("an orphaned/unresolvable report id is simply absent (caller degrades gracefully)",
         "report-does-not-exist" not in labels, labels)


def main():
    store = FakeStore()
    app.db.fetch_advertisers = store.fetch_advertisers
    app.db.create_advertiser = store.create_advertiser
    app.db.fetch_proposals = store.fetch_proposals
    app.db.fetch_proposal = store.fetch_proposal
    app.db.link_proposal_advertiser = store.link_proposal_advertiser
    app.db.log_attribution_report = store.log_attribution_report
    app.db.fetch_attribution_reports = store.fetch_attribution_reports
    app.db.set_advertiser_vertical = store.set_advertiser_vertical
    # Phase 6 additions -- the Report history tab now renders on every run
    # of this page (Streamlit tabs render all content regardless of which
    # is visually active), and Generate now stores the built deck, so all
    # three have to be stubbed here too or a real, unmocked db.get_client()
    # would fire against production the moment any test clicks Generate.
    app.db.upload_report_file = store.upload_report_file
    app.db.report_file = store.report_file
    app.db.delete_attribution_report = store.delete_attribution_report
    # One-slide-summary/case-study work: the case-study save path is the
    # only NEW write call this feature adds (the summary build only reads/
    # renders locally) -- stubbed for the same reason as the three above.
    app.db.upload_case_study = store.upload_case_study

    check_upload_first_new_advertiser_no_proposal(store)
    check_rfpid_confirm_gate(store)
    check_conversions_toggle(store)
    check_mw_has_no_conversions_toggle(store)
    check_polk_upload_and_projection_toggle(store)
    check_polk_phase8_wiring(store)
    check_polk_phase8_persistence(store)
    check_clear_button(store)
    check_prelinked_door(store)
    check_phase5_proposal_link(store)
    check_report_history_tab_and_followup(store)
    check_build_follow_up_proposal(store)
    check_goals_prefill_via_search_and_pick(store)
    check_proposal_picker_shows_identifying_details(store)
    check_pixel_issue_window(store)
    check_vertical_conversion_definition_and_plan_toggle(store)
    check_optimization_controls(store)
    check_optimization_checklist_accept_edit_decline(store)
    check_summary_and_case_study_buttons(store)
    check_vault_browser_report_source_label(store)
    check_dev_warnings_reach_feedback_export()

    print()
    print(f"{len(failures)} failure(s)" if failures else "All checks passed")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
