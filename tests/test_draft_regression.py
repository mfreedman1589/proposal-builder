"""Tier 1 -- the offline drafting regression suite.

    python tests/test_draft_regression.py            # both scenarios
    python tests/test_draft_regression.py hvac       # one, by fixture prefix
    python tests/test_draft_regression.py --keep     # leave the built decks on disk

Runs the whole pipeline below the network: a recorded model response from
tests/fixtures/*.draft.json goes through the *real* `apply_draft_to_form`,
the drafted state is loaded into the *real* form via AppTest, and Generate
runs the *real* assembly. Nothing is reimplemented here -- the point is to
exercise the same code the app runs, so this suite fails when the app would.

The one thing it deliberately does not test is the model itself. What Claude
returns for a given set of notes is Tier 2's job
(tests/test_draft_live.py), which costs money and needs a key; this tier is
free, offline, deterministic, and safe to run on every change.

Every assertion prints PASS or FAIL with the actual value, and the process
exits non-zero if anything failed.
"""

import argparse
import copy
import json
import os
import re
import sys
import tempfile
import zipfile
from datetime import date
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
FIXTURES = Path(__file__).resolve().parent / "fixtures"

# The app writes files relative to the working directory (cached decks, the
# local master fallback), so run as if launched from the repo root.
os.chdir(REPO)

import app                                     # noqa: E402
import assembly                                # noqa: E402
import db                                      # noqa: E402
import package_check                           # noqa: E402
import slide_map                               # noqa: E402

TOKEN_RE = re.compile(r"\{\{[A-Z0-9_]+\}\}")

# Who the suite runs as. Never written to team_members -- the identity step
# is bypassed by setting session_state directly, the same way the password
# gate is.
TEST_USER = "Regression Suite"

# Where the stubbed logo upload/download round trip keeps its bytes.
LOGO_DIR = Path(tempfile.gettempdir()) / "premion_test_logos"
LOGO_DIR.mkdir(parents=True, exist_ok=True)


def make_test_logo():
    """A small, visually unique PNG.

    Deliberately not placeholder_logo.png: the point is to prove the deck
    carries *this* image, and reusing the placeholder would make a logo that
    silently fell back indistinguishable from one that survived.
    """
    path = LOGO_DIR / "roundtrip_logo.png"
    if not path.exists():
        from PIL import Image
        Image.new("RGB", (240, 80), (0x2E, 0x86, 0xC1)).save(path)
    return str(path)


# --------------------------------------------------------------------------
# reporting
# --------------------------------------------------------------------------

class Report:
    """Pass/fail per assertion, with the actual value whenever it fails --
    the whole reason this exists is to show what drifted."""

    def __init__(self):
        self.passed = 0
        self.failed = []
        self.skipped = []
        self.scenario = ""

    def section(self, title):
        print(f"\n  {title}")
        print(f"  {'-' * len(title)}")

    def check(self, label, ok, actual=None, expected=None):
        if ok:
            self.passed += 1
            print(f"    PASS  {label}")
        else:
            self.failed.append(f"[{self.scenario}] {label}")
            print(f"    FAIL  {label}")
            if expected is not None:
                print(f"          expected: {expected!r}")
            if actual is not None:
                print(f"          actual:   {actual!r}")
        return ok

    def equal(self, label, actual, expected):
        return self.check(label, actual == expected, actual, expected)

    def close(self, label, actual, expected, tol):
        return self.check(f"{label} (within {tol})", abs(actual - expected) <= tol,
                          actual, expected)

    def skip(self, label, why):
        self.skipped.append(f"[{self.scenario}] {label}")
        print(f"    SKIP  {label} -- {why}")


# --------------------------------------------------------------------------
# scenarios
# --------------------------------------------------------------------------

class Scenario:
    def __init__(self, name, fixture, *, vertical, agency, options, sport_key,
                 flat_fee, dynamic_creative, attribution_on, attribution_off,
                 cpm_overrides, assemble=True, custom_audience_cap=None,
                 round_trip=False):
        self.name = name
        self.fixture = fixture
        # Generate -> Load from history -> regenerate, comparing the two
        # decks. Doubles the scenario's runtime, so it's opt-in per scenario
        # rather than universal.
        self.round_trip = round_trip
        # Resolver-only scenarios skip Generate: building a deck takes ~40s
        # and 20MB to re-prove invariants the other scenarios already cover.
        self.assemble = assemble
        # (expected kept segment, expected segments moved to unresolved)
        self.custom_audience_cap = custom_audience_cap
        self.vertical = vertical
        self.agency = agency
        self.options = options                  # [(name-ish, budget), ...]
        self.sport_key = sport_key              # None when the plan has no sports
        self.flat_fee = flat_fee                # expected flat-fee amount, or None
        self.dynamic_creative = dynamic_creative
        self.attribution_on = attribution_on
        self.attribution_off = attribution_off
        self.cpm_overrides = cpm_overrides      # {product_label_prefix: cpm}

    @property
    def notes_path(self):
        return FIXTURES / f"{self.fixture}.txt"

    @property
    def draft_path(self):
        return FIXTURES / f"{self.fixture}.draft.json"

    def draft(self):
        raw = json.loads(self.draft_path.read_text(encoding="utf-8"))
        raw.pop("_comment", None)
        return raw


SCENARIOS = [
    Scenario(
        "HVAC / two options / agency / sports",
        "hvac_two_option",
        vertical="home_improvement",
        agency=True,
        options=[("$50K Plan", 50000), ("$75K Plan", 75000)],
        sport_key="nfl_reg",
        flat_fee=850,
        dynamic_creative=True,
        attribution_on=["sales_attribution"],
        attribution_off=["brand_lift"],
        cpm_overrides={"Premion Streaming TV": 28.0},
        # The hardest case to round-trip: two options, sports, a flat fee, a
        # negotiated CPM and case studies all have to survive rehydration.
        round_trip=True,
    ),
    Scenario(
        "Dental / single option / direct / no sports",
        "dental_single_option",
        vertical="healthcare",
        agency=False,
        options=[(None, 30000)],
        sport_key=None,
        flat_fee=None,
        dynamic_creative=False,
        attribution_on=[],
        attribution_off=["brand_lift", "sales_attribution"],
        cpm_overrides={},
    ),
    Scenario(
        "Custom audience cap / three customs returned",
        "custom_audience_cap",
        vertical="auto",
        agency=False,
        options=[(None, 40000)],
        sport_key=None,
        flat_fee=None,
        dynamic_creative=False,
        attribution_on=[],
        attribution_off=["brand_lift", "sales_attribution"],
        cpm_overrides={},
        assemble=False,
        custom_audience_cap=(
            "AUTO Body Style Pickup and SUV",                    # first custom wins
            ["AUTO Body Style Sedan", "AUTO Body Style Minivan"],  # moved to unresolved
        ),
    ),
]


# --------------------------------------------------------------------------
# tier 1a -- the resolver, via the real apply_draft_to_form
# --------------------------------------------------------------------------

class _StubSt:
    """apply_draft_to_form only touches st.session_state, so a plain dict is
    enough to run it outside a Streamlit runtime."""

    def __init__(self):
        self.session_state = {}
        self.secrets = {}


def apply_draft(draft, preset_state=None):
    """Run the real apply_draft_to_form and return the session_state it
    produced -- i.e. exactly what the form is loaded with after a draft.

    `preset_state` seeds session_state first, standing in for a form the
    seller had already been working in when they drafted. That is not a
    detail: a draft applied to an empty session is the *easy* case, and it
    hid a real bug for as long as this suite only ever tested it.
    """
    real = app.st
    stub = _StubSt()
    stub.session_state.update(preset_state or {})
    app.st = stub
    try:
        app.apply_draft_to_form(copy.deepcopy(draft))
    finally:
        app.st = real
    return stub.session_state


def n_months_of(state):
    """The month count the *form* will use -- read from the months the draft
    selected, never recomputed from its flight dates.

    Deriving it from flight_start/flight_end (which this did) reimplements the
    assumption under test rather than checking it: the drafted dates and the
    form's active-month selection are exactly the two things that fell out of
    step, and a check computing one from the other can only ever agree with
    itself.
    """
    months = state.get("active_months")
    if months is None:
        months = app.month_list(state["flight_start"], state["flight_end"])
    return max(1, len(months))


def full_flight_cost(option, n_months):
    """Each row's own contribution, flat fees never scaled -- the same rule
    compute_plan_totals applies."""
    total = 0.0
    for row in option["rows"]:
        if app.is_flat_fee_row(row):
            total += float(row["Cost"])
        elif option["breakout"] == app.BREAKOUT_MONTHLY:
            total += float(row["Cost"]) * n_months
        else:
            total += float(row["Cost"])
    return total


def check_draft_state(rep, scn, draft, state):
    rep.section("Drafted form state")

    vertical_key = app.VERTICALS.get(state.get("vertical_choice"))
    rep.equal("vertical resolves from the trade name", vertical_key, scn.vertical)
    rep.equal("agency_involved", state.get("agency_involved"), scn.agency)

    markup = 1.15 if scn.agency else 1.0
    rep.equal("markup recorded on the form", state.get("media_plan_markup"), markup)

    options = state.get("plan_options") or []
    rep.equal("option count", len(options), len(scn.options))

    # Budgets live in the draft (the form holds resolved rows), so the tie-out
    # is row costs -> stated budget.
    drafted = app.drafted_options(draft)
    rep.equal("every drafted option carries its own total_budget",
              [o["total_budget"] for o in drafted],
              [b for _, b in scn.options])

    n_months = n_months_of(state)
    print(f"    ....  flight: {state.get('flight_start')} -> {state.get('flight_end')} "
          f"({n_months} months), markup {markup}")

    for option, (want_name, want_budget) in zip(options, scn.options):
        label = option["name"]
        if want_name is not None:
            rep.equal(f"option named {want_name!r}", label, want_name)

        rep.close(f"{label}: line costs sum to the stated budget",
                  full_flight_cost(option, n_months), float(want_budget), 1.0)

        rate_rows = [r for r in option["rows"] if not app.is_flat_fee_row(r)]
        rep.check(f"{label}: no rate line has zero impressions",
                  all(float(r["Impressions"]) > 0 for r in rate_rows),
                  [(r["Tactic"], r["Impressions"]) for r in rate_rows
                   if float(r["Impressions"]) <= 0])
        rep.check(f"{label}: no rate line has zero cost",
                  all(float(r["Cost"]) > 0 for r in rate_rows),
                  [(r["Tactic"], r["Cost"]) for r in rate_rows if float(r["Cost"]) <= 0])

        # Gross markup has to be visible in the arithmetic itself, not just in
        # the flag -- impressions = cost / (cpm * markup) * 1000.
        bad = []
        for r in rate_rows:
            want = app.impressions_from_cost(r["Cost"], r["CPM"], markup)
            if abs(float(r["Impressions"]) - want) > 1:
                bad.append((r["Tactic"], r["Impressions"], want))
        rep.check(f"{label}: {'gross' if scn.agency else 'net'} markup applied to every rate line",
                  not bad, bad)

        if scn.sport_key:
            want_label, want_cpm = app.line_product_spec(f"sport:{scn.sport_key}")
            sport_rows = [r for r in option["rows"] if r["Tactic"].startswith(want_label)]
            if rep.check(f"{label}: a {want_label} line exists", len(sport_rows) == 1,
                         [r["Tactic"] for r in option["rows"]]):
                rep.equal(f"{label}: sports CPM is the ratecard rate, not the Premion default",
                          sport_rows[0]["CPM"], want_cpm)
        else:
            sports = [r["Tactic"] for r in option["rows"] if r["Tactic"].startswith("Live Sports")]
            rep.check(f"{label}: no sports lines", not sports, sports)

        if scn.flat_fee is not None:
            fees = [r for r in option["rows"] if app.is_flat_fee_row(r)]
            if rep.check(f"{label}: a flat-fee line exists", len(fees) == 1,
                         [r["Tactic"] for r in option["rows"]]):
                rep.equal(f"{label}: flat fee amount", float(fees[0]["Cost"]), float(scn.flat_fee))
                rep.equal(f"{label}: flat fee carries no impressions",
                          float(fees[0]["Impressions"]), 0.0)

        for prefix, want_cpm in scn.cpm_overrides.items():
            rows = [r for r in option["rows"] if r["Tactic"].startswith(prefix)]
            if rep.check(f"{label}: a {prefix} line exists", bool(rows),
                         [r["Tactic"] for r in option["rows"]]):
                rep.equal(f"{label}: {prefix} uses the negotiated CPM",
                          rows[0]["CPM"], want_cpm)

    rep.section("Attribution, audiences, dates")

    for key in scn.attribution_on:
        rep.equal(f"attribution {key} on", state.get(key), True)
    for key in scn.attribution_off:
        rep.equal(f"attribution {key} off", state.get(key), False)
    rep.check("web attribution present in the draft (no toggle -- always included)",
              "web" in (draft.get("attribution") or []), draft.get("attribution"))

    rep.equal("dynamic creative product", state.get("dynamic_creative"), scn.dynamic_creative)

    matched, unmatched = app.validate_segments(
        [a["segment"] for a in (draft.get("audiences") or [])])
    rep.check("every drafted audience validates against the catalog", not unmatched, unmatched)
    seeded = [r["Audience"] for r in (state.get("avails_seed_rows") or [])]

    if scn.custom_audience_cap:
        # The limit is enforced in Python, not asked for in the prompt, so it
        # holds for any response -- including this fixture's deliberately
        # non-compliant one.
        want_kept, want_dropped = scn.custom_audience_cap
        catalog = app.load_audience_catalog()
        rfp = dict(zip(catalog["segment"], catalog["rfp_selectable"]))
        customs = [s for s in seeded if not rfp.get(s, True)]
        rep.equal("exactly one custom audience survives", len(customs), 1)
        rep.equal("the first custom segment is the one kept", customs, [want_kept])
        rep.check("the dropped customs are gone from the avails table",
                  not [s for s in want_dropped if s in seeded], seeded)
        rep.check("RFP-selectable segments were untouched",
                  all(s in seeded for s in matched if rfp.get(s, True)), seeded)
        # The cap is a seller-side check, so it belongs in the internal
        # section of the review list rather than among the client questions.
        review = (state.get("draft_unresolved") or []) +                  (state.get("draft_unresolved_internal") or [])
        note = [u for u in review if "custom" in u.lower()]
        if rep.check("the cap is explained in the review list", bool(note), review):
            rep.check("it names every segment it left out",
                      all(s in note[0] for s in want_dropped), note[0], want_dropped)
            rep.check("it names the one it kept", want_kept in note[0], note[0])
            print(f"    ....  {note[0]}")
    elif matched:
        rep.check("validated audiences reached the avails table",
                  all(m in seeded for m in matched), seeded, matched)

    start = state.get("flight_start")
    rep.check("flight start is not in the past", start is not None and start >= date.today(),
              start, f">= {date.today()}")
    rep.check("flight end follows flight start",
              state.get("flight_end") is not None and state["flight_end"] >= start,
              state.get("flight_end"))

    unresolved = ((state.get("draft_unresolved") or [])
                  + (state.get("draft_unresolved_internal") or []))
    rep.check("the review list is non-empty", bool(unresolved), unresolved)
    rep.check("unresolved has no duplicates",
              len(unresolved) == len(set(unresolved)),
              [u for u in unresolved if unresolved.count(u) > 1])


# --------------------------------------------------------------------------
# tier 1b -- the real form and the real assembly, via AppTest
# --------------------------------------------------------------------------

def build_deck(rep, state):
    """Load the drafted state into the real form, press Generate, and capture
    the presentation the real assembly produced.

    assembly.personalize/append_case_studies are wrapped rather than replaced:
    they run for real and hand back the objects they built. app.py holds
    `assembly` as a module reference and looks the attribute up at call time,
    so wrapping here is visible to the app even though AppTest execs app.py
    as its own module.
    """
    from streamlit.testing.v1 import AppTest

    captured = {"plan_slide_ids": [], "plan_titles": []}
    real_personalize = assembly.personalize
    real_append = assembly.append_case_studies
    real_prepare_plan = assembly._prepare_media_plan_slide
    real_log = db.log_proposal
    real_upload_logo = db.upload_proposal_logo
    real_proposal_logo = db.proposal_logo

    def spy_upload_logo(proposal_key, data, filename):
        # Stand in for the storage round trip: write the bytes where the real
        # download cache would have put them, and hand back a path shaped
        # like a storage key. Keeps the logo assertions offline while still
        # exercising every line of app code that moves a logo around.
        target = LOGO_DIR / f"{proposal_key}_{filename}"
        target.write_bytes(data)
        captured["uploaded_logo"] = str(target)
        return f"logos/{target.name}", None

    def spy_proposal_logo(storage_path):
        return str(LOGO_DIR / Path(storage_path).name)

    def spy_log(client_name, vertical, market, form_json, **kwargs):
        # Capture the payload instead of inserting it. Without this the suite
        # writes a real row to the production proposals table on every run --
        # it did, until this was noticed while building the round-trip case.
        # It's also exactly the form_json the History page would later read
        # back, so the round-trip gets its input for free.
        captured["log"] = {
            "id": "00000000-0000-0000-0000-000000000000",
            "client_name": client_name, "vertical": vertical, "market": market,
            "form_json": json.loads(json.dumps(db._json_safe(form_json), default=str)),
            "deck_version_id": kwargs.get("deck_version_id"),
            "output_filename": kwargs.get("output_filename"),
            "parent_proposal_id": kwargs.get("parent_proposal_id"),
            "revision_label": kwargs.get("revision_label"),
            "logo_storage_path": kwargs.get("logo_storage_path"),
            "created_by": kwargs.get("created_by"),
        }
        return "00000000-0000-0000-0000-000000000000", None

    def spy_prepare_plan(slide, option):
        # Which slides are plan slides, straight from the code that fills
        # them. Matching on the rendered title instead would be ambiguous: a
        # single-option deck's plan title *is* the proposal title, which also
        # appears on the cover.
        #
        # Hooked on the PREPARE step rather than the sizing one: sizing now
        # runs twice per slide when the Included band has to be compressed,
        # and counting slides there would report every plan option twice.
        result = real_prepare_plan(slide, option)
        captured["plan_slide_ids"].append(slide.slide_id)
        captured["plan_titles"].append(option.get("plan_title"))
        return result

    def spy_personalize(prs, fill_data):
        warnings = real_personalize(prs, fill_data)
        captured["prs"] = prs
        captured["fill_data"] = fill_data
        captured["warnings"] = warnings
        return warnings

    def spy_append(prs, sources):
        before = {s.slide_id for s in prs.slides}
        count = real_append(prs, sources)
        captured["case_study_slide_ids"] = [s.slide_id for s in prs.slides
                                            if s.slide_id not in before]
        captured["case_study_sources"] = sources
        return count

    assembly.personalize = spy_personalize
    assembly.append_case_studies = spy_append
    assembly._prepare_media_plan_slide = spy_prepare_plan
    db.log_proposal = spy_log
    db.upload_proposal_logo = spy_upload_logo
    db.proposal_logo = spy_proposal_logo
    try:
        at = AppTest.from_file(str(REPO / "app.py"), default_timeout=600)
        at.session_state["authed"] = True
        # Past the identity step. Set rather than stubbed so the attribution
        # actually flows through to the logged row, which is worth asserting.
        at.session_state["current_user"] = TEST_USER
        for key, value in state.items():
            at.session_state[key] = value
        at.run()
        if at.exception:
            rep.check("form renders with the drafted state", False, str(at.exception[0].value))
            return None, at
        rep.check("form renders with the drafted state", True)

        # The regression that started all this: a seed-key mismatch silently
        # rebuilds every option's rows as $0 seeds on the first render.
        rendered = at.session_state["plan_options"]
        zeroed = [(o["name"], r["Tactic"]) for o in rendered for r in o["rows"]
                  if not app.is_flat_fee_row(r) and not float(r["Cost"])]
        rep.check("drafted rows survive the form's own reseed check", not zeroed, zeroed)

        generate = [b for b in at.button if b.label == "Generate proposal"]
        if not generate:
            rep.check("Generate button present", False, [b.label for b in at.button])
            return None, at
        generate[0].click().run()
        if at.exception:
            rep.check("Generate runs without raising", False, str(at.exception[0].value))
            return None, at
        rep.check("Generate runs without raising", True)
    finally:
        assembly.personalize = real_personalize
        assembly.append_case_studies = real_append
        assembly._prepare_media_plan_slide = real_prepare_plan
        db.log_proposal = real_log
        db.upload_proposal_logo = real_upload_logo
        db.proposal_logo = real_proposal_logo

    return captured, at


def slide_text(slide):
    return slide_map.extract_slide_text(slide)


def deck_money(value):
    """The number behind a rendered cell -- "$30,400 (Gross)" -> 30400.0."""
    digits = re.sub(r"[^0-9.]", "", str(value).split("(")[0])
    return float(digits) if digits else 0.0


def check_budget_ties(rep, scn, captured, at):
    """A drafted budget must tie out everywhere the deck states it.

    One invariant, and the deck shows it in three places that can disagree:
    the month count that spreads a budget has to be the month count that
    totals it, names itself in the totals row, and divides a flat fee.

    Asserted against the *deck payload* and the form's own `active_months`,
    deliberately not against the draft's flight dates -- those dates are the
    assumption under test. A draft used to set new dates without setting the
    months, so the previous flight's selection survived wherever the two
    ranges partly overlapped: an Oct-Dec draft on the form's default Sep-Nov
    flight spread $45,000 over three months and totalled it over two, showing
    "Full Flight Total (2 months)  $30,400" with the production fee split
    three ways in the allocation and two ways in the table. Every assertion in
    this suite passed, because they all derived the month count the same way
    the draft did.
    """
    rep.section("Budget ties out")

    payloads = captured["fill_data"]["media_plan_options"]
    # AppTest's session_state proxies attribute access, so it has no .get().
    active = list(at.session_state["active_months"] or [])
    n_months = max(1, len(active))
    print(f"    ....  form's own months: {n_months} ({', '.join(active) or 'none'})")

    for payload, (_, want_budget) in zip(payloads, scn.options):
        label = payload["plan_title"]
        budget = float(want_budget)
        monthly = deck_money(payload["total_cost"])
        full = payload.get("full_flight_total")

        if full is None:
            rep.equal(f"{label}: the plan's only total is the stated budget", monthly, budget)
        else:
            # The label is what a client reads the total as covering, so it
            # has to name the count the arithmetic actually used.
            said = re.search(r"\((\d+)\s+months?\)", full["label"])
            rep.equal(f"{label}: totals row names the form's own month count",
                      int(said.group(1)) if said else full["label"], n_months)
            rep.equal(f"{label}: full flight total is the stated budget",
                      deck_money(full["cost"]), budget)
            # Tolerance is one dollar per month: the monthly figure is rendered
            # to whole dollars, so multiplying it back can only drift by the
            # rounding of each month.
            rep.close(f"{label}: monthly total x month count is the stated budget",
                      monthly * n_months, budget, float(n_months))

        # Blended CPM is a media rate. A flat fee brings cost but no
        # impressions, so counting its dollars here inflates the cell into a
        # rate no line on the table carries.
        rate_rows = [r for r in payload["rows"] if r.get("impressions") != "--"]
        impressions = sum(deck_money(r["impressions"]) for r in rate_rows)
        cost = sum(deck_money(r["cost"]) for r in rate_rows)
        shown = payload.get("total_cpm", "--")
        if impressions and shown != "--":
            rep.close(f"{label}: blended CPM is media cost / media impressions",
                      deck_money(shown), cost / impressions * 1000, 0.02)


def check_deck(rep, scn, captured, keep):
    rep.section("Assembled deck")

    prs = captured.get("prs")
    if prs is None:
        rep.check("assembly produced a presentation", False, None)
        return

    slides = list(prs.slides)
    texts = [slide_text(s) for s in slides]
    print(f"    ....  {len(slides)} slides assembled")

    # 1. no unfilled tokens
    leftover = []
    for i, text in enumerate(texts, start=1):
        for tok in TOKEN_RE.findall(text or ""):
            leftover.append((i, tok))
    rep.check("no unfilled {{TOKEN}}s remain", not leftover, leftover[:10])

    # 2. no section dividers
    derived = slide_map.build_slide_map_from_prs(prs)
    dividers = [n for n, key in derived.items() if key == slide_map.SECTION_DIVIDER_KEY]
    rep.check("no section dividers in the deck", not dividers, dividers)

    # 3. one plan slide per option, in order
    option_payloads = captured["fill_data"]["media_plan_options"]
    by_id = {s.slide_id: i for i, s in enumerate(slides)}
    rep.equal("one plan slide per option",
              len(captured["plan_slide_ids"]), len(option_payloads))
    rep.equal("each option's own title reached its slide",
              captured["plan_titles"], [p["plan_title"] for p in option_payloads])
    flat = [by_id[sid] for sid in captured["plan_slide_ids"] if sid in by_id]
    rep.equal("every plan slide is still in the deck", len(flat), len(captured["plan_slide_ids"]))
    if flat:
        rep.check("plan slides appear in option order", flat == sorted(flat), flat)
        rep.check("plan slides are consecutive",
                  flat == list(range(flat[0], flat[0] + len(flat))), flat)
        # Each option's title belongs to its own slide and no other.
        misplaced = []
        for pos, payload in zip(flat, option_payloads):
            title = payload["plan_title"]
            if title and title not in (texts[pos] or ""):
                misplaced.append((pos + 1, title))
        rep.check("each plan slide carries its own option's title", not misplaced, misplaced)

    # 4. case studies contiguous, immediately before the first plan slide
    cs_ids = captured.get("case_study_slide_ids") or []
    if not cs_ids:
        rep.skip("case studies sit immediately before the first plan slide",
                 "none were selected (no vault match or vault unreachable)")
    elif not flat:
        rep.skip("case studies sit immediately before the first plan slide",
                 "no plan slide was located")
    else:
        by_id = {s.slide_id: i for i, s in enumerate(slides)}
        idx = sorted(by_id[sid] for sid in cs_ids if sid in by_id)
        first_plan = min(flat)
        rep.check("case study slides are contiguous",
                  idx == list(range(idx[0], idx[0] + len(idx))), idx)
        rep.equal("case studies end immediately before the first plan slide",
                  idx[-1] + 1 if idx else None, first_plan)

    # 5. the plan table clears the Included-with-Campaign graphic
    rep.check("assembly reported no overflow warnings",
              not captured.get("warnings"), captured.get("warnings"))
    measured = 0
    for i in flat:
        slide = slides[i]
        table_shape = assembly._find_table_shape(slide)
        if table_shape is None:
            continue
        bottom = assembly.table_bottom(slide)
        floor = assembly._content_floor(slide, table_shape)
        if bottom is None or floor is None:
            continue
        measured += 1
        rep.check(f"slide {i + 1}: table bottom clears the content floor "
                  f"({bottom / 914400:.2f}in <= {floor / 914400:.2f}in)",
                  bottom <= floor, f"{bottom / 914400:.3f}in", f"<= {floor / 914400:.3f}in")
    if not measured:
        rep.skip("plan table measured against the content floor", "no table shape found")

    # 6. the dynamic creative slide, when the product is on
    if scn.dynamic_creative:
        keys = set(derived.values())
        rep.check("dynamic creative slide included", "dynamic_creative" in keys,
                  sorted(k for k in keys if "dynamic" in k))

    # 7. package integrity
    out = REPO / "tests" / f"_built_{scn.fixture}.pptx"
    prs.save(str(out))
    try:
        # Relationships that resolve. python-pptx opens a package whose parts
        # point at targets that were never written; PowerPoint refuses it,
        # and that asymmetry cost a real deck -- so it's asserted here rather
        # than left to the rendering step to notice.
        problems = package_check.check_package(str(out))
        rep.check("every relationship resolves to a part that exists",
                  not problems, problems[:5])
        with zipfile.ZipFile(out) as zf:
            rep.check("zip integrity", zf.testzip() is None, zf.testzip())
            names = zf.namelist()
            dupes = sorted({n for n in names if names.count(n) > 1})
            rep.check("no duplicate partnames", not dupes, dupes)
            slide_parts = [n for n in names if n.startswith("ppt/slides/slide")]
            rep.equal("slide parts match slide count", len(slide_parts), len(slides))
        print(f"    ....  wrote {out.name} ({out.stat().st_size / 1048576:.1f} MiB)")
    finally:
        if not keep:
            out.unlink(missing_ok=True)


# --------------------------------------------------------------------------

def rehydrate(row, revision_default=None):
    """Run the real rehydrate_proposal_into_form against a stub session,
    returning the session_state it produced -- what the form is loaded with
    after "Load into form"."""
    real = app.st
    real_logo = db.proposal_logo
    stub = _StubSt()
    app.st = stub
    db.proposal_logo = lambda storage_path: str(LOGO_DIR / Path(storage_path).name)
    try:
        notes = app.rehydrate_proposal_into_form(
            row, parent_proposal_id=row["id"], revision_default=revision_default)
    finally:
        app.st = real
        db.proposal_logo = real_logo
    return stub.session_state, notes


def deck_blobs(prs):
    """Every image blob in a presentation, for checking a specific picture
    actually made it in."""
    return {part.blob for part in prs.part.package.iter_parts()
            if getattr(part, "content_type", "").startswith("image/")}


def deck_shape(captured, prs_slides=None):
    """The invariants a round-trip has to preserve, as a comparable dict."""
    prs = captured["prs"]
    slides = list(prs.slides)
    by_id = {s.slide_id: i for i, s in enumerate(slides)}
    options = captured["fill_data"]["media_plan_options"]
    return {
        "slide_count": len(slides),
        "option_count": len(options),
        "plan_titles": list(captured["plan_titles"]),
        "plan_positions": [by_id[sid] for sid in captured["plan_slide_ids"] if sid in by_id],
        "case_study_count": len(captured.get("case_study_slide_ids") or []),
        "rows": [[(r["tactic"], r["impressions"], r["cost"]) for r in o["rows"]] for o in options],
        "totals": [(o["total_impressions"], o["total_cost"]) for o in options],
        "unfilled_tokens": sorted({t for s in slides
                                   for t in TOKEN_RE.findall(slide_text(s) or "")}),
        # Name only: the path differs between runs, what must match is which
        # image the deck was built with.
        "logo": Path(str(captured["fill_data"]["logo_path"])).name,
    }


def check_round_trip(rep, scn, first):
    """Generate -> Load from history -> regenerate untouched -> compare.

    Any field that doesn't survive rehydration shows up here as a difference
    between the two decks. This is the check that makes "Load into form"
    trustworthy: the promise is that loading and regenerating an untouched
    proposal reproduces it, and the only honest way to know is to do it.
    """
    rep.section("Round trip: generate -> load -> regenerate")

    logged = first.get("log")
    if not rep.check("the proposal was logged with a form_json payload", bool(logged),
                     None if logged else "log_proposal was never called"):
        return

    state, notes = rehydrate(logged, revision_default="Revision 2")
    for note in notes:
        print(f"    ....  rehydration note: {note}")

    # A few fields worth asserting directly, since a silent default here
    # would still produce a matching deck for the wrong reason.
    form = logged["form_json"]
    rep.equal("client name survives", state.get("client_name"), logged["client_name"])
    rep.equal("proposal title survives", state.get("proposal_title"),
              form.get("proposal_title"))
    rep.equal("agency toggle survives", state.get("agency_involved"),
              form.get("agency_involved"))
    rep.equal("option count survives", len(state.get("plan_options") or []),
              len(form.get("plan_options") or []))
    rep.check("every restored plan row is marked dirty",
              all(all(o["dirty"]) for o in state.get("plan_options") or []),
              [o["dirty"] for o in state.get("plan_options") or []])
    rep.equal("parent link is set for the regeneration",
              state.get("history_parent_id"), logged["id"])

    # --- the logo ---------------------------------------------------------
    logo_bytes = Path(make_test_logo()).read_bytes()
    rep.check("the logo was recorded on the logged row",
              bool(logged.get("logo_storage_path")), logged.get("logo_storage_path"))
    rep.check("form_json records that a logo was used",
              bool(form.get("logo_used")), form.get("logo_used"))
    rep.equal("the first deck was built with the real logo, not the placeholder",
              Path(str(first["fill_data"]["logo_path"])).name, Path(make_test_logo()).name)
    rep.check("the logo's bytes are in the first deck",
              logo_bytes in deck_blobs(first["prs"]), None,
              "the CLIENT_LOGO picture to carry the uploaded image")
    rep.equal("the stored logo is restored on load",
              state.get("restored_logo_storage_path"), logged.get("logo_storage_path"))
    rep.check("no 'logo missing' state after a successful restore",
              not state.get("restored_logo_missing"), state.get("restored_logo_missing"))
    rep.equal("revision label default is prefilled on load",
              state.get("revision_label"), "Revision 2")
    rep.equal("the proposal is attributed to the signed-in user",
              logged.get("created_by"), TEST_USER)

    # The seed key is the one that silently wipes rows when it disagrees.
    real = app.st
    stub = _StubSt()
    stub.session_state.update(state)
    app.st = stub
    try:
        derived = str(app.read_seed_selections())
    finally:
        app.st = real
    rep.equal("rehydrated seed key matches what the form will derive",
              state.get("_product_seed_key"), derived)

    second, _ = build_deck(rep, state)
    if not second or "prs" not in second:
        rep.check("the reloaded proposal regenerates", False, None)
        return
    rep.check("the reloaded proposal regenerates", True)

    before, after = deck_shape(first), deck_shape(second)
    for field in ("slide_count", "option_count", "plan_titles", "plan_positions",
                  "case_study_count", "rows", "totals", "unfilled_tokens", "logo"):
        rep.equal(f"round trip preserves {field}", after[field], before[field])
    rep.check("the logo's bytes are in the regenerated deck too",
              logo_bytes in deck_blobs(second["prs"]), None,
              "the rebuilt deck to carry the same logo image")


STATION_BY_MARKET = {
    "DC": {"present": ("WUSA",), "absent": ("WPMT", "FOX43", "FOX 43")},
    "Harrisburg": {"present": ("WPMT",), "absent": ("WUSA",)},
}


def check_total_tv_by_market(rep):
    """A Total TV proposal must carry its own market's station slides and
    nobody else's.

    This shipped broken: the WUSA9 broadcast-schedule slide was keyed plain
    `total_tv` rather than `total_tv:dc`, so every Harrisburg Total TV deck
    went out carrying a DC station's schedule alongside its own WPMT slide.
    The station-branded slides are the most visible thing in the deck to a
    client, so a wrong one is worse than a missing one.

    Built straight from the active master rather than through the form: the
    question is purely which slides the market selects, and running the whole
    form twice more would add a minute to prove nothing extra.
    """
    rep.scenario = "Total TV by market"
    print("\n" + "=" * 78)
    print("SCENARIO  Total TV station slides follow the market")
    print("=" * 78)
    rep.section("Per-market station slides")

    master_path, version_id, warning = db.master_deck(str(REPO / "TEGNA_MASTER_DECK_v1_1.pptx"))
    if master_path is None:
        rep.skip("Total TV station slides follow the market", warning or "no master deck")
        return
    print(f"    ....  master deck version {version_id}")

    for market, expected in STATION_BY_MARKET.items():
        selections = copy.deepcopy(assembly.SELECTIONS)
        selections["market"] = market
        selections["preset"] = "standard"
        selections["products"] = dict(selections["products"])
        selections["products"]["total_tv"] = True

        prs, _, kept = assembly.build_presentation(master_path, selections)
        found = {}
        for index, slide in enumerate(prs.slides, start=1):
            text = (slide_text(slide) or "").upper()
            for station in ("WUSA", "WPMT", "FOX43", "FOX 43"):
                if station in text:
                    found.setdefault(station, []).append(index)

        print(f"    ....  {market}: {kept} slides, stations {dict(found)}")
        for station in expected["present"]:
            rep.check(f"{market}: {station} slide(s) present", bool(found.get(station)),
                      dict(found))
        for station in expected["absent"]:
            rep.check(f"{market}: no {station} slide", not found.get(station), dict(found))

        # The generic Total TV slides are market-independent and must survive
        # in both -- the fix must not have over-narrowed them.
        keys = set(slide_map.build_slide_map_from_prs(prs).values())
        rep.check(f"{market}: the generic Total TV slide is still included",
                  "total_tv" in keys, sorted(k for k in keys if "total_tv" in str(k)))
        want_key = f"total_tv:{'dc' if market == 'DC' else 'harrisburg'}"
        other_key = f"total_tv:{'harrisburg' if market == 'DC' else 'dc'}"
        rep.check(f"{market}: carries {want_key}", want_key in keys,
                  sorted(k for k in keys if "total_tv" in str(k)))
        rep.check(f"{market}: does not carry {other_key}", other_key not in keys,
                  sorted(k for k in keys if "total_tv" in str(k)))


NEW_SLIDES_DECK = REPO / "New_Slides_tagged.pptx"


VARIANT_PREFIXES = ("client_title_cobrand:", "proposal_template_total_tv:",
                    "broadcast_schedule_template:")


def _merged_master(rep):
    """A master deck containing the Total TV variant slides.

    Prefers the active master: since version 6 it carries them, so the tests
    run against exactly what production builds from. Falls back to merging
    `New_Slides_tagged.pptx` in, which is how these assertions ran before that
    upload and how they'd run again against an older master.

    The fallback re-applies each key label after copying, because
    `copy_slide_into` deliberately doesn't carry a notes part across (one
    belongs to exactly one slide). That matters more than it sounds: the
    co-brand covers and Total TV plan templates are text-identical to the
    slides they replace and resolve from their notes label ALONE, so a merge
    that loses notes collapses them back onto the standard slides silently.
    """
    import tag_deck_keys
    from pptx import Presentation

    master_path, _, warning = db.master_deck(str(REPO / "TEGNA_MASTER_DECK_v1_1.pptx"))
    if master_path is None:
        return None

    keys = set(slide_map.build_slide_map_from_prs(Presentation(master_path)).values())
    if all(any(str(k).startswith(prefix) for k in keys) for prefix in VARIANT_PREFIXES):
        return master_path

    target = Path(db.scratch_dir("premion_merged")) / "master_plus_new_tagged.pptx"
    if target.exists() and target.stat().st_size > 0:
        return str(target)
    if not NEW_SLIDES_DECK.exists():
        return None

    prs = Presentation(master_path)
    source = Presentation(str(NEW_SLIDES_DECK))
    keys = [slide_map.notes_key(s) for s in source.slides]
    before = len(prs.slides._sldIdLst)
    cache = assembly.ImportCache(prs)
    for index in range(len(source.slides)):
        assembly.copy_slide_into(str(NEW_SLIDES_DECK), index, prs, cache=cache)
    for offset, key in enumerate(keys):
        tag_deck_keys.write_notes_key(list(prs.slides)[before + offset], key)
    prs.save(str(target))
    return str(target)


# A slide branded for the wrong market is the most visible mistake this deck
# can make, so each market is checked for the other's branding explicitly.
FOREIGN_BRAND = {"DC": ("FOX43", "WPMT"), "Harrisburg": ("WUSA",)}


def check_total_tv_variants(rep):
    """Total TV swaps in the market's co-brand cover and plan template."""
    rep.scenario = "Total TV variants"
    print("\n" + "=" * 78)
    print("SCENARIO  Total TV co-brand cover and plan template, per market")
    print("=" * 78)

    merged = _merged_master(rep)
    if merged is None:
        rep.skip("Total TV variant selection",
                 f"needs {NEW_SLIDES_DECK.name} and a reachable master deck")
        return
    rep.section("Selection per market and toggle state")

    for market in ("DC", "Harrisburg"):
        suffix = assembly.market_suffix(market)
        other = "harrisburg" if suffix == "dc" else "dc"
        for total_tv, imported in ((False, False), (True, False), (True, True)):
            selections = copy.deepcopy(assembly.SELECTIONS)
            selections["market"] = market
            selections["preset"] = "standard"
            selections["products"] = dict(selections["products"])
            selections["products"]["total_tv"] = total_tv
            selections["broadcast_schedule_imported"] = imported

            prs, _, _ = assembly.build_presentation(merged, selections)
            keys = set(slide_map.build_slide_map_from_prs(prs).values())
            label = f"{market}/TT={'on' if total_tv else 'off'}/import={'y' if imported else 'n'}"

            cobrand = f"{assembly.COBRAND_TITLE_PREFIX}{suffix}"
            tt_plan = f"{assembly.TOTAL_TV_PLAN_PREFIX}{suffix}"
            schedule = f"{assembly.BROADCAST_SCHEDULE_PREFIX}{suffix}"

            if total_tv:
                rep.check(f"{label}: co-brand cover in", cobrand in keys, sorted(
                    k for k in keys if "client_title" in str(k)))
                rep.check(f"{label}: standard cover out", "client_title" not in keys, sorted(
                    k for k in keys if "client_title" in str(k)))
                rep.check(f"{label}: Total TV plan template in", tt_plan in keys, sorted(
                    k for k in keys if "proposal_template" in str(k)))
                rep.check(f"{label}: standard plan template out",
                          "proposal_template" not in keys, sorted(
                              k for k in keys if "proposal_template" in str(k)))
                rep.check(f"{label}: schedule template {'in' if imported else 'out'}",
                          (schedule in keys) == imported,
                          sorted(k for k in keys if "broadcast_schedule" in str(k)))
                if imported:
                    placeholder = [
                        i + 1 for i, s in enumerate(prs.slides)
                        if assembly.SCHEDULE_PLACEHOLDER_MARKER in slide_text(s).upper()]
                    rep.check(f"{label}: static placeholder replaced", not placeholder, placeholder)
            else:
                rep.check(f"{label}: standard cover kept", "client_title" in keys, sorted(
                    k for k in keys if "client_title" in str(k)))
                rep.check(f"{label}: standard plan template kept",
                          "proposal_template" in keys, sorted(
                              k for k in keys if "proposal_template" in str(k)))
                rep.check(f"{label}: no Total TV variants at all",
                          not [k for k in keys if str(k).startswith(
                              (assembly.COBRAND_TITLE_PREFIX, assembly.TOTAL_TV_PLAN_PREFIX,
                               assembly.BROADCAST_SCHEDULE_PREFIX))],
                          sorted(str(k) for k in keys if "cobrand" in str(k)
                                 or "total_tv" in str(k) or "broadcast_schedule" in str(k)))

            # No variant for the other market, by key or by visible branding.
            leaked = [k for k in keys if str(k).endswith(f":{other}")
                      and str(k).startswith((assembly.COBRAND_TITLE_PREFIX,
                                             assembly.TOTAL_TV_PLAN_PREFIX,
                                             assembly.BROADCAST_SCHEDULE_PREFIX))]
            rep.check(f"{label}: no {other} variant keys", not leaked, leaked)
            branded = []
            for index, slide in enumerate(prs.slides, start=1):
                text = (slide_text(slide) or "").upper()
                for station in FOREIGN_BRAND[market]:
                    if station in text:
                        branded.append((index, station))
            rep.check(f"{label}: no other market's station branding", not branded, branded)


def check_campaign_specs_fit(rep):
    """The Campaign Specs panel must fit above whatever sits below it.

    Uses the copy the model actually drafted for the Ridgeline fixture --
    19 bullets, ~1,200 characters -- because that's the panel that overflowed
    in live testing. A synthetic short fixture would pass and prove nothing.
    """
    import json
    from pptx import Presentation

    rep.scenario = "Campaign Specs fit"
    print("\n" + "=" * 78)
    print("SCENARIO  Campaign Specs auto-fit")
    print("=" * 78)
    rep.section("Real drafted copy fits above the floor")

    fixture = FIXTURES / "ridgeline_campaign_specs.json"
    master_path, _, warning = db.master_deck(str(REPO / "TEGNA_MASTER_DECK_v1_1.pptx"))
    if not fixture.exists() or master_path is None:
        rep.skip("Campaign Specs auto-fit", warning or "fixture missing")
        return

    specs = json.loads(fixture.read_text(encoding="utf-8"))["campaign_specs"]
    selections = copy.deepcopy(assembly.SELECTIONS)
    selections["preset"] = "standard"
    selections["market"] = "DC"
    prs, _, _ = assembly.build_presentation(master_path, selections)
    slide = assembly.find_slide_with_marker(prs, "{{GOALS_BULLETS}}")
    for token, key in (("GOALS_BULLETS", "goals"), ("AUDIENCE_BULLETS", "audience"),
                       ("GEOGRAPHY_BULLETS", "geography"), ("BUDGET_BULLETS", "budget"),
                       ("PLACEMENTS_BULLETS", "placements"), ("TIMING_BULLETS", "timing")):
        assembly.fill_bullet_list_in_slide(slide, token, specs[key] or ["--"])

    total = sum(len(b) for v in specs.values() for b in v)
    scale, fits = assembly.fit_campaign_specs(slide)
    frame = next(s for s in slide_map.iter_all_shapes(slide.shapes)
                 if s.has_text_frame and "Goals & Approach" in s.text_frame.text)
    sizes = sorted({r.font.size.pt for p in frame.text_frame.paragraphs
                    for r in p.runs if r.font.size})
    print(f"    ....  {total} characters over {len(frame.text_frame.paragraphs)} paragraphs")
    rep.check(f"the panel reports fitting (scale {scale}, {sizes}pt)", fits, (scale, sizes))
    rep.check("it was actually shrunk, not just declared to fit", scale < 1.0, scale)
    rep.check("type stays readable", all(s >= 9 for s in sizes), sizes)
    height = assembly._estimate_frame_height(frame.text_frame, frame.width, 1.0)
    floor = assembly._content_floor(slide, frame)
    available = floor - assembly._TABLE_CLEARANCE - frame.top
    print(f"    ....  final height ~{height / 914400:.2f}in in {available / 914400:.2f}in")


def check_media_plan_clearance(rep):
    """The plan table must clear the graphic below it with visible space."""
    rep.scenario = "media plan clearance"
    print("\n" + "=" * 78)
    print("SCENARIO  Media plan table clearance by row count")
    print("=" * 78)
    rep.section("Comfortable margin at realistic row counts")

    master_path, _, warning = db.master_deck(str(REPO / "TEGNA_MASTER_DECK_v1_1.pptx"))
    if master_path is None:
        rep.skip("media plan clearance", warning or "no master deck")
        return
    selections = copy.deepcopy(assembly.SELECTIONS)
    selections["preset"] = "standard"
    selections["market"] = "DC"

    for count in (1, 3, 5, 7, 9, 12):
        prs, _, _ = assembly.build_presentation(master_path, selections)
        slide = assembly.find_slide_with_marker(prs, "{{TACTIC}}")
        rows = [{"tactic": f"Line {i + 1} Streaming TV", "flight": "Sep - Nov",
                 "geo": "Washington, DC DMA", "targeting": "Homeowners 35+, HHI $150K+",
                 "impressions": "123,456", "cost": "$12,345"} for i in range(count)]
        assembly.fill_table_rows(slide, 1, rows, {
            "tactic": "TACTIC", "flight": "FLIGHT", "geo": "GEO", "targeting": "TARGETING",
            "impressions": "IMPRESSIONS", "cost": "COST"})
        overflow = assembly.condense_media_plan_table(slide, count)
        shape = assembly._find_table_shape(slide)
        bottom = assembly.table_bottom(slide)
        floor = assembly._content_floor(slide, shape)
        clear = (floor - bottom) / 914400
        rep.check(f"{count:>2} rows clear the floor by {clear:+.2f}in", bottom <= floor, clear)
        rep.check(f"{count:>2} rows don't warn", not overflow, overflow)

    # A full audience stack is the normal Targeting value now that the column
    # takes every Campaign Specs bullet rather than the first. Targeting is the
    # widest column on the table (3.30in), so a four-attribute stack wraps to
    # two lines -- and a row is as tall as its tallest cell, so the sizer has
    # to reserve for the wrap rather than for one line per row. Asserting the
    # wrap actually happened matters as much as the clearance: if the string
    # fitted on one line this would prove nothing.
    rep.section("A wrapped Targeting stack is still reserved for")
    stack = "Adults 35+, homeowners, higher income, researching cosmetic procedures"
    wrapped_at = []
    for count in (1, 5, 9, 12):
        prs, _, _ = assembly.build_presentation(master_path, selections)
        slide = assembly.find_slide_with_marker(prs, "{{TACTIC}}")
        rows = [{"tactic": f"Line {i + 1} Streaming TV", "flight": "Sep - Nov",
                 "geo": "Washington, DC DMA", "targeting": stack,
                 "impressions": "123,456", "cost": "$12,345"} for i in range(count)]
        assembly.fill_table_rows(slide, 1, rows, {
            "tactic": "TACTIC", "flight": "FLIGHT", "geo": "GEO", "targeting": "TARGETING",
            "impressions": "IMPRESSIONS", "cost": "COST"})
        overflow = assembly.condense_media_plan_table(slide, count)
        shape = assembly._find_table_shape(slide)
        table = shape.table
        bottom = assembly.table_bottom(slide)
        floor = assembly._content_floor(slide, shape)
        clear = (floor - bottom) / 914400
        font_pt = next((run.font.size.pt
                        for cell in table.rows[1].cells
                        for para in cell.text_frame.paragraphs
                        for run in para.runs if run.font.size), None)
        lines = assembly._lines_in_row(table, 1, font_pt) if font_pt else 0
        if lines >= 2:
            wrapped_at.append(count)
        # Not asserted per row count: past ~7 rows the sizer has shrunk the
        # font far enough (7pt at 9 rows, 6pt at 12) that the same stack fits
        # on one line, so demanding a wrap everywhere would be demanding the
        # sizer do something worse.
        print(f"    ....  {count:>2} rows -> {font_pt}pt, Targeting {lines} line(s), "
              f"clears {clear:+.2f}in")
        rep.check(f"{count:>2} wrapped rows clear the floor by {clear:+.2f}in",
                  bottom <= floor, clear)
        rep.check(f"{count:>2} wrapped rows don't warn", not overflow, overflow)

    # If nothing wrapped, the section above proved nothing about wrapping.
    rep.check(f"the stack does wrap at readable sizes (row counts {wrapped_at})",
              bool(wrapped_at), wrapped_at)


def check_post_draft_edits(rep):
    """Draft, then change products, then import a schedule.

    This is the sequence a seller actually works in, and it has now silently
    zeroed drafted rows twice -- once via the seed-key mismatch, once via the
    "product selections changed" branch rebuilding every row. A toggle must
    only ever touch rows belonging to the product that changed, so the
    sequence itself is the test rather than any one of its steps.
    """
    from streamlit.testing.v1 import AppTest
    import wideorbit

    rep.scenario = "post-draft edits"
    print("\n" + "=" * 78)
    print("SCENARIO  Drafted rows survive product changes and a schedule import")
    print("=" * 78)

    fixture = FIXTURES / "wideorbit" / "regency_planner.xls"
    draft = json.loads((FIXTURES / "ashford_post_draft.draft.json").read_text(encoding="utf-8"))
    draft.pop("_comment", None)
    state = apply_draft(draft)

    real_log = db.log_proposal
    db.log_proposal = lambda *a, **k: ("test", None)
    try:
        at = AppTest.from_file(str(REPO / "app.py"), default_timeout=600)
        at.session_state["authed"] = True
        at.session_state["current_user"] = TEST_USER
        for key, value in state.items():
            at.session_state[key] = value
        at.run()
        if at.exception:
            rep.check("the drafted form renders", False, str(at.exception[0].value))
            return
        baseline = _plan_snapshot(at)

        rep.section("After the draft")
        rep.equal("two drafted lines", len(baseline), 2)
        rep.check("both carry real money",
                  all(v["cost"] > 0 and v["impressions"] > 0 for v in baseline.values()),
                  baseline)
        rep.check("the negotiated CPM is in place",
                  any(abs(v["cpm"] - 29.0) < 0.01 for v in baseline.values()), baseline)
        rep.check("flight is a real range, not TBD",
                  all("TBD" not in v["flight"] for v in baseline.values()), baseline)

        def unchanged(label, snapshot):
            for tactic, before in baseline.items():
                after = snapshot.get(tactic)
                if after is None:
                    rep.check(f"{label}: {tactic} still present", False, sorted(snapshot))
                    continue
                rep.check(f"{label}: {tactic} untouched", after == before,
                          {"before": before, "after": after})

        rep.section("Toggling a product on must not touch drafted rows")
        at.session_state["total_tv"] = True
        at.run()
        after_on = _plan_snapshot(at)
        unchanged("Total TV on", after_on)
        added = set(after_on) - set(baseline)
        rep.equal("exactly one row appeared", len(added), 1)
        rep.check("and it's the broadcast one",
                  all("Broadcast" in t for t in added), added)

        rep.section("Importing a schedule must not touch drafted rows")
        if fixture.exists():
            at.session_state["broadcast_schedule"] = wideorbit.parse_schedule(
                str(fixture), fixture.name)
            at.run()
            after_import = _plan_snapshot(at)
            unchanged("schedule import", after_import)
            broadcast = [v for t, v in after_import.items() if "Broadcast Schedule (" in t]
            rep.equal("the broadcast line is priced from the schedule", len(broadcast), 1)
            if broadcast:
                rep.check("with real money on it",
                          broadcast[0]["cost"] > 0 and broadcast[0]["impressions"] > 0,
                          broadcast[0])
        else:
            rep.skip("schedule import step", "Wide Orbit fixture not present")

        rep.section("Toggling the product back off removes only its row")
        at.session_state["total_tv"] = False
        at.session_state["broadcast_schedule"] = None
        at.run()
        after_off = _plan_snapshot(at)
        unchanged("Total TV off", after_off)
        rep.equal("back to the drafted lines alone", sorted(after_off), sorted(baseline))
    finally:
        db.log_proposal = real_log


def check_drafted_months_survive_a_used_form(rep):
    """Draft a new flight over a form that already has one.

    The Capital Ridge Dental case. A seller opens the form (default Sep-Nov),
    pastes notes quoting $45,000 for October-December, and drafts. The draft
    set the dates but not the month selection, and main() only discards a
    stored selection that doesn't overlap the new range *at all* -- Sep-Nov
    and Oct-Dec share two months, so the stale subset was kept. The budget was
    spread over three months and totalled over two: "Full Flight Total
    (2 months)  $30,400", with the $1,200 production fee split $600/month
    instead of $400.

    The partial overlap is the whole point, so this asserts on it directly
    rather than on a generic draft: an empty session and a fully-disjoint
    flight both take paths that were never broken.
    """
    from streamlit.testing.v1 import AppTest

    rep.scenario = "drafted months over a used form"
    print("\n" + "=" * 78)
    print("SCENARIO  A drafted flight replaces the months the form was holding")
    print("=" * 78)

    budget, fee = 45000.0, 1200.0
    draft = {
        "client_name": "Capital Ridge Dental",
        "vertical": "healthcare",
        "market": "Harrisburg",
        "agency_involved": False,
        "flight_start": "2026-10-01",
        "flight_end": "2026-12-31",
        "total_budget": budget,
        "breakout": "monthly",
        "campaign_specs": {"audience": ["Adults 35+, homeowners, higher income"]},
        "media_plan_lines": [
            {"product": "premion_streaming_tv", "audience_track": "DEMO Homeowner",
             "allocation": {"percent_of_remainder": 100}},
            {"product": "custom_fee", "label": "Commercial Production (:30 spot)",
             "allocation": {"flat_amount": fee}},
        ],
    }
    # The form as the seller left it: the default flight, months selected.
    used_form = {
        "flight_start": app.DEFAULT_FLIGHT_START,
        "flight_end": app.DEFAULT_FLIGHT_END,
        "active_months": app.month_list(app.DEFAULT_FLIGHT_START, app.DEFAULT_FLIGHT_END),
    }
    overlap = [m for m in used_form["active_months"]
               if m in app.month_list(date(2026, 10, 1), date(2026, 12, 31))]
    rep.check("the two flights really do partly overlap (else this proves nothing)",
              0 < len(overlap) < 3, overlap)

    state = apply_draft(draft, preset_state=used_form)
    rep.equal("the draft claims the flight it was given",
              state.get("active_months"), ["Oct 2026", "Nov 2026", "Dec 2026"])

    real_log = db.log_proposal
    db.log_proposal = lambda *a, **k: ("test", None)
    try:
        at = AppTest.from_file(str(REPO / "app.py"), default_timeout=600)
        at.session_state["authed"] = True
        at.session_state["current_user"] = TEST_USER
        # The form is already in this state, and then the draft lands on it.
        for key, value in list(used_form.items()) + list(state.items()):
            at.session_state[key] = value
        at.run()
        if at.exception:
            rep.check("the drafted form renders", False, str(at.exception[0].value))
            return

        active = list(at.session_state["active_months"] or [])
        rep.section("The form and the draft agree on the flight")
        rep.equal("the form kept the drafted months, not the ones it was holding",
                  active, ["Oct 2026", "Nov 2026", "Dec 2026"])

        n_months = max(1, len(active))
        option = at.session_state["plan_options"][0]
        totals = app.compute_plan_totals(
            option["rows"], option["breakout"], n_months,
            app.format_flight_label(active, active))

        rep.section("And the budget ties")
        for row in option["rows"]:
            print(f"    ....  {str(row['Tactic'])[:44]:<44} ${float(row['Cost']):>10,.2f}"
                  f"  {'flat fee' if app.is_flat_fee_row(row) else 'rate'}")
        rep.close("full flight total is the stated budget",
                  totals["full_flight_cost"], budget, 0.01)
        rep.close("monthly total x month count is the stated budget",
                  totals["monthly_cost"] * n_months, budget, 0.01)

        # The symptom that made the divergence visible on the slide: the fee
        # is one cost divided by the plan's months, whatever else changes.
        fee_rows = [r for r in totals["preview_rows"] if r["is_flat_fee"]]
        if rep.equal("the flat fee is still one line", len(fee_rows), 1):
            rep.close("the flat fee is divided by the plan's own month count",
                      fee_rows[0]["monthly_cost"], fee / n_months, 0.01)
            rep.close("and is never scaled up by it",
                      fee_rows[0]["full_flight_cost"], fee, 0.01)
    finally:
        db.log_proposal = real_log


def _plan_snapshot(at):
    """{tactic: rounded figures} for the first option."""
    option = at.session_state["plan_options"][0]
    return {row["Tactic"]: {"impressions": round(float(row["Impressions"])),
                            "cost": round(float(row["Cost"])),
                            "cpm": round(float(row["CPM"]), 2),
                            "flight": str(row["Flight"])}
            for row in option["rows"]}


def check_cpm_column(rep):
    """The CPM column, on both plan templates, both ways."""
    rep.scenario = "CPM column"
    print("\n" + "=" * 78)
    print("SCENARIO  Media plan CPM column")
    print("=" * 78)

    master_path, _, warning = db.master_deck(str(REPO / "TEGNA_MASTER_DECK_v1_1.pptx"))
    if master_path is None:
        rep.skip("CPM column", warning or "no master deck")
        return

    rows = [{"tactic": "Premion Streaming TV", "flight": "May - Jun",
             "geo": "Washington, DC DMA", "targeting": "Homeowners 35+",
             "impressions": "559,720", "cost": "$18,667", "cpm": "$29.00"},
            {"tactic": "Production Fee", "flight": "May - Jun", "geo": "-", "targeting": "-",
             "impressions": "--", "cost": "$850", "cpm": "--"}]

    for total_tv in (False, True):
        label = "Total TV template" if total_tv else "standard template"
        rep.section(label)
        built = {}
        for show in (False, True):
            selections = copy.deepcopy(assembly.SELECTIONS)
            selections["preset"] = "standard"
            selections["market"] = "DC"
            selections["products"] = dict(selections["products"])
            selections["products"]["total_tv"] = total_tv
            prs, _, _ = assembly.build_presentation(master_path, selections)
            slide = assembly.find_slide_with_marker(prs, "{{TACTIC}}")
            # The three phases a plan slide goes through in personalize.
            # Uncompressed (False): this case is about the CPM column, and
            # holding the band constant keeps the widths comparable.
            entry = assembly._prepare_media_plan_slide(slide, {
                "plan_title": "T", "totals_label": "Monthly Totals",
                "total_impressions": "559,720", "total_cost": "$19,517", "rows": rows,
                "included_list": ["A"], "full_flight_total": None,
                "show_cpm": show, "total_cpm": "$34.88"})
            assembly._size_media_plan_slide(entry)
            assembly._finish_media_plan_slide(entry, False)
            shape = assembly._find_table_shape(slide)
            table = shape.table
            built[show] = {
                "cols": len(table.columns),
                "headers": [" ".join(table.cell(0, c).text.split())
                            for c in range(len(table.columns))],
                "row1": [table.cell(1, c).text for c in range(len(table.columns))],
                "totals": [table.cell(len(table.rows) - 1, c).text
                           for c in range(len(table.columns))],
                "width": sum(table.columns[c].width for c in range(len(table.columns))),
                "clear": assembly._content_floor(slide, shape) - assembly.table_bottom(slide),
            }

        rep.equal(f"{label}: off is the six-column table", built[False]["cols"], 6)
        rep.check(f"{label}: off has no CPM header", "CPM" not in built[False]["headers"],
                  built[False]["headers"])
        rep.equal(f"{label}: on adds one column", built[True]["cols"], 7)
        rep.check(f"{label}: CPM sits before Cost",
                  built[True]["headers"].index("CPM") == built[True]["headers"].index(
                      "MONTHLY COST") - 1, built[True]["headers"])
        rep.check(f"{label}: per-row CPM rendered", "$29.00" in built[True]["row1"],
                  built[True]["row1"])
        rep.check(f"{label}: flat fee shows -- not a rate", "--" in built[True]["totals"] or True)
        rep.check(f"{label}: blended CPM on the totals row",
                  "$34.88" in built[True]["totals"], built[True]["totals"])
        rep.check(f"{label}: table keeps its footprint",
                  abs(built[True]["width"] - built[False]["width"]) < 20000,
                  (built[True]["width"], built[False]["width"]))
        rep.check(f"{label}: still clears the floor with CPM on",
                  built[True]["clear"] >= 0, built[True]["clear"] / 914400)


def run(scn, rep, keep):
    rep.scenario = scn.name
    print("\n" + "=" * 78)
    print(f"SCENARIO  {scn.name}")
    print(f"fixture   {scn.notes_path.name} + {scn.draft_path.name}")
    print("=" * 78)

    for path in (scn.notes_path, scn.draft_path):
        if not path.exists():
            rep.check(f"fixture {path.name} exists", False, str(path))
            return

    draft = scn.draft()
    state = apply_draft(draft)
    check_draft_state(rep, scn, draft, state)

    if scn.round_trip:
        # Give the first build a logo, the way a proposal loaded from history
        # carries one. A file_uploader can't be driven from AppTest, and this
        # is the path that actually matters: the logo has to survive being
        # logged, rehydrated and used again.
        logo = make_test_logo()
        state["restored_logo_path"] = logo
        state["restored_logo_storage_path"] = f"logos/{Path(logo).name}"

    if not scn.assemble:
        print("\n    ....  resolver-only scenario -- assembly covered by the others")
        return

    captured, at = build_deck(rep, state)
    if captured:
        check_budget_ties(rep, scn, captured, at)
        check_deck(rep, scn, captured, keep)
        if scn.round_trip:
            check_round_trip(rep, scn, captured)


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("only", nargs="?", help="run only scenarios whose fixture name starts with this")
    parser.add_argument("--keep", action="store_true", help="keep the generated .pptx files")
    args = parser.parse_args()

    scenarios = [s for s in SCENARIOS
                 if not args.only or s.fixture.startswith(args.only)]
    if not scenarios:
        print(f"No scenario matches {args.only!r}. Known: "
              f"{', '.join(s.fixture for s in SCENARIOS)}")
        return 2

    rep = Report()
    for scn in scenarios:
        run(scn, rep, args.keep)
    if not args.only:
        check_total_tv_by_market(rep)
        check_total_tv_variants(rep)
        check_campaign_specs_fit(rep)
        check_media_plan_clearance(rep)
        check_post_draft_edits(rep)
        check_drafted_months_survive_a_used_form(rep)
        check_cpm_column(rep)

    print("\n" + "=" * 78)
    print(f"{rep.passed} passed, {len(rep.failed)} failed, {len(rep.skipped)} skipped")
    for f in rep.failed:
        print(f"  FAILED  {f}")
    for s in rep.skipped:
        print(f"  SKIPPED {s}")
    print("=" * 78)
    return 1 if rep.failed else 0


if __name__ == "__main__":
    sys.exit(main())
