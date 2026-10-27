"""The St. James attribution report review (2026-10-09) -- the report rules it
settled, each guarded here.

    python tests/test_attribution_report_rules.py

    1  The report period comes from the delivery file's flight dates, each
       flight shown and any dark gap named -- never the week-bucket labels.
    2  Recency bucket labels come from the source file; a drafted day range
       the export doesn't report is caught and removed.
    4  A drafted response claim about a dimension with no attribution data
       (daypart) is rejected after the draft.
    5  The ZIP slide's table and subtitle agree: only ZIPs at 1.2x the average
       or better on real volume, sorted by multiple; the heat map's title
       names the metric its legend shades by.
    8  Trend: the attributed rate by real-dated delivery week (dark weeks
       marked) and by flight; it qualifies on 3+ consecutive rises or a
       material flight-to-flight swing, then leads the one-sheet's bottom
       line and is always a highlight thread.
   10  One percentage per fact.
   11  Page tables show single pages with their own visitor counts, labelled
       "Visitors" -- never summed across pages, never a percentage.
   12  CPV: a tile on the highlights slide and the one-sheet with the toggle
       on; nowhere with it off. No Unique Visitor Rate tile.
   13  Recency and referral are shares only -- bar charts, a coverage
       footnote, and no count from either tab anywhere in the deck text.
   14  Page groups come from the goal's own entities plus confirmation pages;
       "Other" is never the top row.
   15  The traffic-mix table is a goal path, counts only.
   16  Confirmation pages are a standard highlight, worded "reached a
       confirmation page" -- never "converted" or "signed up".

Follow-ups (same day): page mapping (contact, party, kids' programs and
utility pages excluded; memberships are Pricing & purchase), the rep's
"Page grouping" overrides, no flow verbs between page-visit steps,
humanized creative names, per-location threads only with Compare
locations, the goal path as bars, one size for one-sheet takeaway heads.

Most checks use SYNTHETIC exports (never SKIP). Deck checks need the
gitignored report template; the real St. James pair ("Premion Website
Attribution and Reach Extension STJ.xlsx" / "Premion OTT STJ.xlsx") adds a
real-data pass and SKIPs without it.
"""
import copy
import sys
from datetime import date, timedelta
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

import db                                  # noqa: E402
db.fetch_audiences = lambda: (None, "stubbed for test isolation")

import market_lookup                       # noqa: E402
from pptx import Presentation              # noqa: E402

import app                                 # noqa: E402
import attribution_import as ai            # noqa: E402
import package_check                       # noqa: E402
import report_assembly as ra               # noqa: E402
import slide_map                           # noqa: E402
from attribution_import import AttributionRow, DateSeriesPoint  # noqa: E402

TEMPLATE = REPO / "REPORT_MASTER_v0_15.pptx"
STJ_ATTRIBUTION = REPO / "Premion Website Attribution and Reach Extension STJ.xlsx"
STJ_DELIVERY = REPO / "Premion OTT STJ.xlsx"
SCRATCH = REPO / "tests" / "_manual_output"
GOALS = ["Promote the Springfield and Bethesda performance clubs; the goal is new membership"]

market_lookup.install()


class Report:
    def __init__(self):
        self.passed, self.failed, self.skipped = 0, [], []

    def check(self, label, ok, detail=""):
        if ok:
            print(f"  PASS  {label}")
            self.passed += 1
        else:
            print(f"  FAIL  {label}{'  -- ' + str(detail) if detail else ''}")
            self.failed.append(label)

    def skip(self, reason):
        print(f"  SKIP  {reason}")
        self.skipped.append(reason)


# ---------------------------------------------------------------------------
# A synthetic two-flight campaign shaped like St. James: Aug 1-14 and Aug
# 23-Sep 5 2026, dark in between; 10,000 impressions a day; weekly buckets
# labelled the Monday after their first day (the export's own habit), the
# attributed rate rising every on-air week.
# ---------------------------------------------------------------------------

FLIGHT_1 = (date(2026, 8, 1), date(2026, 8, 14))
FLIGHT_2 = (date(2026, 8, 23), date(2026, 9, 5))
DAILY = 10_000


def _days(start, end):
    return [start + timedelta(days=i) for i in range((end - start).days + 1)]


def make_delivery(two_flights=True):
    days = _days(*FLIGHT_1) + (_days(*FLIGHT_2) if two_flights else [])
    flights = [{"start": FLIGHT_1[0], "end": FLIGHT_1[1], "impressions": 14 * DAILY,
                "uniques": 30_000, "booked": 140_000}]
    if two_flights:
        flights.append({"start": FLIGHT_2[0], "end": FLIGHT_2[1], "impressions": 14 * DAILY,
                        "uniques": 28_000, "booked": 140_000})
    return ai.DeliveryExport(
        delivered_impressions=DAILY * len(days), vcr=0.97, frequency=4.6, uniques=58_000,
        top_publishers=[("Hulu", 100_000, 0.36), ("Pluto TV", 80_000, 0.29)],
        channel_vcr={"Hulu": 0.98, "Pluto TV": 0.97},
        daily_delivery=[(d, DAILY) for d in days],
        ctv_impressions=DAILY * len(days),          # 100% CTV -> delivered vs. booked tile
        by_daypart=[("6AM - 9AM", 100_000), ("7PM - MID", 180_000)],
        by_geo=[("Springfield Zips", 200_000), ("Bethesda Zips", 80_000)],
        geo_vcr={"Springfield Zips": 0.97, "Bethesda Zips": 0.97},
        booked_impressions=280_000,
        flights=ai.merge_flights(flights),
        creative_flights=[
            {"name": "Main Spot", "start": FLIGHT_1[0], "end": FLIGHT_1[1], "impressions": 140_000},
            {"name": "Main Spot", "start": FLIGHT_2[0], "end": FLIGHT_2[1], "impressions": 136_000},
            {"name": "Anniversary", "start": FLIGHT_2[0], "end": FLIGHT_2[1], "impressions": 4_000}],
    )


WEEKS = [  # (label Monday, delivered, attributed rate) -- Sun-Sat buckets
    (date(2026, 7, 27), 1 * DAILY, 0.005),    # Jul 26-Aug 1: Aug 1 only
    (date(2026, 8, 3), 7 * DAILY, 0.010),     # Aug 2-8
    (date(2026, 8, 10), 6 * DAILY, 0.020),    # Aug 9-15: Aug 9-14
    (date(2026, 8, 24), 7 * DAILY, 0.030),    # Aug 23-29 (Aug 16-22 dark)
    (date(2026, 8, 31), 7 * DAILY, 0.040),    # Aug 30-Sep 5
]

URLS = {
    "https://club.example.com/springfield": 900,
    "https://www.example.com/": 600,
    "https://club.example.com/": 200,
    "https://marketing.example.com/bethesda-free-trial": 400,
    "https://club.example.com/bethesda": 350,
    "https://club.example.com/join": 220,
    "https://www.example.com/memberships/purchase": 120,
    "https://club.example.com/thank-you": 40,
    "https://www.example.com/thank-you/tour-scheduled": 25,
    "https://www.example.com/account/logout/success": 6,
    "https://marketing.example.com/hs/preferences-center/en/confirm": 9,
    "https://www.example.com/academy": 150,
}


def make_attribution():
    weekly = [DateSeriesPoint(day=d, delivered_impressions=n, attributed_impressions=round(n * r),
                              attributed_rate=r) for d, n, r in WEEKS]
    delivered = sum(n for _d, n, _r in WEEKS)
    attributed = sum(round(n * r) for _d, n, r in WEEKS)
    return ai.AttributionExport(
        client_name="Example Club",
        delivered_impressions=delivered, attributed_impressions=attributed,
        attributed_rate=attributed / delivered,
        attributed_unique_visitors=12_000, attributed_unique_visitor_rate=0.0429,
        by_audience=[AttributionRow("DEMO Age A25-64,HH Income 150K Plus", delivered, attributed,
                                    attributed / delivered)],
        by_market=[AttributionRow("WASHINGTON, DC (HAGRSTWN)", 263_000, attributed - 100,
                                  (attributed - 100) / 263_000),
                   AttributionRow("BALTIMORE", 17_000, 100, 100 / 17_000)],
        by_creative=[AttributionRow("Main Spot", 276_000, attributed - 400,
                                    (attributed - 400) / 276_000),
                     AttributionRow("Anniversary", 4_000, 400, 0.1)],
        by_zip=[AttributionRow("20783", 6_000, 1_000, 1_000 / 6_000),       # 2.1% share, high
                AttributionRow("20011", 8_000, 400, 0.05),                  # 2.9% share
                AttributionRow("20019", 9_000, 150, 150 / 9_000),           # below 1.2x
                AttributionRow("20010", 1_500, 300, 0.2),                   # high rate, 0.5% share
                AttributionRow("22407", 30_000, 400, 400 / 30_000)],
        by_url=dict(URLS),
        # Distinctive counts, so check 13 can prove none of them reaches the deck.
        by_recency={"00 - 03 DAYS": 1651, "04 - 07 DAYS": 1483, "08 - 11 DAYS": 1577,
                    "12 - 15 DAYS": 2003},
        by_referral_domain={"Direct": 531, "External Website Referral": 1571,
                            "Organic Search": 973, "Paid Search / Display": 4489, "Social": 37},
        by_device={"smartphone": 900, "desktop": 100},
        by_day_of_week=[AttributionRow(day, 40_000, round(40_000 * r), r) for day, r in
                        (("Mon", 0.025), ("Tue", 0.029), ("Wed", 0.029), ("Thu", 0.03),
                         ("Fri", 0.034), ("Sat", 0.02), ("Sun", 0.026))],
        weekly_trend=weekly,
        flight_start=date(2026, 7, 27), flight_end=date(2026, 8, 31),
    )


def _slide(prs, key):
    return next(s for s in prs.slides if slide_map.notes_key(s) == key)


def _slide_text(slide):
    parts = []
    for shape in slide_map.iter_all_shapes(slide.shapes):
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if shape.has_table:
            parts += [c.text_frame.text for r in shape.table.rows for c in r.cells]
    return "\n".join(parts)


def _table(slide, name):
    shape = next(sh for sh in slide.shapes if sh.name == name)
    return [[c.text_frame.text.strip() for c in r.cells] for r in shape.table.rows]


def build_deck(attribution, delivery, name, cost_per_visit=None, goals=None, **kwargs):
    SCRATCH.mkdir(exist_ok=True)
    out = SCRATCH / f"{name}.pptx"
    goals = goals or GOALS
    ra.build_report_deck(str(TEMPLATE), attribution, delivery, str(out), client_name="Example Club",
                         goals_bullets=goals, whats_next_bullets=["Keep going."],
                         cost_per_visit=cost_per_visit,
                         goal_keywords=ra.extract_goal_keywords(goals), **kwargs)
    return Presentation(out), out


def build_summary(attribution, delivery, name, threads=None, cost_per_visit=None):
    SCRATCH.mkdir(exist_ok=True)
    out = SCRATCH / f"{name}.pptx"
    ra.build_summary_slide(str(TEMPLATE), str(out), attribution=attribution, delivery=delivery,
                           client_name="Example Club", threads=threads or [],
                           accepted_optimizations=[], cost_per_visit=cost_per_visit, goals=GOALS)
    return Presentation(out), out


def _draft(**fields):
    draft = {"threads": [], "goal_alignment_notes": []}
    draft.update(fields)
    return draft


# ---------------------------------------------------------------------------
# 1 -- the report period
# ---------------------------------------------------------------------------

def check_period(rep):
    print("\n1. The report period comes from flight dates, flights and dark gap shown")
    attribution, delivery = make_attribution(), make_delivery()
    rep.check("two flights -> each shown, with the year",
             ra.report_period_label(attribution, delivery) == "Aug 1–14 · Aug 23–Sep 5, 2026",
             ra.report_period_label(attribution, delivery))
    rep.check("the dark gap is named", ra.dark_gap_note(delivery) == "Dark Aug 15–22",
             ra.dark_gap_note(delivery))
    rep.check("report_period spans the flights, not the week labels (Jul 27 - Aug 31)",
             ra.report_period(attribution, delivery) == (FLIGHT_1[0], FLIGHT_2[1]),
             ra.report_period(attribution, delivery))
    merged = ai.merge_flights([
        {"start": date(2026, 4, 1), "end": date(2026, 4, 30), "impressions": 10, "uniques": 1},
        {"start": date(2026, 5, 1), "end": date(2026, 5, 31), "impressions": 10, "uniques": 1},
        {"start": date(2026, 5, 1), "end": date(2026, 5, 31), "impressions": 5, "uniques": 1}])
    rep.check("contiguous month rows merge into one flight (no false dark gap)",
             len(merged) == 1 and merged[0]["impressions"] == 25, merged)
    rep.check("without a delivery file the export's own span is used",
             ra.report_period_label(attribution, None) == "Jul 27 - Aug 31, 2026",
             ra.report_period_label(attribution, None))
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- recap/summary period checks")
        return
    prs, _out = build_deck(attribution, delivery, "rules_period")
    recap = _slide_text(_slide(prs, "report:recap"))
    rep.check("the recap shows both flights", "Aug 1–14 · Aug 23–Sep 5, 2026" in recap, recap)
    rep.check("...and names the dark gap", "Dark Aug 15–22" in recap, recap)
    rep.check("...and never the week-label span", "Jul 27" not in recap, recap)
    summary, _o = build_summary(attribution, delivery, "rules_period_summary")
    text = _slide_text(summary.slides[0])
    rep.check("the one-sheet subtitle shows both flights",
             "Aug 1–14 · Aug 23–Sep 5, 2026" in text, text[:300])


# ---------------------------------------------------------------------------
# 2 -- recency bucket labels
# ---------------------------------------------------------------------------

def check_recency_labels(rep):
    print("\n2. Recency bucket labels come from the source file")
    rep.check("\"12 - 15 DAYS\" reads \"12–15 days\"",
             ra.recency_bucket_label("12 - 15 DAYS") == "12–15 days")
    rep.check("zero padding is dropped, the bounds kept",
             ra.recency_bucket_label("00 - 03 DAYS") == "0–3 days")
    attribution = make_attribution()
    facts = ra.build_facts_payload(attribution, make_delivery(), goals=GOALS)
    labels = [b["label"] for b in facts["response_profile"]["recency"]["buckets"]]
    rep.check("the model's facts carry exactly the source buckets",
             labels == ["0–3 days", "4–7 days", "8–11 days", "12–15 days"], labels)
    draft = _draft(response_profile_narrative=(
        "About 3 in 10 visits with timing data came 12-18 days after exposure. "
        "About 1 in 4 came within 0–3 days."))
    bad = app._recency_range_violations(draft, facts, attribution)
    rep.check("a drafted \"12-18 days\" (not a source bucket) is caught",
             [t for _f, t in bad] == ["12-18 days"], bad)
    cleaned = app._strip_rule_violations(copy.deepcopy(draft), facts, attribution)
    rep.check("...and its sentence removed, the real-bucket sentence kept",
             cleaned["response_profile_narrative"] == "About 1 in 4 came within 0–3 days.",
             cleaned["response_profile_narrative"])
    good = _draft(response_profile_narrative="About 3 in 10 responded 12–15 days after exposure.")
    rep.check("a source bucket range passes", not app._recency_range_violations(good, facts, attribution))


# ---------------------------------------------------------------------------
# 4 -- no claims about dimensions without attribution data
# ---------------------------------------------------------------------------

def check_unsupported_dimension_claims(rep):
    print("\n4. Post-draft fact check: no claim about a dimension with no attribution data")
    attribution = make_attribution()
    facts = ra.build_facts_payload(attribution, make_delivery(), goals=GOALS)
    draft = _draft(threads=[{
        "head": "Weight Toward Strong Areas", "anchor": "goal", "goal_ref": GOALS[0],
        "finding": "Springfield pages drew strong interest.",
        "meaning": "The clubs are resonating.",
        "action": "Continue Premion Streaming TV with targeting weighted toward the ZIP codes and "
                  "dayparts showing the highest attributed rates."}],
        delivery_narrative="Most impressions ran in the 7PM–midnight daypart.")
    hits = app._unsupported_dimension_violations(draft, facts)
    rep.check("the St. James daypart action is rejected",
             len(hits) == 1 and hits[0][2] == "daypart" and "dayparts" in hits[0][1], hits)
    rep.check("a delivery-only daypart sentence (no response claim) passes",
             all("7PM" not in sentence for _f, sentence, _d in hits), hits)
    enforced = app._enforce_draft_rules(copy.deepcopy(draft), facts, attribution)
    action = enforced["threads"][[t["head"] for t in enforced["threads"]].index(
        "Weight Toward Strong Areas")]["action"]
    rep.check("the claim never ships -- the action is removed after the retry", action is None, action)
    rep.check("...and the removal is named for review",
             any("no attribution data" in n for n in enforced.get("_review_notes") or []),
             enforced.get("_review_notes"))
    no_device = dict(facts, device=None)
    device_draft = _draft(zip_narrative="Smartphone visitors responded at the highest rate.")
    rep.check("device claims are rejected when the export has no device data",
             app._unsupported_dimension_violations(device_draft, no_device), None)
    rep.check("...and allowed when it does",
             not app._unsupported_dimension_violations(device_draft, facts), None)


# ---------------------------------------------------------------------------
# 5 -- the ZIP slide
# ---------------------------------------------------------------------------

def check_zip_slide(rep):
    print("\n5. ZIP slide: table and subtitle agree; heat-map title matches its legend")
    attribution = make_attribution()
    rows, _dropped, qualified = ra.zip_slide_rows(attribution)
    zips = [r["zip"] for r in rows]
    rep.check("only ZIPs at 1.2x+ on real volume, sorted by multiple",
             qualified and zips == ["20783", "20011"], zips)
    rep.check("a high-rate ZIP under the volume floor (20010, 0.5%) is left out",
             "20010" not in zips, zips)
    facts = ra.build_facts_payload(attribution, make_delivery(), goals=GOALS)
    rep.check("the model's zip rows are exactly the slide's rows",
             [r["zip"] for r in facts["zip"]["rows"]] == zips, facts["zip"]["rows"])
    flat = make_attribution()
    flat.by_zip = [AttributionRow("22407", 30_000, 300, 0.01)]
    _rows, _d, flat_qualified = ra.zip_slide_rows(flat)
    rep.check("with none qualifying, qualified is False (volume rows, subtitle says so)",
             not flat_qualified)
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- zip slide deck checks")
        return
    prs, _out = build_deck(attribution, make_delivery(), "rules_zip",
                           headline_notes={"zip": "ZIP codes above the campaign average."})
    slide = _slide(prs, "report:zip_analysis")
    table = _table(slide, "TopZipTable")
    rep.check("the table holds exactly the qualifying ZIPs, by multiple",
             [r[0] for r in table[1:]] == zips, table)
    text = _slide_text(slide)
    rep.check("the subtitle is Python's own, describing that set (a drafted note isn't used)",
             ra.ZIP_SUBTITLE_QUALIFIED in text and "above the campaign average." not in text, text)
    rep.check("the heat map's title names the attributed rate (its legend's metric)",
             ra.ZIP_MAP_HEADER in text and "ATTRIBUTED IMPRESSIONS" not in text, text)


# ---------------------------------------------------------------------------
# 8 -- trend detection
# ---------------------------------------------------------------------------

def check_trend(rep):
    print("\n8. Trend: real-dated weeks, dark weeks, flights; qualification and its uses")
    attribution, delivery = make_attribution(), make_delivery()
    weeks = ra.delivery_weeks(attribution, delivery)
    rep.check("weeks are labelled with the days they delivered on, the dark week marked",
             [(w["label"], w["dark"]) for w in weeks] == [
                 ("Aug 1", False), ("Aug 2–8", False), ("Aug 9–14", False),
                 ("Aug 15–22", True), ("Aug 23–29", False), ("Aug 30–Sep 5", False)],
             [(w["label"], w["dark"]) for w in weeks])
    trend = ra.trend_facts(attribution, delivery)
    rep.check("four consecutive rises -> qualifies on rising weeks",
             trend["consecutive_rises"] == 4 and "rising_weeks" in trend["reasons"], trend)
    flights = trend["flights"]
    rep.check("rate by flight comes from the weeks inside each flight",
             len(flights) == 2 and abs(flights[0]["attributed_rate"] - 1950 / 140_000) < 1e-9
             and abs(flights[1]["attributed_rate"] - 4900 / 140_000) < 1e-9, flights)
    rep.check("the flight swing is material, stated as a multiple",
             trend["flight_swing"]["material"] and trend["flight_swing"]["multiple"] == 2.5,
             trend["flight_swing"])
    sentence = ra.trend_sentence(trend)
    rep.check("the trend sentence carries one percentage",
             sentence and sentence.count("%") == 1, sentence)

    flat = make_attribution()
    flat.weekly_trend = [DateSeriesPoint(d, n, round(n * 0.02), 0.02) for d, n, _r in WEEKS]
    flat_trend = ra.trend_facts(flat, delivery)
    rep.check("a flat series doesn't qualify", not flat_trend["qualifies"], flat_trend["reasons"])

    one_flight = make_delivery(two_flights=False)
    short = make_attribution()
    short.weekly_trend = [DateSeriesPoint(d, n, round(n * r), r) for d, n, r in WEEKS[:3]]
    short_trend = ra.trend_facts(short, one_flight)
    rep.check("two rises in one flight doesn't qualify (needs 3+, or a flight swing)",
             not short_trend["qualifies"] and short_trend["consecutive_rises"] == 2, short_trend)

    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS)
    enforced = app._enforce_draft_rules(_draft(threads=[
        {"head": "Springfield Leads", "anchor": "goal", "goal_ref": GOALS[0],
         "finding": "900 visitors viewed the Springfield page.", "meaning": "Location pages work.",
         "action": None}]), facts, attribution)
    trend_threads = [t for t in enforced["threads"] if t.get("protected") and "week" in t["finding"]]
    rep.check("a qualifying trend the draft left out is added as a protected highlight thread",
             len(trend_threads) == 1, enforced["threads"])
    rep.check("...with the keep-it-continuous action (rising, across a dark gap)",
             trend_threads and "continuously" in (trend_threads[0].get("action") or ""),
             trend_threads)
    highlights, _t, whats_next = ra.distribute_threads(enforced["threads"])
    rep.check("the trend is a highlight and drives a What's Next item",
             any(h == trend_threads[0]["head"] for h, _f in highlights)
             and any("continuously" in w for w in whats_next), (highlights, whats_next))
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- one-sheet bottom line check")
        return
    summary, _o = build_summary(attribution, delivery, "rules_trend_summary",
                                threads=enforced["threads"])
    bottom = next(sh.text_frame.text for sh in summary.slides[0].shapes if sh.name == "BottomLine")
    rep.check("the one-sheet bottom line leads with the trend", bottom.startswith(sentence), bottom)


# ---------------------------------------------------------------------------
# 10 -- one percentage per fact
# ---------------------------------------------------------------------------

def check_percent_discipline(rep):
    print("\n10. At most one percentage per fact")
    double = _draft(threads=[{"head": "Bethesda Free Trial", "anchor": "goal", "meaning": "m",
                              "finding": "The Bethesda Free Trial page accounted for 10% of attributed "
                                         "page visits and reached 21.7% of attributed visitors."}],
                    url_intent_narrative="Springfield (23%) led, with Bethesda at 10% (10.3%).",
                    response_profile_narrative="About 3 in 10 (29.9%) responded 12–15 days out.")
    hits = [field for field, _s in app._double_percent_violations(double)]
    rep.check("share-of-visits + share-of-visitors for one page is caught", "thread 1 finding" in hits,
             hits)
    rep.check("a percentage repeated in parentheses is caught", "url_intent_narrative" in hits, hits)
    rep.check("a share phrase restated as a percentage is caught",
             "response_profile_narrative" in hits, hits)
    fine = _draft(threads=[{"head": "Intent", "anchor": "goal", "meaning": "m",
                            "finding": "Purchase pages drew 3.7% of page visits, and lead pages 2.0%."}],
                  attribution_narrative="Flight 2 ran at 4.04%, up from 1.49% in flight 1.")
    rep.check("two different items, or a from/to comparison, pass",
             not app._double_percent_violations(fine), app._double_percent_violations(fine))
    # A live St. James draft wrote Friday's 3.43% rate as "0.03%"; every sub-1
    # token rounded to "0" and "traced" to a 0% conversion rate.
    facts = ra.build_facts_payload(make_attribution(), make_delivery(), goals=GOALS)
    bad = app._attr_draft_number_violations(
        [("x", "Friday led at 0.03%, Fri at 3.40%, the campaign at 2.42%.")], facts)
    rep.check("a mis-scaled sub-1% rate doesn't trace to the facts",
             [t for _f, t in bad] == ["0.03%"], bad)
    rep.check("...while the real rates do", not app._attr_draft_number_violations(
        [("x", "Fri at 3.40%.")], facts))
    zip_double = _draft(zip_narrative="College Park (20783) posted the highest attributed rate at "
                                      "20.98% -- 7.59x the average -- on 1.8% of attributed "
                                      "impressions. Takoma Park ran at 10.70% (3.87x).")
    rep.check("a ZIP's rate plus its impression share in one sentence is caught (live draft)",
             [f for f, _s in app._double_percent_violations(zip_double)] == ["zip_narrative"],
             app._double_percent_violations(zip_double))
    spill = _draft(threads=[{"head": "DC Leads Markets", "anchor": "goal", "meaning": "m",
                             "finding": "Washington, DC ran at 2.88%, well above Baltimore at 0.92%."},
                            {"head": "Pages", "anchor": "goal", "finding": "900 visitors viewed "
                             "Springfield.", "meaning": "Location pages work."}])
    spill_facts = ra.build_facts_payload(make_attribution(), make_delivery(), goals=GOALS)
    rep.check("Baltimore (6% of impressions) is listed as spillover",
             spill_facts["market"]["spillover"] == ["Baltimore"], spill_facts["market"]["spillover"])
    stripped = app._strip_rule_violations(copy.deepcopy(spill), spill_facts)
    rep.check("a thread about a spillover market is dropped whole, the rest kept",
             [t["head"] for t in stripped["threads"]] == ["Pages"], stripped["threads"])
    shares = _draft(response_profile_narrative="TV prompted search: about 7 in 10 visits with "
                                               "referral data came from paid or organic search.")
    rep.check("\"about 7 in 10 visits\" is a share, not a count",
             not app._recency_referral_count_violations(shares))


# ---------------------------------------------------------------------------
# 11 / 14 / 15 -- pages, page groups, goal path
# ---------------------------------------------------------------------------

def check_pages_groups_path(rep):
    print("\n11/14/15. Single-page visitor counts, goal page groups, the goal path")
    attribution = make_attribution()
    rows = ra.top_url_rows(attribution, limit=8)
    labels = [r["label"] for r in rows]
    rep.check("two sites' homepages are two rows, never summed",
             "Homepage" in labels and any("homepage" in l and l != "Homepage" for l in labels)
             and not any(r["visits"] == "800" for r in rows), labels)
    terms = [t["label"] for t in ra.goal_terms(GOALS, None, attribution)]
    rep.check("goal terms are the goal's own entities the site has pages for",
             terms[:3] == ["Springfield", "Bethesda", "Performance Club"] or
             {"Springfield", "Bethesda"} <= set(terms), terms)
    groups = ra.goal_page_groups(attribution, GOALS)["classes"]
    rep.check("page groups lead with goal entities and include confirmation pages",
             groups[0]["label"] in ("Springfield", "Bethesda")
             and any(g["intent"] == "confirmation" for g in groups), [g["label"] for g in groups])
    rep.check("\"Other\" is never the top row (always last)",
             all(g["intent"] != "other" for g in groups[:-1]), [g["label"] for g in groups])
    rep.check("each group's pages carry their own visitor counts",
             all(isinstance(p["visitors"], int) for g in groups for p in g["pages"]), None)
    other_heavy = make_attribution()
    other_heavy.by_url = {"https://ex.com/members-area/benefits": 500, "https://ex.com/news": 200,
                          "https://ex.com/": 100}
    regrouped = ra.goal_page_groups(other_heavy, ["Grow membership"])["classes"]
    rep.check("with Other over 40%, pages are re-read against the goal terms",
             regrouped[0]["label"] == "Membership" and regrouped[0]["visits"] == 500,
             [(g["label"], g["visits"]) for g in regrouped])
    # URL redesign (2026-10-09), items 1-3: one combined path, steps by page
    # purpose, each step the SUM of its pages -- "page visits".
    path = ra.goal_path_combined(attribution, GOALS, delivery=make_delivery())
    rep.check("one combined path, grouped by page purpose regardless of location",
             [st["key"] for st in path] == ["explore", "signup", "purchase", "confirmation"]
             and [st["step"] for st in path][1:] == ["Free trial / Join", "Pricing & purchase",
                                                    "Confirmation pages"], path)
    rep.check("each step is the sum of its pages' attributed visitors",
             [st["page_visits"] for st in path] == [900 + 350 + 200, 400 + 220, 120, 40 + 25],
             [(st["step"], st["page_visits"]) for st in path])
    rep.check("Springfield and Bethesda pages share one row (no location split by default)",
             {pg["label"] for pg in path[0]["pages"]} >= {"Springfield", "Bethesda"},
             path[0]["pages"])
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- URL slide checks")
        return
    prs, _out = build_deck(attribution, make_delivery(), "rules_url")
    slide = _slide(prs, "report:url_report")
    # Follow-ups item 7: the combined path is horizontal bars, one per step,
    # the count at each bar's end -- no table unless locations are compared.
    shapes = {sh.name: sh for sh in slide.shapes}
    labels = [shapes[f"GoalPathLabel{i}"].text_frame.text for i in range(1, 5)
              if f"GoalPathLabel{i}" in shapes]
    counts = [shapes[f"GoalPathCount{i}"].text_frame.text for i in range(1, 5)
              if f"GoalPathCount{i}" in shapes]
    rep.check("one bar per step, labelled with the step, the count at the bar's end",
             labels == ["Location pages", "Free trial / Join", "Pricing & purchase",
                        "Confirmation pages"]
             and counts == ["1,450", "620", "120", "65"]
             and all(f"GoalPathBar{i}" in shapes for i in range(1, 5)), (labels, counts))
    bars = [shapes[f"GoalPathBar{i}"] for i in range(1, 5)]
    rep.check("bar lengths follow the counts, and each count sits at its bar's end",
             bars[0].width > bars[1].width > bars[2].width > bars[3].width
             and all(abs(int(shapes[f"GoalPathCount{i}"].left) - (int(b.left) + int(b.width)))
                     < 914400 * 0.2 for i, b in enumerate(bars, 1)), [b.width for b in bars])
    footnote = shapes.get("UrlReachFootnote")
    rep.check("the bars fill the left panel down to the footnote",
             footnote is not None and int(bars[-1].top) + int(bars[-1].height) > 914400 * 6
             and int(footnote.top) > int(bars[-1].top), None)
    rep.check("no table on the slide unless Compare locations is on",
             "IntentSummaryTable" not in shapes and "TopUrlTable" not in shapes, list(shapes))
    text = _slide_text(slide)
    rep.check("the page-visits footnote explains the sum", ra.PAGE_VISITS_FOOTNOTE in text, text)
    rep.check("no percentage anywhere in the bars", not any("%" in c for c in labels + counts),
             labels + counts)
    rep.check("the subtitle is the headline finding, in the reached form",
             "Attributed visitors reached every step of the membership path, including 65 "
             "confirmation-page visits." in text, text)
    on, _o = build_deck(attribution, make_delivery(), "rules_url_compare", compare_locations=True)
    compare = _table(_slide(on, "report:url_report"), "IntentSummaryTable")
    rep.check("Compare locations on -> a column per location beside the combined count",
             compare[0] == ["Step", "Page visits", "Springfield", "Bethesda"]
             and compare[1][2:] == ["900", "350"], compare)


# ---------------------------------------------------------------------------
# 12 -- CPV
# ---------------------------------------------------------------------------

def check_cpv(rep):
    print("\n12. CPV is a tile with the toggle on, absent with it off")
    attribution, delivery = make_attribution(), make_delivery()
    cpv = ra.compute_cost_per_visit(227_040.0, attribution.attributed_unique_visitors, 0, None)
    facts_off = ra.build_facts_payload(attribution, delivery, goals=GOALS, budget=20_000.0)
    rep.check("off: no cost per visitor reaches the model, even with a budget",
             facts_off["cost_per_visit"] is None
             and facts_off["budget"]["cost_per_attributed_visit"] is None, facts_off["budget"])
    leak = _draft(zip_narrative="Cost per site visitor came to $18.92. 20783 led.")
    stripped = app._strip_rule_violations(copy.deepcopy(leak), facts_off, attribution)
    rep.check("off: a drafted CPV sentence is removed", stripped["zip_narrative"] == "20783 led.",
             stripped["zip_narrative"])
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- CPV tile checks")
        return
    on_deck, _o = build_deck(attribution, delivery, "rules_cpv_on", cost_per_visit=cpv)
    off_deck, _o2 = build_deck(attribution, delivery, "rules_cpv_off")
    on_summary, _s = build_summary(attribution, delivery, "rules_cpv_on_summary", cost_per_visit=cpv)
    off_summary, _s2 = build_summary(attribution, delivery, "rules_cpv_off_summary")
    hi_on = _slide_text(_slide(on_deck, "report:highlights"))
    hi_off = _slide_text(_slide(off_deck, "report:highlights"))
    sum_on, sum_off = _slide_text(on_summary.slides[0]), _slide_text(off_summary.slides[0])
    rep.check("on: the highlights slide has the tile \"$18.92 / Cost per site visitor\"",
             "$18.92" in hi_on and ra.CPV_TILE_LABEL in hi_on, hi_on[:400])
    rep.check("on: the one-sheet has the same tile", "$18.92" in sum_on and ra.CPV_TILE_LABEL in sum_on,
             sum_on[:400])
    names_on = {sh.name for sh in _slide(on_deck, "report:highlights").shapes}
    rep.check("on: no CPV text line under the tiles", "CostPerVisitNote" not in names_on, names_on)
    rep.check("off: absent from the highlights slide",
             "$18.92" not in hi_off and ra.CPV_TILE_LABEL not in hi_off, hi_off[:400])
    rep.check("off: absent from the one-sheet",
             "$18.92" not in sum_off and ra.CPV_TILE_LABEL not in sum_off, sum_off[:400])
    rep.check("no Unique Visitor Rate tile on either one-sheet",
             "Unique Visitor Rate" not in sum_on + sum_off, None)
    for label, path in (("on deck", _o), ("on summary", _s)):
        rep.check(f"{label} passes package_check", package_check.check_package(str(path)) == [],
                 package_check.check_package(str(path)))


# ---------------------------------------------------------------------------
# 13 -- recency and referral are shares only
# ---------------------------------------------------------------------------

_RECENCY_REFERRAL_COUNTS = ("1651", "1,651", "1483", "1,483", "1577", "1,577", "2003", "2,003",
                            "531", "1571", "1,571", "973", "4489", "4,489", "37",
                            "6714", "6,714", "7601", "7,601")   # the two tabs' totals too


def check_recency_referral_shares(rep):
    print("\n13. Recency and referral: shares only, bar charts, a footnote, no counts")
    attribution = make_attribution()
    facts = ra.build_facts_payload(attribution, make_delivery(), goals=GOALS)
    profile = facts["response_profile"]
    has_counts = any("visitors" in b or "total" in profile[k]
                     for k in ("recency", "referral") for b in (profile[k].get("buckets")
                                                                or profile[k].get("sources")))
    rep.check("the model's facts carry no visitor counts from either tab", not has_counts, profile)
    rep.check("search share >= 50% -> the TV-prompted-search flag is set",
             profile["referral"]["search_prompted"], profile["referral"])
    counted = _draft(response_profile_narrative=(
        "Of the 764 visitors with referral data, most came from paid search. "
        "About 3 in 10 responded 12–15 days after exposure."))
    rep.check("a count in a referral sentence is caught",
             app._recency_referral_count_violations(counted), None)
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- response profile slide checks")
        return
    prs, out = build_deck(attribution, make_delivery(), "rules_response")
    slide = _slide(prs, "report:response_profile")
    names = {sh.name for sh in slide.shapes}
    text = _slide_text(slide)
    rep.check("the Source/Visitors/Share table is replaced by a chart",
             "ReferralTable" not in names
             and sum(1 for sh in slide.shapes if sh.shape_type == 13) >= 2, names)
    rep.check("the coverage footnote sits under the charts", ra.RECENCY_REFERRAL_FOOTNOTE in text, text)
    rep.check("the TV-prompted-search line is in the narrative", "TV prompted search" in text, text)
    deck_text = "\n".join(_slide_text(s) for s in prs.slides)
    summary, _o = build_summary(attribution, make_delivery(), "rules_response_summary")
    deck_text += "\n" + _slide_text(summary.slides[0])
    import re
    numbers = set(re.findall(r"\b\d[\d,]*\b", deck_text))
    leaked = [n for n in _RECENCY_REFERRAL_COUNTS if n in numbers]
    rep.check("no count from either tab appears anywhere in the deck or one-sheet text",
             not leaked, leaked)


# ---------------------------------------------------------------------------
# 16 -- confirmation pages
# ---------------------------------------------------------------------------

def check_confirmation_pages(rep):
    print("\n16. Confirmation pages: a standard highlight, \"reached a confirmation page\"")
    attribution = make_attribution()
    pages = ra.confirmation_pages(attribution)
    rep.check("thank-you and tour-scheduled pages count; logout 'success' and preferences "
             "'confirm' don't", [p["visitors"] for p in pages] == [40, 25], pages)
    facts = ra.build_facts_payload(attribution, make_delivery(), goals=GOALS)
    enforced = app._enforce_draft_rules(_draft(threads=[]), facts, attribution)
    confirm = [t for t in enforced["threads"] if t.get("kind") == "confirmation"]
    rep.check("missing from the draft -> added as a protected highlight thread",
             len(confirm) == 1 and confirm[0]["protected"], enforced["threads"])
    rep.check("...leading with the confirmation-page visits total, each page with its own "
             "visitors (the labeling rule)",
             confirm and confirm[0]["finding"].startswith("65 confirmation-page visits")
             and "40 visitors" in confirm[0]["finding"] and "25" in confirm[0]["finding"], confirm)
    stripped_meaning = _draft(threads=[{
        "head": "Confirmation Pages Reached", "anchor": "goal", "goal_ref": GOALS[0],
        "finding": "40 visitors reached the Thank You confirmation page.",
        "meaning": "Bethesda is the more active conversion funnel."}])
    kept = app._enforce_draft_rules(copy.deepcopy(stripped_meaning), facts, attribution)
    confirm_thread = next(t for t in kept["threads"] if t.get("kind") == "confirmation")
    rep.check("a confirmation thread whose meaning a rule removed gets a specific one back "
             "(a protected thread never silently drops out)",
             confirm_thread["meaning"] and "conversion" not in confirm_thread["meaning"],
             confirm_thread)
    wording = _draft(threads=[{"head": "Sign-ups", "anchor": "goal", "meaning": "Members signed up.",
                               "finding": "40 visitors converted on the thank-you page."}])
    rep.check("\"converted\" / \"signed up\" are caught with no conversion data",
             len(app._conversion_word_violations(wording, facts)) == 2,
             app._conversion_word_violations(wording, facts))
    cleaned = app._enforce_draft_rules(copy.deepcopy(wording), facts, attribution)
    text = " ".join(str(t.get(k) or "") for t in cleaned["threads"] for k in ("finding", "meaning"))
    rep.check("...and never ship", "converted" not in text and "signed up" not in text, text)


# ---------------------------------------------------------------------------
# Round 2 (St. James review, 2026-10-09): 2, 4, 6, 9, 12, 14
# ---------------------------------------------------------------------------

def check_labeling_rule(rep):
    print("\nURL redesign 2. One page's count is visitors; any sum is page visits")
    attribution, delivery = make_attribution(), make_delivery()
    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS)
    confirm_step = next(st for st in facts["goal_path"] if st["step"] == "Confirmation pages")
    rep.check("the model gets each step's sum and its single pages",
             confirm_step["page_visits"] == 65 and confirm_step["top_pages"][0]["visitors"] == 40,
             confirm_step)
    draft = _draft(threads=[{"head": "Confirmations", "anchor": "goal", "meaning": "Bethesda led.",
                             "finding": "65 visitors reached a confirmation page; 40 visitors "
                                        "reached the Thank You page."}],
                   url_intent_narrative="1,450 visitors explored the location pages.")
    fixed = app._fix_group_count_labels(copy.deepcopy(draft), facts)
    rep.check("a summed confirmation step reads \"confirmation-page visits\"",
             fixed["threads"][0]["finding"].startswith("65 confirmation-page visits;"),
             fixed["threads"][0]["finding"])
    rep.check("...a single page's count stays \"visitors\"",
             "40 visitors reached the Thank You page" in fixed["threads"][0]["finding"],
             fixed["threads"][0]["finding"])
    rep.check("any other summed step reads \"page visits\"",
             fixed["url_intent_narrative"] == "1,450 page visits explored the location pages.",
             fixed["url_intent_narrative"])
    enforced = app._enforce_draft_rules(_draft(threads=[]), facts, attribution)
    confirm = next(t for t in enforced["threads"] if t.get("kind") == "confirmation")
    rep.check("the confirmation highlight leads with the confirmation-page visits total",
             confirm["finding"].startswith("65 confirmation-page visits"), confirm["finding"])
    filler = _draft(threads=[{"head": "Momentum", "anchor": "goal", "kind": "trend",
                              "protected": True, "finding": "The rate rose.",
                              "meaning": "Sustained exposure is compounding."}])
    refilled = app._enforce_draft_rules(copy.deepcopy(filler), facts, attribution)
    trend_thread = next(t for t in refilled["threads"] if t.get("kind") == "trend")
    rep.check("a story thread's filler takeaway is replaced by its own specific one, never "
             "dropped (a live draft lost the trend from Takeaways this way)",
             trend_thread["meaning"] and "Aug" in trend_thread["meaning"], trend_thread["meaning"])
    timed = _draft(url_headline_note="Visitors spent the most time on the Springfield pages.")
    rep.check("time-on-site claims are rejected (the export counts page visits only)",
             app._time_on_site_violations(timed), None)
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- one-sheet check")
        return
    summary, _o = build_summary(attribution, delivery, "rules_labels_summary",
                                threads=enforced["threads"])
    text = _slide_text(summary.slides[0])
    rep.check("the one-sheet sidebar shows the combined path in page visits",
             "Goal path (page visits)" in text and "1,450" in text and "65" in text, text[:600])


def check_no_page_visit_shares(rep):
    print("\nR2-4. No \"% of attributed page visits\" -- in highlights, takeaways or the one-sheet")
    attribution, delivery = make_attribution(), make_delivery()
    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS)
    rep.check("page groups give the model counts, never a page-visit share",
             all("visit_share" not in c and "visits" not in c for c in facts["intent"]["classes"]),
             facts["intent"]["classes"][:2])
    draft = _draft(threads=[
        {"head": "Both Clubs Engaging", "anchor": "goal", "goal_ref": GOALS[0],
         "finding": "28% of attributed page visits went to Springfield pages. 900 visitors viewed "
                    "the Springfield page.", "meaning": "Springfield leads exploring."}],
        url_intent_narrative=("Springfield and Bethesda pages captured more than half of all "
                              "attributed page visits. 400 visitors started the Bethesda trial."))
    enforced = app._enforce_draft_rules(copy.deepcopy(draft), facts, attribution)
    thread = next(t for t in enforced["threads"] if t["head"] == "Both Clubs Engaging")
    rep.check("the percentage sentence is removed from the finding, the count kept",
             thread["finding"] == "900 visitors viewed the Springfield page.", thread["finding"])
    rep.check("\"more than half of page visits\" is removed from the narrative",
             enforced["url_intent_narrative"] == "400 visitors started the Bethesda trial.",
             enforced["url_intent_narrative"])
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- one-sheet check")
        return
    summary, _o = build_summary(attribution, delivery, "rules_no_page_share",
                                threads=enforced["threads"])
    text = _slide_text(summary.slides[0])
    rep.check("the one-sheet carries no page-visit percentage",
             not app._PAGE_SHARE_CLAIM_RE.search(text), text[:500])


def check_no_repeated_stats(rep):
    print("\nR2-6. No figure repeats across highlight bullets")
    threads = [
        {"head": "Confirmations", "anchor": "goal", "meaning": "Bethesda led completions.",
         "finding": "Visitors reached a confirmation page: 40 on Thank You and 25 on Tour Scheduled."},
        {"head": "Funnel", "anchor": "goal", "meaning": "The Springfield path is full.",
         "finding": "120 visitors reached the purchase page and 40 reached the Thank You page, with "
                    "an additional 25 on Tour Scheduled."},
        {"head": "Echo", "anchor": "goal", "meaning": "Springfield again.",
         "finding": "40 visitors reached the Thank You page."}]
    highlights, takeaways, _wn = ra.distribute_threads(threads)
    by_head = dict(highlights)
    rep.check("the later bullet keeps only its new figure",
             by_head.get("Funnel") == "120 visitors reached the purchase page.", highlights)
    rep.check("a bullet with nothing new is dropped from Highlights (its takeaway stays)",
             "Echo" not in by_head and any(m == "Springfield again." for _h, m in takeaways),
             (highlights, takeaways))
    seen, repeats = set(), []
    for _h, finding in highlights:
        stats = ra._stats(finding)
        repeats += sorted(stats & seen)
        seen |= stats
    rep.check("no figure appears in two bullets", not repeats, repeats)
    rep.check("dates aren't treated as figures (Aug 30 isn't 30)",
             ra._stats("Aug 1 to Aug 30–Sep 5 at 4.46%") == {"4.46%"})


def check_zip_action_names_zips(rep):
    print("\nR2-9. The ZIP What's Next item names the high-response ZIPs")
    attribution = make_attribution()
    facts = ra.build_facts_payload(attribution, make_delivery(), goals=GOALS)
    vague = _draft(threads=[{"head": "Clubs", "anchor": "goal", "finding": "f", "meaning": "Springfield.",
                             "action": "Concentrate weight in the ZIPs surrounding each club."}])
    fixed = app._fix_zip_actions(copy.deepcopy(vague), facts)
    action = fixed["threads"][0]["action"]
    rep.check("a vague ZIP action is replaced by one naming the qualifying ZIPs/areas",
             "20783" in action and "20011" in action and "surrounding" not in action, action)
    named = _draft(threads=[{"head": "Z", "anchor": "goal", "finding": "f", "meaning": "m",
                             "action": "Shift weight toward 20783."}])
    rep.check("an action already naming a qualifying ZIP is kept",
             app._fix_zip_actions(copy.deepcopy(named), facts)["threads"][0]["action"]
             == "Shift weight toward 20783.")
    tied = _draft(threads=[{"head": "Z", "anchor": "goal", "finding": "f", "meaning": "m",
                            "action": "Shift weight toward Springfield zip codes, targeting "
                                      "College Park (20783) and Takoma Park (20011)."}])
    retied = app._fix_zip_actions(copy.deepcopy(tied), facts)["threads"][0]["action"]
    rep.check("a ZIP action tying ZIPs to a location (nothing in the data says so) is replaced",
             "Springfield" not in retied and "20783" in retied, retied)
    flat = make_attribution()
    flat.by_zip = [AttributionRow("22407", 30_000, 300, 0.01)]
    flat_facts = ra.build_facts_payload(flat, make_delivery(), goals=GOALS)
    rep.check("with no qualifying ZIP, a ZIP action is removed rather than guessed",
             app._fix_zip_actions(copy.deepcopy(vague), flat_facts)["threads"][0]["action"] is None)


def check_dark_week_label(rep):
    print("\nR2-12. The dark week carries the real gap dates")
    attribution, delivery = make_attribution(), make_delivery()
    dark = [w for w in ra.delivery_weeks(attribution, delivery) if w["dark"]]
    rep.check("Aug 15–22 (the day after flight 1 to the day before flight 2), not Aug 16–22",
             [w["label"] for w in dark] == ["Aug 15–22"], [w["label"] for w in dark])
    facts = ra.trend_facts(attribution, delivery)
    rep.check("the trend facts and the chart use the same label",
             [w["label"] for w in facts["weeks"] if w["dark"]] == ["Aug 15–22"]
             and facts["dark_gaps"] == ["Aug 15–22"], facts["weeks"])


def check_creative_line(rep):
    print("\nR2-14. Creative comparison always earns a line, plus a rotate-evenly test")
    attribution, delivery = make_attribution(), make_delivery()
    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS)
    rep.check("the breakdown slide is dropped for this campaign (so the line must come from a thread)",
             not facts["breakdown"]["applies"])
    rep.check("the comparison qualifies for a test",
             facts["creative_comparison"]["qualifies_for_test"], facts["creative_comparison"])
    enforced = app._enforce_draft_rules(_draft(threads=[]), facts, attribution)
    line = [t for t in enforced["threads"] if "Anniversary" in (t.get("finding") or "")]
    rep.check("a creative thread is added when the draft leaves it out",
             len(line) == 1 and line[0].get("protected"), enforced["threads"])
    rep.check("...stating the shared dates and both rates, flagged directional",
             line and "Aug 23–Sep 5" in line[0]["finding"] and "10.00%" in line[0]["finding"]
             and "directional" in line[0]["finding"], line)
    rep.check("...with a rotate-evenly test as its What's Next action",
             line and "Rotate" in (line[0].get("action") or "") and "evenly" in line[0]["action"], line)
    again = app._enforce_draft_rules(copy.deepcopy(enforced), facts, attribution)
    rep.check("the post-draft rules are idempotent -- a processed draft gains no second "
             "creative, trend or confirmation thread",
             [t.get("kind") for t in again["threads"]] == [t.get("kind") for t in enforced["threads"]],
             ([t.get("kind") for t in enforced["threads"]], [t.get("kind") for t in again["threads"]]))
    highlights, takeaways, whats_next = ra.distribute_threads(enforced["threads"])
    rep.check("it reaches the deck: a highlight or takeaway line, and What's Next",
             (any("Anniversary" in f for _h, f in highlights)
              or any("Anniversary" in m for _h, m in takeaways))
             and any("Rotate" in w for w in whats_next), (highlights, takeaways, whats_next))



# ---------------------------------------------------------------------------
# St. James follow-ups (2026-10-09): mapping, rep overrides, wording,
# creative names, thread selection, layout
# ---------------------------------------------------------------------------

MAPPING_URLS = {
    "https://www.example.com/contact-us": 36,
    "https://marketing.example.com/birthday-party-request": 1,
    "https://www.example.com/performance-club/laugh-learn-play": 14,
    "https://www.example.com/camps": 3,
    "https://www.example.com/academy-admissions": 52,
    "https://www.example.com/memberships": 53,
    "https://www.example.com/memberships/recovery": 16,
    "https://www.example.com/memberships/corporate": 1,
    "https://www.example.com/performance-club/personal-training": 23,
    "https://www.example.com/performance-club/group-classes": 29,
    "https://www.example.com/performance-club/recovery": 21,
    "https://www.example.com/performance-club/training": 104,
    "https://www.example.com/performance-club/group-class-policies": 1,
    "https://www.example.com/sports/ice-skating": 35,
}


def _mapping_attribution():
    attribution = make_attribution()
    attribution.by_url = {**URLS, **MAPPING_URLS}
    return attribution


def _steps_by_url(path):
    return {pg["url"]: st["key"] for st in path for pg in st["pages"]}


def check_page_mapping(rep):
    print("\nFollow-ups 1. Page mapping rules")
    attribution = _mapping_attribution()
    steps = _steps_by_url(ra.goal_path_combined(attribution, GOALS, delivery=make_delivery()))
    for url, why in (("https://www.example.com/contact-us", "Contact Us"),
                     ("https://marketing.example.com/birthday-party-request", "a party request"),
                     ("https://www.example.com/performance-club/laugh-learn-play",
                      "a kids' program, even under the club's own path"),
                     ("https://www.example.com/camps", "camps"),
                     ("https://www.example.com/academy-admissions", "the academy"),
                     ("https://www.example.com/performance-club/group-class-policies",
                      "a policy page (ambiguous)"),
                     ("https://www.example.com/sports/ice-skating",
                      "a page matching no goal term (ambiguous)")):
        rep.check(f"excluded: {why}", url not in steps, steps.get(url))
    for url in ("https://www.example.com/memberships", "https://www.example.com/memberships/recovery",
                "https://www.example.com/memberships/corporate"):
        rep.check(f"Memberships page -> Pricing & purchase: {url.split('.com')[1]}",
                 steps.get(url) == "purchase", steps.get(url))
    for url in ("https://www.example.com/performance-club/personal-training",
                "https://www.example.com/performance-club/group-classes",
                "https://www.example.com/performance-club/recovery",
                "https://www.example.com/performance-club/training"):
        rep.check(f"stays in the club pages: {url.split('.com')[1]}",
                 steps.get(url) == "explore", steps.get(url))
    rep.check("Tour Scheduled stays a confirmation page",
             steps.get("https://www.example.com/thank-you/tour-scheduled") == "confirmation",
             steps.get("https://www.example.com/thank-you/tour-scheduled"))
    rep.check("a membership page names the step \"Pricing & purchase\"",
             [st["step"] for st in ra.goal_path_combined(attribution, GOALS)
              if st["key"] == "purchase"] == ["Pricing & purchase"])


def check_page_overrides(rep):
    print("\nFollow-ups 2. The rep's page grouping wins over the rules")
    attribution = _mapping_attribution()
    contact, academy = "https://www.example.com/contact-us", "https://www.example.com/academy-admissions"
    join, thanks = "https://club.example.com/join", "https://club.example.com/thank-you"
    overrides = {contact: "signup", join: ra.PATH_STEP_EXCLUDED, thanks: "explore"}
    path = ra.goal_path_combined(attribution, GOALS, delivery=make_delivery(), overrides=overrides)
    steps = _steps_by_url(path)
    rep.check("an excluded page moved into a step lands there", steps.get(contact) == "signup",
             steps.get(contact))
    rep.check("a page moved to Excluded leaves the path", join not in steps, steps.get(join))
    rep.check("an override beats even the confirmation rule", steps.get(thanks) == "explore",
             steps.get(thanks))
    rep.check("pages the rep didn't touch keep their rule", academy not in steps)
    rep.check("the step totals follow the overrides",
             {st["key"]: st["page_visits"] for st in path}["signup"] == 400 + 36,
             [(st["key"], st["page_visits"]) for st in path])
    confirmations = ra.confirmation_pages(attribution, GOALS, None, None, overrides)
    rep.check("the confirmation pages follow the overrides too",
             [c["label"] for c in confirmations] == ["Thank You / Tour Scheduled"], confirmations)
    facts = ra.build_facts_payload(attribution, make_delivery(), goals=GOALS,
                                   page_overrides=overrides)
    rep.check("the facts the model drafts from carry the overridden path",
             {st["step"]: st["page_visits"] for st in facts["goal_path"]}.get("Confirmation pages") == 25,
             facts["goal_path"])
    # The expander: a dropdown change becomes an override; picking the rule's
    # own step back removes it.
    rows = ra.page_rows(attribution)
    contact_index = next(i for i, r in enumerate(rows) if r["url"] == contact)
    for chosen, expected in (("Sign-up / free trial", {contact: "signup"}), ("Excluded", {})):
        stub = _EditorStub(lambda df, i=contact_index, c=chosen: df.assign(
            Step=[c if j == i else v for j, v in enumerate(df["Step"])]))
        real = app.st
        app.st = stub
        try:
            got = app._render_page_grouping(attribution, "\n".join(GOALS), "", None)
        finally:
            app.st = real
        rep.check(f"Page grouping: choosing \"{chosen}\" for Contact Us -> {expected}",
                 got == expected and stub.session_state["attr_page_overrides"] == expected, got)
        rep.check("...collapsed by default, every page listed with its step (incl. Excluded)",
                 stub.expanded is False and len(stub.frame) == len(rows)
                 and "Excluded" in set(stub.frame["Step"]), stub.expanded)
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- override deck check")
        return
    prs, _out = build_deck(attribution, make_delivery(), "rules_overrides", page_overrides=overrides)
    shapes = {sh.name: sh for sh in _slide(prs, "report:url_report").shapes}
    counts = [shapes[f"GoalPathCount{i}"].text_frame.text for i in range(1, 6)
              if f"GoalPathCount{i}" in shapes]
    rep.check("the slide's bars use the overridden totals", "436" in counts and "25" in counts,
             counts)


class _EditorStub:
    """Just enough of `st` for `_render_page_grouping`: an expander, a
    caption and a data_editor that applies one edit."""

    def __init__(self, edit):
        self.edit, self.session_state, self.expanded, self.frame = edit, {}, None, None
        self.column_config = app.st.column_config

    def expander(self, _label, expanded=True):
        self.expanded = expanded
        stub = self

        class _Ctx:
            def __enter__(self):
                return stub

            def __exit__(self, *exc):
                return False
        return _Ctx()

    def caption(self, *_a, **_k):
        return None

    def data_editor(self, frame, **_kwargs):
        self.frame = frame
        return self.edit(frame.copy())


def check_flow_wording(rep):
    print("\nFollow-ups 3. Page-visit steps never lead to, drive or move through another")
    flagged = ["2,894 page visits to club and location pages led to 114 confirmation-page visits.",
               "Visitors are moving from the club page toward pricing and trial pages.",
               "Bethesda visitors reached the free trial page and progressed to confirmation pages.",
               "Club page visits drove 1,050 page visits to the free trial and join pages.",
               "Attributed visitors are moving through the full funnel."]
    kept = ["Attributed visitors reached every step of the membership path, including 114 "
            "confirmation-page visits.",
            "Streaming TV drove 2,885 attributed site visitors.",
            "34 of the 114 confirmation-page visits were tour bookings."]
    hits = {s for _f, s in app._flow_verb_violations(_draft(url_intent_bullets=flagged + kept))}
    rep.check("causal/flow verbs between steps, and \"full funnel\", are caught",
             set(flagged) <= hits, set(flagged) - hits)
    rep.check("...a reached-form sentence, or a verb with no second step, is not",
             not (set(kept) & hits), set(kept) & hits)
    attribution, delivery = make_attribution(), make_delivery()
    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS)
    draft = _draft(url_headline_note=flagged[0], url_intent_bullets=[flagged[1], kept[2]],
                   threads=[{"head": "Confirmation Pages Reached", "anchor": "goal",
                             "finding": "65 confirmation-page visits.",
                             "meaning": "Attributed visitors are moving through the full funnel."}])
    enforced = app._enforce_draft_rules(draft, facts, attribution)
    text = " ".join([enforced.get("url_headline_note") or ""] + enforced["url_intent_bullets"]
                    + [str(t.get("meaning") or "") for t in enforced["threads"]])
    rep.check("after the draft: none survives in the subtitle, bullets or takeaways",
             not app._flow_verb_violations(enforced) and "funnel" not in text.lower(), text)
    rep.check("the confirmation takeaway is refilled from the facts, not left empty",
             all(str(t.get("meaning") or "").strip() for t in enforced["threads"]
                 if t.get("kind") == "confirmation"), enforced["threads"])
    rep.check("the subtitle falls back to the reached form",
             ra.goal_path_headline(ra.goal_path_combined(attribution, GOALS), ra.path_noun(GOALS))
             == "Attributed visitors reached every step of the membership path, including 65 "
                "confirmation-page visits.")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- one-sheet bottom-line check")
        return
    summary, _o = build_summary(attribution, delivery, "rules_flow_summary",
                                threads=enforced["threads"])
    bottom = next(sh.text_frame.text for sh in summary.slides[0].shapes if sh.name == "BottomLine")
    rep.check("the one-sheet bottom line carries no flow verb or funnel phrasing",
             not app._flow_verb_violations(_draft(x_narrative=bottom))
             and "funnel" not in bottom.lower(), bottom)


def check_creative_names(rep):
    print("\nFollow-ups 4. Creative names: humanized, rep-editable; the heading states the finding")
    rep.check("\"TSJV26302 TSJ Anniversary 30\" -> \"the Anniversary :30\"",
             ra.humanize_creative_name("TSJV26302 TSJ Anniversary 30") == "the Anniversary :30",
             ra.humanize_creative_name("TSJV26302 TSJ Anniversary 30"))
    rep.check("a code-only name has no human form (the caller names it)",
             ra.humanize_creative_name("TSJV26301_Rev") is None)
    names = ra.default_creative_names([("TSJV26301_Rev", 700_000),
                                       ("TSJV26302 TSJ Anniversary 30", 14_000),
                                       ("ABC12345", 5_000)])
    rep.check("defaults: the biggest code-only creative is \"the main spot\", another \"spot 2\"",
             names == {"TSJV26301_Rev": "the main spot",
                       "TSJV26302 TSJ Anniversary 30": "the Anniversary :30",
                       "ABC12345": "spot 2"}, names)
    rep.check("a leading trafficking code never reaches a table",
             ra.clean_creative_name("TSJV26302 TSJ Anniversary 30") == "TSJ Anniversary 30")
    attribution, delivery = make_attribution(), make_delivery()
    comparison = ra.creative_comparison(attribution, delivery,
                                        creative_names={"Anniversary": "the Birthday :15"})
    thread = ra.creative_comparison_thread(comparison)
    rep.check("the rep's own name wins, in the finding and the action",
             "the Birthday :15" in thread["finding"] and "the Birthday :15" in thread["action"],
             thread)
    rep.check("the heading states the finding, and \"directional\" stays in the detail",
             thread["head"] == "Birthday Spot Outperformed"
             and "directional" not in thread["head"].lower()
             and "directional" in thread["finding"], thread)
    rep.check("creative_heading: \"the Anniversary :30\" -> \"Anniversary Spot Outperformed\"",
             ra.creative_heading("the Anniversary :30") == "Anniversary Spot Outperformed")
    facts = {"creative_comparison": {
        "lifetime_rates_comparable": False, "window_label": "Aug 23–Sep 5", "directional": True,
        "creatives": [
            {"name": "TSJ Anniversary 30", "display": "the Anniversary :30", "impressions": 14_544,
             "attributed_rate": 0.087, "share_of_window": 0.02},
            {"name": "TSJV26301_Rev", "display": "the main spot", "impressions": 711_952,
             "attributed_rate": 0.0394, "share_of_window": 0.98}]}}
    draft = _draft(threads=[{"head": "Creative Comparison: Directional Signal", "kind": "creative",
                             "anchor": "signal",
                             "finding": "TSJV26302 TSJ Anniversary 30 drew 8.70% to "
                                        "TSJV26301_Rev's 3.94%.",
                             "meaning": "The Anniversary 30 creative showed a higher rate."}],
                   url_intent_narrative="TSJ Anniversary 30 ran in flight 2.")
    swapped = app._swap_creative_names(copy.deepcopy(draft), facts)
    text = " ".join([swapped["threads"][0]["finding"], swapped["threads"][0]["meaning"],
                     swapped["url_intent_narrative"]])
    rep.check("every raw, code-prefixed or short form becomes the display name",
             "TSJ" not in text and "Anniversary 30" not in text
             and text.count("Anniversary :30") == 3 and "the main spot's 3.94%" in text, text)
    rep.check("the drafted creative heading is replaced by the finding heading",
             swapped["threads"][0]["head"] == "Anniversary Spot Outperformed",
             swapped["threads"][0]["head"])
    # A live draft (2026-10-09) named the creative by its display name; the
    # story-thread check missed it and added a second creative thread.
    attribution, delivery = make_attribution(), make_delivery()
    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS)
    display = next(c["display"] for c in facts["creative_comparison"]["creatives"]
                   if c["impressions"] < 10_000)
    drafted = _draft(threads=[{"head": f"{display.title()} Outperformed Directionally",
                               "anchor": "signal",
                               "finding": f"Over Aug 23–Sep 5, {display} posted a 10.00% rate.",
                               "meaning": f"{display} is worth testing.", "action": None}])
    enforced = app._enforce_draft_rules(drafted, facts, attribution)
    creative = [t for t in enforced["threads"] if t.get("kind") == "creative"
                or display.lower() in str(t.get("finding") or "").lower()]
    rep.check("a drafted thread naming the creative by its display name IS the creative thread "
             "(no second one)", len(creative) == 1 and creative[0].get("kind") == "creative",
             [(t["head"], t.get("kind")) for t in enforced["threads"]])


def check_location_threads(rep):
    print("\nFollow-ups 6. Per-location threads only with Compare locations on")
    attribution, delivery = make_attribution(), make_delivery()
    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS)
    springfield = {"head": "Springfield Driving Strong Intent", "anchor": "goal",
                   "finding": "900 visitors viewed the Springfield page.",
                   "meaning": "Springfield is the busiest club page.", "action": None}
    bethesda = {"head": "Bethesda Free Trial Page Active", "anchor": "goal",
                "finding": "400 visitors reached the Bethesda Free Trial page.",
                "meaning": "Bethesda's free trial page is busy.", "action": None}
    zips = {"head": "Springfield-Area ZIPs Lead", "anchor": "signal",
            "finding": "College Park (20783) led the ZIPs.", "meaning": "ZIPs matter.",
            "action": None}
    draft = _draft(threads=[copy.deepcopy(springfield), copy.deepcopy(bethesda)])
    off = app._drop_location_threads(copy.deepcopy(draft), facts)
    heads = [t["head"] for t in off["threads"]]
    rep.check("off: per-location threads are removed", not any(
        "Springfield" in h or "Bethesda" in h for h in heads), heads)
    rep.check("...the first one's slot goes to the geography/ZIP finding",
             heads == ["Top ZIPs Outperformed"] and off["threads"][0]["kind"] == "geo"
             and "20783" in off["threads"][0]["finding"], off["threads"])
    kept_zip = app._drop_location_threads(_draft(threads=[copy.deepcopy(zips),
                                                          copy.deepcopy(springfield)]), facts)
    rep.check("a ZIP thread naming a location isn't a location thread (and no second ZIP thread)",
             [t["head"] for t in kept_zip["threads"]] == ["Springfield-Area ZIPs Lead"],
             [t["head"] for t in kept_zip["threads"]])
    both = {"head": "Both Clubs Drawing Visitors", "anchor": "goal",
            "takeaway_head": "Pages Engaged",
            "finding": "900 visitors viewed the Springfield page and 400 reached the Bethesda "
                       "Free Trial page.", "meaning": "Both clubs drew visitors.", "action": None}
    rep.check("a thread setting two locations' pages side by side is a location thread too "
             "(a live draft's \"Both Clubs Drawing Visitors\")",
             ra.is_location_thread(both, ["Springfield", "Bethesda"]))
    rep.check("...the confirmation story thread naming both clubs is not",
             not ra.is_location_thread({**both, "kind": "confirmation"},
                                       ["Springfield", "Bethesda"]))
    details = ["34 of the 114 confirmation-page visits were tour bookings (the Tour page).",
               "Join was the most-visited sign-up page, with 220 visitors."]
    topped = app._fill_url_bullets(
        _draft(url_intent_bullets=["34 of the 114 confirmation-page visits were tour bookings "
                                   "(the Tour page), and 29 reached the Bethesda thank-you page."]),
        {"goal_path_details": details})["url_intent_bullets"]
    rep.check("URL bullets are topped up with page specifics, never repeating a drafted one",
             len(topped) == 2 and topped[1] == details[1], topped)
    compare = ra.build_facts_payload(attribution, delivery, goals=GOALS, compare_locations=True)
    on = app._drop_location_threads(copy.deepcopy(draft), compare)
    rep.check("on: per-location threads stay",
             [t["head"] for t in on["threads"]] == [springfield["head"], bethesda["head"]],
             [t["head"] for t in on["threads"]])
    weak_facts = {"goal_path": [{"step": "Free trial / Join", "page_visits": 1011, "top_pages": [
        {"label": "Bethesda Free Trial", "visitors": 627}, {"label": "Join", "visitors": 352},
        {"label": "Springfield Free Trial", "visitors": 32}]}]}
    weak = _draft(threads=[{"head": "Free Trials", "anchor": "goal",
                            "finding": "32 visitors reached the Springfield Free Trial page, "
                                       "showing strong intent.",
                            "meaning": "627 visitors reached the Bethesda Free Trial page, the "
                                       "strongest sign-up page."}])
    hits = [s for _f, s in app._weak_evidence_violations(weak, weak_facts)]
    rep.check("a weak number cited as strength is caught; the step's leader isn't",
             len(hits) == 1 and hits[0].startswith("32 visitors"), hits)
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- one-sheet thread checks")
        return
    threads = [copy.deepcopy(springfield), copy.deepcopy(bethesda)]
    summary, _o = build_summary(attribution, delivery, "rules_location_off", threads=threads)
    heads_off = [sh.text_frame.text for sh in summary.slides[0].shapes
                 if sh.name.startswith("SummaryTakeaway") and sh.name.endswith("Head")]
    rep.check("one-sheet, off: takeaways use the combined path, no location takeaway",
             heads_off and heads_off[0] == "Every Membership Step Reached"
             and not any("Springfield" in h or "Bethesda" in h for h in heads_off), heads_off)
    out = SCRATCH / "rules_location_on.pptx"
    ra.build_summary_slide(str(TEMPLATE), str(out), attribution=attribution, delivery=delivery,
                           client_name="Example Club", threads=threads, accepted_optimizations=[],
                           goals=GOALS, compare_locations=True)
    heads_on = [sh.text_frame.text for sh in Presentation(out).slides[0].shapes
                if sh.name.startswith("SummaryTakeaway") and sh.name.endswith("Head")]
    # (Springfield leads the sidebar here, so it's never also a takeaway.)
    rep.check("one-sheet, on: the location takeaways stay, no combined-path stand-in",
             any("Bethesda" in h for h in heads_on)
             and "Every Membership Step Reached" not in heads_on, heads_on)


def check_takeaway_heading_sizes(rep):
    print("\nFollow-ups 8. One-sheet takeaway headings share one font size")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present")
        return
    threads = [{"head": "Short", "anchor": "goal", "finding": "900 visitors viewed Springfield.",
                "meaning": "x"},
               {"head": "A Much Longer Heading That Has To Wrap Twice", "anchor": "goal",
                "finding": "400 visitors reached the Bethesda Free Trial page.", "meaning": "y"},
               {"head": "Mid Length Heading", "anchor": "signal",
                "finding": "65 confirmation-page visits.", "meaning": "z"}]
    summary, _o = build_summary(make_attribution(), make_delivery(), "rules_head_sizes",
                                threads=threads)
    sizes = {sh.name: {r.font.size for p in sh.text_frame.paragraphs for r in p.runs}
             for sh in summary.slides[0].shapes
             if sh.name.startswith("SummaryTakeaway") and sh.name.endswith("Head")}
    rep.check("every takeaway heading renders at the same size",
             len({frozenset(v) for v in sizes.values()}) == 1, sizes)


# ---------------------------------------------------------------------------
# St. James review decisions (Matt's calls) -- 4 through 8, plus the relaxed
# rules 1-3 staying relaxed
# ---------------------------------------------------------------------------

def _zip_thread_draft(takeaway_head, meaning):
    return _draft(threads=[{
        "head": "Top ZIPs Outperformed", "anchor": "goal", "kind": "geo",
        "finding": "College Park (20783) and Takoma Park (20011) each drew an attributed rate at "
                   "2.04x the campaign average or better.",
        "meaning": meaning, "takeaway_head": takeaway_head,
        "action": "Shift Streaming TV weight toward the highest-response ZIPs — College Park "
                  "(20783) and Takoma Park (20011), each at 2.04x the campaign average or better."}])


def check_takeaways_not_actions(rep):
    print("\nDecisions 4. A takeaway is a conclusion; the action stays in What's Next")
    attribution, delivery = make_attribution(), make_delivery()
    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS)
    made = ra.zip_thread(facts["zip"])
    rep.check("Python's own ZIP thread: a meaning-level takeaway heading, not \"Weight Toward\"",
             made["takeaway_head"].startswith("Response Concentrated in")
             and not app._INSTRUCTION_HEAD_RE.match(made["takeaway_head"]), made["takeaway_head"])
    rep.check("...and its meaning interprets the cluster without recommending",
             "cluster in" in made["meaning"] and not app._RECOMMENDING_RE.search(made["meaning"]),
             made["meaning"])
    enforced = app._enforce_draft_rules(_zip_thread_draft(
        "Weight Toward Top ZIPs",
        "Response concentrated in College Park and a handful of nearby ZIPs, the clearest place "
        "to add weight."), facts, attribution)
    thread = next(t for t in enforced["threads"] if t.get("kind") == "geo")
    rep.check("a drafted action heading (\"Weight Toward Top ZIPs\") is replaced",
             thread["takeaway_head"].startswith("Response Concentrated in"), thread["takeaway_head"])
    rep.check("...the recommending meaning is replaced with the cluster reading",
             "add weight" not in thread["meaning"] and "cluster in" in thread["meaning"],
             thread["meaning"])
    rep.check("...the action itself still feeds What's Next",
             thread.get("action", "").startswith("Shift Streaming TV weight"), thread.get("action"))
    rep.check("...and the replacement is named for review",
             any("Weight Toward Top ZIPs" in n for n in enforced.get("_review_notes") or []))
    dup = app._enforce_draft_rules(_draft(threads=[{
        "head": "Anniversary Spot Outperformed", "anchor": "signal", "kind": "creative",
        "finding": "The Anniversary :30 ran at 10.00% against 2.25% -- directional.",
        "meaning": "The Anniversary :30 led over the shared dates.",
        "takeaway_head": "Rotate Creatives More Evenly",
        "action": "Rotate the two creatives more evenly in the next flight."}]), facts, attribution)
    creative = next(t for t in dup["threads"] if t.get("kind") == "creative")
    rep.check("a takeaway heading that repeats its What's Next item falls back to the head",
             creative["takeaway_head"] == "Anniversary Spot Outperformed", creative["takeaway_head"])
    kept = app._enforce_draft_rules(_draft(threads=[{
        "head": "Campaign Momentum Is Rising", "anchor": "goal", "kind": "trend",
        "finding": "The attributed rate rose four weeks in a row.",
        "meaning": "Flight 2's 2.7x lift shows the audience responding more strongly with "
                   "repeated exposure.",
        "takeaway_head": "Continuous Flights Sustain the Lift",
        "action": "Keep the next flight on air continuously."}]), facts, attribution)
    trend = next(t for t in kept["threads"] if t.get("kind") == "trend")
    rep.check("a conclusion heading is left alone (\"Continuous Flights Sustain the Lift\")",
             trend["takeaway_head"] == "Continuous Flights Sustain the Lift", trend["takeaway_head"])


def check_directional_headings(rep):
    print("\nDecisions 5. A directional result: confident heading, no overclaiming word")
    attribution, delivery = make_attribution(), make_delivery()
    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS)
    finding = ("Over the shared dates the Anniversary :30 ran at 10.00% against 2.25% for the "
               "main spot -- a directional comparison.")
    enforced = app._enforce_draft_rules(_draft(threads=[{
        "head": "Anniversary Spot Outperformed", "anchor": "signal", "kind": "creative",
        "finding": finding, "meaning": "The Anniversary :30 led over the shared dates.",
        "takeaway_head": "The Anniversary :30 Shows a Real Rate Advantage",
        "action": "Rotate the two creatives more evenly in the next flight."}]), facts, attribution)
    thread = next(t for t in enforced["threads"] if t.get("kind") == "creative")
    rep.check("\"Real\" is dropped from a directional takeaway heading",
             thread["takeaway_head"] == "The Anniversary :30 Shows a Rate Advantage",
             thread["takeaway_head"])
    rep.check("...the head stays confident (\"Anniversary Spot Outperformed\")",
             thread["head"] == "Anniversary Spot Outperformed", thread["head"])
    rep.check("...and \"directional\" stays in the detail", "directional" in thread["finding"],
             thread["finding"])
    for word in ("Proven", "Clear", "Conclusive"):
        test = {"head": f"A {word} Winner", "finding": "directional, on a small share",
                "takeaway_head": f"{word} Creative Advantage"}
        app._fix_directional_heads(_draft(threads=[test]), facts)
        rep.check(f"\"{word}\" is dropped from both headings",
                 word.lower() not in (test["head"] + test["takeaway_head"]).lower(), test)
    solid = {"head": "A Clear Winner", "finding": "Flight 2 ran at 2.7x flight 1's rate.",
             "takeaway_head": "Clear Momentum"}
    app._fix_directional_heads(_draft(threads=[solid]), {})
    rep.check("a result that isn't directional keeps its wording", solid["head"] == "A Clear Winner",
             solid)


def check_referral_chart_label(rep):
    print("\nDecisions 6. The referral chart keeps the source label; sentences say "
          "\"search or display\"")
    rep.check("the chart label is \"Paid search / display\"",
             ra.referral_chart_label("Paid Search / Display") == "Paid search / display")
    attribution, delivery = make_attribution(), make_delivery()
    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS)
    sources = [s["source"] for s in facts["response_profile"]["referral"]["sources"]]
    rep.check("the narrative facts still say \"Search or display\"", "Search or display" in sources,
             sources)
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- referral chart label")
        return
    seen, real = [], ra.report_charts.render_bar_chart

    def spy(labels, *args, **kwargs):
        seen.append(list(labels))
        return real(labels, *args, **kwargs)
    ra.report_charts.render_bar_chart = spy
    try:
        build_deck(attribution, delivery, "rules_referral_label")
    finally:
        ra.report_charts.render_bar_chart = real
    rep.check("the rendered referral chart's bar reads \"Paid search / display\"",
             any("Paid search / display" in labels for labels in seen), seen)


def check_zip_narrative(rep):
    print("\nDecisions 7. The ZIP narrative interprets the cluster -- at most 2 numbers")
    rows = [{"zip": "20783", "area": "College Park", "multiple": "7.59x"},
            {"zip": "20011", "area": "Takoma Park", "multiple": "3.87x"},
            {"zip": "20782", "area": "Hyattsville", "multiple": "3.73x"},
            {"zip": "20001", "area": "Washington", "multiple": "2.99x"},
            {"zip": "20002", "area": "Washington", "multiple": "1.82x"}]
    rep.check("the St. James ZIPs cluster in DC and Prince George's County (from the crosswalk)",
             ra.zip_cluster(rows) == "DC and Prince George's County", ra.zip_cluster(rows))
    rep.check("ZIPs spread across many counties make no cluster claim",
             ra.zip_cluster([{"zip": z} for z in ("20783", "22150", "21201", "19103", "10001")])
             is None)
    recap = ("College Park (20783) led all ZIP codes with a 20.98% attributed rate at 7.59x the "
             "campaign average, followed by Takoma Park (20011) at 10.70% (3.87x) and Hyattsville "
             "(20782) at 10.31% (3.73x).")
    facts = {"zip": {"rows": rows, "qualified": True}}
    out = app._fix_zip_narrative(_draft(zip_narrative=recap), facts)
    rep.check("a row-by-row recap is replaced by the cluster reading",
             out["zip_narrative"].startswith("The ZIPs responding above the campaign average "
                                             "cluster in DC and Prince George's County"),
             out["zip_narrative"])
    rep.check("...which carries at most 2 numbers (ZIP codes not counted)",
             len(app._zip_narrative_numbers(out["zip_narrative"])) <= 2, out["zip_narrative"])
    fine = ("Response clustered in close-in DC and Prince George's County, led by College Park "
            "(20783) at 7.59x the campaign average.")
    kept = app._fix_zip_narrative(_draft(zip_narrative=fine), facts)
    rep.check("an interpreting sentence with 1-2 numbers is kept", kept["zip_narrative"] == fine,
             kept["zip_narrative"])


def check_tile_rounding(rep):
    print("\nDecisions 8. A stat a tile shows reads at the tile's rounding everywhere")
    attribution, delivery = make_attribution(), make_delivery()
    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS)
    tiles = facts["tile_values"]
    rep.check("the facts carry the tiles as displayed", tiles.get("frequency") == "4.6"
             and tiles.get("attributed_rate") == "2.45%", tiles)
    draft = _draft(delivery_narrative="The campaign reached households at an average frequency "
                                      "of 4.57 per household.",
                   attribution_narrative="The audience ran at a 2.4512% attributed rate.",
                   zip_narrative="College Park ran at 4.57x the average.")
    rate = attribution.attributed_rate * 100
    draft["attribution_narrative"] = f"The audience ran at a {rate:.4f}% attributed rate."
    out = app._match_tile_rounding(draft, facts)
    rep.check("frequency 4.57 reads 4.6, as on the tile",
             "frequency of 4.6 per household" in out["delivery_narrative"], out["delivery_narrative"])
    rep.check("an over-precise attributed rate reads as the tile's 2.45%",
             "2.45% attributed rate" in out["attribution_narrative"], out["attribution_narrative"])
    rep.check("...a different figure that isn't a tile stat (\"4.57x\") is left alone",
             "4.57x" in out["zip_narrative"], out["zip_narrative"])


def check_one_sheet_path_dedupe(rep):
    print("\nOne-sheet: a drafted whole-path thread doesn't repeat the built-in path takeaway")
    rep.check("a whole-path thread is recognized", ra.is_path_thread(
        {"head": "Membership Path Reached at Every Step"}))
    rep.check("...a story thread or a ZIP thread is not",
             not ra.is_path_thread({"head": "Confirmation Pages Reached", "kind": "confirmation"})
             and not ra.is_path_thread({"head": "Top ZIPs Outperformed"}))
    rep.check("a word-capped heading never ends on a dangling word",
             ra._cap_words("Membership Path Reached at Every Step") == "Membership Path Reached")
    if not TEMPLATE.exists():
        rep.skip(f"{TEMPLATE.name} not present -- one-sheet path dedupe")
        return
    attribution, delivery = make_attribution(), make_delivery()
    threads = [{"head": "Membership Path Reached at Every Step", "anchor": "goal",
                "finding": "Attributed visitors reached every step of the membership path.",
                "meaning": "Visitors reached the club, trial, pricing and confirmation pages.",
                "takeaway_head": "Visitors Reached Every Membership Step"}]
    summary, _o = build_summary(attribution, delivery, "rules_path_dedupe", threads=threads)
    heads = [sh.text_frame.text for sh in summary.slides[0].shapes
             if sh.name.startswith("SummaryTakeaway") and sh.name.endswith("Head")]
    path_heads = [h for h in heads if ra.is_path_thread({"head": h})
                  or h.startswith("Every Membership Step")]
    rep.check("the one-sheet carries one path takeaway, not two", len(path_heads) == 1, heads)


def check_relaxed_rules(rep):
    print("\nDecisions 1-3. Relaxed: both-clubs claims, interpretation, day-of-week wording")
    attribution, delivery = make_attribution(), make_delivery()
    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS)
    flat = dict(facts)
    flat["response_profile"] = {**facts["response_profile"], "day_of_week": {
        **(facts["response_profile"].get("day_of_week") or {}), "material_swing": False}}
    narrative = ("TV prompted search: about 6 in 10 visits with referral data came from search or "
                 "display, the campaign moving viewers to look up the clubs. Friday led at 3.40%, "
                 "a meaningful swing across the week.")
    meaning = ("Flight 2's lift shows the audience responding more strongly with repeated "
               "exposure, reaching a confirmation page at both clubs.")
    out = app._enforce_draft_rules(_draft(response_profile_narrative=narrative, threads=[{
        "head": "Campaign Momentum Is Rising", "anchor": "goal", "kind": "trend",
        "finding": "Flight 2 ran at 2.7x flight 1's rate.", "meaning": meaning,
        "takeaway_head": "Momentum Built Across Flights"}]), flat, attribution)
    rep.check("day-of-week \"meaningful swing\" and \"moving viewers to look up\" survive",
             "meaningful swing" in (out.get("response_profile_narrative") or "")
             and "look up the clubs" in (out.get("response_profile_narrative") or ""),
             out.get("response_profile_narrative"))
    trend = next(t for t in out["threads"] if t.get("kind") == "trend")
    rep.check("\"repeated exposure\" and \"at both clubs\" survive in a takeaway",
             "repeated exposure" in str(trend.get("meaning")) and "both clubs" in str(
                 trend.get("meaning")), trend.get("meaning"))
    rep.check("...while a flow verb between page-visit steps is still caught",
             bool(app._flow_verb_violations(_draft(x_narrative=(
                 "2,894 club-page visits led to 114 confirmation-page visits.")))))


# ---------------------------------------------------------------------------
# The real St. James pair
# ---------------------------------------------------------------------------

def check_st_james(rep):
    print("\nReal St. James export (Aug 2026)")
    if not (STJ_ATTRIBUTION.exists() and STJ_DELIVERY.exists()):
        rep.skip("St. James fixtures not present")
        return
    attribution = ai.parse_attribution_export(str(STJ_ATTRIBUTION))
    delivery = ai.parse_delivery_export(str(STJ_DELIVERY))
    rep.check("period: Aug 1–14 · Aug 23–Sep 5, 2026",
             ra.report_period_label(attribution, delivery) == "Aug 1–14 · Aug 23–Sep 5, 2026",
             ra.report_period_label(attribution, delivery))
    trend = ra.trend_facts(attribution, delivery)
    rep.check("trend qualifies on both counts (4 rises; flight 2 at 2.7x)",
             trend["consecutive_rises"] == 4 and trend["flight_swing"]["multiple"] == 2.7, trend)
    comparison = ra.creative_comparison(attribution, delivery)
    rates = {c["display"]: round(c["attributed_rate"] * 100, 1) for c in comparison["creatives"]}
    rep.check("creatives compared over flight 2 only: the Anniversary :30 8.7% vs the main spot "
             "3.9%, directional",
             rates.get("the Anniversary :30") == 8.7 and rates.get("the main spot") == 3.9
             and comparison["directional"], comparison)
    rep.check("Baltimore (6% spillover) isn't a breakdown -- the slide drops",
             not ra.attribution_breakdown_applies(attribution, delivery))
    rep.check("bought geos label the Geography tile",
             ra.bought_geo_labels(delivery) == ["Springfield ZIPs", "Bethesda ZIPs"],
             ra.bought_geo_labels(delivery))
    rep.check("the recency tab's own buckets: 12–15 days, never 12–18",
             "12–15 days" in [ra.recency_bucket_label(b) for b in attribution.by_recency])
    path = ra.goal_path_combined(attribution, GOALS, vertical="healthcare", delivery=delivery)
    rep.check("combined path (follow-ups mapping): 2,809 -> 1,013 -> 373 -> 114 page visits",
             [st["page_visits"] for st in path] == [2809, 1013, 373, 114],
             [(st["step"], st["page_visits"]) for st in path])
    by_label = {pg["label"]: st["key"] for st in path for pg in st["pages"]}
    rep.check("Contact Us, Birthday Party Request and Laugh Learn Play are excluded",
             not {"Contact Us", "Birthday Party Request", "Performance Club / Laugh Learn Play",
                  "Academy"} & set(by_label), {k: v for k, v in by_label.items()})
    rep.check("Memberships, Memberships / Recovery and / Corporate are Pricing & purchase",
             all(by_label.get(k) == "purchase" for k in
                 ("Memberships", "Memberships / Recovery", "Memberships / Corporate")),
             {k: by_label.get(k) for k in ("Memberships", "Memberships / Recovery")})
    rep.check("Training, Recovery, Group Classes, Personal Training stay club pages",
             all(by_label.get(k) == "explore" for k in
                 ("Training", "Recovery", "Performance Club / Group Classes",
                  "Performance Club / Personal Training")), None)
    rep.check("Tour Scheduled stays a confirmation page",
             by_label.get("Thank You / Tour Scheduled") == "confirmation")
    rep.check("the subtitle: every step of the membership path, 114 confirmation-page visits",
             ra.goal_path_headline(path, ra.path_noun(GOALS)) ==
             "Attributed visitors reached every step of the membership path, including 114 "
             "confirmation-page visits.")
    facts = ra.build_facts_payload(attribution, delivery, goals=GOALS, vertical="healthcare")
    rep.check("the top ZIPs cluster in DC and Prince George's County; frequency tile 4.6",
             facts["zip"]["cluster"] == "DC and Prince George's County"
             and facts["tile_values"]["frequency"] == "4.6",
             (facts["zip"]["cluster"], facts["tile_values"]))
    bullets = ra.goal_path_bullets(path)
    rep.check("the first URL bullet is the confirmation detail: 34 of 114 were tour bookings",
             bullets and bullets[0].startswith("34 of the 114 confirmation-page visits were tour "
                                               "bookings"), bullets)



MW_ATTRIBUTION = REPO / "MW attribution excel.xlsx"
CARDINAL_ATTRIBUTION = REPO / "Premion Website Attribution Cardinal.xlsx"


def check_brand_and_store_pages(rep):
    """A page group named after the client never takes an intent label
    (Matt, 2026-10-27). Mattress Warehouse's goal "Drive in-store visits to
    Mattress Warehouse locations" put every page on the site (its domain IS
    the name) under "In-Store pages" -- 7,349 page visits of homepage, sale
    and brand pages. Only pages that signal a visit (store locator,
    locations, directions, hours, booking) are store pages; pages carrying
    only the client's name are "Site pages", outside the goal path."""
    print("\n16. The client's own name is never a goal step; store pages signal a visit")
    att = make_attribution()
    att.by_url = {"https://acmebeds.com/": 900, "https://acmebeds.com/sale/summer-sale": 400,
                  "https://acmebeds.com/acme-beds-giveaway": 50,
                  "https://acmebeds.com/store-locator": 150,
                  "https://acmebeds.com/store-locator/acme-beds-of-cary": 30,
                  "https://acmebeds.com/blogs/buying-in-store-vs-online": 1,
                  "https://acmebeds.com/cart": 20}
    goals = ["Drive in-store visits to Acme Beds locations"]
    terms = {t["label"]: t for t in ra.goal_terms(goals, None, att)}
    rep.check("the client's name is a brand term; 'in-store' matches by visit signals",
             terms.get("Acme Beds", {}).get("brand") and terms.get("In-Store", {}).get("signals"),
             terms)
    path = ra.goal_path_combined(att, goals)
    explore = next((st for st in path if st["key"] == "explore"), None)
    rep.check("the goal step is the store pages only, labelled for what they are",
             explore and explore["step"] == ra.IN_STORE_STEP_LABEL and explore["page_visits"] == 180
             and {p["label"] for p in explore["pages"]} == {"Store Locator",
                                                           "Store Locator / Acme Beds Of Cary"},
             explore)
    groups = {g["label"]: g for g in ra.goal_page_groups(att, goals)["classes"]}
    rep.check("homepage, sale and brand pages are 'Site pages', never an intent label; the cart "
             "keeps its own class; the in-store blog post is a site page, not a store page",
             groups.get(ra.SITE_PAGES_LABEL, {}).get("visits") == 900 + 400 + 50 + 1
             and groups.get("In-Store", {}).get("visits") == 180
             and groups.get("Purchase intent", {}).get("visits") == 20,
             {k: g["visits"] for k, g in groups.items()})
    plumber = make_attribution()
    plumber.by_url = {"https://acmeplumbing.com/": 10, "https://acmeplumbing.com/plumbing/drain": 5}
    plumber_terms = ra.goal_terms(["Service calls for plumbing"], None, plumber)
    rep.check("a one-word product the domain carries still counts on its own pages, never the "
             "homepage", ra._goal_term_for("https://acmeplumbing.com/plumbing/drain", plumber_terms)
             == "Plumbing" and ra._goal_term_for("https://acmeplumbing.com/", plumber_terms) is None,
             plumber_terms)
    if MW_ATTRIBUTION.exists():
        mw = ai.parse_attribution_export(str(MW_ATTRIBUTION))
        mw_path = ra.goal_path_combined(mw, ["Drive in-store visits to Mattress Warehouse locations"])
        rep.check("real MW: 'Store / location pages' 667 page visits (was 'In-Store pages' 7,349 "
                 "with the homepage), then Pricing & purchase 35",
                 [(st["step"], st["page_visits"]) for st in mw_path]
                 == [(ra.IN_STORE_STEP_LABEL, 667), ("Pricing & purchase", 35)],
                 [(st["step"], st["page_visits"]) for st in mw_path])
    else:
        rep.skip(f"{MW_ATTRIBUTION.name} not present")
    if CARDINAL_ATTRIBUTION.exists():
        card = ai.parse_attribution_export(str(CARDINAL_ATTRIBUTION))
        card_path = ra.goal_path_combined(card, ["Generate service calls for plumbing and HVAC work"])
        rep.check("real Cardinal: 'Service pages' no longer carries the homepage/Shop/Offers "
                 "(65 page visits, was 297)",
                 card_path[0]["step"] == "Service pages" and card_path[0]["page_visits"] == 65
                 and "Homepage" not in {p["label"] for p in card_path[0]["pages"]},
                 [(st["step"], st["page_visits"]) for st in card_path])
    else:
        rep.skip(f"{CARDINAL_ATTRIBUTION.name} not present")
    if STJ_ATTRIBUTION.exists() and STJ_DELIVERY.exists():
        stj = ai.parse_attribution_export(str(STJ_ATTRIBUTION))
        stj_path = ra.goal_path_combined(stj, GOALS, delivery=ai.parse_delivery_export(str(STJ_DELIVERY)))
        rep.check("real St. James: unchanged (Club / location pages 2,809, Free trial / Join 1,013, "
                 "Pricing & purchase 373, Confirmation pages 114)",
                 [(st["step"], st["page_visits"]) for st in stj_path]
                 == [("Club / location pages", 2809), ("Free trial / Join", 1013),
                     ("Pricing & purchase", 373), ("Confirmation pages", 114)],
                 [(st["step"], st["page_visits"]) for st in stj_path])


def main():
    rep = Report()
    check_period(rep)
    check_recency_labels(rep)
    check_unsupported_dimension_claims(rep)
    check_zip_slide(rep)
    check_trend(rep)
    check_percent_discipline(rep)
    check_pages_groups_path(rep)
    check_brand_and_store_pages(rep)
    check_cpv(rep)
    check_recency_referral_shares(rep)
    check_confirmation_pages(rep)
    check_labeling_rule(rep)
    check_no_page_visit_shares(rep)
    check_no_repeated_stats(rep)
    check_zip_action_names_zips(rep)
    check_dark_week_label(rep)
    check_creative_line(rep)
    check_page_mapping(rep)
    check_page_overrides(rep)
    check_flow_wording(rep)
    check_creative_names(rep)
    check_location_threads(rep)
    check_takeaway_heading_sizes(rep)
    check_takeaways_not_actions(rep)
    check_directional_headings(rep)
    check_referral_chart_label(rep)
    check_zip_narrative(rep)
    check_tile_rounding(rep)
    check_one_sheet_path_dedupe(rep)
    check_relaxed_rules(rep)
    check_st_james(rep)
    total = rep.passed + len(rep.failed)
    print(f"\n{rep.passed} passed, {len(rep.failed)} failed, {len(rep.skipped)} skipped "
          f"out of {total}")
    return 1 if rep.failed else 0


if __name__ == "__main__":
    sys.exit(main())
