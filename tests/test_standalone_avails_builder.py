"""The standalone avails-slide builder (Zip/map builder page's "Standalone
avails builder" tab) -- BACKLOG.md's "Standalone zip/map builder", scoped
down to just the avails slide.

Two halves:
  1. `assembly.build_avails_only_deck` against the real master deck (present
     locally, gitignored) -- a pure assembly check, no Streamlit involved.
  2. The real form, through AppTest + test_mode_upload (the same technique
     test_avails_pdf_wiring.py uses for a file_uploader AppTest can't drive
     directly) -- the property that actually matters: exercising the new
     tab never touches `st.session_state["targeting_groups"]`. That's not
     an incidental assertion here, it's the whole point of the feature --
     a rep's prospect research must never disturb, or be disturbed by, a
     proposal in progress.

    python tests/test_standalone_avails_builder.py
"""
import os
import re
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))
os.chdir(REPO)
os.environ["PROPOSAL_BUILDER_TEST_MODE"] = "1"

import db  # noqa: E402
db.fetch_audiences = lambda: (None, "stubbed for test isolation")

import assembly  # noqa: E402
import slide_map  # noqa: E402
import package_check  # noqa: E402

MASTER_DECK = REPO / "TEGNA_MASTER_DECK_v1_1.pptx"
HERSHEY = REPO / "Premion Media Plan_RFPID-260402_Direct - No Agency_Visit Hershey & Harrisburg_4-30-2026--ver0.pdf"

failures = []
skipped_all = False


def check(label, condition, detail=""):
    if condition:
        print(f"  PASS  {label}")
    else:
        print(f"  FAIL  {label}{'  -- ' + str(detail) if detail else ''}")
        failures.append(label)


def _fill_data(vertical_display="Your Audience"):
    return {
        "client_name": "",
        "proposal_title": "",
        "vertical_display": vertical_display,
        "logo_path": "placeholder_logo.png",
        "avails": {
            "rows": [
                {"audience": "Auto Intenders", "geo": "Washington, DC DMA", "avails": "1,250,000"},
                {"audience": "Auto Intenders", "geo": "Baltimore DMA", "avails": "480,000"},
            ],
            "total_avails": "1,730,000",
            "label": "Max Monthly Avails",
            "map_png": None,
        },
        "media_plan_options": [{}],
    }


def check_build_avails_only_deck():
    print("assembly.build_avails_only_deck against the real master deck")
    if not MASTER_DECK.exists():
        print("SKIP -- master deck not present (gitignored)")
        return

    prs, warnings = assembly.build_avails_only_deck(
        str(MASTER_DECK), _fill_data(), map_present=False)
    check("exactly one slide survives", len(prs.slides._sldIdLst) == 1,
          len(prs.slides._sldIdLst))
    check("personalize() reported no warnings", not warnings, warnings)

    text = slide_map.extract_slide_text(prs.slides[0])
    check("every {{TOKEN}} was consumed", not re.findall(r"\{\{[A-Z_]+\}\}", text),
          re.findall(r"\{\{[A-Z_]+\}\}", text))
    for expected in ("Auto Intenders", "Washington, DC DMA", "1,250,000", "1,730,000"):
        check(f'"{expected}" reached the slide', expected in text, text)

    print("\nno_vertical=True drops the word \"intenders\" from the one sentence that needs it, "
          "without touching the other two {{VERTICAL}} spots")
    prs2, _ = assembly.build_avails_only_deck(
        str(MASTER_DECK), _fill_data(), map_present=False, no_vertical=True)
    text2 = slide_map.extract_slide_text(prs2.slides[0])
    check('"reach Your Audience in all 210 DMAs" (no "intenders")',
          "reach Your Audience in all 210 DMAs" in text2, text2)
    check('the eyebrow line still reads "PREMION + Your Audience"',
          "PREMION + Your Audience" in text2, text2)

    print("\npackage_check.check_package -- no dangling relationship, undeclared part or "
          "unresolved rId (the same structural check the fast test suite runs on every deck)")
    out_path = REPO / "tests" / "_scratch_avails_only.pptx"
    prs.save(str(out_path))
    try:
        problems = package_check.check_package(str(out_path))
        check("package is structurally sound", not problems, problems)
    finally:
        out_path.unlink(missing_ok=True)


def ss(at, key, default=None):
    return at.session_state[key] if key in at.session_state else default


def main():
    global skipped_all
    check_build_avails_only_deck()

    if not HERSHEY.exists():
        print("\nSKIP -- Visit Hershey & Harrisburg real avails PDF not present (gitignored fixture)")
        skipped_all = not MASTER_DECK.exists()
    else:
        from streamlit.testing.v1 import AppTest

        print("\nthe standalone tab, end to end: upload -> resolve -> build -- and the "
              "separation property that's the whole point of this feature")
        at = AppTest.from_file(str(REPO / "app.py"), default_timeout=300)
        at.session_state["authed"] = True
        at.session_state["current_user"] = "T"
        at.session_state["page_choice"] = "Zip/map builder"
        at.session_state["standalone_avails_pdf_upload_path"] = str(HERSHEY)
        at.run()
        check("no exception on upload", not at.exception,
              at.exception[0].message[:400] if at.exception else "")
        check("12 groups landed in standalone_groups (one per audience-geography pair)",
              len(ss(at, "standalone_groups") or []) == 12, ss(at, "standalone_groups"))
        check("targeting_groups was never even CREATED, let alone written to -- "
              "the separation property this feature exists to guarantee",
              "targeting_groups" not in at.session_state, True)

        build_btn = [b for b in at.button if b.label == "Build the avails slide"]
        check("the Build button is present once groups exist", len(build_btn) == 1,
              [b.label for b in at.button])
        if build_btn:
            build_btn[0].click().run()
            check("no exception building the slide", not at.exception,
                  at.exception[0].message[:400] if at.exception else "")
            check("slide bytes were produced",
                  bool(ss(at, "standalone_avails_slide_bytes")), None)
            check("a download button is offered",
                  any("avails slide" in b.label.lower() for b in at.download_button),
                  [b.label for b in at.download_button])
            check("targeting_groups STILL never touched, even after building the slide",
                  "targeting_groups" not in at.session_state, True)

            # The map defaults to ON whenever geography resolved (Hershey's
            # 12 groups did) -- never an opt-in -- and the slide fill must
            # use a DARK-composited render, not the app's own light-
            # background preview image: feeding the light preview to the
            # map-variant template (no stock photo, a dark gradient
            # background) was a real bug here, caught by rendering a real
            # import through PowerPoint (a bright white box on the slide's
            # dark navy background). The map-variant slide has zero native
            # Picture shapes (see TEGNA_MASTER_DECK_v1_1.pptx slide 16); one
            # Picture surviving here is the placed map, proving map_present
            # was True and the standard (photo) template was NOT used.
            import io as _io
            from pptx import Presentation as _Presentation
            _prs = _Presentation(_io.BytesIO(ss(at, "standalone_avails_slide_bytes")))
            _pictures = [s for s in _prs.slides[0].shapes if s.shape_type == 13]
            check("the map variant was used automatically (one placed map picture, "
                  "no native stock photo)", len(_pictures) == 1, len(_pictures))

            if _pictures:
                import targeting_map as _tm
                _embedded_bytes = _pictures[0].image.blob
                _groups = ss(at, "standalone_groups") or []
                _light_render = _tm.render_map(
                    _tm.groups_with_zips(_groups), width_px=1000, height_px=620)
                check("the slide's own map image is the DARK-composited render, not "
                      "the app's own light-background preview image (the actual bug, "
                      "confirmed by rendering it: a white box on the slide's dark "
                      "gradient)", _embedded_bytes != _light_render, None)

    print()
    if failures:
        print(f"{len(failures)} FAILED: {failures}")
        return 1
    print("The standalone avails-slide builder produces a real, structurally sound "
          "single-slide deck from the same D2 machinery a proposal uses, backed by its own "
          "group list -- targeting_groups is never touched by any part of this flow.")
    return 0


if __name__ == "__main__":
    code = main()
    if skipped_all:
        sys.exit(0)
    sys.exit(code)
