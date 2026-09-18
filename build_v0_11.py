"""build_v0_11.py -- inventory check for REPORT_MASTER_v0_11.pptx.

Same role as build_v0_10.py: v0_11 was hand-built by Matt directly from
his own handoff spec (one change -- RateTileDelta/VisitorsTileDelta on
report:highlights, month-over-month delta captions, cross-month evidence
work item 5). This script doesn't build the file -- it VERIFIES the
delivered file actually matches the spec, and that nothing else on the
slide moved from v0_10.

    python build_v0_11.py
"""
import sys

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL_TYPE

import slide_map
from pptx import Presentation

V0_10 = "REPORT_MASTER_v0_10.pptx"
V0_11 = "REPORT_MASTER_v0_11.pptx"

FAILURES = []


def check(label, condition, detail=None):
    status = "PASS" if condition else "FAIL"
    print(f"  {status}  {label}" + (f" -- {detail}" if detail and not condition else ""))
    if not condition:
        FAILURES.append(label)


def _slide_by_key(prs):
    keys = {}
    for i, slide in enumerate(prs.slides):
        if slide.has_notes_slide:
            key = slide_map.notes_key(slide)
            if key:
                keys[key] = i
    return keys


def _shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def _in(value_emu, expected_in, tol=0.02):
    return abs(value_emu / 914400 - expected_in) < tol


def main():
    prs10 = Presentation(V0_10)
    prs11 = Presentation(V0_11)
    keys10, keys11 = _slide_by_key(prs10), _slide_by_key(prs11)

    check("v0_11 has the same slide count as v0_10", len(list(prs11.slides)) == len(list(prs10.slides)),
         (len(list(prs11.slides)), len(list(prs10.slides))))

    hl10 = prs10.slides[keys10["report:highlights"]]
    hl11 = prs11.slides[keys11["report:highlights"]]

    print("\n1. RateTileDelta / VisitorsTileDelta -- new shapes, new tokens")
    rate_tile = _shape(hl11, "RateTile")
    visitors_tile = _shape(hl11, "VisitorsTile")
    rate_delta = _shape(hl11, "RateTileDelta")
    visitors_delta = _shape(hl11, "VisitorsTileDelta")
    check("RateTileDelta exists", rate_delta is not None)
    check("VisitorsTileDelta exists", visitors_delta is not None)
    if rate_delta is not None:
        check("RateTileDelta carries the {{RATE_DELTA}} token",
             "{{RATE_DELTA}}" in rate_delta.text_frame.text, rate_delta.text_frame.text)
        check("RateTileDelta is centered under RateTile (same left, same width)",
             rate_delta.left == rate_tile.left and rate_delta.width == rate_tile.width,
             (rate_delta.left, rate_tile.left, rate_delta.width, rate_tile.width))
    if visitors_delta is not None:
        check("VisitorsTileDelta carries the {{VISITORS_DELTA}} token",
             "{{VISITORS_DELTA}}" in visitors_delta.text_frame.text, visitors_delta.text_frame.text)
        check("VisitorsTileDelta is centered under VisitorsTile (same left, same width)",
             visitors_delta.left == visitors_tile.left and visitors_delta.width == visitors_tile.width,
             (visitors_delta.left, visitors_tile.left, visitors_delta.width, visitors_tile.width))

    print("\n2. Position -- y=3.38in, the gap between the tiles (bottom 3.35) and the card (top 3.70)")
    for shape, label in ((rate_delta, "RateTileDelta"), (visitors_delta, "VisitorsTileDelta")):
        if shape is None:
            continue
        check(f"{label} top is 3.38in", _in(shape.top, 3.38), shape.top / 914400)
    tile_bottom_in = (rate_tile.top + rate_tile.height) / 914400
    check("the tiles' own bottom edge is 3.35in (the floor this sits above)",
         abs(tile_bottom_in - 3.35) < 0.02, tile_bottom_in)
    card_header = _shape(hl11, "TrendChartHeader")
    check("TrendChartHeader's own top is 3.70in (the ceiling this sits below)",
         card_header is not None and _in(card_header.top, 3.70),
         card_header.top / 914400 if card_header else None)

    print("\n3. Style -- small, muted, no fill")
    for shape, label in ((rate_delta, "RateTileDelta"), (visitors_delta, "VisitorsTileDelta")):
        if shape is None:
            continue
        check(f"{label} has no fill", shape.fill.type == MSO_FILL_TYPE.BACKGROUND, shape.fill.type)
        run = shape.text_frame.paragraphs[0].runs[0]
        check(f"{label}'s own run is small (<= 11pt)",
             run.font.size is not None and run.font.size.pt <= 11, run.font.size)
        check(f"{label}'s own run is muted (not pure black)",
             run.font.color is not None and run.font.color.type is not None
             and run.font.color.rgb != RGBColor(0, 0, 0), run.font.color.rgb)

    print("\n4. Nothing else on report:highlights moved from v0_10")
    unchanged_names = ("ImpressionsTile", "VisitorsTile", "RateTile", "ConversionsTile",
                       "TrendChartHeader", "TrendChartRegion", "TrendChartRegionLabel")
    for name in unchanged_names:
        s10, s11 = _shape(hl10, name), _shape(hl11, name)
        check(f"{name} geometry unchanged", s10 is not None and s11 is not None
             and (s10.left, s10.top, s10.width, s10.height) == (s11.left, s11.top, s11.width, s11.height),
             None if (s10 and s11) else (s10, s11))

    print("\n5. report_assembly._fill_momentum_captions's own fill/delete code "
         "already targets these exact names")
    import report_assembly as ra
    check("RateTileDelta is in _MOMENTUM_CAPTION_SHAPES", "RateTileDelta" in ra._MOMENTUM_CAPTION_SHAPES)
    check("VisitorsTileDelta is in _MOMENTUM_CAPTION_SHAPES",
         "VisitorsTileDelta" in ra._MOMENTUM_CAPTION_SHAPES)
    check("the tokens match exactly (bare names -- _fill_tokens wraps them itself)",
         ra._MOMENTUM_CAPTION_SHAPES["RateTileDelta"][0] == "RATE_DELTA"
         and ra._MOMENTUM_CAPTION_SHAPES["VisitorsTileDelta"][0] == "VISITORS_DELTA")

    print(f"\n{'ALL PASSED' if not FAILURES else f'{len(FAILURES)} FAILURE(S): ' + ', '.join(FAILURES)}")
    return 1 if FAILURES else 0


if __name__ == "__main__":
    sys.exit(main())
