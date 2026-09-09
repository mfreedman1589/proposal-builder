import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.enum.text import PP_ALIGN

SRC, OUT = "REPORT_MASTER_v0_5.pptx", "REPORT_MASTER_v0_6.pptx"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
prs = Presentation(SRC)

# ---------------- helpers ----------------
def key_of(slide):
    return slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""

def find_slide(key):
    for s in prs.slides:
        if f"key: {key}" in key_of(s): return s
    raise KeyError(key)

def clone_after(src, after_key, notes):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    rid = {}
    for rel in src.part.rels.values():
        if rel.reltype == RT.IMAGE: rid[rel.rId] = new.part.relate_to(rel.target_part, RT.IMAGE)
    for sh in src.shapes:
        el = copy.deepcopy(sh._element)
        for node in el.iter():
            for attr in ("{%s}embed" % NS["r"], "{%s}link" % NS["r"]):
                if attr in node.attrib and node.attrib[attr] in rid: node.attrib[attr] = rid[node.attrib[attr]]
        new.shapes._spTree.append(el)
    new.notes_slide.notes_text_frame.text = notes
    lst = prs.slides._sldIdLst; ids = list(lst); last = ids[-1]
    idx = [i for i, s in enumerate(prs.slides) if f"key: {after_key}" in key_of(s)][0]
    lst.remove(last); lst.insert(idx + 1, last)
    return new

class S:
    def __init__(self, slide): self.s = slide
    def shape(self, name):
        for sh in self.s.shapes:
            if sh.name == name: return sh
        raise KeyError(name)
    def drop(self, *names):
        for n in names:
            try: sh = self.shape(n); sh._element.getparent().remove(sh._element)
            except KeyError: pass
    def rename(self, old, new): self.shape(old).name = new
    def clone(self, name, new_name):
        el = copy.deepcopy(self.shape(name)._element); self.s.shapes._spTree.append(el)
        sh = self.s.shapes[-1]; sh.name = new_name; return sh
    @staticmethod
    def text(sh, t):
        tf = sh.text_frame; p = tf.paragraphs[0]
        if p.runs:
            p.runs[0].text = t
            for r in p.runs[1:]: r._r.getparent().remove(r._r)
        else: p.add_run().text = t
        for extra in tf.paragraphs[1:]: extra._p.getparent().remove(extra._p)
    @staticmethod
    def place(sh, x, y, w=None, h=None):
        sh.left, sh.top = Inches(x), Inches(y)
        if w is not None: sh.width = Inches(w)
        if h is not None: sh.height = Inches(h)
    @staticmethod
    def cell(c, t):
        p = c.text_frame.paragraphs[0]
        if p.runs:
            p.runs[0].text = t
            for r in p.runs[1:]: r._r.getparent().remove(r._r)
        elif t: p.add_run().text = t
    @staticmethod
    def recolumn(tbl_shape, heads, widths, token, left_cols):
        tbl = tbl_shape.table; el = tbl._tbl; grid = el.find("a:tblGrid", NS)
        cur = len(grid.findall("a:gridCol", NS)); want = len(heads)
        while cur < want:
            grid.append(copy.deepcopy(grid.findall("a:gridCol", NS)[-1]))
            for tr in el.findall("a:tr", NS): tr.append(copy.deepcopy(tr.findall("a:tc", NS)[-1]))
            cur += 1
        while cur > want:
            grid.remove(grid.findall("a:gridCol", NS)[-1])
            for tr in el.findall("a:tr", NS): tr.remove(tr.findall("a:tc", NS)[-1])
            cur -= 1
        for gc, w in zip(grid.findall("a:gridCol", NS), widths): gc.set("w", str(Inches(w)))
        tbl_shape.width = Inches(sum(widths))
        for c, h in zip(list(tbl.rows[0].cells), heads): S.cell(c, h)
        row = list(tbl.rows[1].cells)
        S.cell(row[0], token)
        for c in row[1:]: S.cell(c, "")
        for r in tbl.rows:
            for i, c in enumerate(r.cells):
                c.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT if i < left_cols else PP_ALIGN.RIGHT
    def tile(self, old, new, token, label):
        for suf in ("", "Value", "Label"): self.rename(old + suf, new + suf)
        S.text(self.shape(new + "Value"), token); S.text(self.shape(new + "Label"), label)
    def space_tiles(self, names, x0=0.50, span=12.33, gap=0.25, y=2.10, h=0.85):
        tw = (span - gap * (len(names) - 1)) / len(names)
        for i, n in enumerate(names):
            x = x0 + i * (tw + gap)
            S.place(self.shape(n), x, y, tw, h)
            S.place(self.shape(n + "Value"), x + 0.15, y + 0.08, tw - 0.30, 0.48)
            S.place(self.shape(n + "Label"), x + 0.15, y + 0.56, tw - 0.30, 0.25)

# =====================================================================
# 1. response_profile  (clone attribution_breakdown, insert after it)
# =====================================================================
rp = S(clone_after(find_slide("report:attribution_breakdown"), "report:attribution_breakdown",
                   "key: report:response_profile"))
S.text(rp.shape("Text 1"), "Response Timing & Source")
S.text(rp.shape("Text 2"), "{{RECENCY_STAT}}")
L, R, CW = 0.50, 6.87, 5.97
# left: recency chart, then conditional day-of-week
hdr_l = rp.shape("Text 3"); S.text(hdr_l, "DAYS FROM EXPOSURE TO VISIT"); S.place(hdr_l, L, 2.10, CW, 0.26)
S.place(rp.shape("ChartRegion"), L, 2.45, CW, 2.10)
S.place(rp.shape("ChartRegionLabel"), L, 2.45, CW, 2.10)
S.text(rp.shape("ChartRegionLabel"), "ChartRegion — attributed visitors by days since exposure")
dow_h = rp.clone("Text 3", "DayOfWeekHeader"); S.text(dow_h, "ATTRIBUTED RATE BY DAY OF WEEK"); S.place(dow_h, L, 4.75, CW, 0.26)
dow_t = rp.clone("BreakdownTable", "DayOfWeekTable")
S.recolumn(dow_t, ["Day", "Attributed rate"], [3.97, 2.00], "{{DAY_OF_WEEK_ROWS}}", 1)
S.place(dow_t, L, 5.10)
# right: referral table, then narrative card
hdr_r = rp.shape("Text 4"); S.text(hdr_r, "HOW VISITORS ARRIVED"); S.place(hdr_r, R, 2.10, CW, 0.26)
ref = rp.shape("BreakdownTable"); ref.name = "ReferralTable"
S.recolumn(ref, ["Source", "Visitors", "Share"], [3.17, 1.40, 1.40], "{{REFERRAL_ROWS}}", 1)
S.place(ref, R, 2.45)
card = rp.shape("Shape 7"); S.place(card, R, 4.55, CW, 2.35)
nar = rp.shape("AttributionNarrative"); nar.name = "ResponseProfileNarrative"
S.text(nar, "{{RESPONSE_PROFILE_NARRATIVE}}"); S.place(nar, R + 0.30, 4.75, CW - 0.60, 1.95)

# =====================================================================
# 2. ott_retargeting  (clone live_sports, insert after zip_analysis)
# =====================================================================
ot = S(clone_after(find_slide("report:live_sports"), "report:zip_analysis", "key: report:ott_retargeting"))
S.text(ot.shape("Text 1"), "OTT Retargeting")
S.text(ot.shape("Text 2"), "Display banners served to households exposed to the streaming campaign")
ot.tile("SportsImpressionsTile", "DisplayImpressionsTile", "{{OTT_IMPRESSIONS}}", "Display impressions")
ot.tile("SportsVcrTile", "DisplayClicksTile", "{{OTT_CLICKS}}", "Clicks")
ot.tile("SportsPacingTile", "DisplayCtrTile", "{{OTT_CTR}}", "Click-through rate")
ot.space_tiles(["DisplayImpressionsTile", "DisplayClicksTile", "DisplayCtrTile"])
ot.drop("SportsRfpidCaption", "SportsByLeagueHeader", "SportsByLeagueTable")
LW = 7.70
# creative table (conditional) then ad-size table
ch = ot.shape("Text 18"); ch.name = "CreativeHeader"; S.text(ch, "BY CREATIVE"); S.place(ch, L, 3.15, LW, 0.26)
ct = ot.shape("SportsEventTable"); ct.name = "CreativeTable"
S.recolumn(ct, ["Creative", "Impressions", "Clicks", "CTR"], [3.50, 1.60, 1.30, 1.30], "{{OTT_CREATIVE_ROWS}}", 1)
S.place(ct, L, 3.45)
ah = ot.clone("CreativeHeader", "AdSizeHeader"); S.text(ah, "BY UNIT SIZE"); S.place(ah, L, 4.60, LW, 0.26)
at = ot.clone("CreativeTable", "AdSizeTable")
S.recolumn(at, ["Unit", "Impressions", "Clicks", "CTR"], [3.50, 1.60, 1.30, 1.30], "{{OTT_AD_SIZE_ROWS}}", 1)
S.place(at, L, 4.95)
# right column: two stats + narrative
RX, RW = 8.45, 4.38
sh1 = ot.clone("CreativeHeader", "ScreenHeader"); S.text(sh1, "WHERE CLICKS CAME FROM"); S.place(sh1, RX, 3.15, RW, 0.26)
st1 = ot.clone("LiveSportsNarrative", "ScreenStat"); S.text(st1, "{{OTT_SCREEN_STAT}}"); S.place(st1, RX, 3.45, RW, 0.55)
sh2 = ot.clone("CreativeHeader", "BlendedHeader"); S.text(sh2, "STREAMING + RETARGETING COMBINED"); S.place(sh2, RX, 4.15, RW, 0.26)
st2 = ot.clone("LiveSportsNarrative", "BlendedStat"); S.text(st2, "{{OTT_BLENDED_STAT}}"); S.place(st2, RX, 4.45, RW, 0.55)
nar2 = ot.shape("LiveSportsNarrative"); nar2.name = "OttRetargetingNarrative"
S.text(nar2, "{{OTT_RETARGETING_NARRATIVE}}"); S.place(nar2, RX, 5.20, RW, 1.65)

prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides))
print([key_of(s).split("\n")[0].replace("key: report:", "") for s in prs.slides])
