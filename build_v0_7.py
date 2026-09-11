import copy
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

SRC, OUT = "REPORT_MASTER_v0_6.pptx", "REPORT_MASTER_v0_7.pptx"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
prs = Presentation(SRC)

def slide(key):
    for s in prs.slides:
        if f"key: {key}" in s.notes_slide.notes_text_frame.text: return s
    raise KeyError(key)

def shape(s, name):
    for sh in s.shapes:
        if sh.name == name: return sh
    raise KeyError(name)

def cell_set(c, t):
    p = c.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = t
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
    elif t: p.add_run().text = t

# ---------- 1. widen DeliveryByGeoTable 3 -> 5 columns ----------
bd = slide("report:delivery_breakdown")
geo = shape(bd, "DeliveryByGeoTable")
tbl, el = geo.table, geo.table._tbl
grid = el.find("a:tblGrid", NS)
for _ in range(2):
    grid.append(copy.deepcopy(grid.findall("a:gridCol", NS)[-1]))
    for tr in el.findall("a:tr", NS):
        tr.append(copy.deepcopy(tr.findall("a:tc", NS)[-1]))

widths = [2.25, 1.30, 1.30, 1.15, 1.00]   # = 7.00, unchanged total
for gc, w in zip(grid.findall("a:gridCol", NS), widths):
    gc.set("w", str(Inches(w)))
geo.width = Inches(sum(widths))

heads = ["Geography", "Planned", "Delivered", "% of plan", "VCR"]
for c, h in zip(list(tbl.rows[0].cells), heads):
    cell_set(c, h)
row = list(tbl.rows[1].cells)
cell_set(row[0], "{{DELIVERY_BY_GEO_ROWS}}")
for c in row[1:]:
    cell_set(c, "")
for r in tbl.rows:
    for i, c in enumerate(r.cells):
        c.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT if i < 1 else PP_ALIGN.RIGHT

# ---------- 2. PLAN_VS_ACTUAL_NOTE on delivery_recap ----------
dr = slide("report:delivery_recap")
nar = shape(dr, "DeliveryNarrative")
note_el = copy.deepcopy(nar._element)
dr.shapes._spTree.append(note_el)
note = dr.shapes[-1]
note.name = "PlanVsActualNote"
p = note.text_frame.paragraphs[0]
p.runs[0].text = "{{PLAN_VS_ACTUAL_NOTE}}"
for extra in p.runs[1:]: extra._r.getparent().remove(extra._r)
for extra in note.text_frame.paragraphs[1:]: extra._p.getparent().remove(extra._p)
# narrative moves up slightly; note sits beneath it
nar.top, nar.height = Inches(6.10), Inches(0.50)
note.left, note.top = Inches(0.50), Inches(6.60)
note.width, note.height = Inches(12.33), Inches(0.35)

prs.save(OUT)
print("saved", OUT)
