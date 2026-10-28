"""A generated deck's slide-by-slide text -- titles, text boxes, tables --
so a tester can check slide content without opening PowerPoint. Pure: no
Streamlit, no database. Image-only slides (an appended Analyst deck's charts)
show whatever text they carry and say when they have none."""
from pptx import Presentation

import slide_map


def _table_block(table):
    rows = [[cell.text.strip().replace("|", "\\|").replace("\n", " ")
             for cell in row.cells] for row in table.rows]
    if not rows:
        return None
    return "\n".join(["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
                     + ["| " + " | ".join(r) + " |" for r in rows[1:]])


def deck_text(pptx_path):
    """Markdown text, one section per slide, in deck order. Every text box is
    its own paragraph and every line in it its own line (a markdown hard
    break), so a tile's value stays beside its own label and nothing past a
    box's first line is dropped (QA, 2026-10-28: run together, Muse couldn't
    tell the Executive Summary's tiles apart)."""
    prs = Presentation(pptx_path)
    out = []
    for number, slide in enumerate(prs.slides, start=1):
        key = slide_map.notes_key(slide) if slide.has_notes_slide else None
        blocks = []  # (kind, content): ("text", [lines]) or ("table", markdown)
        for shape in slide_map.iter_all_shapes(slide.shapes):
            if getattr(shape, "has_table", False) and shape.has_table:
                table = _table_block(shape.table)
                if table:
                    blocks.append(("table", table))
            elif shape.has_text_frame and shape.text_frame.text.strip():
                lines = [line.strip() for line in shape.text_frame.text.splitlines()
                         if line.strip()]
                blocks.append(("text", lines))
        title = "(no text on this slide)"
        if blocks and blocks[0][0] == "text":
            title = blocks[0][1][0]
            rest = blocks[0][1][1:]
            blocks = ([("text", rest)] if rest else []) + blocks[1:]
        header = f"### Slide {number}: {title}"
        if key:
            header += f"  `{key}`"
        out += [header, ""]
        for kind, content in blocks:
            out += [content if kind == "table" else "  \n".join(content), ""]
    return "\n".join(out)
