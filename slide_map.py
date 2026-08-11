"""
slide_map.py -- scan a master deck and derive slide_number -> condition_key.

Per master_v1_1_changelog.md: "Don't hand-patch offsets -- regenerate
SLIDE_MAP by scanning v1.1 (extract per-slide text, match section anchors)."
This is also called out as reusable infrastructure for the Vault versioning
flow (build spec section 8), so it's written as a standalone, deck-agnostic
utility rather than a one-off script: point it at any deck built on the same
content anchors and it re-derives the map without touching slide numbers.

Approach: extract concatenated text per slide, then run each slide's text
through an ordered set of anchor rules (most specific first). A few rules
mark section-divider slides and update a "current section" carry-forward
default; that default is what genuinely blank/undifferentiated slides fall
back to. Every other slide is resolved by a directly-matched anchor.
"""

import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# An explicit key label in a slide's speaker notes: "key:" at the start of its
# own line, then the condition_key. Written by tag_deck_keys.py and editable
# by hand in PowerPoint (View -> Notes Page). Matched case-insensitively and
# tolerant of whitespace, since it's hand-edited.
NOTES_KEY_LINE = re.compile(r"^\s*key\s*:\s*(\S+)\s*$", re.IGNORECASE)


def iter_all_shapes(shapes):
    for shape in shapes:
        yield shape
        # python-pptx raises NotImplementedError for a <p:sp> that is neither
        # a placeholder nor an autoshape, which is a legal thing for a deck to
        # contain -- imported case study slides have their <p:ph> stripped, and
        # anything without preset geometry lands here. A shape this walk can't
        # classify is simply not a group; refusing to enumerate the slide at
        # all is a far worse answer than assuming that.
        try:
            is_group = shape.shape_type == MSO_SHAPE_TYPE.GROUP
        except NotImplementedError:
            is_group = False
        if is_group:
            yield from iter_all_shapes(shape.shapes)


def extract_slide_text(slide):
    """Concatenate all run/cell text on a slide into one fingerprint string."""
    parts = []
    for shape in iter_all_shapes(slide.shapes):
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                parts.append(t)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    t = cell.text_frame.text.strip()
                    if t:
                        parts.append(t)
    return " || ".join(parts)


def _contains_all(text, *substrings):
    return all(s in text for s in substrings)


def _classify_sport_package(text):
    """Disambiguate which sport a 'PREMION | SPORTS PACKAGES' slide covers.

    Ordered so that more specific families (WNBA before NBA, Home Team
    before Regular/Playoffs) are checked first -- several of these strings
    are substrings of each other (e.g. "WNBA" contains "NBA").
    """
    rules = [
        (("NHL", "Regular"), "sport:nhl_reg"),
        (("NHL", "Playoffs"), "sport:nhl_playoffs"),
        (("WNBA", "Playoffs"), "sport:wnba_playoffs"),
        (("WNBA", "Regular"), "sport:wnba_reg"),
        (("NCAA", "Basketball"), "sport:ncaa_basketball"),
        (("NBA", "Playoffs"), "sport:nba_playoffs"),
        (("NBA", "Regular"), "sport:nba_reg"),
        # "Home Team Game Inventory Only" is the actual home-team-package
        # marker; bare "Home Team" also appears in regular-season slides'
        # coverage bullet ("League-Wide & Home Team Games"), so that alone
        # would misclassify them.
        (("Home Team Game Inventory", "NFL"), "sport:nfl_home_team"),
        (("NFL", "Playoffs"), "sport:nfl_playoffs"),
        (("NFL", "Regular"), "sport:nfl_reg"),
        (("Home Team Game Inventory", "NCAA"), "sport:ncaaf_home_team"),
        (("NCAA Football", "Playoffs"), "sport:ncaaf_playoffs"),
        (("NCAA Football", "Conferences"), "sport:ncaaf_conference"),
        (("NCAA", "Football", "Regular"), "sport:ncaaf_reg"),
        (("MLB", "Playoffs"), "sport:mlb_playoffs"),
        (("MLB", "Regular"), "sport:mlb_reg"),
        (("Professional", "Soccer"), "sport:soccer_pro"),
        # PGA Majors is its own package/slide, distinct from the general
        # Live Golf package -- check the more specific one first.
        (("PGA", "Majors"), "sport:pga_majors"),
        (("Professional", "Golf"), "sport:golf"),
        (("Prestige", "Sports"), "sport:prestige_sports"),
        (("All Live", "Sports"), "sport:all_live_sports"),
    ]
    for needles, key in rules:
        if _contains_all(text, *needles):
            return key
    return None


VERTICAL_ANCHORS = [
    ("PREMION + RETAIL", "vertical:retail"),
    ("PREMION + TRAVEL", "vertical:travel"),
    ("PREMION + HOME IMPROVEMENT", "vertical:home_improvement"),
    ("PREMION + BANKING", "vertical:banking"),
    ("PREMION + ENTERTAINMENT", "vertical:entertainment"),
    ("PREMION + EDUCATION", "vertical:education"),
    ("PREMION + CASUAL DINING", "vertical:dining_qsr"),
    ("PREMION + AUTOMOTIVE", "vertical:auto"),
    ("POLK AUDIENCES", "vertical:auto"),
    ("POLK SIGNALS", "vertical:auto"),
]

# Every section-divider slide classifies as this one key, and no preset ever
# includes it -- dividers are dropped from generated proposals globally.
# Keeping it a single key (rather than one per section) is what lets
# assembly.build_presentation exclude them with one rule instead of having to
# know the name of every section.
SECTION_DIVIDER_KEY = "section_divider"

# Exact-text section dividers: (literal text, the carry-forward default the
# slides *after* it adopt). The divider slide itself is always
# SECTION_DIVIDER_KEY; the second element is only about what follows it, and
# still matters even though the divider is dropped -- the slides it
# introduces are resolved by that default.
SECTION_DIVIDERS = [
    ("Content || Premium || © PREMION 2024", "full_deck"),
    ("Targeting || Precision || © PREMION 2024", "full_deck"),
    ("Measurement || Attribution +", "full_deck"),
    ("Specialties || Vertical", "any_vertical"),
    ("Media Group || TEGNA", "tegna_positioning"),
    ("Marketplace || Audience", "am"),
    ("Audience Marketplace", "am"),
    ("Live Sports", "sports"),
    ("Total TV", "total_tv"),
    ("Proposal Slides", "proposal_divider"),
]


def _classify_sport_viewership(upper_text):
    """Disambiguate which sport an 'Ad-Supported Streaming TV Viewer' slide
    covers, mirroring _classify_sport_package. Order matters: WNBA before
    NBA (substring), NCAAF before bare NCAA, etc.; PGA and PRESTIGE before
    bare GOLF (the Prestige Sports Viewers slide's own MRI footnote lists
    "Streams Golf, Tennis, Horse Racing", which would otherwise misclassify
    it). A handful of viewership slides ("Motorsports Viewers", "World Cup
    Viewers", "March Madness") have no 1:1 package counterpart in this deck
    version and are intentionally left unmatched -- they're never
    selectable so never included.
    """
    rules = [
        (("WNBA", "PLAYOFF"), "sport_viewership:wnba_playoffs"),
        (("WNBA",), "sport_viewership:wnba_reg"),
        (("NBA", "PLAYOFF"), "sport_viewership:nba_playoffs"),
        (("NBA",), "sport_viewership:nba_reg"),
        (("NFL", "PLAYOFF"), "sport_viewership:nfl_playoffs"),
        (("NFL",), "sport_viewership:nfl_reg"),
        (("NHL", "PLAYOFF"), "sport_viewership:nhl_playoffs"),
        (("NHL",), "sport_viewership:nhl_reg"),
        (("MLB", "PLAYOFF"), "sport_viewership:mlb_playoffs"),
        (("MLB",), "sport_viewership:mlb_reg"),
        # "CONF" alone also matches "confidential" in every slide's MRI
        # footnote boilerplate; "CONF." (with the period, as in "Conf.
        # Viewers") does not.
        (("NCAAF", "CONF."), "sport_viewership:ncaaf_conference"),
        (("NCAAF", "PLAYOFF"), "sport_viewership:ncaaf_playoffs"),
        (("NCAAF",), "sport_viewership:ncaaf_reg"),
        (("NCAA BASKETBALL",), "sport_viewership:ncaa_basketball"),
        (("PGA",), "sport_viewership:pga_majors"),
        (("PRESTIGE",), "sport_viewership:prestige_sports"),
        # "SOCCER VIEWERS" (the heading, two words together) rather than
        # bare "SOCCER" -- the World Cup slide's own footnote lists
        # "Soccer-World Cup" among included leagues, which would otherwise
        # misclassify it as the Soccer slide.
        (("SOCCER VIEWERS",), "sport_viewership:soccer_pro"),
        # Bare "Golf Viewers" -- must come after PGA and PRESTIGE (see
        # docstring: Prestige's own footnote mentions "Golf").
        (("GOLF",), "sport_viewership:golf"),
    ]
    for needles, key in rules:
        if _contains_all(upper_text, *needles):
            return key
    return None


def _classify_slide_raw(text, current_default):
    """Return (condition_key, new_default) for one slide's text."""
    upper = text.upper()

    # --- Priority 0: Total TV market variants --------------------------
    # These must be tested before the generic token rules below, because the
    # co-brand covers also carry {{PROPOSAL_TITLE}} and the Total TV plan
    # templates also carry {{PLAN_TITLE}} -- a generic match would collapse
    # each variant onto the standard slide it's meant to replace.
    #
    # Only the broadcast schedule templates are distinguishable by text: they
    # name their station. The co-brand covers and the Total TV plan templates
    # are byte-identical between markets apart from a logo image, so **they
    # resolve from their notes `key:` labels only**. That's not a gap to fix
    # with a cleverer anchor -- there is nothing in the text to match on --
    # and it's a concrete example of why the labels live in the deck.
    if "{{BROADCAST_PLAN_DESC}}" in text:
        if "WUSA" in upper:
            return "broadcast_schedule_template:dc", current_default
        if "FOX43" in upper or "WPMT" in upper:
            return "broadcast_schedule_template:harrisburg", current_default

    # --- Priority 1: template token slides -----------------------------
    if "{{PROPOSAL_TITLE}}" in text:
        return "client_title", current_default
    if "{{GOALS_BULLETS}}" in text:
        return "campaign_specs", current_default
    if "{{AVAILS}}" in text and "{{VERTICAL}}" in text:
        return "targeting_avails_template", current_default
    if "{{PLAN_TITLE}}" in text:
        return "proposal_template", current_default

    # --- Priority 2: distinctive single-slide markers -------------------
    if "BILINGUAL-SPANISH" in upper:
        return "spanish", current_default
    if "CROSSIX" in upper or "PREMION + HEALTHCARE" in upper:
        return "vertical:healthcare", current_default
    if "1ST PARTY DATA TARGETING" in upper:
        return "first_party", current_default
    if "STREAMING TV RETARGETING" in upper:
        return "streaming_retargeting", current_default
    if "Dynamic Video Ad" in text:
        return "dynamic_creative", current_default
    if "Linear TV + PREMION" in text:
        return "linear_reach_ext", current_default
    if "PREMION BRAND LIFT STUDY" in upper or "Brand Lift Campaign Results" in text:
        return "brand_lift", current_default
    if "CRM data" in text:
        return "sales_attribution", current_default
    if "Regional Furniture Store" in text:
        return "case_study:retail", current_default
    if "ADVANCED ATTRIBUTION INSIGHTS" in upper:
        return "vertical:travel", current_default

    # --- Priority 2c: the core content slides -----------------------------
    # The handful of slides the "Standard" preset includes on top of the
    # always-on ones. They'd otherwise fall through to the ambient
    # "full_deck" default, which is shared by eleven slides and so can't
    # single them out -- and singling them out by *number* is exactly what
    # breaks the moment the master deck is updated.
    if _contains_all(text, "Over 125", "Media Brands"):
        return "core:premium_content", current_default
    if "Solutions for Every Step" in text:
        return "core:attribution_overview", current_default
    # "Transparent Reporting" alone also appears in the Live Sports "Why
    # Choose PREMION" copy, which would misclassify a sports slide -- the
    # same anchor-collision hazard as "CONF" matching "confidential".
    if _contains_all(text, "Transparent Reporting", "DETAILED REPORTING"):
        return "core:reporting", current_default
    if "Exposed Visitors" in text:
        return "core:web_attribution", current_default

    # --- Priority 2b: generic per-vertical anchor -----------------------
    for needle, key in VERTICAL_ANCHORS:
        if needle in upper:
            return key, current_default

    # --- Priority 3: section dividers + subsection markers --------------
    stripped = text.strip()
    for divider_text, default_key in SECTION_DIVIDERS:
        if stripped == divider_text:
            return SECTION_DIVIDER_KEY, default_key

    if "puts your brand alongside trusted journalism" in text:
        return "tegna_positioning", "tegna_positioning"
    if "product solution suite was designed" in text:
        return "tegna_positioning", "tegna_positioning"
    if "AWARENESS || ENGAGEMENT || CONVERSION" in text:
        return "tegna_positioning", "tegna_positioning"
    if "utilizes precision audience targeting, geofencing and retargeting" in text:
        return "am", "am"
    if "AUDIENCE MARKETPLACE ///" in text:
        return "am", "am"
    if "HOW WE DO IT" in upper and "RETARGETING" in upper:
        return "am:retargeting", current_default
    if "HOW WE DO IT" in upper and "AUDIENCE TARGETING" in upper:
        return "am:audience", current_default
    if "HOW WE DO IT" in upper and "GEOFENCING" in upper:
        return "am:geofencing", current_default
    if "Formats & Screens" in text:
        return "am", current_default
    if "HOW WE DO IT" in upper and "REPORTING" in upper:
        return "am", current_default
    if "Reach Sports Fans" in text:
        return "sports", "sports"
    if "SPORTS CONTENT PACKAGES" in upper:
        return "sports", "sports"

    # --- Priority 4: sports viewership / package blocks ------------------
    if "LIVE SPORTS VIEWERSHIP" in upper:
        if "Live Sports Viewers" in text:
            # the one generic overview slide, not tied to a specific sport
            return "sports_viewership_intro", current_default
        sv_key = _classify_sport_viewership(upper)
        return (sv_key or "sports_viewership_unmapped"), current_default
    if "SPORTS PACKAGES" in upper:
        sport_key = _classify_sport_package(text)
        if sport_key:
            return sport_key, current_default

    # --- Priority 5: Total TV specifics ----------------------------------
    if "TV SCHEDULE PLACEHOLDER" in upper:
        return "total_tv", "total_tv"
    if "Total TV is Still King" in text:
        return "total_tv", "total_tv"
    if "WUSA + PREMION" in text:
        return "total_tv:dc", current_default
    if "WPMT + PREMION" in text:
        return "total_tv:harrisburg", current_default

    # Nothing matched. The caller substitutes the carry-forward ambient
    # section, but returns None here so it can tell the difference between
    # "this slide identified itself" and "we guessed from its neighbours" --
    # which is exactly what the deck update page's gate turns on.
    return None, current_default


def classify_slide(text, current_default):
    """Wraps _classify_slide_raw, appending ':targeting' to any 'vertical:X'
    result whose slide text reads as the vertical's own "PRECISION
    TARGETING" page (as opposed to its stats/attribution/case-study
    slides). This lets the caller swap that one slide out in favor of the
    personalized targeting_avails_template, per the vertical specialty
    targeting/avails mutual-exclusion rule.
    """
    key, new_default = _classify_slide_raw(text, current_default)
    if key and key.startswith("vertical:") and ":targeting" not in key:
        upper = text.upper()
        if "PRECISION" in upper and "TARGETING" in upper:
            key = key + ":targeting"
    return key, new_default


def notes_key(slide):
    """The condition_key labelled in this slide's speaker notes, or None.

    An explicit `key: <condition_key>` line always beats the text anchors:
    anchors infer a slide's role from its prose, which is exactly what
    changes when someone retitles a slide or rewrites a heading. A label
    travels with the slide instead. Kept tolerant of spacing and case
    because it's hand-edited in PowerPoint, and the last such line wins so a
    correction appended below an old line does the obvious thing.
    """
    if not slide.has_notes_slide:
        return None
    found = None
    for line in slide.notes_slide.notes_text_frame.text.splitlines():
        match = NOTES_KEY_LINE.match(line)
        if match:
            found = match.group(1)
    return found


def build_slide_map_from_prs(prs):
    """Scan an already-open Presentation and return {slide_number: condition_key}.

    A slide's own `key:` notes label wins; text anchors resolve anything
    untagged, so a partly-tagged deck (a new slide dropped into a tagged
    master, say) still works throughout.

    Slides that resolve to None are printed as warnings -- they need either a
    notes label or a new anchor rule added above.
    """
    result = {key: value for key, (value, _) in _classify_all(prs).items()}
    return result


SOURCE_NOTES = "notes"          # an explicit key: label in the speaker notes
SOURCE_ANCHOR = "anchor"        # a text-anchor rule matched this slide directly
SOURCE_CARRY_FORWARD = "carry"  # neither -- inherited from the current section


def _classify_all(prs, verbose=True):
    """{slide_number: (condition_key, source)} for a deck.

    The source matters as much as the key. A slide resolved by
    SOURCE_CARRY_FORWARD didn't identify itself at all -- it inherited
    whatever section it happens to sit in, which is a guess that reads as an
    answer. That's fine for a deck built on the v1.1 conventions, and
    unacceptable for a slide someone just added, so the deck update page
    treats carry-forward as "needs a key".
    """
    result = {}
    current_default = "always"
    guessed = []

    for i, slide in enumerate(prs.slides, start=1):
        text = extract_slide_text(slide)
        # Anchors run even when a label is present, because the carry-forward
        # default has to keep tracking the current section for any *untagged*
        # slide further down the deck.
        anchor_key, current_default = classify_slide(text, current_default)
        label = notes_key(slide)

        if label:
            result[i] = (label, SOURCE_NOTES)
        elif anchor_key:
            result[i] = (anchor_key, SOURCE_ANCHOR)
        else:
            result[i] = (current_default, SOURCE_CARRY_FORWARD)
            guessed.append((i, text[:80]))

    if guessed and verbose:
        print(f"WARNING: {len(guessed)} slide(s) have no key: label and match no anchor "
              f"-- they inherited the surrounding section:")
        for n, snippet in guessed:
            print(f"  slide {n}: {snippet!r}")

    return result


def build_slide_map(pptx_path):
    """Scan pptx_path (opening it fresh) and return {slide_number: condition_key}."""
    return build_slide_map_from_prs(Presentation(pptx_path))


def slide_fingerprints(prs):
    """{slide_number: (key, source, text_fingerprint, title)} for one deck.

    The fingerprint is the slide's whole text, which is what lets two deck
    versions be compared without relying on slide numbers -- the very thing
    an update changes.
    """
    result = {}
    classified = _classify_all(prs, verbose=False)
    for number, slide in enumerate(prs.slides, start=1):
        text = extract_slide_text(slide)
        key, source = classified[number]
        result[number] = (key, source, text, _slide_label(slide, text))
    return result


def _slide_label(slide, text):
    """A short human label for a slide -- the largest explicitly-sized text on
    it, which is its heading far more reliably than shape order is (most
    slides here have no title placeholder and open with a stray bullet)."""
    biggest, biggest_size = None, 0
    for shape in iter_all_shapes(slide.shapes):
        if not shape.has_text_frame:
            continue
        shape_text = " ".join(shape.text_frame.text.split())
        if not shape_text:
            continue
        sizes = [run.font.size.pt
                 for para in shape.text_frame.paragraphs
                 for run in para.runs if run.font.size]
        if sizes and max(sizes) > biggest_size:
            biggest, biggest_size = shape_text, max(sizes)
    return (biggest or " ".join(text.split()) or "(no text)")[:80]


def diff_decks(old_prs, new_prs):
    """Compare two deck versions. Returns a dict describing what changed.

    Slides are matched on their text, not their position, because position is
    exactly what a deck update churns. A slide that kept its text but moved
    is "moved"; one whose text is gone is "removed"; new text is "added".

    `unresolved` is the list that gates activation: slides that didn't
    identify themselves at all and merely inherited the section around them.
    That inheritance is a guess dressed as an answer -- fine for the slides
    the v1.1 anchors were written against, wrong for one somebody just added,
    which would otherwise be silently assembled into decks under whatever
    key its neighbour happened to have.
    """
    old = slide_fingerprints(old_prs) if old_prs is not None else {}
    new = slide_fingerprints(new_prs)

    old_by_text = {}
    for number, (key, source, text, label) in old.items():
        old_by_text.setdefault(text, []).append(number)

    added, moved, retagged, matched_old = [], [], [], set()
    for number, (key, source, text, label) in sorted(new.items()):
        candidates = old_by_text.get(text)
        if not candidates:
            added.append({"number": number, "key": key, "label": label})
            continue
        old_number = candidates.pop(0)
        matched_old.add(old_number)
        old_key = old[old_number][0]
        if old_number != number:
            moved.append({"number": number, "was": old_number, "key": key, "label": label})
        if old_key != key:
            retagged.append({"number": number, "label": label,
                             "was_key": old_key, "key": key})

    removed = [{"number": number, "key": old[number][0], "label": old[number][3]}
               for number in sorted(set(old) - matched_old)]

    unresolved = [{"number": number, "label": label, "guessed_key": key}
                  for number, (key, source, text, label) in sorted(new.items())
                  if source == SOURCE_CARRY_FORWARD]

    return {
        "old_count": len(old), "new_count": len(new),
        "added": added, "removed": removed, "moved": moved,
        "retagged": retagged, "unresolved": unresolved,
        "unchanged": len(matched_old) - len(moved),
    }


if __name__ == "__main__":
    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else "TEGNA_MASTER_DECK_v1_1.pptx"
    result = build_slide_map(path)
    for n in sorted(result):
        print(n, "->", result[n])
