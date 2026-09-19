import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

SRC, OUT = "REPORT_MASTER_v0_11.pptx", "REPORT_MASTER_v0_12.pptx"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
prs = Presentation(SRC)
def key_of(s): return s.notes_slide.notes_text_frame.text
def find(key): return [s for s in prs.slides if f"key: {key}" in key_of(s)][0]

def clone_after(src, after_key, notes):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    rid = {r.rId: new.part.relate_to(r.target_part, RT.IMAGE) for r in src.part.rels.values() if r.reltype == RT.IMAGE}
    for sh in src.shapes:
        el = copy.deepcopy(sh._element)
        for node in el.iter():
            for attr in ("{%s}embed" % NS["r"], "{%s}link" % NS["r"]):
                if attr in node.attrib and node.attrib[attr] in rid: node.attrib[attr] = rid[node.attrib[attr]]
        new.shapes._spTree.append(el)
    new.notes_slide.notes_text_frame.text = notes
    lst = prs.slides._sldIdLst; last = list(lst)[-1]
    idx = [i for i, s in enumerate(prs.slides) if f"key: {after_key}" in key_of(s)][0]
    lst.remove(last); lst.insert(idx + 1, last)
    return new

s = clone_after(find("report:live_sports"), "report:takeaways", "key: report:automotive_registrations")
def shape(n):
    for sh in s.shapes:
        if sh.name == n: return sh
    raise KeyError(n)
def drop(*ns):
    for n in ns:
        try: sh = shape(n); sh._element.getparent().remove(sh._element)
        except KeyError: pass
def set_text(sh, t):
    tf = sh.text_frame; p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = t
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
    else: p.add_run().text = t
    for e in tf.paragraphs[1:]: e._p.getparent().remove(e._p)
def cell_set(c, t):
    p = c.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = t
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
    elif t: p.add_run().text = t
def place(sh, x, y, w=None, h=None):
    sh.left, sh.top = Inches(x), Inches(y)
    if w is not None: sh.width = Inches(w)
    if h is not None: sh.height = Inches(h)
def clone(n, new):
    el = copy.deepcopy(shape(n)._element); s.shapes._spTree.append(el)
    sh = s.shapes[-1]; sh.name = new; return sh

set_text(shape("Text 1"), "Automotive Registrations")
set_text(shape("Text 2"), "Polk new-vehicle registrations matched to households exposed to the campaign")

# ---- four tiles: rename three, clone a fourth ----
ren = [("SportsImpressionsTile", "PolkHouseholdsTile", "{{POLK_MATCHED_HOUSEHOLDS}}", "Matched households"),
       ("SportsVcrTile",         "PolkSalesTile",      "{{POLK_TARGET_DEALER_SALES}}", "Target dealer sales"),
       ("SportsPacingTile",      "PolkBuyRateTile",    "{{POLK_BUY_RATE}}",            "Buy rate")]
for old, new, tok, lab in ren:
    for suf in ("", "Value", "Label"): shape(old + suf).name = new + suf
    set_text(shape(new + "Value"), tok); set_text(shape(new + "Label"), lab)
for suf in ("", "Value", "Label"): clone("PolkBuyRateTile" + suf, "PolkLiftTile" + suf)
set_text(shape("PolkLiftTileValue"), "{{POLK_CAMPAIGN_LIFT}}"); set_text(shape("PolkLiftTileLabel"), "Campaign lift")
names = ["PolkHouseholdsTile", "PolkSalesTile", "PolkBuyRateTile", "PolkLiftTile"]
x0, span, gap, y, h = 0.50, 12.33, 0.25, 2.10, 0.85
tw = (span - 3 * gap) / 4
for i, n in enumerate(names):
    x = x0 + i * (tw + gap)
    place(shape(n), x, y, tw, h)
    place(shape(n + "Value"), x + 0.10, y + 0.08, tw - 0.20, 0.48)
    place(shape(n + "Label"), x + 0.10, y + 0.56, tw - 0.20, 0.25)
# room for "49,461 (projected)" — 18 chars in a 2.9in tile: 18pt fits, 20pt doesn't
for n in ("PolkHouseholdsTileValue", "PolkSalesTileValue"):
    for p in shape(n).text_frame.paragraphs:
        for r in p.runs: r.font.size = Pt(18)

# ---- match-rate note under the tiles ----
# SportsRfpidCaption's own color (C9CBE8, a light lavender) is meant for
# where it sits on report:live_sports; carried over unchanged here it read
# as near-invisible on white. Matched to PlanVsActualNote's dark 1A1A2E
# instead -- the same caption-under-a-KPI-row precedent this shape is
# built to match, and the one real color already proven readable on white
# in this exact position on another slide. Found by rendering, not by
# reading the fill code -- see REPORT_MASTER_README.md.
note = clone("SportsRfpidCaption", "PolkMatchRateNote")
set_text(note, "{{POLK_MATCH_RATE_NOTE}}"); place(note, 0.50, 3.05, 12.33, 0.28)
for p in note.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
drop("SportsRfpidCaption")

# ---- dealer table, left ----
hdr = shape("Text 18"); hdr.name = "PolkTargetDealersHeader"
set_text(hdr, "TARGET DEALER PERFORMANCE"); place(hdr, 0.50, 3.55, 7.70, 0.26)
tbl_sh = shape("SportsEventTable"); tbl_sh.name = "PolkTargetDealersTable"
tbl, el = tbl_sh.table, tbl_sh.table._tbl
grid = el.find("a:tblGrid", NS)
for gc in grid.findall("a:gridCol", NS)[3:]: grid.remove(gc)
for tr in el.findall("a:tr", NS):
    for tc in tr.findall("a:tc", NS)[3:]: tr.remove(tc)
widths = [4.30, 1.70, 1.70]
for gc, w in zip(grid.findall("a:gridCol", NS), widths): gc.set("w", str(Inches(w)))
place(tbl_sh, 0.50, 3.85, sum(widths))
for c, t in zip(list(tbl.rows[0].cells), ["Dealer", "Market rank", "Campaign rank"]): cell_set(c, t)
row = list(tbl.rows[1].cells); cell_set(row[0], "{{POLK_TARGET_DEALER_ROWS}}")
for c in row[1:]: cell_set(c, "")
for r in tbl.rows:
    for i, c in enumerate(r.cells):
        c.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT if i == 0 else PP_ALIGN.RIGHT

# ---- narrative, right ----
drop("SportsByLeagueHeader", "SportsByLeagueTable")
nar = shape("LiveSportsNarrative"); nar.name = "PolkNarrative"
set_text(nar, "{{POLK_NARRATIVE}}"); place(nar, 8.45, 3.55, 4.38, 3.10)

prs.save(OUT)
print("saved", OUT, [key_of(x).split("\n")[0].replace("key: report:", "") for x in prs.slides][10:])
