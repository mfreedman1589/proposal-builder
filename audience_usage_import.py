"""audience_usage_import.py -- parse Premion's YTD audience-usage workbook
and derive catalog updates from it.

Pure module: no Streamlit, no DB -- `db.py`'s admin-page functions call
this, and it's testable offline the same way `wideorbit.py` is. `parse_workbook`
raises `UsageImportError` (a message aimed at whoever's uploading the file,
not a developer) rather than guessing at a malformed row, the same
philosophy `wideorbit.ScheduleParseError` uses for a broadcast schedule --
guessing wrong here silently changes which audiences look popular.

The workbook is one sheet, two columns (Segment Name / Delivered
Impressions), one row per BOOKED STACK -- the segment cell is a
comma-joined AND-stack, exactly the shape `audience_usage_ytd.csv` already
has (that CSV was hand-derived from an earlier cut of this same workbook).
"No Data Targeting" is a delivery-log label for untargeted impressions, not
an audience -- dropped here, before it can reach any downstream table.

**CLT exclusion.** A `CLT`-prefixed component is a client's own first-party
data list -- usable only for the client it belongs to, never selectable for
anyone else. `derive_catalog_updates` drops CLT components entirely before
normalization, before categorization, before anything else, so no catalog
row is ever built from one. The raw usage rows handed to
`db.replace_audience_usage` keep the full original stack text (CLT and
all) -- that table is the historical transcript of what was actually
booked, and CLT exclusion is enforced at every CONSUMING boundary instead
(this module's catalog derivation, and `audience_evidence.parse_stack`),
not by rewriting the transcript.

**Merge, don't replace.** `derive_catalog_updates` takes the EXISTING
catalog (brochure-derived, or already merged from an earlier workbook) and
only ever ADDS to it: a component that normalizes to an existing catalog
segment keeps that segment's exact spelling, category and
`rfp_selectable` (the brochure is authoritative for RFP-selectable, which
`targeting_groups.custom_segment_count` depends on) and only gets its
usage numbers refreshed; a component the catalog lacks becomes a new row,
`rfp_selectable=False` by default (the safe default -- an unknown
component must never silently count as RFP-selectable), category inferred
from its own prefix or, for a CUSTOM-prefixed one, from a second prefix
embedded in the name ("CUSTOM FOOD Grocery Delivery Services" -> FOOD).
Nothing is guessed past that: a component with no recognizable prefix at
all is reported uncategorized rather than forced into one of the 14.

**Client-retargeting exclusion, mechanical and manual.** A second class of
component is excluded the same way CLT is, for the same reason -- it's one
advertiser's own retargeting or address-list pool, not a reusable
audience -- but has no CLT prefix to key off. Most of these follow one of
four structural patterns (`is_client_pattern`: `WEB RT`, `LOCATION RT`,
`..._RFPID-..._RT`, an "Address List" suffix), checked automatically on
every component so a future workbook needs no manual pass for THESE. The
rest -- a bare client/campaign name with no structural marker at all
("CUSTOM Out West", "CUSTOM City of Mesa") -- can't be told apart from a
legitimately reusable custom segment by pattern alone, and live in
`audience_component_overrides.csv` (`action=exclude_client`) instead,
alongside the other three manual actions the same registry holds:
`categorize` (a category this module couldn't derive, e.g. "Custom
Personal Injury Lawyer Intender" -> LEGAL), `split` (an OR-joined compound
component that's really two -- see `_apply_split` below), and
`leave_blank` (confirmed uncategorizable, so it stops re-surfacing in
every future uncategorized report as if it were new). `load_overrides`
reads that file once; `derive_catalog_updates` calls it automatically
when no `overrides` argument is given, so this is baked into the pipeline
permanently rather than a one-time script.
"""
import csv
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from audience_catalog import CATEGORY_PREFIXES
from market_profiles import _normalize as normalize

NO_DATA_TARGETING = "no data targeting"

DEFAULT_OVERRIDES_PATH = Path(__file__).parent / "audience_component_overrides.csv"

# A component is CLT/CUSTOM by its own leading word, matched case-
# insensitively -- the workbook isn't consistent about case ("CUSTOM
# Blue Bell Ice Cream Shoppers" vs "Custom Personal Injury Lawyer
# Intender"), and treating those as two different prefixes would let a
# title-cased custom segment slip through as if it were a real catalog
# category. The lookahead is `[_\s]` or end-of-string, not `\b` -- a plain
# word boundary doesn't fire between "m" and "_" (both are word
# characters), which missed "Custom_Doctors and Nurses" entirely on a real
# pass over the workbook and left it uncategorized instead of recognized
# as custom.
_CLT_RE = re.compile(r"^CLT(?=[_\s]|$)", re.IGNORECASE)
_CUSTOM_RE = re.compile(r"^CUSTOM(?=[_\s]|$)", re.IGNORECASE)
_CUSTOM_STRIP_RE = re.compile(r"^CUSTOM[_\s]*", re.IGNORECASE)

# Four structural forms of "this is one advertiser's own retargeting/address
# pool" -- not anchored to start/end, since the real workbook has both
# "WEB RT- Trane LMG" (prefix) and "Rutter Mills _Web RT" (suffix), and a
# plain \b boundary doesn't fire before an underscore (same trap the
# CLT/CUSTOM regexes above had, and for the same reason: underscore is a
# \w character). Verified against the real markup: 0 false positives across
# all 106 previously-uncategorized components, 29 of 41 confirmed
# exclude_client rows caught without a single manual entry.
_CLIENT_PATTERN_RE = re.compile(
    r"(WEB\s*RT|LOCATION\s*RT|_RFPID-[^_]*_RT|ADDRESS\s*LIST)", re.IGNORECASE)


class UsageImportError(Exception):
    """Raised with a message aimed at whoever's uploading the workbook."""


@dataclass
class UsageRow:
    segment: str
    impressions: int


@dataclass
class WorkbookImport:
    rows: list                     # [UsageRow, ...] -- "No Data Targeting" already dropped
    dropped_no_data_targeting: int
    sheet_name: str


@dataclass
class ComponentUpdate:
    """One row to upsert into `audiences`. `segment` is either the EXISTING
    catalog's own exact spelling (a brochure/prior-merge collision) or a
    representative display spelling picked from the workbook's own raw
    variants (a net-new component) -- never the workbook's raw text
    verbatim when a catalog spelling already exists, which is exactly the
    trap that would split one segment's history across two near-identical
    rows.
    """
    segment: str
    category: str
    rfp_selectable: bool
    times_used: int                # distinct stacks this component appears in
    impressions: int                # total delivered impressions attributed to it
    is_new: bool                    # not already in the catalog handed in
    source: str                     # "existing" | "prefix" | "custom-prefix" |
                                     # "override-categorize" | "override-blank" | "uncategorized"


@dataclass
class MergeReport:
    updates: list                        # [ComponentUpdate, ...] to upsert -- CLT/client-excluded already dropped
    gained: int                          # net-new segments (is_new)
    collided: int                        # normalized matches to an existing catalog segment
    uncategorized: int                   # updates with source == "uncategorized" (NOT override-blank --
                                          # that's a confirmed decision, not a gap to keep reporting)
    uncategorized_examples: list         # display segment strings, capped
    unusual_prefixes: dict               # {first-token: count} among the uncategorized ones
    excluded_clt: list                   # every distinct CLT-prefixed raw component, sorted
    excluded_clt_stack_appearances: int  # total (row, CLT component) occurrences
    clt_aliases: list                    # [(clt_raw, matching_non_clt_segment), ...] -- should be empty
    excluded_client_pattern: list        # raw components matched by is_client_pattern, sorted
    excluded_client_override: list       # raw components excluded via the manual registry, sorted --
                                          # the pattern couldn't catch these
    custom_case_variants: list           # raw CUSTOM-prefixed strings not spelled exactly "CUSTOM ..."
    overrides_applied: dict              # {action: count} across categorize/split/leave_blank/exclude_client
    usage_rows: list                     # [{"segment_string", "delivered_impressions"}, ...] for
                                          # db.replace_audience_usage -- the FULL transcript, CLT/client-
                                          # excluded and pre-split text all included verbatim


def is_clt(raw_component):
    return bool(_CLT_RE.match(raw_component.strip()))


def is_custom(raw_component):
    return bool(_CUSTOM_RE.match(raw_component.strip()))


def is_client_pattern(raw_component):
    """One of the four structural markers of a single advertiser's own
    retargeting or address-list pool -- see this module's own docstring
    and `_CLIENT_PATTERN_RE`'s comment for what's covered and why it's a
    search, not an anchored match."""
    return bool(_CLIENT_PATTERN_RE.search(raw_component))


def _first_token(text):
    parts = text.split(None, 1)
    return parts[0] if parts else ""


def load_overrides(path=DEFAULT_OVERRIDES_PATH):
    """{normalized component: {"action", "category", "split_into"}} from
    the committed overrides registry -- exact-name decisions a human made
    because they can't be derived mechanically (a category with no
    recognizable prefix, an OR-joined component that's really two things,
    a client name with no structural marker to catch it). Missing file
    degrades to no overrides at all, same "never raise on a missing local
    fallback" policy every other loader in this app follows -- an empty
    registry just means every component falls through to mechanical
    categorization, which is exactly what happened before this existed.
    """
    path = Path(path)
    if not path.exists():
        return {}
    overrides = {}
    try:
        with open(path, encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                component = (row.get("component") or "").strip()
                if not component:
                    continue
                split_into = [p.strip() for p in (row.get("split_into") or "").split("|") if p.strip()]
                overrides[normalize(component)] = {
                    "action": (row.get("action") or "").strip(),
                    "category": (row.get("category") or "").strip(),
                    "split_into": split_into,
                }
    except (OSError, csv.Error):
        return {}
    return overrides


def parse_workbook(source):
    """`source` is a path or a file-like/bytes buffer openpyxl can open.
    Returns a `WorkbookImport`, or raises `UsageImportError`.
    """
    import openpyxl

    try:
        book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise UsageImportError(f"Couldn't open that file as an Excel workbook ({exc}).") from exc

    if not book.sheetnames:
        raise UsageImportError("That workbook has no sheets.")
    sheet_name = book.sheetnames[0]
    ws = book[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise UsageImportError("That workbook's sheet is empty.")

    # A leading BOM shows up on the SHEET NAME in the real file, not (so
    # far) on any cell -- stripped here too, defensively, since Excel
    # exports are exactly the kind of thing that varies without warning.
    header_cells = [str(c).strip().lstrip("﻿") if c is not None else "" for c in header]
    if (len(header_cells) < 2 or "segment" not in header_cells[0].lower()
            or "impression" not in header_cells[1].lower()):
        raise UsageImportError(
            f"Expected two columns headed something like \"Segment Name\" / "
            f"\"Delivered Impressions\" -- got {header_cells!r}. This doesn't look like "
            f"the usage workbook.")

    rows, dropped = [], 0
    for i, raw in enumerate(rows_iter, start=2):
        if raw is None or all(cell is None for cell in raw):
            continue  # a genuinely blank row -- Excel exports carry these sometimes
        segment = raw[0] if len(raw) > 0 else None
        impressions = raw[1] if len(raw) > 1 else None
        if segment is None or str(segment).strip() == "":
            raise UsageImportError(f"Row {i} has no segment name.")
        segment = str(segment).strip()
        if segment.lower() == NO_DATA_TARGETING:
            dropped += 1
            continue
        if impressions is None or (isinstance(impressions, str) and not impressions.strip()):
            raise UsageImportError(f"Row {i} ({segment!r}) has no impressions figure.")
        try:
            impressions = int(impressions)
        except (TypeError, ValueError):
            raise UsageImportError(
                f"Row {i} ({segment!r})'s impressions figure ({impressions!r}) isn't a number.")
        rows.append(UsageRow(segment=segment, impressions=impressions))

    if not rows:
        raise UsageImportError("No usable rows found in that workbook -- every row was blank "
                                "or \"No Data Targeting\".")
    return WorkbookImport(rows=rows, dropped_no_data_targeting=dropped, sheet_name=sheet_name)


def _infer_category(raw_variants, existing_by_norm, overrides):
    """(category, rfp_selectable, existing_segment_or_None, source) for one
    normalized group of raw component spellings. `existing_by_norm` is
    {normalized: {"segment","category","rfp_selectable"}} for the catalog
    handed in -- checked FIRST, since an existing catalog entry is
    authoritative over anything this function would otherwise infer (and
    self-reinforcing across successive workbook uploads: the first upload
    applies an override and writes the category; the NEXT upload sees it
    as "existing" and never needs the override again). A manual
    `categorize`/`leave_blank` override is checked next, before mechanical
    prefix inference -- it exists specifically for what the mechanical
    rules below get wrong or can't reach.
    """
    key = normalize(raw_variants[0])
    existing = existing_by_norm.get(key)
    if existing is not None:
        return existing["category"], existing["rfp_selectable"], existing["segment"], "existing"

    for raw in raw_variants:
        override = overrides.get(normalize(raw))
        if override and override["action"] == "categorize":
            return override["category"], False, None, "override-categorize"
        if override and override["action"] == "leave_blank":
            return "", False, None, "override-blank"

    for raw in raw_variants:
        token = _first_token(raw).upper()
        if token in CATEGORY_PREFIXES:
            return token, False, None, "prefix"

    for raw in raw_variants:
        if is_custom(raw):
            rest = _CUSTOM_STRIP_RE.sub("", raw.strip()).strip()
            second = _first_token(rest).upper()
            if second in CATEGORY_PREFIXES:
                return second, False, None, "custom-prefix"

    return "", False, None, "uncategorized"


def _pick_display(raw_variants, stack_counts_by_raw):
    """The raw spelling to show for a net-new segment -- the one seen in
    the most distinct stacks, alphabetical as the tiebreak so the choice is
    deterministic rather than dependent on dict/set iteration order.
    """
    return sorted(raw_variants, key=lambda r: (-len(stack_counts_by_raw[r]), r))[0]


def clt_alias_matches(clt_raws, non_clt_raws):
    """[(clt_raw, matching_non_clt_raw), ...] -- a CLT component whose text,
    with the leading "CLT" / "CLT 1P" stripped and normalized, matches some
    OTHER, non-CLT component's normalized text exactly. Checked so CLT
    exclusion can be verified rather than assumed: if a client's first-party
    list is really just a renamed copy of a real catalog segment, dropping
    it outright would be the wrong call. Empty when nothing aliases, which
    is the expected, checked-for-real case (see this module's own tests).
    """
    non_clt_by_norm = {}
    for raw in non_clt_raws:
        non_clt_by_norm.setdefault(normalize(raw), raw)
    hits = []
    for clt_raw in clt_raws:
        rest = re.sub(r"^CLT[_\s]*(1P[_\s]*)?", "", clt_raw.strip(), flags=re.IGNORECASE).strip()
        key = normalize(rest)
        if key and key in non_clt_by_norm:
            hits.append((clt_raw, non_clt_by_norm[key]))
    return hits


def derive_catalog_updates(workbook_import, existing_catalog, overrides=None):
    """The full merge: catalog rows to upsert, plus a report of what
    happened. `existing_catalog` is an iterable of PLAIN DICTS carrying at
    least `segment`, `category`, `rfp_selectable` -- `db.fetch_audiences()`
    rows already are; a DataFrame needs `.to_dict("records")` first.
    `overrides` defaults to `load_overrides()` (the committed registry) --
    pass `{}` explicitly to see what the mechanical rules alone would do.
    """
    if overrides is None:
        overrides = load_overrides()

    existing_by_norm = {}
    for row in existing_catalog or []:
        seg = str(row.get("segment") or "").strip()
        if not seg:
            continue
        existing_by_norm[normalize(seg)] = {
            "segment": seg,
            "category": row.get("category") or "",
            "rfp_selectable": bool(row.get("rfp_selectable")),
        }

    raw_stacks = defaultdict(set)          # raw component -> {stack index, ...}
    raw_impressions = Counter()            # raw component -> total impressions
    clt_stacks = defaultdict(set)          # CLT raw component -> {stack index, ...}
    client_pattern_stacks = defaultdict(set)   # mechanically excluded -> {stack index, ...}
    client_override_stacks = defaultdict(set)  # manually excluded (registry) -> {stack index, ...}
    usage_rows = []
    overrides_applied = Counter()

    for i, row in enumerate(workbook_import.rows):
        usage_rows.append({"segment_string": row.segment, "delivered_impressions": row.impressions})
        parts = [p.strip() for p in row.segment.split(",") if p.strip()]
        for part in parts:
            if is_clt(part):
                clt_stacks[part].add(i)
                continue
            if is_client_pattern(part):
                client_pattern_stacks[part].add(i)
                continue
            override = overrides.get(normalize(part))
            action = override["action"] if override else None
            if action == "exclude_client":
                client_override_stacks[part].add(i)
                overrides_applied["exclude_client"] += 1
                continue
            if action == "split":
                # Attributed exactly like a real comma-stack: EACH split
                # piece gets the FULL stack membership and impressions of
                # the compound phrase it replaces, not a share of it --
                # the same rule two components sitting in one real,
                # comma-separated stack already follow. The compound
                # phrase itself never becomes its own raw_stacks entry, so
                # it can't ALSO produce a row for the un-split text.
                for piece in override["split_into"]:
                    raw_stacks[piece].add(i)
                    raw_impressions[piece] += row.impressions
                overrides_applied["split"] += 1
                continue
            raw_stacks[part].add(i)
            raw_impressions[part] += row.impressions

    groups = defaultdict(list)             # normalized -> [raw, raw, ...]
    for raw in raw_stacks:
        groups[normalize(raw)].append(raw)

    updates = []
    gained = collided = uncategorized = 0
    uncategorized_examples = []
    unusual_prefixes = Counter()
    custom_case_variants = []

    for key, raws in groups.items():
        stacks_union = set()
        for raw in raws:
            stacks_union |= raw_stacks[raw]
        times_used = len(stacks_union)
        impressions = sum(raw_impressions[raw] for raw in raws)

        category, rfp_selectable, existing_segment, source = _infer_category(
            raws, existing_by_norm, overrides)
        if source == "existing":
            segment = existing_segment
            collided += 1
            is_new = False
        else:
            segment = _pick_display(raws, raw_stacks)
            gained += 1
            is_new = True
            if source == "override-categorize":
                overrides_applied["categorize"] += 1
            elif source == "override-blank":
                overrides_applied["leave_blank"] += 1
            elif source == "uncategorized":
                uncategorized += 1
                uncategorized_examples.append(segment)
                unusual_prefixes[_first_token(segment).upper()] += 1

        for raw in raws:
            if is_custom(raw) and not raw.strip().startswith("CUSTOM"):
                custom_case_variants.append(raw)

        updates.append(ComponentUpdate(
            segment=segment, category=category, rfp_selectable=rfp_selectable,
            times_used=times_used, impressions=impressions, is_new=is_new, source=source,
        ))

    excluded_clt = sorted(clt_stacks.keys())
    excluded_clt_appearances = sum(len(v) for v in clt_stacks.values())
    clt_aliases = clt_alias_matches(excluded_clt, raw_stacks.keys())

    return MergeReport(
        updates=updates, gained=gained, collided=collided,
        uncategorized=uncategorized, uncategorized_examples=sorted(uncategorized_examples)[:200],
        unusual_prefixes=dict(unusual_prefixes.most_common()),
        excluded_clt=excluded_clt, excluded_clt_stack_appearances=excluded_clt_appearances,
        clt_aliases=clt_aliases,
        excluded_client_pattern=sorted(client_pattern_stacks.keys()),
        excluded_client_override=sorted(client_override_stacks.keys()),
        custom_case_variants=sorted(set(custom_case_variants)),
        overrides_applied=dict(overrides_applied),
        usage_rows=usage_rows,
    )
