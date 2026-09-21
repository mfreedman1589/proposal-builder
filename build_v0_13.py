import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SRC, OUT = "REPORT_MASTER_v0_12.pptx", "REPORT_MASTER_v0_13.pptx"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
DARK = RGBColor(0x1A, 0x1A, 0x2E)
prs = Presentation(SRC)

def slide(key):
    return [s for s in prs.slides if f"key: {key}" in s.notes_slide.notes_text_frame.text][0]
def shape(s, n):
    for sh in s.shapes:
        if sh.name == n: return sh
    raise KeyError(n)
def set_text(sh, t):
    tf = sh.text_frame; p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = t
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
    else: p.add_run().text = t
    for e in tf.paragraphs[1:]: e._p.getparent().remove(e._p)
def cell_set(c, t):
    p = c.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = t
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
    elif t: p.add_run().text = t
def place(sh, x, y, w=None, h=None):
    sh.left, sh.top = Inches(x), Inches(y)
    if w is not None: sh.width = Inches(w)
    if h is not None: sh.height = Inches(h)
def clone(src_slide, src_shape, dst_slide, new):
    el = copy.deepcopy(src_shape._element); dst_slide.shapes._spTree.append(el)
    sh = dst_slide.shapes[-1]; sh.name = new; return sh
def style_note(sh, size=10, align=PP_ALIGN.LEFT):
    for p in sh.text_frame.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size); r.font.bold = False; r.font.color.rgb = DARK

# ===== report:automotive_registrations =====
a = slide("report:automotive_registrations")
# Households tile -> ROI, BuyRate tile -> MSRP (reuse geometry/styling), then reorder
for suf in ("", "Value", "Label"):
    shape(a, "PolkHouseholdsTile" + suf).name = "PolkRoiTile" + suf
    shape(a, "PolkBuyRateTile" + suf).name = "PolkMsrpTile" + suf
spec = [("PolkSalesTile", "{{POLK_TARGET_DEALER_SALES}}", "Target dealer sales", 18),
        ("PolkMsrpTile",  "{{POLK_MSRP_SOLD}}",           "Total MSRP sold",     16),
        ("PolkLiftTile",  "{{POLK_CAMPAIGN_LIFT}}",       "Campaign lift",       20),
        ("PolkRoiTile",   "{{POLK_ROI}}",                 "ROI",                 20)]
x0, span, gap, y, h = 0.50, 12.33, 0.25, 2.10, 0.85
tw = (span - 3 * gap) / 4
for i, (n, tok, lab, size) in enumerate(spec):
    x = x0 + i * (tw + gap)
    place(shape(a, n), x, y, tw, h)
    v, l = shape(a, n + "Value"), shape(a, n + "Label")
    place(v, x + 0.10, y + 0.08, tw - 0.20, 0.48); place(l, x + 0.10, y + 0.56, tw - 0.20, 0.25)
    set_text(v, tok); set_text(l, lab)
    for p in v.text_frame.paragraphs:
        for r in p.runs: r.font.size = Pt(size)
# match-rate note: longer sentence now, room for two lines
place(shape(a, "PolkMatchRateNote"), 0.50, 3.03, 12.33, 0.42)
style_note(shape(a, "PolkMatchRateNote"), 10)  # re-apply Claude Code's v0_12 colour fix: 1A1A2E, not the inherited C9CBE8
# dealer table 3 -> 4 columns
set_text(shape(a, "PolkTargetDealersHeader"), "WHERE EXPOSED HOUSEHOLDS BOUGHT")
t = shape(a, "PolkTargetDealersTable"); tbl, el = t.table, t.table._tbl
grid = el.find("a:tblGrid", NS)
grid.append(copy.deepcopy(grid.findall("a:gridCol", NS)[-1]))
for tr in el.findall("a:tr", NS): tr.append(copy.deepcopy(tr.findall("a:tc", NS)[-1]))
widths = [4.10, 1.10, 1.60, 0.90]   # 7.70
for gc, w in zip(grid.findall("a:gridCol", NS), widths): gc.set("w", str(Inches(w)))
t.width = Inches(sum(widths))
for c, hd in zip(list(tbl.rows[0].cells), ["Dealer", "Sales", "MSRP sold", ""]): cell_set(c, hd)
row = list(tbl.rows[1].cells); cell_set(row[0], "{{POLK_TARGET_DEALER_ROWS}}")
for c in row[1:]: cell_set(c, "")
for r in tbl.rows:
    for i, c in enumerate(r.cells):
        c.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT if i == 0 else (PP_ALIGN.CENTER if i == 3 else PP_ALIGN.RIGHT)

# ===== report:recap — PolkSalesWindowNote in the gap under the tiles =====
rc = slide("report:recap")
n1 = clone(a, shape(a, "PolkMatchRateNote"), rc, "PolkSalesWindowNote")
set_text(n1, "{{POLK_SALES_WINDOW_NOTE}}"); place(n1, 0.50, 4.43, 12.33, 0.27); style_note(n1, 10)

# ===== report:highlights — CostPerVisitNote under the tiles/deltas =====
hl = slide("report:highlights")
# push the card and trend column down 0.25 to make a caption row
for n, dy, dh in (("Shape 11", 0.25, -0.25), ("Text 12", 0.25, 0), ("HIGHLIGHTBullets", 0.25, -0.25),
                  ("TrendChartHeader", 0.25, 0), ("TrendChartRegion", 0.25, -0.25), ("TrendChartRegionLabel", 0.25, -0.25)):
    sh = shape(hl, n); sh.top += Inches(dy); sh.height += Inches(dh)
n2 = clone(a, shape(a, "PolkMatchRateNote"), hl, "CostPerVisitNote")
set_text(n2, "{{COST_PER_VISIT_NOTE}}"); place(n2, 0.50, 3.63, 12.33, 0.26); style_note(n2, 10)

prs.save(OUT); print("saved", OUT)
