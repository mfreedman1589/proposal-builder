"""Builds the Harrisburg rep user guide (.docx) by driving the real,
running app with Playwright and capturing screenshots for every step, then
assembling them with python-docx -- the same "build it from a script, not
by hand" treatment REPORT_MASTER's own build_v0_N.py scripts get.

    python build_user_guide.py [--task proposal|report|all]

Re-run this after any change to the proposal or report flow -- same
standing treatment as the "walk the page as a rep" rule (CLAUDE.md). The
guide is never hand-edited: change a step's wording in STEPS_PROPOSAL /
STEPS_REPORT below, and the next run picks it up. Screenshots are always
re-captured, never hand-pasted -- there is no way to edit one without
re-running the capture that produced it.

Requires:
  - playwright + its chromium browser: `pip install playwright &&
    playwright install chromium`. Test tooling only, not a runtime
    dependency of the app itself (same footing as pywin32 in
    requirements.txt).
  - The test-mode app, started automatically here via
    tests/app_process.py on port 8517 (PROPOSAL_BUILDER_TEST_MODE=1) --
    the sidebar identity is "Test Mode", never a real rep's name, and
    nothing this script does writes to Supabase under any other identity.
  - Two real fixtures already in the repo root (gitignored, same as every
    other real client file this project tests against): the Visit
    Hershey & Harrisburg avails PDF for Task 1 (the Harrisburg market's
    own real avail), and WAEPA's real attribution export
    ("...Reach Extension (14).xlsx") plus a real logged WAEPA proposal
    already in Supabase for Task 2.
  - Draft-from-notes and the report narrative both make real, live
    Anthropic API calls -- a few cents, covered by this project's
    standing authorization for verifying drafting behaviour (CLAUDE.md).

Output: Harrisburg_Rep_User_Guide.docx at the repo root (gitignored, same
as every other generated deck) plus its screenshots under
user_guide_assets/ (also gitignored -- regenerated every run).
"""
import argparse
import os
import sys
import time
from pathlib import Path

REPO = Path(__file__).resolve().parent
ASSETS_DIR = REPO / "user_guide_assets"
OUTPUT_DOCX = REPO / "Harrisburg_Rep_User_Guide.docx"

HERSHEY_PDF = REPO / ("Premion Media Plan_RFPID-260402_Direct - No Agency_"
                       "Visit Hershey & Harrisburg_4-30-2026--ver0.pdf")
WAEPA_ATTRIBUTION_XLSX = REPO / "Premion Website Attribution and Reach Extension (14).xlsx"

HERSHEY_NOTES = """Met with the Visit Hershey & Harrisburg tourism board about their
early-summer push. They want CTV in front of two kinds of travelers:
younger travel-minded couples and friends looking for a quick getaway,
and families planning a summer trip with kids. Feeder markets are the
usual out-of-market drive audiences -- Philadelphia, New York, Baltimore,
and DC.

No agency involved, this is direct with the tourism board. Budget
follows whatever the avails price out to at standard rate -- no
negotiated CPM discussed. They'd like a web-visit attribution report
after the flight wraps, nothing beyond that for now. If we have a
travel-related case study, include it; otherwise leave that open for
them to see what we've got."""

sys.path.insert(0, str(REPO / "tests"))
import app_process  # noqa: E402  (tests/app_process.py -- start/status the local test-mode app)

PORT = app_process.DEFAULT_PORT
BASE_URL = app_process.base_url(PORT)


# ===========================================================================
# STEP TEXT -- data, not formatting code. Edit a step by editing its string
# here; nothing about layout or screenshot mechanics lives in this section.
# Each step: shot (screenshot filename stem, matches a capture_* call below),
# text (the numbered instruction, one action), callout (optional -- one of
# the "things that trip people," rendered as a highlighted note under the
# step it belongs to).
# ===========================================================================
STEPS_PROPOSAL = [
    {
        "shot": "p01_gate_closed",
        "text": "Open the app. “Build a proposal” is the home page, and the "
                "Setup band at the top is the first thing to fill in.",
        "callout": "The gate: nothing below the Setup band appears until the "
                   "originating market and both flight dates are set. This is "
                   "deliberate -- it stops a rep from filling out a form that "
                   "has nowhere to put the numbers yet.",
    },
    {
        "shot": "p02_avails_uploaded",
        "text": "Upload the avail. Drop the Salesforce avails PDF onto “Avails "
                "PDF (from Salesforce).” The client name, the flight dates, "
                "and every audience/geography group on the document are filled "
                "in for you.",
        "callout": None,
    },
    {
        "shot": "p03_form_open",
        "text": "Pick the originating market -- Harrisburg, for this walkthrough. "
                "With the market and both flight dates set, the rest of the "
                "form opens.",
        "callout": None,
    },
    {
        "shot": "p04_drafted",
        "text": "Paste the meeting or discovery notes into “Meeting / discovery "
                "notes” and click “Draft proposal from notes.” Claude fills in "
                "the plan -- budget, products, audiences -- and lists anything "
                "it wasn’t sure about for you to confirm.",
        "callout": "Draft from notes is optional. Skip it and build the plan by "
                   "hand instead -- but it’s what fills the plan in for you, "
                   "so most reps use it, and everything it fills in stays "
                   "editable afterward.",
    },
    {
        "shot": "p05_media_plan",
        "text": "Review the plan Claude built, in “E. Proposal / media plan.” "
                "Every row -- Tactic, Geo, Targeting, Cost -- is editable.",
        "callout": None,
    },
    {
        "shot": "p06_avails_table",
        "text": "In the avails table (“D2. Audiences & avails”), check the "
                "“Plan” column for each row you actually want on the deck.",
        "callout": "The Plan checkbox is the real on/off switch for the deck. A "
                   "row can sit in this table as research an audience or "
                   "geography you looked at but decided not to sell -- without "
                   "ever appearing in the proposal, unless its Plan box is "
                   "checked.",
    },
    {
        "shot": "p07_case_studies",
        "text": "Pick case studies, if any fit. Open “Case studies” near the "
                "bottom of the page and check the ones to include -- matches "
                "for this vertical are suggested first.",
        "callout": None,
    },
    {
        "shot": "p08_generated",
        "text": "Click Generate, then download the finished PowerPoint.",
        "callout": None,
    },
]

STEPS_REPORT = [
    {
        "shot": "r01_upload_door",
        "text": "Open “Attribution reports” from the sidebar. You land on the "
                "“New report” tab.",
        "callout": None,
    },
    {
        "shot": "r02_export_uploaded",
        "text": "Upload the Website Attribution export.",
        "callout": "The Delivery export is optional. Upload it too and the "
                   "report gains a Delivery Recap slide; without it, that "
                   "slide is simply left out -- nothing else about the report "
                   "changes.",
    },
    {
        "shot": "r03_confirm_client",
        "text": "Confirm the client match the app found.",
        "callout": None,
    },
    {
        "shot": "r04_link_proposal",
        "text": "Link a proposal, if one exists for this client -- the app "
                "finds the best match and offers it as the default. Choose "
                "“No proposal” instead to build the report standalone.",
        "callout": None,
    },
    {
        "shot": "r05_goals",
        "text": "Fill in Goals, unless a linked proposal already filled them in "
                "for you.",
        "callout": "Goals are the one field this page always needs. Linking a "
                   "proposal pre-fills them from it, so there’s usually "
                   "nothing to type here -- but with no proposal linked, "
                   "Goals has to be typed by hand before Generate will run.",
    },
    {
        "shot": "r06_optimization",
        "text": "If the optimization engine has a recommendation, an "
                "“Optimization recommendations” checklist appears here -- "
                "accept, edit, or decline each one before Generate. If it "
                "doesn’t appear, there’s nothing to do.",
        "callout": None,
    },
    {
        "shot": "r07_preview",
        "text": "“Preview narrative” lets you read Claude’s draft before it "
                "becomes part of the deck.",
        "callout": "Preview narrative is optional. Skip it and Generate drafts "
                   "the narrative automatically -- either way, once a draft "
                   "has been previewed, Generate uses it exactly as it stands "
                   "rather than drafting a second time.",
        },
    {
        "shot": "r08_summary_toggle",
        "text": "Decide whether to also build the one-slide summary.",
        "callout": "“Also build a one-slide summary” is a toggle, not a "
                   "separate task -- check it before Generate to get an "
                   "abbreviated, single-slide version alongside the full "
                   "report deck.",
    },
    {
        "shot": "r09_generated",
        "text": "Click “Generate report deck,” then download it.",
        "callout": None,
    },
]


# ===========================================================================
# CAPTURE -- drives the real app with Playwright. Nothing here decides what
# a step says; it only produces the screenshot file that step's "shot" name
# points to.
# ===========================================================================
VIEWPORT = {"width": 1440, "height": 1000}
APP_CONTAINER = '[data-testid="stAppViewContainer"]'


def _wait_settled(page, ms=1200):
    page.wait_for_timeout(ms)


def _wait_for_rerun(page, settle_ms=1500, max_wait_ms=90_000):
    """Best-effort wait for a Streamlit rerun (including a live Claude call)
    to finish: wait for the top-right running indicator to appear, then for
    it to go away, and always add a fixed settle on top since Streamlit
    paints in more than one pass. Never raises -- a script driving a real
    running app for documentation screenshots should degrade to "waited the
    fixed amount" rather than abort the whole run over a timing guess.
    """
    try:
        page.wait_for_selector('[data-testid="stStatusWidget"]', timeout=4000)
    except Exception:
        pass
    try:
        page.wait_for_selector('[data-testid="stStatusWidget"]', state="detached",
                               timeout=max_wait_ms)
    except Exception:
        pass
    _wait_settled(page, settle_ms)


def _visible_matches(page, text, exact=True, retries=6, retry_wait=500):
    """Streamlit renders more than one DOM node for some headings (an
    sr-only/hidden duplicate alongside the real, visible one) -- matching by
    text alone can silently grab the hidden copy, whose zero-size bounding
    box then produces a garbage crop with no error anywhere. Return only the
    matches that actually have a paintable box, in document order. Retries
    briefly before giving up, since right after a rerun the real node can
    still be mid-layout the first time this is checked."""
    loc = None
    for _ in range(retries):
        loc = page.get_by_text(text, exact=exact)
        visible = []
        for i in range(loc.count()):
            item = loc.nth(i)
            box = item.bounding_box()
            if box and box["width"] > 0 and box["height"] > 0:
                visible.append(item)
        if visible:
            return visible
        page.wait_for_timeout(retry_wait)
    return [loc.first] if loc is not None else []


def _crop_locator(page, locator, out_path, height=500, top_pad=20):
    """Scroll a resolved locator into view and screenshot a clipped region
    below it -- the "crop to the relevant area" this guide is supposed to
    use everywhere, rather than a full, un-cropped page.

    Streamlit's main content area scrolls WITHIN a fixed-height container
    (the outer window never grows), so `scroll_into_view_if_needed` can park
    the element anywhere from the top to the bottom edge of the viewport --
    never past it. When it lands too low for the requested crop height to
    fit below it, nudge the scroll further with a wheel event and
    re-measure, rather than silently clipping to whatever was already
    on-screen above it.
    """
    locator.scroll_into_view_if_needed(timeout=15_000)
    _wait_settled(page, 400)
    box = locator.bounding_box()
    vp_h = VIEWPORT["height"]
    if box["y"] + height > vp_h:
        extra = (box["y"] + height) - vp_h + 40
        page.mouse.move(VIEWPORT["width"] * 0.6, vp_h * 0.5)
        page.mouse.wheel(0, extra)
        _wait_settled(page, 400)
        box = locator.bounding_box()
    y0 = max(box["y"] - top_pad, 0)
    clip = {"x": 0, "y": y0, "width": VIEWPORT["width"],
            "height": min(height, vp_h - y0)}
    page.screenshot(path=str(out_path), clip=clip)


def _crop_below(page, heading_text, out_path, extra_scroll=0, height=700, exact=True):
    """Like `_crop`, but for a section whose interesting content (a
    st.data_editor grid, in particular) is a canvas -- glide-data-grid,
    which st.data_editor renders on, exposes no visible DOM text Playwright
    can locate, so column headers/cells can't be used as scroll anchors the
    way a real heading can. Scrolls to the section heading, then nudges an
    extra fixed amount to bring the grid itself into frame, and screenshots
    whatever is now at the top of the viewport."""
    matches = _visible_matches(page, heading_text, exact=exact)
    heading = matches[0]
    heading.scroll_into_view_if_needed(timeout=15_000)
    _wait_settled(page, 300)
    if extra_scroll:
        page.mouse.move(VIEWPORT["width"] * 0.6, VIEWPORT["height"] * 0.5)
        page.mouse.wheel(0, extra_scroll)
        _wait_settled(page, 500)
    clip = {"x": 0, "y": 0, "width": VIEWPORT["width"],
            "height": min(height, VIEWPORT["height"])}
    page.screenshot(path=str(out_path), clip=clip)


def _crop(page, heading_text, out_path, height=500, exact=True, top_pad=20, nth=0):
    """`_crop_locator`, resolving the target by exact/substring text match
    among only the VISIBLE matches (see `_visible_matches`) -- the right
    tool for a real st.header/st.subheader, which Streamlit renders once.
    Do not use this for a checkbox/button/widget LABEL that also gets
    echoed elsewhere on the page (a review panel, a caption) -- resolve
    those with a role-based locator and call `_crop_locator` directly, the
    way capture_proposal_flow's "Draft from notes" steps do; text alone
    can't tell two genuine, visible, identically-worded elements apart."""
    matches = _visible_matches(page, heading_text, exact=exact)
    heading = matches[min(nth, len(matches) - 1)]
    _crop_locator(page, heading, out_path, height=height, top_pad=top_pad)


def _goto_page(page, nav_text):
    page.get_by_text(nav_text, exact=True).first.click()
    _wait_settled(page, 2500)


def _upload(page, index, file_path):
    inputs = page.locator('[data-testid="stFileUploaderDropzoneInput"]')
    inputs.nth(index).set_input_files(str(file_path))
    _wait_settled(page, 3500)


def capture_proposal_flow(page, out_dir):
    shots = {s["shot"]: out_dir / f"{s['shot']}.png" for s in STEPS_PROPOSAL}

    page.goto(BASE_URL, timeout=30_000)
    page.wait_for_selector(APP_CONTAINER, timeout=20_000)
    _wait_settled(page, 3000)
    _crop(page, "\U0001F9ED Setup", shots["p01_gate_closed"], height=560, exact=False)

    _upload(page, 0, HERSHEY_PDF)
    _crop(page, "\U0001F9ED Setup", shots["p02_avails_uploaded"], height=560, exact=False)

    page.get_by_text("Harrisburg", exact=True).first.click()
    _wait_settled(page, 2000)
    draft_from_notes = page.get_by_role("checkbox", name="Draft from notes")
    _crop_locator(page, draft_from_notes, shots["p03_form_open"], height=520)

    notes_box = page.get_by_label("Meeting / discovery notes")
    notes_box.click()
    notes_box.fill(HERSHEY_NOTES)
    _wait_settled(page, 300)
    page.get_by_role("button", name="Draft proposal from notes", exact=True).click()
    _wait_for_rerun(page, settle_ms=2000, max_wait_ms=120_000)
    _crop_locator(page, draft_from_notes, shots["p04_drafted"], height=760)

    _crop_below(page, "E. Proposal / media plan", shots["p05_media_plan"],
               extra_scroll=650, height=760)

    _crop_below(page, "D2. Audiences & avails", shots["p06_avails_table"],
               extra_scroll=220, height=760)

    cs_expander = page.get_by_text("Case studies (", exact=False).first
    cs_expander.scroll_into_view_if_needed(timeout=15_000)
    cs_expander.click()
    _wait_settled(page, 800)
    _crop_locator(page, cs_expander, shots["p07_case_studies"], height=560, top_pad=140)

    generate_heading = page.get_by_text("Generate", exact=True).first
    generate_heading.scroll_into_view_if_needed(timeout=15_000)
    _wait_settled(page, 500)
    page.get_by_role("button", name="Generate", exact=False).first.click()
    _wait_for_rerun(page, settle_ms=2000, max_wait_ms=120_000)
    _crop(page, "Generate", shots["p08_generated"], height=480, exact=True)


def capture_report_flow(page, out_dir):
    shots = {s["shot"]: out_dir / f"{s['shot']}.png" for s in STEPS_REPORT}

    page.goto(BASE_URL, timeout=30_000)
    page.wait_for_selector(APP_CONTAINER, timeout=20_000)
    _wait_settled(page, 3000)
    _goto_page(page, "Attribution reports")
    _crop(page, "1. Upload the export(s)", shots["r01_upload_door"], height=440, exact=True)

    _upload(page, 0, WAEPA_ATTRIBUTION_XLSX)
    _crop(page, "1. Upload the export(s)", shots["r02_export_uploaded"], height=520, exact=True)

    _crop(page, "2. Confirm the client", shots["r03_confirm_client"], height=220, exact=True)
    page.get_by_role("button", name="Confirm", exact=True).first.click()
    _wait_settled(page, 2000)

    _crop(page, "3. Link a proposal", shots["r04_link_proposal"], height=260, exact=False)
    page.get_by_role("button", name="Confirm", exact=True).first.click()
    _wait_settled(page, 2500)

    _crop(page, "4. Goals, notes, what's next", shots["r05_goals"], height=420, exact=False)

    _crop(page, "Optimization recommendations", shots["r06_optimization"], height=380,
          exact=False)

    page.get_by_role("button", name="Preview narrative", exact=False).first.click()
    _wait_for_rerun(page, settle_ms=2000, max_wait_ms=120_000)
    _crop(page, "Preview narrative", shots["r07_preview"], height=520, exact=False, nth=1)

    _crop(page, "Also build a one-slide summary", shots["r08_summary_toggle"], height=260,
          exact=False)

    page.get_by_role("button", name="Generate report deck", exact=False).first.click()
    _wait_for_rerun(page, settle_ms=2000, max_wait_ms=120_000)
    _crop(page, "5. Generate", shots["r09_generated"], height=420, exact=False)


def run_captures(tasks):
    from playwright.sync_api import sync_playwright

    ok, msg = app_process.ensure_running(PORT, wait_seconds=60)
    print(f"[app] {msg}")
    if not ok:
        raise SystemExit(f"Test-mode app didn't come up: {msg}")

    ASSETS_DIR.mkdir(exist_ok=True)

    with sync_playwright() as p:
        browser = p.chromium.launch()
        try:
            if "proposal" in tasks:
                if not HERSHEY_PDF.exists():
                    raise SystemExit(f"Missing fixture: {HERSHEY_PDF}")
                page = browser.new_page(viewport=VIEWPORT)
                print("[capture] Task 1 -- Build a proposal ...")
                capture_proposal_flow(page, ASSETS_DIR)
                page.close()
            if "report" in tasks:
                if not WAEPA_ATTRIBUTION_XLSX.exists():
                    raise SystemExit(f"Missing fixture: {WAEPA_ATTRIBUTION_XLSX}")
                page = browser.new_page(viewport=VIEWPORT)
                print("[capture] Task 2 -- Build an attribution report ...")
                capture_report_flow(page, ASSETS_DIR)
                page.close()
        finally:
            browser.close()


# ===========================================================================
# ASSEMBLE -- python-docx. Reads STEPS_PROPOSAL / STEPS_REPORT plus the
# screenshots they name; decides nothing about wording.
# ===========================================================================
def _build_stamp():
    """Read the same build hash+date the app's own sidebar shows, so a rep
    comparing their screen to a screenshot can tell whether they're looking
    at the same version."""
    import subprocess
    try:
        sha = subprocess.check_output(["git", "rev-parse", "--short=7", "HEAD"],
                                      cwd=REPO, text=True).strip()
        ts = subprocess.check_output(["git", "log", "-1", "--format=%cd",
                                       "--date=format:%Y-%m-%d %H:%M"],
                                     cwd=REPO, text=True).strip()
        return f"{sha} · {ts}"
    except Exception:
        return "(unknown -- not a git checkout)"


def _add_step(doc, n, step, assets_dir):
    from docx.shared import Inches, Pt, RGBColor

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    run = p.add_run(f"{n}. ")
    run.bold = True
    run.font.size = Pt(13)
    run2 = p.add_run(step["text"])
    run2.font.size = Pt(13)

    img_path = assets_dir / f"{step['shot']}.png"
    if img_path.exists():
        doc.add_picture(str(img_path), width=Inches(6.2))
    else:
        warn = doc.add_paragraph()
        warn.add_run(f"[screenshot missing: {img_path.name} -- re-run "
                     f"build_user_guide.py]").italic = True

    if step.get("callout"):
        cp = doc.add_paragraph()
        cp.paragraph_format.left_indent = Inches(0.25)
        cp.paragraph_format.space_after = Pt(10)
        r = cp.add_run("⚠ " + step["callout"])
        r.italic = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x8A, 0x5A, 0x00)


def _add_task_page(doc, title, intro, steps, assets_dir, first=False):
    from docx.shared import Pt

    if not first:
        doc.add_page_break()
    h = doc.add_heading(title, level=1)
    if intro:
        ip = doc.add_paragraph(intro)
        ip.runs[0].font.size = Pt(11)
    for n, step in enumerate(steps, start=1):
        _add_step(doc, n, step, assets_dir)


def assemble_docx(tasks, out_path=OUTPUT_DOCX):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("Proposal Builder — Harrisburg Rep Guide", level=0)
    stamp = doc.add_paragraph()
    stamp.add_run(f"App build: {_build_stamp()}").bold = True
    guide_stamp = doc.add_paragraph()
    guide_stamp.add_run(f"Guide generated: {time.strftime('%Y-%m-%d %H:%M')}")
    intro = doc.add_paragraph(
        "This is the shortest path through the app's two real jobs: building a "
        "client proposal, and building an attribution report after a campaign "
        "runs. If you can't get through page one alone, something in here needs "
        "fixing -- tell whoever generated this guide.\n\n"
        "If your screen doesn't match a screenshot below, check the build "
        "stamp in the app's own sidebar against the one above -- the app may "
        "have changed since this guide was generated.")

    if "proposal" in tasks:
        _add_task_page(
            doc, "Task 1 — Build a proposal",
            "The real path: upload the avail, fill the Setup band, paste "
            "notes, draft, review the plan, pick case studies, generate, "
            "download.",
            STEPS_PROPOSAL, ASSETS_DIR, first=True)

    if "report" in tasks:
        _add_task_page(
            doc, "Task 2 — Build an attribution report",
            "The real path: upload the export, confirm the client, link a "
            "proposal if one exists, fill goals, review the optimization "
            "checklist if one appears, generate, download.",
            STEPS_REPORT, ASSETS_DIR, first=("proposal" not in tasks))

    doc.save(str(out_path))
    print(f"[docx] wrote {out_path}")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--task", choices=["proposal", "report", "all"], default="all")
    ap.add_argument("--assemble-only", action="store_true",
                    help="Skip capture; assemble the .docx from whatever screenshots "
                         "already exist in user_guide_assets/.")
    args = ap.parse_args()
    tasks = ["proposal", "report"] if args.task == "all" else [args.task]

    if not args.assemble_only:
        run_captures(tasks)
    assemble_docx(tasks)


if __name__ == "__main__":
    main()
