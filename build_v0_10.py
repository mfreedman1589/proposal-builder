"""build_v0_10.py -- inventory check for REPORT_MASTER_v0_10.pptx.

Unlike build_v0_6/7/8.py (which programmatically widened a table by
python-pptx surgery), v0_10 was hand-built by Matt directly from this
handoff spec (three changes: url_report's reach-basis headers/columns +
footnote, delivery_recap's TopPublishersTable +Attributed Rate, and
highlights' narrowed WHAT STOOD OUT card + new TrendChartRegion). This
script doesn't build the file -- it VERIFIES the delivered file actually
matches the spec Claude Code was given, the same role migrate_report_
master_v0_4.py's own `_inventory_check` plays for a script-built version,
just checking a hand-built one instead. Run it after any future v0_10
template edit to confirm nothing drifted.

    python build_v0_10.py
"""
import sys

from pptx.util import Emu

import slide_map
from pptx import Presentation

V0_9 = "REPORT_MASTER_v0_9.pptx"
V0_10 = "REPORT_MASTER_v0_10.pptx"

FAILURES = []


def check(label, condition, detail=None):
    status = "PASS" if condition else "FAIL"
    print(f"  {status}  {label}" + (f" -- {detail}" if detail and not condition else ""))
    if not condition:
        FAILURES.append(label)


def _slide_by_key(prs):
    keys = {}
    for i, slide in enumerate(prs.slides):
        if slide.has_notes_slide:
            key = slide_map.notes_key(slide)
            if key:
                keys[key] = i
    return keys


def _shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def _table_header(shape):
    return [c.text_frame.text for c in shape.table.rows[0].cells]


def main():
    prs10 = Presentation(V0_10)
    keys10 = _slide_by_key(prs10)

    print("\n1. url_report -- reach-basis headers/columns + footnote")
    url_slide = prs10.slides[keys10["report:url_report"]]
    intent_table = _shape(url_slide, "IntentSummaryTable")
    url_table = _shape(url_slide, "TopUrlTable")
    footnote = _shape(url_slide, "UrlReachFootnote")
    check("IntentSummaryTable header is Intent | Visitors | % of visitors | Converted",
         _table_header(intent_table) == ["Intent", "Visitors", "% of visitors", "Converted"],
         _table_header(intent_table))
    check("TopUrlTable header is Page / section | Visitors | % of visitors | Converted",
         _table_header(url_table) == ["Page / section", "Visitors", "% of visitors", "Converted"],
         _table_header(url_table))
    check("UrlReachFootnote shape exists with the reach-basis wording",
         footnote is not None and "don't sum to 100%" in footnote.text_frame.text,
         footnote.text_frame.text if footnote else None)
    section_headers = [sh.text_frame.text for sh in url_slide.shapes
                       if sh.has_text_frame and sh.text_frame.text in
                       ("VISITORS BY INTENT", "TOP PAGES BY ATTRIBUTED VISITORS")]
    check("both section headers renamed to the reach-basis wording",
         set(section_headers) == {"VISITORS BY INTENT", "TOP PAGES BY ATTRIBUTED VISITORS"},
         section_headers)

    print("\n2. delivery_recap -- TopPublishersTable +Attributed Rate")
    recap_slide = prs10.slides[keys10["report:delivery_recap"]]
    pub_table = _shape(recap_slide, "TopPublishersTable")
    check("TopPublishersTable header is Channel | Impressions | VCR | Attr. rate",
         _table_header(pub_table) == ["Channel", "Impressions", "VCR", "Attr. rate"],
         _table_header(pub_table))
    check("TopPublishersTable width unchanged at 5.97in",
         abs(pub_table.width / 914400 - 5.97) < 0.01, pub_table.width / 914400)

    print("\n3. highlights -- narrowed card + TrendChartRegion")
    hl_slide = prs10.slides[keys10["report:highlights"]]
    bullets = _shape(hl_slide, "HIGHLIGHTBullets")
    trend_header = _shape(hl_slide, "TrendChartHeader")
    trend_region = _shape(hl_slide, "TrendChartRegion")
    trend_label = _shape(hl_slide, "TrendChartRegionLabel")
    card = next((sh for sh in hl_slide.shapes if sh.shape_type == 1  # MSO_SHAPE_TYPE.AUTO_SHAPE
                and bullets is not None
                and sh.left is not None and sh.left <= bullets.left
                and sh.top is not None and sh.top <= bullets.top
                and sh.width is not None and sh.left + sh.width >= bullets.left + bullets.width
                and sh.height > 2000000), None)
    check("HIGHLIGHTBullets narrowed to 7.10in", bullets is not None
         and abs(bullets.width / 914400 - 7.10) < 0.02, bullets.width / 914400 if bullets else None)
    check("the WHAT STOOD OUT card narrowed to 7.70in", card is not None
         and abs(card.width / 914400 - 7.70) < 0.02, card.width / 914400 if card else None)
    check("TrendChartHeader reads ATTRIBUTED RATE BY WEEK", trend_header is not None
         and trend_header.text_frame.text == "ATTRIBUTED RATE BY WEEK",
         trend_header.text_frame.text if trend_header else None)
    check("TrendChartRegion + TrendChartRegionLabel both exist",
         trend_region is not None and trend_label is not None)
    if trend_region is not None:
        check("TrendChartRegion at (8.45, 4.05) sized 4.38 x 2.85",
             abs(trend_region.left / 914400 - 8.45) < 0.02
             and abs(trend_region.top / 914400 - 4.05) < 0.05
             and abs(trend_region.width / 914400 - 4.38) < 0.02
             and abs(trend_region.height / 914400 - 2.85) < 0.05,
             (trend_region.left / 914400, trend_region.top / 914400,
              trend_region.width / 914400, trend_region.height / 914400))

    print("\n4. Restore targets for the no-weekly-series reflow "
         "(_widen_highlights_card in report_assembly.py reads these back from v0_9)")
    prs9 = Presentation(V0_9)
    keys9 = _slide_by_key(prs9)
    hl9 = prs9.slides[keys9["report:highlights"]]
    bullets9 = _shape(hl9, "HIGHLIGHTBullets")
    card9 = next((sh for sh in hl9.shapes if sh.shape_type == 1
                 and sh.left is not None and sh.left <= bullets9.left
                 and sh.top is not None and sh.top <= bullets9.top
                 and sh.width is not None and sh.left + sh.width >= bullets9.left + bullets9.width
                 and sh.height > 2000000), None)
    print(f"  v0_9 card width:    {card9.width} EMU ({card9.width / 914400:.3f}in)")
    print(f"  v0_9 bullets width: {bullets9.width} EMU ({bullets9.width / 914400:.3f}in)")
    check("v0_9's own card width matches the spec's '12.33'",
         abs(card9.width / 914400 - 12.33) < 0.01)

    print(f"\n{'ALL PASSED' if not FAILURES else f'{len(FAILURES)} FAILURE(S): ' + ', '.join(FAILURES)}")
    return 1 if FAILURES else 0


if __name__ == "__main__":
    sys.exit(main())
