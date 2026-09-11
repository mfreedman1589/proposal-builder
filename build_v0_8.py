import copy
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

SRC, OUT = "REPORT_MASTER_v0_7.pptx", "REPORT_MASTER_v0_8.pptx"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
prs = Presentation(SRC)

def cell_set(c, t):
    p = c.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = t
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
    elif t: p.add_run().text = t

def add_col(tbl_shape, new_head, widths):
    tbl, el = tbl_shape.table, tbl_shape.table._tbl
    grid = el.find("a:tblGrid", NS)
    grid.append(copy.deepcopy(grid.findall("a:gridCol", NS)[-1]))
    for tr in el.findall("a:tr", NS):
        tr.append(copy.deepcopy(tr.findall("a:tc", NS)[-1]))
    for gc, w in zip(grid.findall("a:gridCol", NS), widths):
        gc.set("w", str(Inches(w)))
    tbl_shape.width = Inches(sum(widths))
    cell_set(list(tbl.rows[0].cells)[-1], new_head)
    cell_set(list(tbl.rows[1].cells)[-1], "")
    for r in tbl.rows:
        cells = list(r.cells)
        cells[-1].text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

targets = {
    # name: (new header, new widths — total unchanged)
    "BreakdownTable":     ("Conv. rate", [1.95, 1.10, 1.10, 0.86, 0.96]),   # 5.97
    "IntentSummaryTable": ("Converted",  [3.36, 1.28, 1.28, 1.28]),         # 7.20
    "TopUrlTable":        ("Converted",  [3.36, 1.28, 1.28, 1.28]),         # 7.20
}
done = set()
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_table and sh.name in targets and sh.name not in done:
            head, widths = targets[sh.name]
            add_col(sh, head, widths)
            done.add(sh.name)
assert done == set(targets), done
prs.save(OUT)
print("saved", OUT)
