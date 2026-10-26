"""REPORT_MASTER_v0_16 — the Auto-Sales Analyst's High-Engagement Filter
(Analyst app_build 2026.10.26, still facts schema v2).

Run against the CURRENT active template (v0_15):
    python build_v0_16.py REPORT_MASTER_v0_15.pptx REPORT_MASTER_v0_16.pptx
then upload the output on dev only, inactive (the nightly merge activates it).

What changes, all on the native Analyst slides (`analyst_set: true`):
  * Section titles match the Analyst's own deck and PDF: Inventory Movement ->
    "Executive Summary", Missed Opportunities -> "Missed Opportunities: The
    Watch List", Store Scoreboard -> "Auto Group Performance". (The Traffic
    Mix retitle of report:url_report is done in code: that slide also serves
    reports with no Analyst file.)
  * Executive Summary gains AnalystInventoryNote, an accent line in the header
    band under the headline -- the high-engagement callout in highlight mode,
    the "based on N+ campaign visits" caveat in filter mode, deleted for a
    file from before the filter. Its tier table says "Sold", not "Units sold".
  * A new report:analyst_high_engagement slide ("High-Engagement Sales: N+
    Campaign Visits") right after Executive Summary, cloned from it: one tile
    row (high-engagement vehicles sold | Est. value sold | share of all viewed
    vehicles sold), a by-condition table whose two Look-to-Book columns the
    code removes unless the Analyst's comparison_holds is true, a price-tier
    table, a computed narrative and the glossary footnote.
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

SRC = sys.argv[1] if len(sys.argv) > 1 else "REPORT_MASTER_v0_15.pptx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "REPORT_MASTER_v0_16.pptx"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
ACCENT = RGBColor(0x99, 0x99, 0xFF)   # the slide titles' own lavender
prs = Presentation(SRC)


def key_of(s): return s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
def find(key): return [s for s in prs.slides if f"key: {key}" in key_of(s)][0]


def clone_after(src, after_key, notes):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    rid = {r.rId: new.part.relate_to(r.target_part, RT.IMAGE)
           for r in src.part.rels.values() if r.reltype == RT.IMAGE}
    for sh in src.shapes:
        el = copy.deepcopy(sh._element)
        for node in el.iter():
            for attr in ("{%s}embed" % NS["r"], "{%s}link" % NS["r"]):
                if attr in node.attrib and node.attrib[attr] in rid:
                    node.attrib[attr] = rid[node.attrib[attr]]
        new.shapes._spTree.append(el)
    new.notes_slide.notes_text_frame.text = notes
    lst = prs.slides._sldIdLst; last = list(lst)[-1]
    idx = [i for i, s in enumerate(prs.slides) if f"key: {after_key}" in key_of(s)][0]
    lst.remove(last); lst.insert(idx + 1, last)
    return new


class S:
    def __init__(self, s): self.s = s
    def get(self, n):
        for sh in self.s.shapes:
            if sh.name == n: return sh
        raise KeyError(n)
    def drop(self, *ns):
        for n in ns:
            sh = self.get(n); sh._element.getparent().remove(sh._element)
    def clone(self, n, new):
        el = copy.deepcopy(self.get(n)._element); self.s.shapes._spTree.append(el)
        sh = self.s.shapes[-1]; sh.name = new; return sh
    def rename(self, old, new): self.get(old).name = new


def text(sh, t):
    tf = sh.text_frame
    p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = t
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
    else: p.add_run().text = t
    for e in tf.paragraphs[1:]: e._p.getparent().remove(e._p)


def cell(c, t):
    p = c.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = t
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
    elif t: p.add_run().text = t


def place(sh, x, y, w=None, h=None):
    sh.left, sh.top = Inches(x), Inches(y)
    if w is not None: sh.width = Inches(w)
    if h is not None: sh.height = Inches(h)


def recolumn(tsh, heads, widths, token, left_cols=1):
    tbl, el = tsh.table, tsh.table._tbl
    grid = el.find("a:tblGrid", NS)
    while len(grid.findall("a:gridCol", NS)) < len(heads):
        grid.append(copy.deepcopy(grid.findall("a:gridCol", NS)[-1]))
        for tr in el.findall("a:tr", NS): tr.append(copy.deepcopy(tr.findall("a:tc", NS)[-1]))
    while len(grid.findall("a:gridCol", NS)) > len(heads):
        grid.remove(grid.findall("a:gridCol", NS)[-1])
        for tr in el.findall("a:tr", NS): tr.remove(tr.findall("a:tc", NS)[-1])
    for gc, w in zip(grid.findall("a:gridCol", NS), widths): gc.set("w", str(Inches(w)))
    tsh.width = Inches(sum(widths))
    for c, h in zip(list(tbl.rows[0].cells), heads): cell(c, h)
    row = list(tbl.rows[1].cells); cell(row[0], token)
    for c in row[1:]: cell(c, "")
    for r in tbl.rows:
        for i, c in enumerate(r.cells):
            c.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT if i < left_cols else PP_ALIGN.RIGHT


inv_slide = find("report:analyst_inventory")
inv = S(inv_slide)

# ======================================================================
# 1. report:analyst_high_engagement -- cloned BEFORE the Executive Summary
#    gains its note line, so the clone never carries one.
# ======================================================================
he = S(clone_after(inv_slide, "report:analyst_inventory",
                   "key: report:analyst_high_engagement\nanalyst_set: true"))
text(he.get("Text 1"), "High-Engagement Sales: {{ANALYST_HE_THRESHOLD}}+ Campaign Visits")
text(he.get("Text 2"), "{{ANALYST_HE_HEADLINE}}")
for t in ("AnalystLtbTile", "AnalystLtbNewTile", "AnalystLtbUsedTile"):
    he.drop(t, t + "Value", t + "Label")
for old, new, tok, lab in (
        ("AnalystRevenueTile", "AnalystHeSoldTile", "{{ANALYST_HE_SOLD}}",
         "High-engagement vehicles sold"),
        ("AnalystUnitsTile", "AnalystHeValueTile", "{{ANALYST_HE_VALUE}}", "Est. value sold"),
        ("AnalystPipelineTile", "AnalystHeShareTile", "{{ANALYST_HE_SHARE}}",
         "of all viewed vehicles sold")):
    for suf in ("", "Value", "Label"): he.rename(old + suf, new + suf)
    text(he.get(new + "Value"), tok); text(he.get(new + "Label"), lab)
# One tile row instead of two, so the row gets the room: taller tiles,
# larger figures.
for i, name in enumerate(("AnalystHeSoldTile", "AnalystHeValueTile", "AnalystHeShareTile")):
    x = 0.50 + i * 4.195
    place(he.get(name), x, 2.10, 3.94, 0.98)
    place(he.get(name + "Value"), x + 0.10, 2.18, 3.74, 0.50)
    place(he.get(name + "Label"), x + 0.10, 2.70, 3.74, 0.26)
    for r in he.get(name + "Value").text_frame.paragraphs[0].runs: r.font.size = Pt(26)
    for r in he.get(name + "Label").text_frame.paragraphs[0].runs: r.font.size = Pt(11)

# left: by-condition table
he.rename("AnalystModelsHeader", "AnalystHeConditionHeader")
text(he.get("AnalystHeConditionHeader"), "HIGH-ENGAGEMENT SALES BY CONDITION")
place(he.get("AnalystHeConditionHeader"), 0.50, 3.38, 7.70, 0.26)
he.rename("AnalystModelsTable", "AnalystHeConditionTable")
recolumn(he.get("AnalystHeConditionTable"),
         ["Condition", "Viewed", "Sold", "Look-to-Book, 5+ visits", "Look-to-Book, under 5"],
         [1.90, 1.10, 1.10, 1.80, 1.80], "{{ANALYST_HE_CONDITION_ROWS}}")
place(he.get("AnalystHeConditionTable"), 0.50, 3.68)

# right: price-tier table
he.rename("AnalystTierHeader", "AnalystHeTierHeader")
text(he.get("AnalystHeTierHeader"), "HIGH-ENGAGEMENT SALES BY PRICE TIER")
place(he.get("AnalystHeTierHeader"), 8.45, 3.38, 4.38, 0.26)
he.rename("AnalystTierTable", "AnalystHeTierTable")
recolumn(he.get("AnalystHeTierTable"), ["Tier", "Sold", "% of sold"],
         [2.18, 1.00, 1.20], "{{ANALYST_HE_TIER_ROWS}}")
place(he.get("AnalystHeTierTable"), 8.45, 3.68)
for name in ("AnalystHeConditionTable", "AnalystHeTierTable"):
    for row in he.get(name).table.rows: row.height = Inches(0.38)

# bottom: the narrative on the Traffic Mix slide's own "WHAT THIS SIGNALS"
# panel (copied across: plain shapes, no relationships)
url = S(find("report:url_report"))
for src_name, new_name in (("Shape 5", "AnalystHePanel"), ("Text 6", "AnalystHePanelHeader")):
    el = copy.deepcopy(url.get(src_name)._element)
    he.s.shapes._spTree.append(el)
    he.s.shapes[-1].name = new_name
place(he.get("AnalystHePanel"), 0.50, 5.32, 12.33, 1.22)
place(he.get("AnalystHePanelHeader"), 0.80, 5.46, 6.00, 0.26)
text(he.get("AnalystHePanelHeader"), "WHAT THIS SHOWS")
he.rename("AnalystInventoryNarrative", "AnalystHeNarrative")
narrative = he.get("AnalystHeNarrative")
text(narrative, "{{ANALYST_HE_NARRATIVE}}")
place(narrative, 0.80, 5.78, 11.73, 0.66)
for r in narrative.text_frame.paragraphs[0].runs: r.font.size = Pt(15)
he.s.shapes._spTree.append(narrative._element)   # in front of the panel

he.rename("AnalystFootnote", "AnalystHeFootnote")
text(he.get("AnalystHeFootnote"), "{{ANALYST_HE_FOOTNOTE}}")

# ======================================================================
# 2. Executive Summary (was Inventory Movement)
# ======================================================================
text(inv.get("Text 1"), "Executive Summary")
note = inv.clone("Text 2", "AnalystInventoryNote")
text(note, "{{ANALYST_INVENTORY_NOTE}}")
place(note, 0.50, 1.46, 11.03, 0.36)
for p in note.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(14)
        r.font.color.rgb = ACCENT
cell(inv.get("AnalystTierTable").table.rows[0].cells[1], "Sold")

# ======================================================================
# 3. Watch list and scoreboard: the Analyst's own section names
# ======================================================================
text(S(find("report:analyst_watchlist")).get("Text 1"), "Missed Opportunities: The Watch List")
text(S(find("report:analyst_group")).get("Text 1"), "Auto Group Performance")

prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides))
print([key_of(x).split("\n")[0].replace("key: report:", "") for x in prs.slides])
