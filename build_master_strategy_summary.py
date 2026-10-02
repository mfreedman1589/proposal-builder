"""Proposal master deck: add the media plan's one-sentence strategy line.

Adds a text box named `StrategySummary` carrying `{{STRATEGY_SUMMARY}}` under
the plan title on every media plan slide (every slide with `{{PLAN_TITLE}}`:
the standard plan and both Total TV variants), and moves the plan table down
to make room. The box takes the table's old top position, so when the rep's
toggle is off `assembly.apply_strategy_summary` deletes the box and puts the
table back exactly where the template always had it.

Run against the CURRENT active master (version 9 at the time of writing):
    python build_master_strategy_summary.py <active master>.pptx <out>.pptx
then upload the output on dev only (never activate from dev -- the nightly
merge does that).
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

import slide_map

SRC, OUT = sys.argv[1], sys.argv[2]
BOX_HEIGHT = Inches(0.30)
GAP_BELOW_BOX = Inches(0.10)
FONT, SIZE, COLOR = "Proxima Nova Light", Pt(12), RGBColor(0, 0, 0)   # = "Included with Campaign"

prs = Presentation(SRC)
done = 0
for slide in prs.slides:
    shapes = list(slide_map.iter_all_shapes(slide.shapes))
    if not any(sh.has_text_frame and "{{PLAN_TITLE}}" in sh.text_frame.text for sh in shapes):
        continue
    if any(sh.name == "StrategySummary" for sh in slide.shapes):
        continue                                  # idempotent: already built
    table = next(sh for sh in slide.shapes if getattr(sh, "has_table", False) and sh.has_table)
    box = slide.shapes.add_textbox(table.left, table.top, table.width, BOX_HEIGHT)
    box.name = "StrategySummary"
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Emu(0)
    frame.margin_top = frame.margin_bottom = Emu(0)
    run = frame.paragraphs[0].add_run()
    run.text = "{{STRATEGY_SUMMARY}}"
    run.font.name, run.font.size, run.font.bold = FONT, SIZE, False
    run.font.color.rgb = COLOR
    table.top = Emu(table.top + BOX_HEIGHT + GAP_BELOW_BOX)
    done += 1
    print(f"slide {prs.slides.index(slide) + 1} ({slide_map.notes_key(slide)}): "
          f"StrategySummary at {box.top / 914400:.2f}in, table moved to {table.top / 914400:.2f}in")

prs.save(OUT)
print(f"{done} media plan slide(s) updated -> {OUT}")
