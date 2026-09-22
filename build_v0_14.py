"""build_v0_14.py -- Netmaker Communications review (2026-09-22), item 1's
template half: widens BreakdownTable on report:attribution_breakdown from
5 to 6 columns (inserting a VCR column between Delivered/Attributed, and
renaming the Delivered header to Impressions -- see CLAUDE.md's "one slide,
not two" item and report_assembly.py's `attribution_breakdown_display`),
and relabels report:ott_retargeting's DisplayImpressionsTileLabel static
text from "Display impressions" to "Retargeting impressions" (the largest
real unit on a real export is video, not display).

Everything else in that review (delivery_breakdown's creative section,
the OTT blended stat) is handled entirely in code -- report_assembly.py
now unconditionally deletes those shapes at fill time regardless of
whether the template still has them, so no template edit is needed for
either.
"""
from pptx import Presentation
from pptx.util import Inches

import assembly

SRC, OUT = "REPORT_MASTER_v0_13.pptx", "REPORT_MASTER_v0_14.pptx"
prs = Presentation(SRC)


def slide(key):
    return [s for s in prs.slides if f"key: {key}" in s.notes_slide.notes_text_frame.text][0]


def shape(s, n):
    for sh in s.shapes:
        if sh.name == n:
            return sh
    raise KeyError(n)


def cell_set(c, t):
    p = c.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = t
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    elif t:
        p.add_run().text = t


# ===== report:attribution_breakdown — BreakdownTable 5 -> 6 columns =====
# ChartRegion starts at left=6278728 EMU (6.87in); BreakdownTable's own
# left is 457200 (0.5in) -- a 6.37in budget total, only ~0.4in more than
# the original 5-column table's own 5.97in. A full duplicated-width VCR
# column would overrun that budget and overlap the chart, so all six
# columns are resized to fit within it rather than just inserting one at
# its clone-source's own width.
ab = slide("report:attribution_breakdown")
t = shape(ab, "BreakdownTable")
tbl = t.table
# Column 1 ("Delivered") is the clone source -- the copy lands at index 2,
# immediately after it, which is exactly where VCR belongs (Delivered
# already renamed to Impressions below).
assembly.clone_table_column(tbl, 1)
header = list(tbl.rows[0].cells)
data = list(tbl.rows[1].cells)
assert [c.text for c in header] == ["{{BREAKDOWN_DIMENSION_LABEL}}", "Delivered", "Delivered",
                                    "Attributed", "Rate", "Conv. rate"], [c.text for c in header]
cell_set(header[1], "Impressions")
cell_set(header[2], "VCR")
cell_set(data[2], "")  # the cloned data cell copied "" from Delivered's own -- already blank,
                       # stated explicitly so the intent is visible in the diff
widths = [1.75, 0.95, 0.75, 0.95, 0.80, 0.90]  # label, impressions, vcr, attributed, rate, conv rate
for col, w in zip(tbl.columns, widths):
    col.width = Inches(w)
t.width = Inches(sum(widths))
print("BreakdownTable header now:", [c.text for c in tbl.rows[0].cells])
print("BreakdownTable columns:", len(tbl.columns), "total width (in):", sum(widths))

# ===== report:ott_retargeting — DisplayImpressionsTileLabel relabel =====
ott = slide("report:ott_retargeting")
label = shape(ott, "DisplayImpressionsTileLabel")
assert label.text_frame.text == "Display impressions", label.text_frame.text
p = label.text_frame.paragraphs[0]
p.runs[0].text = "Retargeting impressions"
for r in p.runs[1:]:
    r._r.getparent().remove(r._r)
print("DisplayImpressionsTileLabel now:", label.text_frame.text)

prs.save(OUT)
print("saved", OUT)
