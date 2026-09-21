"""polk_import.py -- parse a Polk (IHS Markit) automotive match-back export
into one shape. Pure -- no Streamlit, no DB, mirrors attribution_import.py's
own shape: `parse_polk_export(path)` returns a `PolkExport` dataclass or
raises `PolkParseError` with a message aimed at a seller.

ATTRIBUTION_REPORT_PLAN.md Phase 7. The real fixture on hand (`Polk
Dashboard.xlsx`, project root, gitignored) carries 19 tabs: six single-value
headline tabs (Matched Households, Target Dealer Sales, Buy Rate, Matched
Impressions, Match Rate, Campaign Lift), four sales-cut breakdowns (days
elapsed, gender, age, income), a make/model-per-dealer table, three
audience/creative/publisher PAIRS (one tab is that dimension's share of
TOTAL matched impressions, the other its share of TARGET SALES impressions --
kept as two distinct fields, never conflated into one number), and two
dealer-roster tabs (Target Dealer(s), with market rank vs. campaign rank;
All Dealers, campaign share only).

**Tabs are identified by header row, exactly like attribution_import.py --
never by sheet name.** Unlike the Premion exports that module parses, this
vendor's sheet names are descriptive, not auto-numbered -- but Excel's own
31-character sheet-name limit already truncates one of them ("Target Dealer
Sales by Days Ela[psed]" in the real fixture), which is reason enough to
never trust a name to stay intact. Every tab here happens to carry its own
distinct header shape, so there is no need for attribution_import.py's
"pick the fuller of two identical headers" logic -- a straight header-tuple
index is enough.

**The audience/creative/publisher share-of-impressions and share-of-sales
tabs are a matched PAIR per dimension, never merged.** "AUTO Ford Intenders
drove 65% of matched impressions" and "AUTO Ford Intenders drove 88% of
target sales impressions" are two different, real questions (where the
reach went vs. where the sales came from) -- a caller that wants one number
picks the field for the question it's actually answering.

**`has_target_dealer_sales` is the same "present-and->0" detection
attribution_import.py's `has_conversions` already uses.** All six headline
tabs are ALWAYS present in a Polk export (it's a fixed dashboard shape,
unlike the Premion "Attributed Conversions" widget that's absent entirely
on a non-conversions campaign) -- so presence alone says nothing here. What
distinguishes a real signal from a paid-but-not-yet-converting campaign is
whether Target Dealer Sales is genuinely > 0; every sales-cut breakdown
(days elapsed/gender/age/income) and the Target Dealer(s) rank table can be
real rows summing to zero without this flag, which is a valid report (a
campaign that hasn't driven a matched sale yet), not a parse failure.

**Phase 8 adds `normalize_for_months(export, months)`** -- Polk's dashboard
stacks match rate, campaign lift, and every cut's share percentage across
the months a report covers (see that function's own docstring). Parsing
itself is unaffected; the caller (app.py, once a rep enters "months
covered by this Polk report") normalizes the freshly-parsed export before
handing it to anything downstream.
"""
import dataclasses
from dataclasses import dataclass, field

import openpyxl


class PolkParseError(Exception):
    """Raised with a message intended for a seller, not a developer."""


def _normalize_cell(value):
    return " ".join(str(value if value is not None else "").split()).strip().lower()


def _sheet_header(ws):
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not row:
        return ()
    return tuple(_normalize_cell(c) for c in row)


def _sheet_rows(ws):
    header = _sheet_header(ws)
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and all(c is None for c in row):
            continue
        out.append(dict(zip(header, row)))
    return out


def _index_by_header(wb):
    index = {}
    for ws in wb.worksheets:
        index.setdefault(_sheet_header(ws), []).append(ws)
    return index


def _find_one(index, header):
    sheets = index.get(header)
    return sheets[0] if sheets else None


def _clean_int(value):
    if value is None:
        return 0
    try:
        return int(round(float(str(value).replace(",", "").strip())))
    except (TypeError, ValueError):
        return 0


def _clean_float(value):
    if value is None:
        return 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def _single_value(index, header, cleaner):
    ws = _find_one(index, header)
    if ws is None:
        return None
    rows = _sheet_rows(ws)
    if not rows:
        return None
    return cleaner(rows[0].get(header[0]))


_MATCHED_HOUSEHOLDS = ("matched households",)
_TARGET_DEALER_SALES = ("target dealer sales",)
_BUY_RATE = ("buy rate",)
_MATCHED_IMPRESSIONS = ("matched impressions",)
_MATCH_RATE = ("match rate",)
_CAMPAIGN_LIFT = ("campaign lift",)

_BY_DAYS_ELAPSED = ("days elapsed group", "target dealer sales")
_BY_GENDER = ("target dealer sales", "gender")
_BY_AGE = ("target dealer sales", "age")
_BY_INCOME = ("target dealer sales", "income")
_MAKE_MODEL = ("target dealer name", "make 1", "model 1", "models sold", "average msrp 1")
_AUDIENCE_IMPRESSIONS = ("segment name", "matched impressions", "percentage in total matched impressions")
_AUDIENCE_SALES = ("segment name", "percentage in target sales impressions")
_CREATIVE_IMPRESSIONS = ("creative name", "matched impressions", "percentage in total matched impressions")
_CREATIVE_SALES = ("creative name", "percentage in target sales impressions")
_PUBLISHER_IMPRESSIONS = ("publisher name", "matched impressions", "percentage in total matched impressions")
_PUBLISHER_SALES = ("publisher name", "percentage in target sales impressions")
_TARGET_DEALERS = ("selling dealer name", "selling dealer addr", "new sales", "campaign share",
                   "target market rank", "campaign rank", "rank diff")
_ALL_DEALERS = ("selling dealer name", "selling dealer addr", "new sales", "campaign share")


def _label_count_rows(index, header, label_key, count_key="target dealer sales"):
    ws = _find_one(index, header)
    if ws is None:
        return []
    return [(str(row.get(label_key) or "").strip(), _clean_int(row.get(count_key)))
            for row in _sheet_rows(ws) if str(row.get(label_key) or "").strip()]


@dataclass
class PolkExport:
    source_name: str = ""
    matched_households: int = 0
    target_dealer_sales: int = 0
    buy_rate: float = 0.0
    matched_impressions: int = 0
    match_rate: float = 0.0
    campaign_lift: float = 0.0
    has_target_dealer_sales: bool = False   # target_dealer_sales > 0 -- see module docstring

    sales_by_days_elapsed: list = field(default_factory=list)   # [(label, count)]
    sales_by_gender: list = field(default_factory=list)
    sales_by_age: list = field(default_factory=list)
    sales_by_income: list = field(default_factory=list)

    make_model_rows: list = field(default_factory=list)         # [{"dealer","make","model",
                                                                  #   "models_sold","avg_msrp"}]

    # Share-of-TOTAL-matched-impressions and share-of-TARGET-SALES-impressions
    # are two separate fields per dimension -- see module docstring. Each
    # "_by_impressions" row carries both a raw impressions count and its own
    # share; each "_by_sales" row carries ONLY the sales-impressions share
    # (the export itself has no raw count on that tab).
    audience_by_impressions: list = field(default_factory=list)  # [{"segment","impressions","share"}]
    audience_by_sales: list = field(default_factory=list)        # [{"segment","share"}]
    creative_by_impressions: list = field(default_factory=list)  # [{"creative","impressions","share"}]
    creative_by_sales: list = field(default_factory=list)        # [{"creative","share"}]
    publisher_by_impressions: list = field(default_factory=list) # [{"publisher","impressions","share"}]
    publisher_by_sales: list = field(default_factory=list)       # [{"publisher","share"}]

    target_dealers: list = field(default_factory=list)   # [{"name","address","new_sales",
                                                           #   "campaign_share","market_rank",
                                                           #   "campaign_rank","rank_diff"}]
    all_dealers: list = field(default_factory=list)       # [{"name","address","new_sales","campaign_share"}]
    warnings: list = field(default_factory=list)

    @property
    def top_audience(self):
        """The audience row with the largest share of TOTAL matched
        impressions, or None -- "where the matched impressions went"."""
        return (max(self.audience_by_impressions, key=lambda r: r["impressions"])
               if self.audience_by_impressions else None)

    @property
    def top_creative(self):
        return (max(self.creative_by_impressions, key=lambda r: r["impressions"])
               if self.creative_by_impressions else None)

    @property
    def top_publisher(self):
        return (max(self.publisher_by_impressions, key=lambda r: r["impressions"])
               if self.publisher_by_impressions else None)


def normalize_for_months(export, months):
    """Phase 8 (ATTRIBUTION_REPORT_PLAN.md): Polk's dashboard STACKS certain
    figures across the months a report covers, rather than averaging them --
    confirmed from real multi-month dashboard exports Matt supplied (`Polk
    Dashboard (1).pptx`/`(2).pptx`), which show match rates of 414.95% and
    474.42% and audience-cut shares of 400%/200%/600%: each is the true
    per-month figure times the month count. A single-month file (the only
    shape `parse_polk_export` has a real fixture for) is unaffected --
    dividing by 1 is a no-op.

    What stacks, per Matt's own enumeration -- **match rate, campaign lift,
    and every share percentage in the audience/creative/publisher cuts**
    (both the share-of-matched-impressions and share-of-target-sales
    views). Counts (matched households, target dealer sales, matched
    impressions, days-elapsed/gender/age/income sale counts, make/model
    rows, dealer new_sales/campaign_share) are NOT divided -- they're
    already totals, correct as parsed. **`buy_rate` is deliberately left
    untouched too** -- Matt's own list names match rate/campaign lift/cut
    percentages only, not buy rate, and no real multi-month buy-rate figure
    is on hand to confirm one way or the other; flag this assumption back
    to Matt the first time a real multi-month Excel export is available to
    check it against (the two on-hand files are PowerPoint dashboard
    exports, not parseable Excel -- they prove the STACKING rule exists,
    not every field it touches).

    Returns a new `PolkExport` (the input is never mutated, matching this
    module's existing pure-function convention). Raises `PolkParseError`
    for `months <= 0` -- a required rep input with no sane fallback,
    exactly the "raise a plain-language message" contract every parse
    failure in this module already uses.

    Sanity check (Matt's own spec): if any of match_rate or a cut's share
    is still over 100% after dividing by `months`, the month count is
    probably wrong -- appended to the returned export's own `warnings` list
    (the same list a rep-facing report already surfaces from a parse), not
    raised, since a genuinely unusual campaign could still clear 100% on a
    single dimension without the month count being wrong; the rep decides.
    """
    if months <= 0:
        raise PolkParseError(
            f"\"Months covered\" must be a positive number, not {months}.")

    def _scaled(rows, key):
        return [{**row, key: row[key] / months} for row in rows]

    normalized = dataclasses.replace(
        export,
        match_rate=export.match_rate / months,
        campaign_lift=export.campaign_lift / months,
        audience_by_impressions=_scaled(export.audience_by_impressions, "share"),
        audience_by_sales=_scaled(export.audience_by_sales, "share"),
        creative_by_impressions=_scaled(export.creative_by_impressions, "share"),
        creative_by_sales=_scaled(export.creative_by_sales, "share"),
        publisher_by_impressions=_scaled(export.publisher_by_impressions, "share"),
        publisher_by_sales=_scaled(export.publisher_by_sales, "share"),
        warnings=list(export.warnings),
    )

    over_100 = []
    if normalized.match_rate > 1.0:
        over_100.append(f"match rate ({normalized.match_rate * 100:.1f}%)")
    for label, rows, key in (
        ("an audience", normalized.audience_by_impressions + normalized.audience_by_sales, "share"),
        ("a creative", normalized.creative_by_impressions + normalized.creative_by_sales, "share"),
        ("a publisher", normalized.publisher_by_impressions + normalized.publisher_by_sales, "share"),
    ):
        if any(row[key] > 1.0 for row in rows):
            over_100.append(f"{label} share")
    if over_100:
        normalized.warnings.append(
            f"After dividing by {months} month(s), {', '.join(over_100)} "
            f"{'is' if len(over_100) == 1 else 'are'} still over 100% -- "
            f"double check the month count entered for this Polk report.")

    return normalized


def _impressions_share_rows(index, header, label_key):
    ws = _find_one(index, header)
    if ws is None:
        return []
    out = []
    for row in _sheet_rows(ws):
        label = str(row.get(label_key) or "").strip()
        if not label:
            continue
        out.append({
            label_key.replace(" name", ""): label,
            "impressions": _clean_int(row.get("matched impressions")),
            "share": _clean_float(row.get("percentage in total matched impressions")),
        })
    return out


def _sales_share_rows(index, header, label_key):
    ws = _find_one(index, header)
    if ws is None:
        return []
    out = []
    for row in _sheet_rows(ws):
        label = str(row.get(label_key) or "").strip()
        if not label:
            continue
        out.append({
            label_key.replace(" name", ""): label,
            "share": _clean_float(row.get("percentage in target sales impressions")),
        })
    return out


def parse_polk_export(path, source_name=None):
    """PolkExport, or raises PolkParseError."""
    source_name = source_name or str(path)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:                                              # noqa: BLE001
        raise PolkParseError(
            f"Couldn't open \"{source_name}\" as an Excel file: {exc}") from exc

    index = _index_by_header(wb)
    if _find_one(index, _MATCHED_HOUSEHOLDS) is None:
        raise PolkParseError(
            f"\"{source_name}\" doesn't look like a Polk automotive match-back export -- "
            f"no \"Matched Households\" tab found.")

    result = PolkExport(source_name=source_name)
    result.matched_households = _single_value(index, _MATCHED_HOUSEHOLDS, _clean_int) or 0
    result.target_dealer_sales = _single_value(index, _TARGET_DEALER_SALES, _clean_int) or 0
    result.buy_rate = _single_value(index, _BUY_RATE, _clean_float) or 0.0
    result.matched_impressions = _single_value(index, _MATCHED_IMPRESSIONS, _clean_int) or 0
    result.match_rate = _single_value(index, _MATCH_RATE, _clean_float) or 0.0
    result.campaign_lift = _single_value(index, _CAMPAIGN_LIFT, _clean_float) or 0.0
    result.has_target_dealer_sales = result.target_dealer_sales > 0

    result.sales_by_days_elapsed = _label_count_rows(index, _BY_DAYS_ELAPSED, "days elapsed group")
    result.sales_by_gender = _label_count_rows(index, _BY_GENDER, "gender")
    result.sales_by_age = _label_count_rows(index, _BY_AGE, "age")
    result.sales_by_income = _label_count_rows(index, _BY_INCOME, "income")

    ws = _find_one(index, _MAKE_MODEL)
    if ws is not None:
        for row in _sheet_rows(ws):
            dealer = str(row.get("target dealer name") or "").strip()
            if not dealer:
                continue
            result.make_model_rows.append({
                "dealer": dealer,
                "make": str(row.get("make 1") or "").strip(),
                "model": str(row.get("model 1") or "").strip(),
                "models_sold": _clean_int(row.get("models sold")),
                "avg_msrp": _clean_float(row.get("average msrp 1")),
            })

    result.audience_by_impressions = _impressions_share_rows(index, _AUDIENCE_IMPRESSIONS, "segment name")
    result.audience_by_sales = _sales_share_rows(index, _AUDIENCE_SALES, "segment name")
    result.creative_by_impressions = _impressions_share_rows(index, _CREATIVE_IMPRESSIONS, "creative name")
    result.creative_by_sales = _sales_share_rows(index, _CREATIVE_SALES, "creative name")
    result.publisher_by_impressions = _impressions_share_rows(index, _PUBLISHER_IMPRESSIONS, "publisher name")
    result.publisher_by_sales = _sales_share_rows(index, _PUBLISHER_SALES, "publisher name")

    ws = _find_one(index, _TARGET_DEALERS)
    if ws is not None:
        for row in _sheet_rows(ws):
            name = str(row.get("selling dealer name") or "").strip()
            if not name:
                continue
            result.target_dealers.append({
                "name": name,
                "address": str(row.get("selling dealer addr") or "").strip(),
                "new_sales": _clean_int(row.get("new sales")),
                "campaign_share": _clean_float(row.get("campaign share")),
                "market_rank": _clean_int(row.get("target market rank")),
                "campaign_rank": _clean_int(row.get("campaign rank")),
                "rank_diff": _clean_int(row.get("rank diff")),
            })
    else:
        result.warnings.append("No \"Target Dealer(s)\" tab found -- dealer rank table will be empty.")

    ws = _find_one(index, _ALL_DEALERS)
    if ws is not None:
        for row in _sheet_rows(ws):
            name = str(row.get("selling dealer name") or "").strip()
            if not name:
                continue
            result.all_dealers.append({
                "name": name,
                "address": str(row.get("selling dealer addr") or "").strip(),
                "new_sales": _clean_int(row.get("new sales")),
                "campaign_share": _clean_float(row.get("campaign share")),
            })

    return result
