"""segment_catalog.py -- parse a RANKED audience-usage list into a scored,
role-classified, vertical-tagged segment catalog.

Pure module: no Streamlit, no DB. `db.py` stores what `derive_segment_catalog`
returns as JSON beside the uploaded workbook (one derived catalog per
`audience_usage_versions` row), and `audience_catalog` overlays it on the
`audiences` table at read time.

The ranked list (first seen 2026-10-12, "IMPRESSIONS BY DATA SEGMENT") is one
column, one targeting LINE per row, sorted by usage -- row 1 is the most-used
line. It carries no impressions, so rank is the only volume signal: each
segment scores the sum, over the lines containing it, of 1/log2(row + 1).
A segment on the top line gets 1.0, one on row 1,000 about 0.1, so a segment
booked on a handful of top lines outranks one booked on many lines nobody
reuses.

**Splitting.** A line is an AND-stack of components joined by commas, but a
comma alone is not a separator ("WEB RT- Pella Carolina, Inc"). A component
boundary is a comma followed by a taxonomy prefix (`SPLIT_PREFIXES`), with or
without a space ("DEMO Age A25-64,HH Income 150K Plus"). AND/OR inside a name
never split it -- "CUSTOM Homeowner AND Auto or Home Insurance" is one recipe.

**What never becomes a recommendable segment** (`status`):
- `excluded`: a client's own first-party list (CLT), retargeting or
  campaign-specific pools (RT, RFPID, pixel, address list, WEB RT, LOCATION
  RT), exclusion lines, "No Data Targeting", third-party taxonomy paths (">").
  These are dropped here and only counted -- a CLT name never reaches the
  derived catalog at all.
- `hidden`: a component naming a client or brand (`audience_brand_terms.csv`,
  the overrides registry's `exclude_client` rows, and client names harvested
  from this same list's CLT and retargeting lines). Kept in the derived
  catalog, never shown, so an admin can review it.
- `custom`: a CUSTOM build -- a "proven custom recipe", shown separately and
  labeled custom.
Everything else is `standard`.

**CLT wrapping a standard segment.** "CLT 1P Motor City Buick FIN Business
Industry Agriculture" is a client list narrowed by a real segment: the
standard segment (from the first UPPERCASE taxonomy prefix after the client
text) is extracted and counted; the client text is dropped. Matching is
case-sensitive on purpose -- "Lifestage" or "Food" inside a client's own
list name is client text, and guessing wrong would leak it.

**Role.** A `modifier` narrows any audience (age, gender, ethnicity,
language, education level, income, children, home value, dwelling,
homeowner); a `core` segment is the audience itself. Recommendations rank
core segments and offer each one's most common modifier pairing
(`pairings`, from co-occurrence on the same lines) as an optional add.

**Tags.** Cross-cutting vertical tags (`VERTICAL_TAGS`), several per segment,
assigned by the rules in `TAG_RULES` once, at ingest -- never by a model at
runtime. B2B carries a subgroup (`b2b_group`) and, for decision makers, a
function (`b2b_function`). Admin corrections override per segment at read
time (`apply_overrides`).
"""
import csv
import math
import re
from collections import Counter, defaultdict
from pathlib import Path

from market_profiles import _normalize as normalize

FORMAT = "ranked_v1"

SPLIT_PREFIXES = ["DEMO", "HH", "HLTH", "HLH", "AUTO", "FIN", "RETAIL", "LIFESTYLE",
                  "LIFESTAGE", "ENT", "SPORTS", "FOOD", "TRAVEL", "AFIRST", "POL", "WEB",
                  "B2B", "CUSTOM", "CLT"]
# Prefixes a CLT list can wrap -- the standard taxonomy, not another client
# list, a custom build or a retargeting pool.
STANDARD_PREFIXES = ["DEMO", "HH", "HLTH", "AUTO", "FIN", "RETAIL", "LIFESTYLE",
                     "LIFESTAGE", "ENT", "SPORTS", "FOOD", "TRAVEL", "AFIRST", "POL", "B2B"]

_PREFIX_ALT = "|".join(sorted(SPLIT_PREFIXES, key=len, reverse=True))
# A comma, then a prefix as a whole token (followed by a non-letter or the
# end), case-insensitive -- the list writes "Custom" and "CUSTOM" both.
_SPLIT_RE = re.compile(rf"\s*,\s*(?=(?:{_PREFIX_ALT})(?![A-Za-z]))", re.IGNORECASE)
_STANDARD_ALT = "|".join(sorted(STANDARD_PREFIXES, key=len, reverse=True))
_CLT_INNER_RE = re.compile(rf"(?:^|[\s_])((?:{_STANDARD_ALT})(?![A-Za-z]).*)$")
_CLT_RE = re.compile(r"^CLT(?![A-Za-z])", re.IGNORECASE)
_CLT_HEAD_RE = re.compile(r"^CLT[\s_]*(?:1P(?![A-Za-z]))?[\s_]*", re.IGNORECASE)
_CUSTOM_RE = re.compile(r"^CUSTOM(?![A-Za-z])", re.IGNORECASE)

# Spellings that are the same segment under an older export's naming --
# matched on the cleaned display, mapped to the current name.
KNOWN_ALIASES = {
    "financebankingintenders": "FIN Banking Intenders",
    "educationvocationalschool": "LIFESTYLE Vocational School",
}

TYPO_FIXES = [
    (re.compile(r"^HLH(?![A-Za-z])", re.IGNORECASE), "HLTH"),
    (re.compile(r"Entrepeneur", re.IGNORECASE), "Entrepreneur"),
]

_RETARGETING_RE = re.compile(
    r"(\bRT\b|RFPID|\bpixel\b|address\s*list|retarget|^WEB\b|^LOC(ATION)?\b)", re.IGNORECASE)
_EXCLUSION_RE = re.compile(r"\b(exclusion|exclusions|exclude|excludes|excluding|excl|suppression)\b",
                           re.IGNORECASE)
NO_DATA_TARGETING = "nodatatargeting"

VERTICAL_TAGS = ["Healthcare", "Education", "Auto", "Retail", "Legal", "Home Services",
                 "Financial", "B2B",
                 # Added 2026-10-20 (second pass) so travel, dining, events and
                 # the rest get within-tag recommendations too.
                 "Tourism/Travel", "Restaurants/Food", "Entertainment/Events", "Sports",
                 "Outdoors/Recreation", "Fitness/Wellness", "Insurance", "Real Estate/Movers",
                 "Political/Advocacy", "Military/Government", "Nonprofit/Charity", "Pets",
                 "Telecom/Tech", "Life Stage/Family"]
B2B_GROUPS = ["size_owner", "decision_maker", "industry", "occupation", "business_travel"]
B2B_GROUP_LABELS = {
    "size_owner": "Business size / owners",
    "decision_maker": "Decision makers",
    "industry": "Industry",
    "occupation": "Occupation",
    "business_travel": "Business travel / commercial",
}

BRAND_TERMS_PATH = Path(__file__).parent / "audience_brand_terms.csv"
OVERRIDES_PATH = Path(__file__).parent / "audience_component_overrides.csv"


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------

def clean_display(raw):
    """A component as it should read: typos fixed, underscores and runs of
    whitespace collapsed to single spaces, a leading prefix uppercased. The
    raw spelling is kept separately (`raw_variants`) for traceability."""
    text = raw.strip()
    for pattern, replacement in TYPO_FIXES:
        text = pattern.sub(replacement, text)
    text = re.sub(r"[_\s]+", " ", text).strip(" ,")
    alias = KNOWN_ALIASES.get(normalize(text))
    if alias:
        return alias
    first, _, rest = text.partition(" ")
    if first.upper() in SPLIT_PREFIXES or first.upper() == "POLITICAL":
        text = (first.upper() + (" " + rest if rest else "")).strip()
    return text


def segment_key(raw):
    """The form two spellings are compared in: typos fixed, then letters and
    digits only, lowercased -- the same `normalize` the `audiences` table is
    matched with, so a ranked-list segment lands on its brochure row."""
    return normalize(clean_display(raw))


def _words(text):
    """Lowercase, punctuation to spaces -- for whole-word rule matching."""
    return " " + re.sub(r"[^a-z0-9+]+", " ", clean_display(text).lower()).strip() + " "


def prefix_of(display):
    first = display.split(" ", 1)[0].upper()
    if first == "POLITICAL":
        return "POL"
    return first if first in SPLIT_PREFIXES else ""


def split_line(line):
    """A targeting line's components, in order. Splits only on a comma that
    is followed by a taxonomy prefix."""
    return [p.strip() for p in _SPLIT_RE.split(line or "") if p and p.strip()]


def unwrap_clt(component):
    """(standard_component_or_None, client_text) for a CLT component. The
    client text is returned only so a caller can harvest it as a brand term
    -- it must never be shown."""
    rest = _CLT_HEAD_RE.sub("", component.strip())
    match = _CLT_INNER_RE.search(rest)
    if match:
        client = rest[:match.start(1)].strip(" _-")
        return match.group(1).strip(), client
    return None, rest.strip(" _-")


# ---------------------------------------------------------------------------
# Brand / client terms
# ---------------------------------------------------------------------------

def load_brand_terms(path=BRAND_TERMS_PATH, overrides_path=OVERRIDES_PATH):
    """Brand phrases (lowercase word strings) a shown segment must not
    contain: the committed `audience_brand_terms.csv` plus every
    `exclude_client` row in the overrides registry (the custom part of the
    name, prefix stripped). Missing files degrade to fewer terms."""
    terms = set()
    path = Path(path)
    if path.exists():
        with open(path, encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                term = _words(row.get("term") or "").strip()
                if term:
                    terms.add(term)
    overrides_path = Path(overrides_path)
    if overrides_path.exists():
        with open(overrides_path, encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                if (row.get("action") or "").strip() != "exclude_client":
                    continue
                name = re.sub(r"^custom[\s_]*", "", (row.get("component") or "").strip(),
                              flags=re.IGNORECASE)
                # Retargeting rows are caught structurally already; only a
                # bare client name is worth a phrase.
                if _RETARGETING_RE.search(name):
                    continue
                term = _words(name).strip()
                if term:
                    terms.add(term)
    return terms


def harvest_client_terms(client_texts, generic_text=""):
    """Client names out of CLT and retargeting lines, as brand phrases: the
    first two words of each client text (one word when that's all there
    is, and only if it's distinctive -- four letters or more, not a common
    word). Two words is enough to name a client ("lennar homes", "pima
    medical") without swallowing the descriptive tail ("buy sell home")."""
    common = {"the", "and", "all", "new", "home", "homes", "audience", "audiences", "list",
              "campaign", "custom", "retail", "frequency", "test", "web", "location"}
    terms = set()
    for text in client_texts:
        words = [w for w in _words(text).split() if not w.isdigit()]
        while words and words[0] in common:
            words = words[1:]
        if len(words) >= 2:
            term = " ".join(words[:2])
        elif len(words) == 1 and len(words[0]) >= 4 and words[0] not in common:
            term = words[0]
        else:
            continue
        # A phrase that also turns up in a standard segment is ordinary
        # vocabulary ("credit union", "children in"), not a client's name.
        if len(term.replace(" ", "")) >= 4 and f" {term} " not in generic_text:
            terms.add(term)
    return terms


def contains_brand(display, brand_terms):
    text = _words(display)
    return any(f" {term} " in text for term in brand_terms)


# ---------------------------------------------------------------------------
# Role
# ---------------------------------------------------------------------------

_MODIFIER_RULES = [
    # (prefix or "" for unprefixed, regex on the words after the prefix)
    ("DEMO", r"^ (age|gender|male|males|female|females|men|women|ethnicity|hispanic|"
             r"african american|asian|caucasian|native american|bipoc|language|spanish|"
             r"bilingual|education (level|attainment|high school|hs|some college|college|"
             r"bachelor|graduate|post graduate|associate|less than high school)|"
             r"homeowner|homeowners|renter|renters)\b"),
    ("HH", r"^ (income|hhi|children|child|kids|no children|presence of children|"
           r"with children|home value|dwelling)\b"),
    ("", r"^ (male|males|female|females|men|women|adults?)\b.*\b\d{2}\b"),
]


def classify_role(display):
    """'modifier' or 'core' -- see this module's docstring."""
    prefix = prefix_of(display)
    rest = _words(display[len(display.split(" ", 1)[0]):] if prefix else display)
    rest = " " + rest.strip()
    for rule_prefix, pattern in _MODIFIER_RULES:
        if rule_prefix == prefix and re.search(pattern, rest + " "):
            return "modifier"
    return "core"


# ---------------------------------------------------------------------------
# Tags
# ---------------------------------------------------------------------------

def _rx(*phrases):
    return re.compile(r"\b(" + "|".join(phrases) + r")\b")


# Each vertical: (prefixes that imply it, word patterns that imply it,
# prefixes the word patterns never apply to -- "repair" under AUTO is a car
# repair, "college" under SPORTS is a fan, "insurance" under HLTH is health
# coverage, not a financial product).
_NEVER = {"Home Services": {"AUTO"}, "Education": {"SPORTS"}, "Financial": {"HLTH"}}
TAG_RULES = {
    "Healthcare": (["HLTH"], _rx(
        "health", "healthcare", "medical", "medicare", "medicaid", "doctor", "doctors",
        "nurse", "nurses", "nursing", "dental", "dentist", "hospital", "physician",
        "pharmacy", "prescription", "caregiver", "caregivers", "senior living",
        "home care", "chiropractic", "vision", "eyeglasses", "hearing", "pain",
        "weight loss", "diabetes", "cancer", "heart", "stroke", "depression", "anxiety",
        "therapy", "paramedic", "mesothelioma", "wellness", "fitness", "veins",
        "fibroids", "menopause", "vaccine", "unvaccinated", "uninsured", "wheelchair",
        "mobility", "hair loss", "cosmetic procedure")),
    "Education": ([], _rx(
        "education", "educational", "college", "colleges", "university", "student",
        "students", "school", "schools", "vocational", "tuition", "degree", "teachers",
        "educators", "academic", "graduate", "learning")),
    "Auto": (["AUTO"], _rx(
        "auto", "vehicle", "vehicles", "car", "cars", "truck", "trucks", "suv", "suvs",
        "sedan", "pickup", "dealer", "dealership", "motorcycle", "motorcycles", "rv",
        "tire", "tires", "collision", "body shop", "drivers", "ev", "crossover")),
    "Retail": (["RETAIL"], _rx(
        "shopper", "shoppers", "shopping", "shopaholics", "store", "stores", "furniture",
        "apparel", "clothing", "mattress", "mattresses", "appliance", "appliances",
        "jewelry", "grocery", "retail", "toys", "electronics", "home furnishing",
        "home furnishings", "department", "coupon", "deal seekers", "luxury shopper",
        "luxury shoppers")),
    "Legal": (["LEGAL"], _rx(
        "legal", "lawyer", "lawyers", "attorney", "attorneys", "law", "personal injury",
        "accident", "accidents", "divorce", "divorced", "family law",
        "bankruptcy", "immigration", "mesothelioma", "lawsuit", "estate planning",
        "recently single")),
    "Home Services": ([], _rx(
        "home services", "home service", "home improvement", "renovation", "renovations",
        "remodel", "remodeling", "repair", "hvac", "plumbing", "roofing", "roof",
        "flooring", "windows", "doors", "landscaping", "landscape", "lawn", "pool",
        "pools", "hot tub", "solar", "pest", "homes built", "home built", "home exterior",
        "diy", "diyers", "fixers", "garage", "siding", "gutters", "electrician", "kitchen",
        "bathroom", "home warranty", "security system", "single family", "homeowner",
        "homeowners", "home builder", "new construction home")),
    "Financial": (["FIN"], _rx(
        "bank", "banking", "credit", "credit union", "loan", "loans", "mortgage", "heloc",
        "insurance", "invest", "investor", "investors", "investing", "investment",
        "retirement planning", "wealth", "net worth", "tax", "precious metals",
        "financial", "finance", "savings", "annuity", "annuities", "credit card")),
    # A casino is both a night out and a destination -- it carries both.
    "Tourism/Travel": (["TRAVEL"], _rx(
        "travel", "travels", "traveler", "travelers", "traveller", "travellers", "vacation",
        "vacations", "road trip", "road trips", "resort", "resorts", "hotel", "hotels",
        "staycation", "cruise", "cruises", "visitors", "tourism", "tourist", "tourists",
        "getaway", "getaways", "casino", "casinos", "honeymooners", "sightseers")),
    "Restaurants/Food": (["FOOD"], _rx(
        "foodies", "dining", "diner", "restaurant", "restaurants", "qsr", "bar", "bars",
        "brewery", "breweries", "fast food", "takeout", "food delivery", "wine", "beer")),
    "Entertainment/Events": (["ENT"], _rx(
        "concert", "concerts", "concert goers", "music", "performing arts", "fair", "fairs",
        "festival", "festivals", "comedy", "museum", "museums", "casino", "casinos", "gaming",
        "gambling", "lottery", "sweepstakes", "nightlife", "theater", "theatre", "broadway",
        "movies", "movie", "gamers", "zoo", "waterparks", "amusement", "events")),
    "Sports": (["SPORTS"], _rx(
        "sport", "sports", "golf", "golfer", "golfers", "betting", "nfl", "nba", "wnba", "mlb",
        "nhl", "mls", "nascar", "auto racing", "college football", "college basketball",
        "mavericks", "bills", "seahawks", "spurs", "athlete", "athletes", "super fans")),
    "Outdoors/Recreation": ([], _rx(
        "outdoors", "outdoor", "hunting", "fishing", "camping", "boating", "boats", "boat",
        "rv", "rvs", "skiing", "snowboard", "hiking", "outdoor adventure", "outdoor adventurers",
        "bird watchers", "motocross", "national parks")),
    "Fitness/Wellness": ([], _rx(
        "fitness", "gym", "gyms", "running", "runners", "yoga", "weight loss", "wellness",
        "health fitness", "marathon", "crossfit", "pilates", "fitness fanatics")),
    "Insurance": ([], _rx("insurance")),
    "Real Estate/Movers": ([], _rx(
        "home buyer", "home buyers", "homebuyer", "new movers", "new mover", "movers",
        "renter", "renters", "new construction", "home intender", "home renter intender",
        "buy sell home", "likely to move", "first time home buyer", "real estate")),
    "Political/Advocacy": (["POL"], _rx("voter", "voters", "political", "democratic",
                                        "republican")),
    "Military/Government": ([], _rx(
        "military", "veteran", "veterans", "govt employee", "government", "law and government",
        "gov workers", "federal", "public safety", "first responder")),
    "Nonprofit/Charity": ([], _rx(
        "charity", "charities", "charitable", "volunteerism", "volunteer", "volunteers",
        "donor", "donors", "donation", "causes")),
    "Pets": ([], _rx("pet", "pets", "pet adoption", "dog", "dogs", "cat", "cats", "horses")),
    "Life Stage/Family": ([], _rx(
        "parents", "parent", "mothers", "mother", "moms", "expecting", "expecting mothers",
        "married", "newlyweds", "retired", "retirees", "job seeker", "job seekers",
        "empty nesters", "empty nester", "new parents", "grandparents")),
    "Telecom/Tech": ([], _rx(
        "cord cutters", "cord cutter", "streaming services", "internet", "internet intenders",
        "technology", "tech", "wireless", "cable", "broadband", "cell phone", "smartphone")),
}

# B2B subgroups, checked in order -- the first that matches wins, so a
# business-owner phrase beats the "executives" next to it ("DEMO Career
# Business Owners Executives" is an owner segment), and a decision-maker
# function beats the industry word it sits beside.
_B2B_RULES = [
    ("size_owner", _rx(
        "business owner", "business owners", "small business", "small businesses",
        "business revenue", "revenue", "self employed", "entrepreneur", "entrepreneurs",
        "business size", "smb", "small to medium", "small to midsize",
        "small to midsize business", "mid size business")),
    ("decision_maker", _rx(
        "decision maker", "decision makers", "it decision makers", "human resources",
        "hr", "facility management", "facilities management", "facility maintenance",
        "property project and facility management", "property manager",
        "property management", "insurance prospects and decision makers", "ceo",
        "operations management", "executives", "executive", "exec")),
    ("industry", _rx(
        "construction", "agriculture", "manufacturing", "commercial property",
        "commercial property owner", "real estate commercial", "mid size office",
        "business industry", "b2b")),
    ("business_travel", _rx("business traveler", "business travelers",
                            "owner commercial vehicle", "commercial vehicle")),
]
# Industry phrases specific enough to be B2B under any prefix -- a custom
# build selling to the trades, not a consumer buying a newly built home.
_B2B_ANY_PREFIX = _rx("professional construction supplies", "construction supplies", "b2b")
# Which function a decision-maker segment serves -- recommendations match it
# to the product (IT for VoIP, HR for staffing, facilities for building
# services, insurance for commercial insurance).
_B2B_FUNCTION_RULES = [
    ("it", _rx("it decision makers", "it", "technology", "tech decision")),
    ("hr", _rx("human resources", "hr")),
    ("facilities", _rx("facility", "facilities", "property manager", "property management",
                       "property project", "maintenance")),
    ("insurance", _rx("insurance")),
    ("executive", _rx("ceo", "operations management", "executive", "executives", "exec")),
]
_CAREER_RE = re.compile(r"^ (custom )?(demo )?(career|careers|employment)\b")
# An employment STATUS under DEMO Career is not an occupation a business
# sells to -- "DEMO Career Unemployed" is a consumer, not a trade.
_EMPLOYMENT_STATUS_RE = _rx("unemployed", "employed", "job seeker", "job seekers", "veteran",
                            "caregiver", "retired")


def b2b_group_of(display):
    """(group, function) for a B2B segment, or (None, None).

    Occupation is a DEMO Career segment (or a custom build of one) -- a
    LIFESTYLE "Career Changers" segment is a consumer interest, not an
    occupation. An industry word only makes a segment B2B under a business
    prefix (FIN, B2B) or a phrase specific to selling to a trade: "CUSTOM New
    Construction Home Intender" is a home buyer."""
    prefix = prefix_of(display)
    words = _words(display)
    is_career = bool(_CAREER_RE.search(words))
    business_prefix = prefix in ("FIN", "B2B")
    for group, pattern in _B2B_RULES:
        if not pattern.search(words):
            continue
        if group == "size_owner" and not (business_prefix or is_career or prefix == "CUSTOM"):
            continue
        if group == "decision_maker" and not (business_prefix or is_career
                                               or prefix in ("HLTH", "CUSTOM")):
            continue
        if group == "industry" and not (business_prefix or _B2B_ANY_PREFIX.search(words)):
            continue
        if group == "business_travel" and prefix not in ("TRAVEL", "AUTO"):
            continue
        function = None
        if group == "decision_maker":
            for name, fn_pattern in _B2B_FUNCTION_RULES:
                if fn_pattern.search(words):
                    function = name
                    break
        return group, function
    if _B2B_ANY_PREFIX.search(words):
        return "industry", None
    if prefix == "B2B" or (is_career and not _EMPLOYMENT_STATUS_RE.search(words)):
        return "occupation", None
    return None, None


def tags_for(display):
    """(tags, b2b_group, b2b_function) for one segment name."""
    prefix = prefix_of(display)
    # A CUSTOM build often embeds a second prefix ("CUSTOM AUTO Intenders
    # Expanded") -- it implies the same vertical the bare prefix would.
    inner = ""
    if prefix == "CUSTOM":
        inner = prefix_of(display.split(" ", 1)[1] if " " in display else "")
    words = _words(display)
    tags = []
    for tag, (prefixes, pattern) in TAG_RULES.items():
        never = _NEVER.get(tag, set())
        if prefix in prefixes or inner in prefixes or (
                prefix not in never and inner not in never and pattern.search(words)):
            tags.append(tag)
    group, function = b2b_group_of(display)
    if group:
        tags.append("B2B")
        # A business-function FIN segment is B2B, not consumer finance.
        if prefix == "FIN" and group in ("decision_maker", "industry") and "Financial" in tags:
            tags.remove("Financial")
    return [t for t in VERTICAL_TAGS if t in tags], group, function


# ---------------------------------------------------------------------------
# Parsing a ranked workbook
# ---------------------------------------------------------------------------

class RankedImportError(Exception):
    """Raised with a message aimed at whoever's uploading the list."""


def read_ranked_lines(source):
    """[(row, line), ...] from a one-column ranked workbook, row numbers
    1-based over the data rows in file order (row 1 = most used). Raises
    RankedImportError for anything that isn't that shape."""
    import openpyxl

    try:
        book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise RankedImportError(f"Couldn't open that file as an Excel workbook ({exc}).") from exc
    if not book.sheetnames:
        raise RankedImportError("That workbook has no sheets.")
    rows = book[book.sheetnames[0]].iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        raise RankedImportError("That workbook's sheet is empty.")
    first = str(header[0] or "").strip().lstrip("﻿").lower() if header else ""
    if "segment" not in first:
        raise RankedImportError(
            f"Expected a first column headed \"Segment Name\" -- got {first!r}.")
    lines = []
    for raw in rows:
        if raw is None or not raw or raw[0] is None or not str(raw[0]).strip():
            continue
        lines.append((len(lines) + 1, str(raw[0]).strip()))
    if not lines:
        raise RankedImportError("No targeting lines found in that workbook.")
    return lines


def is_ranked_workbook(source):
    """True when the workbook is the one-column ranked list (no impressions
    column) rather than the older Segment Name / Delivered Impressions
    export."""
    import openpyxl

    try:
        book = openpyxl.load_workbook(source, read_only=True, data_only=True)
        header = next(book[book.sheetnames[0]].iter_rows(values_only=True))
    except Exception:
        return False
    cells = [str(c or "").strip().lower() for c in header]
    return bool(cells) and "segment" in cells[0] and not any("impression" in c for c in cells[1:])


def rank_weight(row):
    return 1.0 / math.log2(row + 1)


def classify_component(raw, brand_terms, client_terms=frozenset(), known_keys=None):
    """(display_or_None, status, flags, client_text) for one raw component.
    display is None when nothing standard survives (a pure CLT list, or one
    wrapping something that isn't a standard segment).

    `brand_terms` (curated) apply to every component; `client_terms`
    (harvested from this list's own CLT/retargeting lines) only to custom
    builds and unprefixed text, since a standard segment's name is Premion's
    taxonomy, not a client's. `known_keys` is the set of segment keys that
    exist outside any client list -- a CLT wrapper's inner text counts only
    when it's one of them ("POL <candidate> 2025" inside a client list is
    that client's own list, not a standard segment)."""
    flags = []
    client_text = None
    text = raw.strip()
    if _CLT_RE.match(text):
        inner, client_text = unwrap_clt(text)
        flags.append("client 1P list")
        if inner is None or (known_keys is not None and segment_key(inner) not in known_keys):
            return None, "excluded", flags, client_text
        text = inner
        flags.append("extracted from client list")
    display = clean_display(text)
    key = segment_key(display)
    if ">" in display:
        return display, "excluded", flags + ["3rd-party taxonomy path"], client_text
    if key == NO_DATA_TARGETING:
        return display, "excluded", flags + ["no targeting"], client_text
    if _RETARGETING_RE.search(display) or _RETARGETING_RE.search(raw):
        return display, "excluded", flags + ["retargeting/campaign-specific"], client_text
    if _EXCLUSION_RE.search(display):
        return display, "excluded", flags + ["exclusion"], client_text
    prefix = prefix_of(display)
    if contains_brand(display, brand_terms) or (
            prefix in ("", "CUSTOM") and contains_brand(display, client_terms)):
        status = "hidden"
        flags.append("names a client or brand")
        if _CUSTOM_RE.match(display):
            flags.append("custom build")
        return display, status, flags, client_text
    if _CUSTOM_RE.match(display):
        return display, "custom", flags + ["custom build"], client_text
    return display, "standard", flags, client_text


def _client_text_of_retargeting(raw):
    """The client name in a retargeting component ("WEB RT - The Skin
    Center" -> "The Skin Center", "Pima Medical Institute_RFPID-..._RT" ->
    "Pima Medical Institute"), for brand harvesting only."""
    text = re.sub(r"^(WEB|LOC(ATION)?)\s*RT\s*-?\s*", "", raw.strip(), flags=re.IGNORECASE)
    text = re.split(r"_RFPID|RFPID|_RT\b|\bRT\b|Address List|Web Retargeting|Website RT|_Web",
                    text, flags=re.IGNORECASE)[0]
    return text.strip(" _-")


def derive_segment_catalog(lines, brand_terms=None, category_overrides=None,
                           existing_keys=(), source_filename=None, pairings_per_segment=5):
    """The derived catalog for a ranked list -- see this module's docstring.
    `lines` is [(row, line), ...] from `read_ranked_lines`. Returns a plain,
    JSON-serializable dict:

        {"format", "source_filename", "lines", "segments": [...], "summary": {...}}

    Each segment: key, segment (display), raw_variants, prefix, status,
    flags, role, tags, b2b_group, b2b_function, score, line_count,
    best_row, pairings ([[modifier display, count], ...]).
    `category_overrides` is {normalized: category} (the registry's
    `categorize` rows) for a component no prefix can categorize.
    `existing_keys` are the current catalog's segment keys (the brochure and
    earlier merges) -- a CLT list wrapping one of them counts it.
    """
    if brand_terms is None:
        brand_terms = load_brand_terms()
    brand_terms = set(brand_terms)
    category_overrides = category_overrides or {}

    # Pass 1: harvest client names from CLT and retargeting components so a
    # custom build carrying the same name is hidden too, and collect which
    # segments exist outside any client list.
    harvested = []
    known_keys = set(existing_keys)
    generic = []
    for _row, line in lines:
        for raw in split_line(line):
            if _CLT_RE.match(raw):
                harvested.append(unwrap_clt(raw)[1])
            elif _RETARGETING_RE.search(raw) and not _CUSTOM_RE.match(raw.strip()):
                harvested.append(_client_text_of_retargeting(raw))
            else:
                known_keys.add(segment_key(raw))
                if prefix_of(clean_display(raw)) not in ("", "CUSTOM", "WEB", "CLT"):
                    generic.append(_words(raw))
    client_terms = harvest_client_terms(harvested, " ".join(generic))

    seg = {}
    line_members = []
    excluded_counts = Counter()
    excluded_keys = set()
    clt_components = 0
    clt_extracted = 0
    for row, line in lines:
        members = []
        for raw in split_line(line):
            display, status, flags, _client = classify_component(
                raw, brand_terms, client_terms, known_keys)
            if "client 1P list" in flags:
                clt_components += 1
                if display is not None:
                    clt_extracted += 1
            if display is None or status == "excluded":
                excluded_keys.add(segment_key(display) if display else "clt:" + segment_key(raw))
                for flag in flags:
                    if flag not in ("client 1P list", "extracted from client list"):
                        excluded_counts[flag] += 1
                if display is None:
                    excluded_counts["client 1P list"] += 1
                continue
            key = segment_key(display)
            entry = seg.setdefault(key, {
                "key": key, "variants": Counter(), "raw_variants": set(), "status": status,
                "flags": set(), "rows": set()})
            entry["variants"][display] += 1
            # Raw spellings for traceability -- a CLT wrapper's client text
            # never goes in, only the extracted standard segment's own text.
            entry["raw_variants"].add(display if "client 1P list" in flags else raw.strip())
            entry["flags"].update(f for f in flags if f != "client 1P list")
            if "client 1P list" not in flags:
                entry["standalone"] = True
            entry["rows"].add(row)
            members.append(key)
        line_members.append((row, list(dict.fromkeys(members))))

    segments = []
    for key, entry in seg.items():
        display = sorted(entry["variants"], key=lambda d: (-entry["variants"][d], d))[0]
        if entry.get("standalone"):
            # Booked on its own too -- a client list wrapping it is no
            # longer the interesting fact about it.
            entry["flags"].discard("extracted from client list")
        role = classify_role(display)
        tags, group, function = tags_for(display)
        rows = sorted(entry["rows"])
        prefix = prefix_of(display)
        category = category_overrides.get(key, "")
        if not category:
            if prefix == "CUSTOM":
                inner = prefix_of(display.split(" ", 1)[1] if " " in display else "")
                category = "" if inner in ("", "CUSTOM", "CLT", "WEB") else inner
            elif prefix not in ("", "WEB", "CLT"):
                category = prefix
        if category == "HLH":
            category = "HLTH"
        segments.append({
            "key": key, "segment": display, "raw_variants": sorted(entry["raw_variants"]),
            "prefix": prefix, "category": category, "status": entry["status"],
            "flags": sorted(entry["flags"]), "role": role, "tags": tags,
            "b2b_group": group, "b2b_function": function,
            "score": round(sum(rank_weight(r) for r in rows), 6),
            "line_count": len(rows), "best_row": rows[0], "pairings": [],
        })
    by_key = {s["key"]: s for s in segments}

    # Co-occurrence: for every core segment and every custom recipe, the
    # modifiers it shares a line with most often.
    pair_counts = defaultdict(Counter)
    for _row, members in line_members:
        modifiers = [m for m in members if by_key[m]["role"] == "modifier"
                     and by_key[m]["status"] == "standard"]
        for member in members:
            entry = by_key[member]
            if entry["role"] == "modifier" and entry["status"] == "standard":
                continue
            for modifier in modifiers:
                pair_counts[member][modifier] += 1
    for key, counts in pair_counts.items():
        ranked = sorted(counts.items(), key=lambda kv: (-kv[1], -by_key[kv[0]]["score"]))
        by_key[key]["pairings"] = [[by_key[m]["segment"], n]
                                   for m, n in ranked[:pairings_per_segment]]

    segments.sort(key=lambda s: (-s["score"], s["segment"]))
    # Each line as the segments a seller may see (standard and custom only)
    # -- what the booking-evidence panel indexes. Client lists, retargeting
    # and client-named builds are already gone, so no client text is kept.
    stacks = []
    for _row, members in line_members:
        shown = [by_key[m]["segment"] for m in members
                 if by_key[m]["status"] in ("standard", "custom")]
        if shown:
            stacks.append(shown)
    return {
        "format": FORMAT,
        "source_filename": source_filename,
        "lines": len(lines),
        "segments": segments,
        "stacks": stacks,
        "summary": summarize(segments, excluded_counts, clt_components, clt_extracted,
                             excluded_unique=len(excluded_keys - set(by_key))),
    }


def summarize(segments, excluded_counts=None, clt_components=0, clt_extracted=0,
              excluded_unique=0):
    status = Counter(s["status"] for s in segments)
    shown = [s for s in segments if s["status"] in ("standard", "custom")]
    standard = [s for s in segments if s["status"] == "standard"]
    tags = Counter(t for s in shown for t in s["tags"])
    b2b = Counter(s["b2b_group"] for s in shown if s["b2b_group"])
    return {
        "unique_segments": len(segments) + excluded_unique,
        "excluded_unique": excluded_unique,
        "recommendable": len(standard),
        "status": dict(status),
        "roles": dict(Counter(s["role"] for s in standard)),
        "categories": dict(Counter(s["category"] or "(none)" for s in shown)),
        "tags": dict(tags),
        "b2b_groups": dict(b2b),
        "untagged_standard": sum(1 for s in standard if not s["tags"]),
        "excluded_components": dict(excluded_counts or {}),
        "clt_components": clt_components,
        "clt_standard_extracted": clt_extracted,
    }


# ---------------------------------------------------------------------------
# Corrections and the review CSV
# ---------------------------------------------------------------------------

def apply_overrides(segment, overrides):
    """One derived segment with an admin correction applied -- `overrides`
    is {key: {"tags": [...], "role": "...", "hidden": bool,
    "b2b_group": "..."}}, any field optional. Returns a new dict."""
    correction = (overrides or {}).get(segment["key"])
    if not correction:
        return segment
    out = dict(segment)
    if "tags" in correction:
        out["tags"] = [t for t in VERTICAL_TAGS if t in (correction["tags"] or [])]
        if "B2B" not in out["tags"]:
            out["b2b_group"] = None
            out["b2b_function"] = None
    if correction.get("b2b_group") in B2B_GROUPS:
        out["b2b_group"] = correction["b2b_group"]
    if correction.get("role") in ("core", "modifier"):
        out["role"] = correction["role"]
    if "hidden" in correction:
        if correction["hidden"]:
            out["status"] = "hidden"
        elif out["status"] == "hidden":
            out["status"] = "custom" if "custom build" in out.get("flags", []) else "standard"
    return out


REVIEW_COLUMNS = ["segment", "status", "role", "category", "tags", "b2b_group",
                  "b2b_function", "score", "line_count", "best_row", "top_pairing",
                  "flags", "raw_variants"]


def write_review_csv(derived, path):
    """The reviewable tagging output: one row per derived segment, best
    score first. A pure-CLT list never appears -- it isn't a segment."""
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=REVIEW_COLUMNS)
        writer.writeheader()
        for s in derived["segments"]:
            writer.writerow({
                "segment": s["segment"], "status": s["status"], "role": s["role"],
                "category": s["category"], "tags": "; ".join(s["tags"]),
                "b2b_group": s["b2b_group"] or "", "b2b_function": s["b2b_function"] or "",
                "score": f"{s['score']:.4f}", "line_count": s["line_count"],
                "best_row": s["best_row"],
                "top_pairing": s["pairings"][0][0] if s["pairings"] else "",
                "flags": "; ".join(s["flags"]), "raw_variants": " | ".join(s["raw_variants"]),
            })


# ---------------------------------------------------------------------------
# One entry point for the admin page and the ingest script
# ---------------------------------------------------------------------------

def load_category_overrides(path=OVERRIDES_PATH):
    """{segment key: category} from the overrides registry's `categorize`
    rows -- a category for a component no prefix can give one. (Its
    `split` rows are deliberately not used: a custom build is one recipe
    here, never split on OR.)"""
    path = Path(path)
    if not path.exists():
        return {}
    out = {}
    with open(path, encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            if (row.get("action") or "").strip() == "categorize" and row.get("category"):
                out[segment_key(row.get("component") or "")] = row["category"].strip()
    return out


def derive_from_workbook(source, existing_segments=(), source_filename=None):
    """Read a ranked workbook and derive its catalog, with the committed
    brand terms and category overrides. `existing_segments` are the current
    catalog's segment names. Raises RankedImportError."""
    lines = read_ranked_lines(source)
    return derive_segment_catalog(
        lines, category_overrides=load_category_overrides(),
        existing_keys={segment_key(s) for s in existing_segments},
        source_filename=source_filename)
