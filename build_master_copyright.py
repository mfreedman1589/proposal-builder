"""Proposal master deck: tidy two stale copyright marks.

1. The media plan slides (every slide with `{{PLAN_TITLE}}`) draw a
   "(c) PREMION 2021" text box from their slide MASTER -- not their layout --
   and it now sits behind the TEGNA terms block. Only the plan slides use
   that master, so the text box is removed from the master itself.
2. "(c) PREMION 2024" text boxes placed directly on slides become
   "(c) PREMION 2026". Only the year changes, run by run, so the formatting
   is untouched.

Run against the CURRENT master (version 11 at the time of writing):
    python build_master_copyright.py <master>.pptx <out>.pptx
then upload the output on dev only (never activate from dev -- the nightly
merge does that). Idempotent: a second run changes nothing.
"""
import re
import sys

from pptx import Presentation

import slide_map

STALE_SLIDE_MARK = re.compile(r"^©\s*PREMION\s+2024$")
NEW_YEAR = "2026"


def plan_slide_masters(prs):
    masters = []
    for slide in prs.slides:
        if any(sh.has_text_frame and "{{PLAN_TITLE}}" in sh.text_frame.text
               for sh in slide_map.iter_all_shapes(slide.shapes)):
            master = slide.slide_layout.slide_master
            if all(master is not m for m in masters):
                masters.append(master)
    return masters


def remove_master_copyright(prs):
    removed = 0
    for master in plan_slide_masters(prs):
        users = [s for s in prs.slides if s.slide_layout.slide_master is master]
        if any("{{PLAN_TITLE}}" not in "".join(sh.text_frame.text for sh in
                                                slide_map.iter_all_shapes(s.shapes)
                                                if sh.has_text_frame) for s in users):
            raise SystemExit("a non-plan slide shares the plan slides' master -- "
                             "removing its copyright would change that slide too")
        for shape in list(master.shapes):
            if shape.has_text_frame and shape.text_frame.text.strip().startswith("©"):
                print(f"master: removed {shape.name!r} ({shape.text_frame.text.strip()!r})")
                shape._element.getparent().remove(shape._element)
                removed += 1
    return removed


def update_slide_marks(prs):
    updated = 0
    for index, slide in enumerate(prs.slides, start=1):
        for shape in slide_map.iter_all_shapes(slide.shapes):
            if not (shape.has_text_frame
                    and STALE_SLIDE_MARK.match(shape.text_frame.text.strip())):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace("2024", NEW_YEAR)
            print(f"slide {index} ({slide_map.notes_key(slide)}): "
                  f"{shape.text_frame.text.strip()!r}")
            updated += 1
    return updated


def main(src, out):
    prs = Presentation(src)
    removed = remove_master_copyright(prs)
    updated = update_slide_marks(prs)
    prs.save(out)
    print(f"{removed} master mark(s) removed, {updated} slide mark(s) updated -> {out}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
